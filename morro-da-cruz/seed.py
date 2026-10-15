"""Carga inicial do banco a partir das planilhas oficiais auditadas.

Uso:  python seed.py [--db dados/morro_da_cruz.db]

Fontes (pasta planilhas/, ver config/fontes.json):
  06.1 REDES REV01 ............ trechos (aba REDE)
  06.2 CONEXÕES REV08 (.xlsm) . quantidades/unidades/avisos AA·AB das peças (aba CONEXÕES)
  EVIDÊNCIAS REV08 ............ estrutura das peças, conjuntos e elementos (enquanto a 06.2 não estiver na pasta)
  BM REV04 IDENTIFICAÇÃO ...... BM Juntada e Vertical
  03.3.2 BM REV03 REAJUSTE .... aba REAJUSTE (data-base, aniversário, índice CEFR)

Ao final confere os totais com a auditoria e sai com código 1 se algo não bater.
"""
import argparse
import json
import math
import os
import sys
from datetime import date, datetime

import openpyxl
from openpyxl.utils import column_index_from_string as col

import regras as R
import servicos as S

BASE = os.path.dirname(os.path.abspath(__file__))


def carregar_config():
    with open(os.path.join(BASE, "config", "fontes.json"), encoding="utf-8") as f:
        return json.load(f)


def caminho(cfg, chave):
    p = os.path.join(BASE, "planilhas", cfg["arquivos"][chave])
    return p if os.path.exists(p) else None


def abrir(p):
    return openpyxl.load_workbook(p, read_only=True, data_only=True)


def v(row, letra):
    i = col(letra) - 1
    return row[i] if i < len(row) else None


def num(x):
    return float(x) if isinstance(x, (int, float)) else None


def txt(x):
    if x is None:
        return None
    s = str(x).strip()
    return s or None


# ------------------------------------------------------------------ UTM (SIRGAS 2000 / UTM 23S) -> lat/lon
def utm_para_latlon(e, n, zona=23, sul=True):
    a, f = 6378137.0, 1 / 298.257222101
    k0, e2 = 0.9996, f * (2 - f)
    ep2 = e2 / (1 - e2)
    x = e - 500000.0
    y = n - (10000000.0 if sul else 0)
    m = y / k0
    mu = m / (a * (1 - e2 / 4 - 3 * e2 ** 2 / 64 - 5 * e2 ** 3 / 256))
    e1 = (1 - math.sqrt(1 - e2)) / (1 + math.sqrt(1 - e2))
    p1 = (mu + (3 * e1 / 2 - 27 * e1 ** 3 / 32) * math.sin(2 * mu) + (21 * e1 ** 2 / 16 - 55 * e1 ** 4 / 32) * math.sin(4 * mu)
          + (151 * e1 ** 3 / 96) * math.sin(6 * mu) + (1097 * e1 ** 4 / 512) * math.sin(8 * mu))
    n1 = a / math.sqrt(1 - e2 * math.sin(p1) ** 2)
    t1, c1 = math.tan(p1) ** 2, ep2 * math.cos(p1) ** 2
    r1 = a * (1 - e2) / (1 - e2 * math.sin(p1) ** 2) ** 1.5
    d = x / (n1 * k0)
    lat = p1 - (n1 * math.tan(p1) / r1) * (d ** 2 / 2 - (5 + 3 * t1 + 10 * c1 - 4 * c1 ** 2 - 9 * ep2) * d ** 4 / 24
                                           + (61 + 90 * t1 + 298 * c1 + 45 * t1 ** 2 - 252 * ep2 - 3 * c1 ** 2) * d ** 6 / 720)
    lon = (d - (1 + 2 * t1 + c1) * d ** 3 / 6 + (5 - 2 * c1 + 28 * t1 - 3 * c1 ** 2 + 8 * ep2 + 24 * t1 ** 2) * d ** 5 / 120) / math.cos(p1)
    lon0 = math.radians((zona - 1) * 6 - 180 + 3)
    return round(math.degrees(lat), 7), round(math.degrees(lon0 + lon), 7)


# ------------------------------------------------------------------ cargas
def carregar_equipes(conn):
    nomes = {f"Zona{n:02d}": f"EQUIPE ZONA {n:02d}" for n in range(1, 12)}
    nomes.update({"ADUTORA": "EQUIPE ADUTORA", "CH": "EQUIPE CH — CAIXAS HIDRÁULICAS"})
    for eid in R.FRENTES:
        conn.execute("INSERT OR REPLACE INTO equipes(id, nome) VALUES (?,?)", (eid, nomes[eid]))


COLUNAS_SERVICO = ["AL", "AM", "AN", "AO", "AP", "AQ", "AR", "AS", "AT", "AU", "AV", "AW", "AX", "AY", "AZ",
                   "BA", "BB", "BC", "BD", "BE", "BF", "BG", "BH", "BI", "BJ", "BK"]


def carregar_trechos(conn, p):
    ws = abrir(p)["REDE"]
    n = 0
    for row in ws.iter_rows(min_row=6, values_only=True):
        pid = txt(v(row, "R"))
        if not pid:
            continue
        quant = {c: num(v(row, c)) for c in COLUNAS_SERVICO if num(v(row, c)) is not None}
        e1, n1, e2, n2 = (num(v(row, c)) for c in ("BS", "BT", "BU", "BV"))
        prof, larg, lastro = num(v(row, "AF")), num(v(row, "AG")), num(v(row, "AH"))
        pend = None
        if any(str(v(row, c) or "").strip().upper() == R.PERFIL_PENDENTE for c in ("AE", "AF", "AG", "AH")):
            pend = R.PERFIL_PENDENTE
        conn.execute(
            """INSERT INTO trechos(id,item,frente_id,equipe,tipo,pb,prancha,area,de_mm,material,solda,metodologia,classe,
                 extensao_m,codigo_bm,recobrimento_m,profundidade_m,largura_m,lastro_m,perfil_status,quantidades,
                 no_ini,no_fim,handle,utm_ini,utm_fim,lat,lon,observacao)
               VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            (pid, v(row, "A"), txt(v(row, "K")), txt(v(row, "BM")), txt(v(row, "BN")), txt(v(row, "BO")), txt(v(row, "U")),
             txt(v(row, "X")), num(v(row, "H")), txt(v(row, "I")), txt(v(row, "G")), txt(v(row, "N")), txt(v(row, "S")),
             num(v(row, "M")), txt(v(row, "AC")), num(v(row, "AE")), prof, larg, lastro, pend,
             json.dumps(quant), txt(v(row, "BP")), txt(v(row, "BQ")), txt(v(row, "BR")),
             json.dumps([e1, n1]) if e1 else None, json.dumps([e2, n2]) if e2 else None,
             num(v(row, "BW")), num(v(row, "BX")), txt(v(row, "AB"))))
        n += 1
    return n


def _unidade_regra(descricao):
    d = (descricao or "").upper()
    return "M" if d.startswith("TUBO") or d.startswith("ESCADA") else "UN"


def numero_projeto(nomenclatura):
    """Zona01_PB01_CH_DF-33 -> DF-33."""
    return nomenclatura.split("_")[-1]


def carregar_conexoes_evidencias(conn, p):
    wb = abrir(p)
    elem = {}
    for row in wb["ELEMENTOS_ZONA"].iter_rows(min_row=2, values_only=True):
        if not row[0]:
            continue
        noms = [s.strip() for s in str(row[1]).split("|")] if row[1] else []
        eid = row[0]   # identificador único da auditoria (SUBDMC1_NÓ1.8); o número no projeto se repete entre PBs
        e, n = num(row[3]), num(row[4])
        lat, lon = utm_para_latlon(e, n) if e and n else (None, None)
        elem[eid] = noms
        conn.execute(
            "INSERT INTO elementos(id,numero,nomenclaturas,tipo,pb,zona,utm_e,utm_n,lat,lon,status_posicao) VALUES (?,?,?,?,?,?,?,?,?,?,?)",
            (eid, numero_projeto(eid), " | ".join(noms), row[2], txt(row[6]), txt(row[5]), e, n, lat, lon, txt(row[8])))
    conj = set()
    npecas = 0
    for row in wb["DE-PARA_CONEXOES"].iter_rows(min_row=2, values_only=True):
        if not row[0]:
            continue
        nom, eid = row[2], row[1]
        if eid not in elem:
            raise SystemExit(f"Elemento {eid} da linha {row[0]} não existe em ELEMENTOS_ZONA.")
        if nom not in conj:
            conn.execute("INSERT INTO conjuntos(id,elemento_id,frente_id) VALUES (?,?,?)", (nom, eid, row[14]))
            conj.add(nom)
        conn.execute(
            """INSERT INTO pecas(conjunto_id,descricao,diametro,unidade,unidade_fonte,quantidade,trechos,linha_origem)
               VALUES (?,?,?,?, 'REGRA', NULL, ?, ?)""",
            (nom, row[6], txt(row[7]), _unidade_regra(row[6]), txt(row[15]), row[0]))
        npecas += 1
    return len(elem), len(conj), npecas


def carregar_conexoes_062(conn, p):
    """Completa quantidade, unidade, código BM e avisos AA/AB a partir da 06.2 xlsm (aba CONEXÕES).

    Colunas localizadas pelo cabeçalho (linha com 'DESCRIÇÃO' e 'QUANT'), casando pela linha da planilha.
    """
    wb = abrir(p)
    ws = wb["CONEXÕES"] if "CONEXÕES" in wb.sheetnames else wb.worksheets[0]
    cab, ini = None, None
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=15, values_only=True), 1):
        h = [str(c or "").strip().upper() for c in row]
        if any("DESCRI" in c for c in h) and any(c.startswith("QUANT") for c in h):
            cab, ini = h, i + 1
            break
    if not cab:
        raise SystemExit("06.2: cabeçalho da aba CONEXÕES não encontrado.")

    def idx(*nomes):
        for n in nomes:
            for j, c in enumerate(cab):
                if c == n or c.startswith(n):
                    return j
        return None

    iq, iu, ibm = idx("QUANT"), idx("UNID"), idx("CÓDIGO BM", "COD. BM", "CÓD. BM")
    iaa, iab = col("AA") - 1, col("AB") - 1
    if iq is None:
        raise SystemExit("06.2: coluna de quantidade não encontrada.")
    un_bm = {r["item_bm"]: r["unidade"] for r in conn.execute("SELECT item_bm, unidade FROM bm_itens WHERE planilha='VERTICAL'")}
    n = 0
    for k, row in enumerate(ws.iter_rows(min_row=ini, values_only=True), ini):
        q = num(row[iq]) if iq < len(row) else None
        if q is None:
            continue
        bm = txt(row[ibm]) if ibm is not None else None
        un = txt(row[iu]) if iu is not None else None
        fonte = "06.2"
        if not un:
            un, fonte = (un_bm.get(bm), "BM") if bm in un_bm else (None, None)
        cur = conn.execute(
            """UPDATE pecas SET quantidade=?, codigo_bm=?, aviso_problema=?, aviso_sugestao=?,
                 unidade=COALESCE(?, unidade), unidade_fonte=COALESCE(?, unidade_fonte)
               WHERE linha_origem=?""",
            (q, bm, txt(row[iaa]) if iaa < len(row) else None, txt(row[iab]) if iab < len(row) else None,
             R.normalizar_unidade(un) if un else None, fonte, k))
        n += cur.rowcount
    return n


def carregar_bm(conn, p):
    wb = abrir(p)
    # Juntada: A item, B:J atributos, K código Caesb, L classificação, M descrição, N unid, O quant, P preço, Q total
    nj = 0
    for k, row in enumerate(wb["BM MORRO DA CRUZ"].iter_rows(min_row=6, values_only=True), 6):
        if not (row[0] and row[1]):
            continue
        atr = dict(zip(["descricao_resumida", "pressao", "resina", "tipo_conexao", "solda", "diametro", "material", "grupo",
                        "grupo_orcamento_ankara"], row[1:10]))
        conn.execute(
            """INSERT INTO bm_itens(planilha,item,item_bm,codigo_caesb,classificacao,descricao,unidade,quantidade,
                 preco_unitario,valor_total,atributos,linha_origem) VALUES ('JUNTADA',?,?,?,?,?,?,?,?,?,?,?)""",
            (str(row[0]), str(row[0]), txt(row[10]), txt(row[11]), row[12], txt(row[13]), num(row[14]), num(row[15]),
             num(row[16]), json.dumps(atr, ensure_ascii=False, default=str), k))
        nj += 1
    # Vertical: B subgrupo (cron. FF), C nome, D item, E item BM, F:N atributos, O cód. Caesb, P classif., Q desc., R unid, S quant, T preço, U total
    nv = 0
    for k, row in enumerate(wb["BM MORRO DA CRUZ - VERTICAL"].iter_rows(min_row=6, values_only=True), 6):
        if not row[3]:
            continue
        atr = dict(zip(["descricao_resumida", "pressao", "resina", "tipo_conexao", "solda", "diametro", "material", "grupo",
                        "grupo_orcamento_ankara"], row[5:14]))
        conn.execute(
            """INSERT INTO bm_itens(planilha,item,item_bm,subgrupo,subgrupo_nome,codigo_caesb,classificacao,descricao,unidade,
                 quantidade,preco_unitario,valor_total,atributos,linha_origem) VALUES ('VERTICAL',?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            (row[3], str(row[4]), txt(row[1]), txt(row[2]), txt(row[14]), txt(row[15]), row[16], txt(row[17]), num(row[18]),
             num(row[19]), num(row[20]), json.dumps(atr, ensure_ascii=False, default=str), k))
        nv += 1
    # rodapé: BDI, K, P0
    rod = {}
    for row in wb["BM MORRO DA CRUZ - VERTICAL"].iter_rows(min_row=6, values_only=True):
        cel = [x for x in row if x not in (None, "")]
        rots = [i for i, x in enumerate(cel) if isinstance(x, str)]
        if not rots:
            continue
        rot = str(cel[rots[-1]]).strip()
        vals = [x for x in cel[rots[-1] + 1:] if isinstance(x, (int, float))]
        if rot.startswith("BDI MATERIAL") and vals:
            rod["bdi_material"] = vals[0]
        elif rot.startswith("BDI SERVI") and vals:
            rod["bdi_servico"] = vals[0]
        elif rot.startswith("COEF. - K") and vals:
            rod["coef_k"] = vals[0]
        elif rot.startswith("TOTAL 01 + 02 + BDI -K") and vals:
            rod["p0"] = vals[0]
    for kk, vv in rod.items():
        conn.execute("INSERT OR REPLACE INTO parametros VALUES (?,?)", (kk, str(vv)))
    return nj, nv, rod


def carregar_reajuste(conn, p):
    ws = abrir(p)["REAJUSTE"]
    achados = {}
    for row in ws.iter_rows(values_only=True):
        rot = str(row[0] or "")
        vals = [x for x in row[1:] if x is not None]
        if rot.startswith("Data-base do reajuste") and vals:
            achados["data_base"] = vals[0]
        elif rot.startswith("Percentual do 1º reajuste") and vals:
            achados["indice"] = vals[0]
        elif rot.startswith("1º aniversário") and vals:
            achados["aniversario"] = vals[0]
        elif rot.startswith("Início da execução") and vals:
            achados["inicio_os"] = vals[0]
    for k in ("data_base", "indice", "aniversario"):
        if k not in achados:
            raise SystemExit(f"REAJUSTE: '{k}' não encontrado na aba.")
    aniv = achados["aniversario"].date() if isinstance(achados["aniversario"], datetime) else achados["aniversario"]
    prox = date(aniv.year + 1, aniv.month, aniv.day)
    conn.execute("INSERT OR REPLACE INTO reajuste_indices VALUES (?,?,?,?)",
                 (aniv.isoformat(), (prox.fromordinal(prox.toordinal() - 1)).isoformat(), float(achados["indice"]),
                  "1º reajuste — Despacho 322/2026-CEFR (ICC-Brasília-DI); a partir de 31/01/2027 pendente de novo índice"))
    for k, val in achados.items():
        conn.execute("INSERT OR REPLACE INTO parametros VALUES (?,?)",
                     (f"reajuste_{k}", val.date().isoformat() if isinstance(val, datetime) else str(val)))
    return achados


def carregar_materiais(conn, cfg):
    """Almoxarifado: tubos PEAD da rede principal (fornecimento Caesb, Seção 6 do SL0327 — não estão
    no BM, valor R$ 0,00; quantidade de projeto = extensão da 06.1 por DE), itens MATERIAL do BM
    Vertical (aquisição Ankara) e materiais de consumo diário."""
    nc = 0
    for r in conn.execute("SELECT de_mm, SUM(extensao_m) ext FROM trechos GROUP BY de_mm ORDER BY de_mm").fetchall():
        de = int(r["de_mm"])
        conn.execute(
            "INSERT INTO materiais(id,descricao,unidade,propriedade,preco_unitario,qtd_projeto) VALUES (?,?,'M','CAESB',0,?)",
            (f"CAESB-PEAD-DE{de}", f"TUBO PEAD DE{de} PN10 — fornecimento Caesb", round(r["ext"], 2)))
        nc += 1
    na = 0
    for b in conn.execute("SELECT * FROM bm_itens WHERE planilha='VERTICAL' AND classificacao='MATERIAL'").fetchall():
        conn.execute(
            "INSERT INTO materiais(id,descricao,unidade,propriedade,preco_unitario,bm_item_id,qtd_projeto) VALUES (?,?,?,'ANKARA',?,?,?)",
            (b["item"], b["descricao"], b["unidade"] or "UN", b["preco_unitario"] or 0, b["id"], b["quantidade"]))
        na += 1
    for m in cfg["materiais_consumo"]:
        conn.execute("INSERT INTO materiais(id,descricao,unidade,propriedade) VALUES (?,?,?,'CONSUMO')",
                     (m["id"], m["descricao"], m["unidade"]))
    return nc, na


# ------------------------------------------------------------------ auditoria
def auditar(conn, cfg):
    esperado = cfg["auditoria"]
    q = lambda sql: conn.execute(sql).fetchone()[0]
    obtido = {
        "trechos": q("SELECT COUNT(*) FROM trechos"),
        "extensao_m": round(q("SELECT SUM(extensao_m) FROM trechos"), 2),
        "elementos": q("SELECT COUNT(*) FROM elementos"),
        "nos": q("SELECT COUNT(*) FROM elementos WHERE tipo='NÓ'"),
        "caixas": q("SELECT COUNT(*) FROM elementos WHERE tipo<>'NÓ'"),
        "conjuntos": q("SELECT COUNT(*) FROM conjuntos"),
        "pecas": q("SELECT COUNT(*) FROM pecas"),
        "bm_vertical": q("SELECT COUNT(*) FROM bm_itens WHERE planilha='VERTICAL'"),
        "bm_juntada": q("SELECT COUNT(*) FROM bm_itens WHERE planilha='JUNTADA'"),
        "p0": round(float(S.parametro(conn, "p0", 0)), 2),
        "sem_posicao": sorted(r[0] for r in conn.execute("SELECT numero FROM elementos WHERE lat IS NULL")),
    }
    ok = True
    print("\nAUDITORIA DA CARGA")
    for k, ev in esperado.items():
        st = "OK " if obtido[k] == ev else "ERRO"
        ok &= obtido[k] == ev
        print(f"  [{st}] {k:<12} esperado {ev!s:<28} obtido {obtido[k]}")
    sem_q = q("SELECT COUNT(*) FROM pecas WHERE quantidade IS NULL")
    if sem_q:
        print(f"  [PEND] {sem_q} linhas de peças sem quantidade — coloque a 06.2 REV08 (.xlsm) em planilhas/ e rode de novo.")
    pend_bm = [s["servico"] for s in json.load(open(os.path.join(BASE, "config", "servicos_rede.json"), encoding="utf-8"))["gaveta1"]
               if s["status"] != "CONFIRMADO"]
    print(f"  [PEND] {len(pend_bm)} serviços da 06.1 sem item de BM confirmado (config/servicos_rede.json).")
    return ok, obtido


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--db", default=S.DB_PADRAO)
    args = ap.parse_args()
    cfg = carregar_config()
    faltando = [k for k in ("redes", "bm", "reajuste") if not caminho(cfg, k)]
    if faltando:
        raise SystemExit("Faltam planilhas obrigatórias em planilhas/: " + ", ".join(cfg["arquivos"][k] for k in faltando))
    if not caminho(cfg, "evidencias"):
        raise SystemExit("Faltam as EVIDÊNCIAS REV08 em planilhas/ (estrutura das conexões).")
    if os.path.exists(args.db):
        os.remove(args.db)
    conn = S.conectar(args.db)
    with S.transacao(conn):
        conn.execute("INSERT INTO parametros VALUES ('contrato', ?)", (cfg["contrato"],))
        carregar_equipes(conn)
        nj, nv, rod = carregar_bm(conn, caminho(cfg, "bm"))
        print(f"BM: {nj} itens Juntada, {nv} itens Vertical, BDI mat {rod.get('bdi_material')}, serv {rod.get('bdi_servico')}, K {rod.get('coef_k')}")
        aj = carregar_reajuste(conn, caminho(cfg, "reajuste"))
        print(f"REAJUSTE: base {aj['data_base']:%d/%m/%Y}, 1º aniversário {aj['aniversario']:%d/%m/%Y}, índice {aj['indice']:.2%}")
        print(f"06.1: {carregar_trechos(conn, caminho(cfg, 'redes'))} trechos")
        ne, nc, npc = carregar_conexoes_evidencias(conn, caminho(cfg, "evidencias"))
        print(f"Conexões (estrutura): {ne} elementos, {nc} conjuntos, {npc} linhas de peças")
        if caminho(cfg, "conexoes"):
            print(f"06.2: quantidades em {carregar_conexoes_062(conn, caminho(cfg, 'conexoes'))} linhas")
        nc, na = carregar_materiais(conn, cfg)
        print(f"Materiais: {nc} tubos PEAD Caesb, {na} itens MATERIAL do BM (Ankara), {len(cfg['materiais_consumo'])} de consumo")
    ok, _ = auditar(conn, cfg)
    sys.exit(0 if ok else 1)


if __name__ == "__main__":
    main()
