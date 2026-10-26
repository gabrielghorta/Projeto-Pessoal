"""Edição de um .xlsx existente por dentro (XML), para gerar documentos no modelo oficial sem perder nada dele.

Abrir e salvar pelo openpyxl perde ou altera partes do modelo (comentários, desenhos, gráficos com formatação própria).
Aqui o arquivo é tratado como pacote: só as partes que mudam são reescritas; logos, gráficos, estilos e configuração de
impressão ficam como estão.

* Aba editada: fórmulas compartilhadas viram fórmulas explícitas (openpyxl Translator), para que mudar uma célula não
  quebre as vizinhas que dependiam dela.
* Valores guardados das fórmulas saem de todas as abas e o arquivo pede recálculo completo ao abrir: Excel e
  LibreOffice mostram os números calculados das células, nunca os do modelo.
"""
import html
import io
import posixpath
import re
import zipfile
from xml.sax.saxutils import escape

from openpyxl.formula.tokenizer import Tokenizer
from openpyxl.formula.translate import Translator
from openpyxl.utils import column_index_from_string, get_column_letter

_ATR = re.compile(r'([\w:]+)="([^"]*)"')
_LINHA = re.compile(r"<row\b([^>]*?)(?:/>|>(.*?)</row>)", re.S)
_CEL = re.compile(r"<c\b([^>]*?)(?:/>|>(.*?)</c>)", re.S)
_F = re.compile(r"<f\b([^>]*?)(?:/>|>(.*?)</f>)", re.S)
_V = re.compile(r"<v>(.*?)</v>", re.S)
_IS = re.compile(r"<is>.*?</is>", re.S)
_REF = re.compile(r"([A-Z]{1,3})(\d+)$")
_CELREF = re.compile(r"(\$?)([A-Z]{1,3})(\$?)(\d+)")
# ordem dos elementos depois de <sheetData> (CT_Worksheet): um elemento novo entra antes do primeiro que vem depois dele
ORDEM = ["sheetCalcPr", "sheetProtection", "protectedRanges", "scenarios", "autoFilter", "sortState", "dataConsolidate",
         "customSheetViews", "mergeCells", "phoneticPr", "conditionalFormatting", "dataValidations", "hyperlinks",
         "printOptions", "pageMargins", "pageSetup", "headerFooter", "rowBreaks", "colBreaks", "customProperties",
         "cellWatches", "ignoredErrors", "smartTags", "drawing", "legacyDrawing", "legacyDrawingHF", "picture",
         "oleObjects", "controls", "webPublishItems", "tableParts", "extLst"]
NS_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def atributos(s):
    return dict(_ATR.findall(s or ""))


def _atr_xml(d):
    return "".join(f' {k}="{v}"' for k, v in d.items())


def ref_partes(ref):
    m = _REF.match(ref)
    return m.group(1), int(m.group(2))


def deslocar_linhas(formula, dl):
    """A mesma fórmula dl linhas abaixo, como numa cópia de bloco: todas as referências desta aba (absolutas ou não)
    andam dl linhas; as de outra aba ('ABA'!A1) ficam."""
    out = []
    for t in Tokenizer("=" + formula).items:
        v = t.value
        if t.type == "OPERAND" and t.subtype == "RANGE" and "!" not in v:
            v = _CELREF.sub(lambda m: f"{m[1]}{m[2]}{m[3]}{int(m[4]) + dl}", v)
        out.append(v)
    return "".join(out)


class Celula:
    __slots__ = ("atr", "f", "fatr", "v", "inline", "extra")

    def __init__(self, atr=None):
        self.atr, self.f, self.fatr, self.v, self.inline, self.extra = atr or {}, None, {}, None, None, ""

    @classmethod
    def ler(cls, atr, dentro):
        c = cls(atributos(atr))
        dentro = dentro or ""
        m = _F.search(dentro)
        if m:
            c.fatr = atributos(m.group(1))
            c.f = html.unescape(m.group(2)) if m.group(2) is not None else ""
        m = _V.search(dentro)
        c.v = html.unescape(m.group(1)) if m else None
        m = _IS.search(dentro)
        c.inline = m.group(0) if m else None
        m = re.search(r"<extLst>.*?</extLst>", dentro, re.S)
        c.extra = m.group(0) if m else ""
        return c

    def copia(self):
        c = Celula(dict(self.atr))
        c.f, c.fatr, c.v, c.inline, c.extra = self.f, dict(self.fatr), self.v, self.inline, self.extra
        return c

    def xml(self, ref):
        a = dict(self.atr)
        a["r"] = ref
        dentro = ""
        if self.f is not None:
            dentro += f"<f{_atr_xml(self.fatr)}>{escape(self.f)}</f>" if self.f else f"<f{_atr_xml(self.fatr)}/>"
        if self.inline is not None:
            dentro += self.inline
        elif self.v is not None:
            dentro += f"<v>{escape(str(self.v))}</v>"
        dentro += self.extra
        a = {"r": a.pop("r"), **a}
        return f"<c{_atr_xml(a)}>{dentro}</c>" if dentro else f"<c{_atr_xml(a)}/>"


class Aba:
    """Uma aba (worksheet XML): as linhas viram {número: [atributos, {coluna: Celula}]}; o resto do XML fica como está."""

    def __init__(self, xml, textos=None):
        i = xml.index("<sheetData")
        fim = xml.find("</sheetData>")
        if fim < 0:                                     # <sheetData/>
            j = xml.index(">", i) + 1
            corpo, self.depois = "", xml[j:]
        else:
            j = xml.index(">", i) + 1
            corpo, self.depois = xml[j:fim], xml[fim + len("</sheetData>"):]
        self.antes = xml[:i]
        self.textos = textos or []
        self.linhas = {}
        for m in _LINHA.finditer(corpo):
            atr = atributos(m.group(1))
            r = int(atr["r"])
            cels = {}
            for c in _CEL.finditer(m.group(2) or ""):
                cel = Celula.ler(c.group(1), c.group(2))
                col, _ = ref_partes(cel.atr.pop("r"))
                cels[col] = cel
            self.linhas[r] = [atr, cels]

    # ------------------------------------------------------------------ leitura
    def celula(self, ref):
        col, r = ref_partes(ref)
        return self.linhas.get(r, [None, {}])[1].get(col)

    def valor(self, ref):
        """Texto ou número guardado na célula (texto compartilhado resolvido); fórmula sem valor → None."""
        c = self.celula(ref)
        if c is None:
            return None
        if c.inline is not None:
            return html.unescape(re.sub(r"<[^>]+>", "", c.inline))
        if c.v is None:
            return None
        t = c.atr.get("t")
        if t == "s":
            return self.textos[int(c.v)]
        if t in ("str", "inlineStr", "e"):
            return c.v
        try:
            n = float(c.v)
            return int(n) if n.is_integer() else n
        except ValueError:
            return c.v

    # ------------------------------------------------------------------ escrita (o estilo da célula fica)
    def _cel(self, ref):
        col, r = ref_partes(ref)
        linha = self.linhas.setdefault(r, [{"r": str(r)}, {}])
        return linha[1].setdefault(col, Celula())

    def _limpar(self, c):
        c.f, c.fatr, c.v, c.inline = None, {}, None, None
        c.atr.pop("t", None)
        return c

    def numero(self, ref, n):
        c = self._limpar(self._cel(ref))
        c.v = repr(float(n)) if not float(n).is_integer() else str(int(n))

    def texto(self, ref, s):
        c = self._limpar(self._cel(ref))
        c.atr["t"] = "inlineStr"
        c.inline = f'<is><t xml:space="preserve">{escape(str(s))}</t></is>'

    def formula(self, ref, f):
        c = self._limpar(self._cel(ref))
        c.f = f[1:] if f.startswith("=") else f

    def vazia(self, ref):
        """Tira o conteúdo e deixa o estilo."""
        c = self.celula(ref)
        if c is not None:
            self._limpar(c)

    def estilo(self, ref, s):
        self._cel(ref).atr["s"] = str(s)

    # ------------------------------------------------------------------ fórmulas
    def expandir_compartilhadas(self):
        mestres = {}
        for r, (_, cels) in self.linhas.items():
            for col, c in cels.items():
                if c.f and c.fatr.get("t") == "shared":
                    mestres[c.fatr["si"]] = (c.f, f"{col}{r}")
        trad = {}
        for r, (_, cels) in self.linhas.items():
            for col, c in cels.items():
                if c.f is not None and c.fatr.get("t") == "shared":
                    si = c.fatr["si"]
                    f, origem = mestres[si]
                    ref = f"{col}{r}"
                    if ref != origem:
                        if si not in trad:
                            trad[si] = Translator("=" + f, origin=origem)
                        f = trad[si].translate_formula(ref)[1:]
                    c.f, c.fatr = f, {}

    def sem_valores_guardados(self):
        for _, cels in self.linhas.values():
            for c in cels.values():
                if c.f is not None:
                    c.v = None
                    if c.atr.get("t") in ("str", "n", "b", "e"):
                        c.atr.pop("t")

    # ------------------------------------------------------------------ linhas, colunas e partes da aba
    def copiar_bloco(self, ini, fim, destino, origem=None):
        """Cópia das linhas ini..fim (de `origem`, outra Aba, ou desta) para destino.., com as fórmulas deslocadas."""
        fonte = origem or self
        dl = destino - ini
        for r in range(ini, fim + 1):
            if r not in fonte.linhas:
                continue
            atr, cels = fonte.linhas[r]
            novo = {}
            for col, c in cels.items():
                n = c.copia()
                if n.f:
                    n.f = deslocar_linhas(n.f, dl)
                novo[col] = n
            self.linhas[r + dl] = [{**atr, "r": str(r + dl)}, novo]

    def apagar_linhas(self, ini, fim):
        for r in [r for r in self.linhas if ini <= r <= fim]:
            del self.linhas[r]

    def ocultar_linha(self, r, oculta=True):
        atr = self.linhas.setdefault(r, [{"r": str(r)}, {}])[0]
        if oculta:
            atr["hidden"] = "1"
        else:
            atr.pop("hidden", None)

    def altura(self, r):
        atr = self.linhas.get(r, [{}])[0]
        return float(atr["ht"]) if "ht" in atr else None

    def elemento(self, nome):
        m = re.search(rf"<{nome}\b[^>]*?(?:/>|>.*?</{nome}>)", self.depois + self.antes, re.S)
        return m.group(0) if m else None

    def por_elemento(self, nome, xml):
        """Troca (ou insere, na ordem do esquema) o elemento `nome` depois de <sheetData>; xml vazio só remove."""
        self.depois = re.sub(rf"<{nome}\b[^>]*?(?:/>|>.*?</{nome}>)", "", self.depois, flags=re.S)
        if not xml:
            return
        seguintes = ORDEM[ORDEM.index(nome) + 1:]
        pos = len(self.depois) - len("</worksheet>")
        for s in seguintes:
            m = re.search(rf"<{s}\b", self.depois)
            if m:
                pos = min(pos, m.start())
        self.depois = self.depois[:pos] + xml + self.depois[pos:]

    def mesclagens(self):
        m = re.search(r"<mergeCells\b[^>]*>(.*?)</mergeCells>", self.depois, re.S)
        return re.findall(r'<mergeCell ref="([^"]+)"', m.group(1)) if m else []

    def por_mesclagens(self, refs):
        self.por_elemento("mergeCells", f'<mergeCells count="{len(refs)}">' +
                          "".join(f'<mergeCell ref="{r}"/>' for r in refs) + "</mergeCells>" if refs else "")

    def ocultar_colunas(self, indices):
        """Oculta as colunas (índices a partir de 1), mantendo largura e estilo de cada uma."""
        m = re.search(r"<cols>(.*?)</cols>", self.antes, re.S)
        cols = {}
        for c in re.finditer(r"<col\b([^>]*)/>", m.group(1) if m else ""):
            a = atributos(c.group(1))
            for i in range(int(a["min"]), int(a["max"]) + 1):
                cols[i] = {k: v for k, v in a.items() if k not in ("min", "max")}
        for i in indices:
            cols.setdefault(i, {"width": "8.6640625", "customWidth": "1"})["hidden"] = "1"
        partes, ordem = [], sorted(cols)
        k = 0
        while k < len(ordem):                           # junta as colunas seguidas com os mesmos atributos
            ini = fim = ordem[k]
            while k + 1 < len(ordem) and ordem[k + 1] == fim + 1 and cols[ordem[k + 1]] == cols[ini]:
                k += 1
                fim = ordem[k]
            partes.append(f'<col min="{ini}" max="{fim}"{_atr_xml(cols[ini])}/>')
            k += 1
        novo = "<cols>" + "".join(partes) + "</cols>"
        if m:
            self.antes = self.antes[:m.start()] + novo + self.antes[m.end():]
        else:
            i = self.antes.index("<sheetData") if "<sheetData" in self.antes else len(self.antes)
            self.antes = self.antes[:i] + novo + self.antes[i:]

    def quebras_de_pagina(self, linhas):
        """Quebra manual depois de cada linha da lista (uma folha por página na impressão)."""
        self.por_elemento("rowBreaks", f'<rowBreaks count="{len(linhas)}" manualBreakCount="{len(linhas)}">' +
                          "".join(f'<brk id="{r}" max="16383" man="1"/>' for r in linhas) + "</rowBreaks>"
                          if linhas else "")

    def fixar_alturas(self, ini, fim, padrao):
        """Altura explícita em cada linha (a que não tem fica com a padrão): o LibreOffice reajusta as linhas sem altura
        depois de posicionar os desenhos, e as fotos e logos escorregam folha após folha."""
        for r in range(ini, fim + 1):
            atr = self.linhas.setdefault(r, [{"r": str(r)}, {}])[0]
            if "ht" not in atr:
                atr["ht"], atr["customHeight"] = str(padrao), "1"

    def uma_pagina_de_largura(self):
        """Impressão com todas as colunas visíveis numa página de largura (a altura segue em quantas páginas precisar)."""
        m = re.search(r"<sheetPr\b([^>]*)(/>|>(.*?)</sheetPr>)", self.antes, re.S)
        if m:
            dentro = re.sub(r"<pageSetUpPr\b[^>]*/>", "", m.group(3) or "") + '<pageSetUpPr fitToPage="1"/>'
            self.antes = self.antes[:m.start()] + f"<sheetPr{m.group(1)}>{dentro}</sheetPr>" + self.antes[m.end():]
        else:
            i = self.antes.index("<dimension") if "<dimension" in self.antes else self.antes.index("<sheetViews")
            self.antes = self.antes[:i] + '<sheetPr><pageSetUpPr fitToPage="1"/></sheetPr>' + self.antes[i:]
        self.depois = re.sub(r"<pageSetup\b([^>]*?)\s*/>", lambda m: "<pageSetup" + re.sub(
            r'\s(?:scale|fitToWidth|fitToHeight)="[^"]*"', "", m.group(1)) + ' fitToWidth="1" fitToHeight="0"/>', self.depois)

    def cabecalho(self, texto):
        """Texto no cabeçalho de todas as páginas impressas (centro, vermelho, negrito)."""
        h = escape(f'&C&"-,Bold"&14&KFF0000{texto}')
        self.por_elemento("headerFooter", f"<headerFooter><oddHeader>{h}</oddHeader></headerFooter>")

    # ------------------------------------------------------------------ saída
    def xml(self):
        partes, cmin, cmax, rmax = [], None, None, 1
        for r in sorted(self.linhas):
            atr, cels = self.linhas[r]
            atr = {"r": str(r), **{k: v for k, v in atr.items() if k not in ("r", "spans")}}
            ordem = sorted(cels, key=column_index_from_string)
            if ordem:
                a, b = column_index_from_string(ordem[0]), column_index_from_string(ordem[-1])
                cmin, cmax = min(cmin or a, a), max(cmax or b, b)
            rmax = max(rmax, r)
            corpo = "".join(cels[c].xml(f"{c}{r}") for c in ordem)
            partes.append(f"<row{_atr_xml(atr)}>{corpo}</row>" if corpo else f"<row{_atr_xml(atr)}/>")
        dim = f"{get_column_letter(cmin or 1)}1:{get_column_letter(cmax or 1)}{rmax}"
        antes = re.sub(r'<dimension ref="[^"]*"/>', f'<dimension ref="{dim}"/>', self.antes, count=1)
        return antes + "<sheetData>" + "".join(partes) + "</sheetData>" + self.depois


class Pacote:
    """O .xlsx como pacote de partes. Partes não tocadas saem byte a byte iguais."""

    def __init__(self, dados):
        z = zipfile.ZipFile(io.BytesIO(dados))
        self.ordem = z.namelist()
        self.partes = {n: z.read(n) for n in self.ordem}
        self._abas_abertas = {}
        ss = self.partes.get("xl/sharedStrings.xml", b"").decode("utf-8")
        self.textos = [html.unescape("".join(re.findall(r"<t\b[^>]*>(.*?)</t>", si, re.S)))
                       for si in re.findall(r"<si>(.*?)</si>", ss, re.S)]

    def texto(self, nome):
        return self.partes[nome].decode("utf-8")

    def por_texto(self, nome, s, tipo=None):
        if nome not in self.partes:
            self.ordem.append(nome)
            if tipo:
                self.por_tipo(nome, tipo)
        self.partes[nome] = s.encode("utf-8") if isinstance(s, str) else s

    def remover(self, nome):
        self.partes.pop(nome, None)
        if nome in self.ordem:
            self.ordem.remove(nome)
        ct = re.sub(rf'<Override PartName="/{re.escape(nome)}"[^>]*/>', "", self.texto("[Content_Types].xml"))
        self.partes["[Content_Types].xml"] = ct.encode("utf-8")

    def por_tipo(self, nome, tipo):
        ct = self.texto("[Content_Types].xml")
        if f'PartName="/{nome}"' not in ct:
            ct = ct.replace("</Types>", f'<Override PartName="/{nome}" ContentType="{tipo}"/></Types>')
            self.partes["[Content_Types].xml"] = ct.encode("utf-8")

    def extensao(self, ext, tipo):
        ct = self.texto("[Content_Types].xml")
        if f'Extension="{ext}"' not in ct:
            ct = ct.replace("<Default ", f'<Default Extension="{ext}" ContentType="{tipo}"/><Default ', 1)
            self.partes["[Content_Types].xml"] = ct.encode("utf-8")

    # ------------------------------------------------------------------ relações
    @staticmethod
    def _rels_de(nome):
        d, b = posixpath.split(nome)
        return posixpath.join(d, "_rels", b + ".rels")

    def relacoes(self, nome):
        """{id: (tipo, alvo absoluto)} da parte."""
        rels = self._rels_de(nome)
        if rels not in self.partes:
            return {}
        out = {}
        for m in re.finditer(r"<Relationship\b([^>]*)/>", self.texto(rels)):
            a = atributos(m.group(1))
            out[a["Id"]] = (a["Type"].rsplit("/", 1)[-1], posixpath.normpath(posixpath.join(posixpath.dirname(nome), a["Target"])))
        return out

    def relacionar(self, nome, tipo, alvo):
        """Nova relação da parte `nome` para `alvo` (caminho absoluto no pacote); devolve o Id."""
        rels = self._rels_de(nome)
        xml = self.texto(rels) if rels in self.partes else (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"></Relationships>')
        usados = set(re.findall(r'Id="rId(\d+)"', xml))
        n = 1
        while str(n) in usados:
            n += 1
        alvo_rel = posixpath.relpath(alvo, posixpath.dirname(nome))
        xml = xml.replace("</Relationships>", f'<Relationship Id="rId{n}" Type="{NS_REL}/{tipo}" Target="{alvo_rel}"/></Relationships>')
        self.por_texto(rels, xml)
        return f"rId{n}"

    def desrelacionar(self, nome, rid):
        rels = self._rels_de(nome)
        self.por_texto(rels, re.sub(rf'<Relationship [^>]*Id="{rid}"[^>]*/>', "", self.texto(rels)))

    # ------------------------------------------------------------------ abas
    def abas(self):
        """[(nome, parte)] na ordem da pasta."""
        rels = self.relacoes("xl/workbook.xml")
        return [(html.unescape(a["name"]), rels[a["r:id"]][1])
                for a in (atributos(m.group(1)) for m in re.finditer(r"<sheet\b([^>]*)/>", self.texto("xl/workbook.xml")))]

    def parte_da_aba(self, nome):
        for n, p in self.abas():
            if n == nome:
                return p
        raise KeyError(nome)

    def aba(self, nome):
        """A aba aberta para edição (a mesma instância até salvar)."""
        p = self.parte_da_aba(nome)
        if p not in self._abas_abertas:
            self._abas_abertas[p] = Aba(self.texto(p), self.textos)
        return self._abas_abertas[p]

    def renomear_aba(self, antigo, novo):
        q_ant, q_novo = f"'{antigo}'!", f"'{novo}'!"
        wb = self.texto("xl/workbook.xml").replace(f'name="{escape(antigo)}"', f'name="{escape(novo)}"')
        self.por_texto("xl/workbook.xml", wb.replace(q_ant, q_novo))
        for n in list(self.partes):
            if n.endswith(".xml") and (n.startswith("xl/worksheets/") or n.startswith("xl/charts/")):
                if n in self._abas_abertas:
                    a = self._abas_abertas[n]
                    for _, cels in a.linhas.values():
                        for c in cels.values():
                            if c.f and q_ant in c.f:
                                c.f = c.f.replace(q_ant, q_novo)
                else:
                    b = self.partes[n]
                    if q_ant.encode() in b:
                        self.partes[n] = b.replace(q_ant.encode(), q_novo.encode())

    def nome_definido(self, nome, aba_idx, valor):
        wb = self.texto("xl/workbook.xml")
        wb = re.sub(rf'(<definedName name="{re.escape(nome)}" localSheetId="{aba_idx}"[^>]*>)[^<]*(</definedName>)',
                    lambda m: m.group(1) + escape(valor) + m.group(2), wb)
        self.por_texto("xl/workbook.xml", wb)

    def recalcular_ao_abrir(self):
        """Sem cadeia de cálculo antiga e sem valores guardados: o Excel e o LibreOffice recalculam tudo ao abrir."""
        if "xl/calcChain.xml" in self.partes:
            self.remover("xl/calcChain.xml")
            self.por_texto("xl/_rels/workbook.xml.rels",
                           re.sub(r"<Relationship [^>]*calcChain[^>]*/>", "", self.texto("xl/_rels/workbook.xml.rels")))
        wb = self.texto("xl/workbook.xml")
        if "fullCalcOnLoad" not in wb:
            wb = re.sub(r"<calcPr\b([^>]*?)/>", r'<calcPr\1 fullCalcOnLoad="1"/>', wb, count=1)
            if "<calcPr" not in wb:
                wb = wb.replace("</workbook>", '<calcPr fullCalcOnLoad="1"/></workbook>')
        self.por_texto("xl/workbook.xml", wb)
        for _, p in self.abas():
            if p in self._abas_abertas:
                self._abas_abertas[p].sem_valores_guardados()
            else:
                b = self.partes[p]
                b = re.sub(rb"(</f>|<f\b[^>]*/>)<v>[^<]*</v>", rb"\1", b).replace(b"<v/>", b"")
                self.partes[p] = b

    def salvar(self):
        for p, a in self._abas_abertas.items():
            self.partes[p] = a.xml().encode("utf-8")
        saida = io.BytesIO()
        with zipfile.ZipFile(saida, "w", zipfile.ZIP_DEFLATED) as z:
            for n in ["[Content_Types].xml"] + [n for n in self.ordem if n != "[Content_Types].xml"]:
                if n in self.partes:
                    z.writestr(n, self.partes[n])
        return saida.getvalue()
