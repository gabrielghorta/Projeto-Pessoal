"""Importação idempotente das planilhas oficiais para o esquema v3.

Princípios:
  * nada é apagado: registro que sai de uma revisão fica com vigente = 0;
  * identidade estável própria (TR-/EL-/AT-/IP-), casada pela identidade FÍSICA, nunca pelo nº da linha
    nem só pelo nome visível:
      - trecho: handle CAD + pontas UTM (o handle sozinho se repete em 123 trechos da 06.1 REV01);
      - elemento: handle CAD; sem handle, nº de projeto (identidade provisória, promovida quando o handle aparecer);
      - item: atividade + assinatura técnica completa (descrição, DE, PN, resina, tipo, solda, material);
  * divergência entre nome e geometria, retirada de objeto com histórico, divisão e união nunca transferem
    histórico por presunção: viram pendência IMPEDITIVA até haver conciliação explícita (config/conciliacoes.json);
  * toda linha importada guarda proveniência (arquivo + SHA-256 + aba + linha + células);
  * reimportar os mesmos arquivos não altera nada (status SEM_ALTERACOES);
  * fonte ausente não apaga dado: campos que dependem dela são preservados e a ausência vira pendência;
  * quantidade/unidade sem fonte ficam NULL (pendentes) — sem regra de descrição, sem equivalência UN↔PÇ;
  * avisos AA/AB/AW entram como texto de origem; decisões ficam em `decisoes`, ligadas só por escopo explícito.
"""
import hashlib
import json
import os
import re
from collections import Counter, defaultdict
from contextlib import contextmanager
from datetime import date, datetime

import openpyxl
from openpyxl.utils import column_index_from_string as col
from openpyxl.utils import get_column_letter

import regras as R
import servicos as S
from backup import sha256

BASE = os.path.dirname(os.path.abspath(__file__))
DIR_PLANILHAS = os.path.join(BASE, "planilhas")

COLUNAS_SERVICO = ["AL", "AM", "AN", "AO", "AP", "AQ", "AR", "AS", "AT", "AU", "AV", "AW", "AX", "AY", "AZ",
                   "BA", "BB", "BC", "BD", "BE", "BF", "BG", "BH", "BI", "BJ", "BK"]

# Cabeçalhos da aba CONEXÕES (06.2) localizados pelo nome; a letra encontrada vai para o relatório.
CAB_CONEXOES = {
    "descricao": "DESCRIÇÃO", "pressao": "PRESSÃO", "resina": "RESINA", "tipo_conexao": "TIPO_CONEXÃO",
    "solda": "SOLDA", "diametro": "DIÂMETRO", "material": "MATERIAL", "grupo": "GRUPO", "quantidade": "QUANT.",
    "nomenclatura": "NOMENCLATURA (OBRA)", "local": "LOCAL", "codigo_bm": "CÓDIGO BM",
    "consideracao": "CONSIDERAÇÃO", "sugestao": "SUGESTÃO DE CORREÇÃO", "equipe": "EQUIPE",
    "tipo_entroncamento": "TIPO ENTRONCAMENTO", "status_conexao": "STATUS CONEXÃO", "handle": "HANDLE CAD",
    "utm_e": "COORD. E (UTM)", "utm_n": "COORD. N (UTM)", "lat": "LATITUDE", "lon": "LONGITUDE",
    "prioridade": "PRIORIDADE DO ERRO", "enquadramento": "ENQUADRAMENTO BM", "obs_enquadramento": "OBS. ENQUADRAMENTO BM",
    "trechos": "TRECHO (06.1)",
}

# Identidade técnica de um item de projeto (o resto — quantidade, código BM, BA/BB/BC, AG/AH — é atributo mutável)
CAMPOS_ASSINATURA = ("descricao", "diametro", "pressao", "resina", "tipo_conexao", "solda", "material")
TOL_POSICAO_M = 0.01


def carregar_config():
    cfg_dir = os.path.join(BASE, "config")
    with open(os.path.join(cfg_dir, "fontes.json"), encoding="utf-8") as f:
        cfg = json.load(f)
    for chave, arq, padrao in (("controles", "controles.json", {"conjuntos": []}),
                               ("conciliacoes", "conciliacoes.json", {"conciliacoes": []})):
        p = os.path.join(cfg_dir, arq)
        dados = padrao
        if os.path.exists(p):
            with open(p, encoding="utf-8") as f:
                dados = json.load(f)
        cfg[chave] = dados["conjuntos"] if chave == "controles" else dados["conciliacoes"]
    cfg.setdefault("perfil", "REAL")          # perfis de teste são explícitos e nunca valem para o perfil real
    return cfg


def caminho(cfg, chave, pasta=DIR_PLANILHAS):
    nome = cfg["arquivos"].get(chave)
    p = os.path.join(pasta, nome) if nome else None
    return p if p and os.path.exists(p) else None


@contextmanager
def planilha(p):
    """Abre a planilha só para leitura e SEMPRE fecha (no Windows, arquivo aberto impede mover/apagar)."""
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    try:
        yield wb
    finally:
        wb.close()


def revisao_do_nome(nome):
    m = re.findall(r"REV\d+", nome.upper())
    return m[0] if m else None


def num(x):
    return float(x) if isinstance(x, (int, float)) and not isinstance(x, bool) else None


def txt(x):
    if x is None:
        return None
    if isinstance(x, datetime):
        return x.date().isoformat()
    s = str(x).strip()
    return s or None


def bruto(x):
    """Valor de célula serializável para a proveniência."""
    if isinstance(x, (datetime, date)):
        return x.isoformat()
    return x


def _hash_texto(t):
    return hashlib.sha1(t.encode("utf-8")).hexdigest()


def numero_projeto(nomenclatura):
    """Zona01_PB01_CH_DF-33 -> DF-33."""
    return nomenclatura.split("_")[-1]


def tipo_elemento(numero):
    return "NÓ" if numero.upper().startswith(("NÓ", "NO")) else numero.split("-")[0].upper()


def assinatura_item(d):
    return "|".join("" if d.get(k) is None else str(d[k]) for k in CAMPOS_ASSINATURA)


def _mesma_posicao(a, b):
    return (None not in (a[0], a[1], b[0], b[1])
            and abs(a[0] - b[0]) <= TOL_POSICAO_M and abs(a[1] - b[1]) <= TOL_POSICAO_M)


_LER = object()


class Importador:
    def __init__(self, conn, cfg, importacao_id, pasta=DIR_PLANILHAS):
        self.c = conn
        self.cfg = cfg
        self.imp = importacao_id
        self.pasta = pasta
        self.stats = defaultdict(Counter)       # entidade -> {novos, alterados, iguais, retirados}
        self.fontes = {}
        self.colunas_06_2 = {}
        self.conferencias = {}
        self.bloqueios = []                     # pendências impeditivas: impedem o COMMIT
        self.decisoes_novas = 0
        self._ultimo_id = {}
        self._av_vistos = set()
        self.concil = self._validar_conciliacoes(cfg.get("conciliacoes", []))

    # -------------------------------------------------------------- utilitários
    def fonte(self, chave, p):
        h = sha256(p)
        nome = os.path.basename(p)
        r = self.c.execute("SELECT id FROM fontes_arquivo WHERE chave=? AND sha256=?", (chave, h)).fetchone()
        if r:
            fid = r[0]
        else:
            fid = self.c.execute(
                "INSERT INTO fontes_arquivo(chave,nome,revisao,sha256,bytes,primeira_importacao_id) VALUES (?,?,?,?,?,?)",
                (chave, nome, revisao_do_nome(nome), h, os.path.getsize(p), self.imp)).lastrowid
        self.fontes[chave] = {"id": fid, "nome": nome, "sha256": h, "revisao": revisao_do_nome(nome)}
        return fid

    def prov(self, entidade, eid, fid, aba, linha, celulas):
        # só reescreve se mudou (INSERT OR REPLACE apagaria e reinseriria a linha a cada reimportação)
        self.c.execute("""INSERT INTO proveniencia VALUES (?,?,?,?,?,?)
                          ON CONFLICT(entidade, entidade_id, fonte_id) DO UPDATE SET aba=excluded.aba,
                            linha=excluded.linha, celulas=excluded.celulas
                          WHERE aba IS NOT excluded.aba OR linha IS NOT excluded.linha OR celulas IS NOT excluded.celulas""",
                       (entidade, eid, fid, aba, linha, json.dumps(celulas, ensure_ascii=False, default=str)))

    def pendencia(self, tipo, ref, detalhe, bloqueante=False):
        self.c.execute("INSERT INTO pendencias_importacao(importacao_id,tipo,referencia,detalhe,bloqueante) VALUES (?,?,?,?,?)",
                       (self.imp, tipo, ref, detalhe, 1 if bloqueante else 0))
        if bloqueante:
            self.bloqueios.append({"tipo": tipo, "referencia": ref, "detalhe": detalhe})

    def nome_hist(self, entidade, eid, nome, fid):
        self.c.execute("INSERT OR IGNORE INTO nomes_historico VALUES (?,?,?,?,?)", (entidade, eid, nome, fid, S._agora()))

    def _alteracao(self, entidade, eid, campo, antes, depois):
        self.c.execute("INSERT INTO alteracoes(importacao_id,entidade,entidade_id,campo,antes,depois) VALUES (?,?,?,?,?,?)",
                       (self.imp, entidade, eid, campo, antes, depois))

    def upsert(self, tabela, entidade, chave_where, campos, atual=_LER):
        """Insere ou atualiza registrando cada campo alterado. `atual` permite comparar com o estado lido
        antes de renomeações temporárias (troca de nomes entre registros)."""
        where = " AND ".join(f"{k}=?" for k in chave_where)
        if atual is _LER:
            atual = self.c.execute(f"SELECT * FROM {tabela} WHERE {where}", tuple(chave_where.values())).fetchone()
        ident = str(next(iter(chave_where.values()))) if len(chave_where) == 1 else json.dumps(chave_where, ensure_ascii=False)
        if atual is None:
            dados = {**chave_where, **campos}
            self.c.execute(f"INSERT INTO {tabela}({','.join(dados)}) VALUES ({','.join('?' * len(dados))})",
                           tuple(dados.values()))
            self._alteracao(entidade, ident, "(novo)", None, None)
            self.stats[entidade]["novos"] += 1
            return "novo"
        dif = {k: v for k, v in campos.items() if atual[k] != v}
        if not dif:
            self.stats[entidade]["iguais"] += 1
            return "igual"
        self.c.execute(f"UPDATE {tabela} SET {', '.join(f'{k}=?' for k in dif)} WHERE {where}",
                       (*dif.values(), *chave_where.values()))
        for k, v in dif.items():
            self._alteracao(entidade, ident, k, None if atual[k] is None else str(atual[k]), None if v is None else str(v))
        self.stats[entidade]["alterados"] += 1
        return "alterado"

    def retirar(self, tabela, entidade, eid):
        self.c.execute(f"UPDATE {tabela} SET vigente=0 WHERE id=?", (eid,))
        self._alteracao(entidade, str(eid), "vigente", "1", "0")
        self.stats[entidade]["retirados"] += 1

    def retirar_nao_vistos(self, tabela, entidade, vistos, filtro="1=1", args=()):
        for (eid,) in self.c.execute(f"SELECT id FROM {tabela} WHERE vigente=1 AND {filtro}", args).fetchall():
            if eid not in vistos:
                self.retirar(tabela, entidade, eid)

    def novo_id(self, prefixo, tabela, largura):
        """IDs novos em sequência (em memória, para reservar vários antes de gravar)."""
        if prefixo not in self._ultimo_id:
            r = self.c.execute(f"SELECT MAX(CAST(SUBSTR(id, {len(prefixo) + 1}) AS INTEGER)) FROM {tabela} WHERE id LIKE ?",
                               (prefixo + "%",)).fetchone()[0]
            self._ultimo_id[prefixo] = r or 0
        self._ultimo_id[prefixo] += 1
        return f"{prefixo}{self._ultimo_id[prefixo]:0{largura}d}"

    def parametro(self, chave, valor, fonte, sobrescrever=True):
        atual = S.parametro(self.c, chave)
        conflito = ("DO UPDATE SET valor=excluded.valor, fonte=excluded.fonte "
                    "WHERE valor IS NOT excluded.valor OR fonte IS NOT excluded.fonte") if sobrescrever else "DO NOTHING"
        self.c.execute(f"INSERT INTO parametros(chave,valor,fonte) VALUES (?,?,?) ON CONFLICT(chave) {conflito}",
                       (chave, str(valor), fonte))
        if sobrescrever and atual is not None and atual != str(valor):
            self._alteracao("parametro", chave, "valor", atual, str(valor))

    def historico(self, alvo_tipo):
        """IDs (TRECHO / ATIVIDADE) com qualquer registro operacional: movimento, check-in, foto, execução, medição,
        soldas, reserva, intervalos, conciliação, cadastro (definição única em servicos.historico_alvos)."""
        return set(S.historico_alvos(self.c, alvo_tipo))

    # -------------------------------------------------------------- conciliações explícitas (de-para)
    def _validar_conciliacoes(self, lista):
        """config/conciliacoes.json: decisões humanas de de-para. Cada uma exige responsável, data e motivo.

        {"tipo": "TRECHO",   "ponto": "...", "handle": "...", "usar_id": "TR-00001", ...}  linha da fonte herda o ID
        {"tipo": "TRECHO",   "retirar_id": "TR-00007", ...}   confirma retirar um trecho com histórico (o histórico fica nele)
        {"tipo": "ELEMENTO", "chave_natural": "H:5FD4F", "usar_id": "EL-0001", ...}
        {"tipo": "ELEMENTO" | "ATIVIDADE", "retirar_id": "...", ...}
        """
        out = {"TRECHO": {"usar": {}, "retirar": {}}, "ELEMENTO": {"usar": {}, "retirar": {}},
               "ATIVIDADE": {"usar": {}, "retirar": {}}}
        self._concil_invalidas = []
        for e in lista:
            falta = [k for k in ("tipo", "responsavel", "data", "motivo") if not e.get(k)]
            if falta or e.get("tipo") not in out:
                self._concil_invalidas.append((e, f"campos obrigatórios ausentes ou tipo inválido: {falta}"))
                continue
            if e.get("retirar_id"):
                out[e["tipo"]]["retirar"][e["retirar_id"]] = e
            elif e.get("usar_id") and e["tipo"] == "TRECHO" and e.get("ponto"):
                out["TRECHO"]["usar"][(e["ponto"], e.get("handle"))] = e
            elif e.get("usar_id") and e["tipo"] == "ELEMENTO" and e.get("chave_natural"):
                out["ELEMENTO"]["usar"][e["chave_natural"]] = e
            else:
                self._concil_invalidas.append((e, "precisa de 'retirar_id' ou de 'usar_id' com a referência da fonte"))
        return out

    def _concil_aplicada(self, e, ref, texto):
        self.pendencia("CONCILIACAO_APLICADA", ref, f"{texto} — {e['responsavel']}, {e['data']}: {e['motivo']}")
        self._alteracao("conciliacao", ref, e["tipo"], None, json.dumps(e, ensure_ascii=False))

    # -------------------------------------------------------------- cargas
    def equipes(self):
        nomes = {f"Zona{n:02d}": f"EQUIPE ZONA {n:02d}" for n in range(1, 12)}
        nomes.update({"ADUTORA": "EQUIPE ADUTORA", "CH": "EQUIPE CH — CAIXAS HIDRÁULICAS"})
        for eid in R.FRENTES:
            self.c.execute("INSERT OR IGNORE INTO equipes(id, nome) VALUES (?,?)", (eid, nomes[eid]))

    def bm(self, p):
        fid = self.fonte("bm", p)
        nomes_atr = ["descricao_resumida", "pressao", "resina", "tipo_conexao", "solda", "diametro", "material", "grupo",
                     "grupo_orcamento_ankara"]
        with planilha(p) as wb:
            for planilha_bm, aba, ini_atr in (("JUNTADA", "BM MORRO DA CRUZ", 1), ("VERTICAL", "BM MORRO DA CRUZ - VERTICAL", 5)):
                vistos = set()
                for k, row in enumerate(wb[aba].iter_rows(min_row=6, values_only=True), 6):
                    if planilha_bm == "JUNTADA":
                        if not (row[0] and row[1]):
                            continue
                        item = str(row[0])
                        campos = dict(item_bm=item, subgrupo=None, subgrupo_nome=None, codigo_caesb=txt(row[10]),
                                      classificacao=txt(row[11]), descricao=row[12], unidade=txt(row[13]),
                                      quantidade=num(row[14]), preco_unitario=num(row[15]), valor_total=num(row[16]))
                    else:
                        if not row[3]:
                            continue
                        item = str(row[3])
                        campos = dict(item_bm=str(row[4]), subgrupo=txt(row[1]), subgrupo_nome=txt(row[2]),
                                      codigo_caesb=txt(row[14]), classificacao=txt(row[15]), descricao=row[16],
                                      unidade=txt(row[17]), quantidade=num(row[18]), preco_unitario=num(row[19]),
                                      valor_total=num(row[20]))
                    campos["atributos"] = json.dumps(dict(zip(nomes_atr, row[ini_atr:ini_atr + 9])), ensure_ascii=False,
                                                     default=str)
                    campos["linha_origem"] = k
                    campos["vigente"] = 1
                    self.upsert("bm_itens", "bm_item", {"planilha": planilha_bm, "item": item}, campos)
                    bid = self.c.execute("SELECT id FROM bm_itens WHERE planilha=? AND item=?", (planilha_bm, item)).fetchone()[0]
                    vistos.add(bid)
                    self.prov("bm_item", str(bid), fid, aba, k, {"item": item, "quantidade": bruto(campos["quantidade"]),
                                                                 "preco": bruto(campos["preco_unitario"])})
                self.retirar_nao_vistos("bm_itens", "bm_item", vistos, "planilha=?", (planilha_bm,))
            # rodapé da Vertical: BDI, K, totais por classe e P0
            rod, no_rodape = {}, False
            for row in wb["BM MORRO DA CRUZ - VERTICAL"].iter_rows(min_row=6, values_only=True):
                no_rodape = no_rodape or any(str(x).strip() == "RESUMO PRINCIPAL" for x in row if x is not None)
                if not no_rodape:
                    continue
                cel = [x for x in row if x not in (None, "")]
                rots = [i for i, x in enumerate(cel) if isinstance(x, str)]
                if not rots:
                    continue
                rot = str(cel[rots[-1]]).strip()
                vals = [x for x in cel[rots[-1] + 1:] if isinstance(x, (int, float))]
                if not vals:
                    continue
                chave = {"MATERIAL": "bm_total_material", "SERVIÇO": "bm_total_servico", "BDI MATERIAL =>": "bdi_material",
                         "BDI SERVIÇO =>": "bdi_servico", "COEF. - K =>": "coef_k", "TOTAL 01 + 02 + BDI -K": "p0_bm"}.get(rot)
                if chave and chave not in rod:
                    rod[chave] = vals[0]
        faltam = [k for k in ("bm_total_material", "bm_total_servico", "bdi_material", "bdi_servico", "coef_k", "p0_bm")
                  if k not in rod]
        if faltam:
            self.pendencia("BM_RODAPE_INCOMPLETO", os.path.basename(p), f"células do rodapé não encontradas: {faltam}",
                           bloqueante=True)
        fonte_rod = f"{os.path.basename(p)} · BM VERTICAL, rodapé"
        for k, v in rod.items():
            self.parametro(k, v, fonte_rod)
        return rod

    def reajuste(self, p):
        self.fonte("reajuste", p)
        achados = {}
        with planilha(p) as wb:
            for row in wb["REAJUSTE"].iter_rows(values_only=True):
                rot = str(row[0] or "")
                vals = [x for x in row[1:] if x is not None]
                if not vals:
                    continue
                if rot.startswith("Data-base do reajuste"):
                    achados["data_base"] = vals[0]
                elif rot.startswith("Percentual do 1º reajuste"):
                    achados["indice"] = vals[0]
                elif rot.startswith("1º aniversário"):
                    achados["aniversario"] = vals[0]
                elif rot.startswith("Início da execução"):
                    achados["inicio_os"] = vals[0]
        for k in ("data_base", "indice", "aniversario"):
            if k not in achados:
                raise RuntimeError(f"REAJUSTE: '{k}' não encontrado na aba.")
        aniv = achados["aniversario"].date() if isinstance(achados["aniversario"], datetime) else achados["aniversario"]
        fim = date.fromordinal(date(aniv.year + 1, aniv.month, aniv.day).toordinal() - 1)
        self.c.execute("""INSERT INTO reajuste_indices VALUES (?,?,?,?) ON CONFLICT(vigencia_ini) DO UPDATE SET
                            vigencia_fim=excluded.vigencia_fim, indice=excluded.indice, descricao=excluded.descricao
                          WHERE vigencia_fim IS NOT excluded.vigencia_fim OR indice IS NOT excluded.indice
                            OR descricao IS NOT excluded.descricao""",
                       (aniv.isoformat(), fim.isoformat(), float(achados["indice"]),
                        "1º reajuste — Despacho 322/2026-CEFR; frações elegíveis e pedido da contratada pendentes"))
        for k, v in achados.items():
            self.parametro(f"reajuste_{k}", v.date().isoformat() if isinstance(v, datetime) else v,
                           f"{os.path.basename(p)} · aba REAJUSTE")
        return achados

    # -------------------------------------------------------------- trechos (06.1)
    def _ler_trechos(self, p):
        linhas, pontos = [], set()
        with planilha(p) as wb:
            ws = wb["REDE"]
            v = lambda row, l: row[col(l) - 1] if col(l) - 1 < len(row) else None
            for k, row in enumerate(ws.iter_rows(min_row=6, values_only=True), 6):
                ponto = txt(v(row, "R"))
                if not ponto:
                    continue
                if ponto in pontos:
                    self.pendencia("TRECHO_DUPLICADO", ponto, f"PONTO repetido na linha {k} da aba REDE", bloqueante=True)
                    continue
                pontos.add(ponto)
                e1, n1, e2, n2 = (num(v(row, c)) for c in ("BS", "BT", "BU", "BV"))
                utm_ini = json.dumps([e1, n1]) if e1 else None
                utm_fim = json.dumps([e2, n2]) if e2 else None
                handle = txt(v(row, "BR"))
                pend = R.PERFIL_PENDENTE if any(str(v(row, c) or "").strip().upper() == R.PERFIL_PENDENTE
                                                 for c in ("AE", "AF", "AG", "AH")) else None
                campos = dict(
                    ponto=ponto, vigente=1, item=v(row, "A"), frente_id=txt(v(row, "K")), equipe=txt(v(row, "BM")),
                    tipo=txt(v(row, "BN")), pb=txt(v(row, "BO")), prancha=txt(v(row, "U")), area=txt(v(row, "X")),
                    de_mm=num(v(row, "H")), material=txt(v(row, "I")), solda=txt(v(row, "G")),
                    metodologia=txt(v(row, "N")), classe=txt(v(row, "S")), extensao_m=num(v(row, "M")),
                    codigo_bm=txt(v(row, "AC")), recobrimento_m=num(v(row, "AE")), profundidade_m=num(v(row, "AF")),
                    largura_m=num(v(row, "AG")), lastro_m=num(v(row, "AH")), perfil_status=pend,
                    quantidades=json.dumps({c: num(v(row, c)) for c in COLUNAS_SERVICO if num(v(row, c)) is not None}),
                    no_ini=txt(v(row, "BP")), no_fim=txt(v(row, "BQ")), handle=handle, utm_ini=utm_ini, utm_fim=utm_fim,
                    lat=num(v(row, "BW")), lon=num(v(row, "BX")), observacao=txt(v(row, "AB")))
                linhas.append({"k": k, "ponto": ponto, "handle": handle, "ini": utm_ini, "fim": utm_fim, "campos": campos,
                               "celulas": {c: bruto(v(row, c)) for c in
                                           ("R", "K", "H", "M", "N", "AC", "AE", "AF", "AG", "AH", "BR", "BP", "BQ")}})
        return linhas

    def trechos(self, p):
        """Identidade = handle CAD + pontas UTM (únicas na 06.1). O nome (PONTO) nunca decide sozinho."""
        fid = self.fonte("redes", p)
        linhas = self._ler_trechos(p)
        exist = {r["id"]: dict(r) for r in self.c.execute("SELECT * FROM trechos")}
        por_fisica, por_pontas = defaultdict(list), defaultdict(list)
        for t in exist.values():
            por_fisica[(t["handle"], t["utm_ini"], t["utm_fim"])].append(t["id"])
            if t["utm_ini"] and t["utm_fim"]:
                por_pontas[(t["utm_ini"], t["utm_fim"])].append(t["id"])
        conc = self.concil["TRECHO"]
        atrib, usados = [None] * len(linhas), set()

        def atribuir(n, tid):
            atrib[n] = tid
            usados.add(tid)

        for n, l in enumerate(linhas):                                   # 1) de-para explícito
            e = conc["usar"].get((l["ponto"], l["handle"]))
            if not e:
                continue
            if e["usar_id"] not in exist or e["usar_id"] in usados:
                self.pendencia("CONCILIACAO_INVALIDA", l["ponto"], f"usar_id {e['usar_id']} inexistente ou já usado",
                               bloqueante=True)
            else:
                atribuir(n, e["usar_id"])
                self._concil_aplicada(e, e["usar_id"], f"{l['ponto']} (handle {l['handle']}) → {e['usar_id']}")
        for n, l in enumerate(linhas):                                   # 2) mesma identidade física
            if atrib[n] is None:
                cands = [i for i in por_fisica.get((l["handle"], l["ini"], l["fim"]), []) if i not in usados]
                if len(cands) == 1:
                    atribuir(n, cands[0])
                elif len(cands) > 1:
                    self.pendencia("TRECHO_AMBIGUO", l["ponto"], f"geometria igual a {cands}", bloqueante=True)
        for n, l in enumerate(linhas):                                   # 3) mesmas pontas, handle redesenhado
            if atrib[n] is None and l["ini"] and l["fim"]:
                cands = [i for i in por_pontas.get((l["ini"], l["fim"]), []) if i not in usados]
                if len(cands) == 1:
                    atribuir(n, cands[0])
                    self.pendencia("TRECHO_HANDLE_ALTERADO", cands[0],
                                   f"{l['ponto']}: handle {exist[cands[0]]['handle']} → {l['handle']}, mesmas pontas UTM")
        novos = [n for n in range(len(linhas)) if atrib[n] is None]
        for n in novos:                                                  # 4) novo ID (sem herança presumida)
            atrib[n] = self.novo_id("TR-", "trechos", 5)
        for n, l in enumerate(linhas):
            t = exist.get(atrib[n])
            if t and t["ponto"] != l["ponto"]:
                self.pendencia("TRECHO_RENOMEADO", atrib[n], f"{t['ponto']} → {l['ponto']} (mesma geometria física: "
                                                             f"handle {l['handle']}, pontas {l['ini']} → {l['fim']})")
        # retirada: o que não tem correspondência física
        hist = self.historico("TRECHO")
        for tid, t in exist.items():
            if not t["vigente"] or tid in usados:
                continue
            suc = [linhas[n]["ponto"] for n in novos
                   if linhas[n]["ponto"] == t["ponto"] or (linhas[n]["handle"] == t["handle"] and
                                                           {linhas[n]["ini"], linhas[n]["fim"]} & {t["utm_ini"], t["utm_fim"]})]
            if tid in conc["retirar"]:
                self._concil_aplicada(conc["retirar"][tid], tid, f"retirada confirmada de {t['ponto']}")
            elif tid in hist:
                self.pendencia("TRECHO_COM_HISTORICO_SEM_CORRESPONDENCIA", tid,
                               f"{t['ponto']} (handle {t['handle']}) tem histórico operacional e não tem correspondência "
                               f"física na nova revisão (traçado alterado, divisão, união ou exclusão). Possíveis "
                               f"sucessores: {suc or 'nenhum'}. Registre em config/conciliacoes.json qual linha herda o "
                               f"ID (usar_id) ou confirme a retirada (retirar_id).", bloqueante=True)
            elif suc:
                self.pendencia("TRECHO_SUBSTITUIDO", tid, f"{t['ponto']} sem histórico, geometria diferente: novo ID para "
                                                          f"{suc}; o ID antigo fica não vigente")
            self.retirar("trechos", "trecho", tid)
        # nomes temporários para permitir troca de nomes entre trechos (índice único do PONTO vigente)
        for n, l in enumerate(linhas):
            t = exist.get(atrib[n])
            if t and t["ponto"] != l["ponto"]:
                self.c.execute("UPDATE trechos SET ponto=? WHERE id=?", ("⁣" + atrib[n], atrib[n]))
        for n, l in enumerate(linhas):
            tid = atrib[n]
            self.upsert("trechos", "trecho", {"id": tid}, l["campos"], atual=exist.get(tid))
            self.nome_hist("trecho", tid, l["ponto"], fid)
            self.prov("trecho", tid, fid, "REDE", l["k"], l["celulas"])

    # -------------------------------------------------------------- conexões (06.2)
    def _ler_conexoes(self, p):
        with planilha(p) as wb:
            ws = wb["CONEXÕES"]
            cab_linha, idx = None, {}
            for i, row in enumerate(ws.iter_rows(min_row=1, max_row=15, values_only=True), 1):
                nomes = [str(c or "").strip().upper() for c in row]
                if CAB_CONEXOES["nomenclatura"] in nomes:
                    cab_linha = i
                    for campo, rotulo in CAB_CONEXOES.items():
                        if rotulo in nomes:
                            idx[campo] = nomes.index(rotulo)
                    break
            faltando = [CAB_CONEXOES[c] for c in CAB_CONEXOES if c not in idx]
            if cab_linha is None or faltando:
                raise RuntimeError(f"06.2: cabeçalhos não encontrados: {faltando or 'linha de cabeçalho'}")
            self.colunas_06_2 = {c: get_column_letter(j + 1) for c, j in idx.items()}
            linhas = []
            for k, row in enumerate(ws.iter_rows(min_row=cab_linha + 1, values_only=True), cab_linha + 1):
                g = lambda campo: row[idx[campo]] if idx[campo] < len(row) else None
                if txt(g("nomenclatura")):
                    linhas.append((k, {c: g(c) for c in idx}))
        return linhas

    def _evidencias(self, p, p_conexoes, p_redes):
        """EVIDÊNCIAS (aba ELEMENTOS_ZONA): conferência do agrupamento + zona, status e ID da auditoria.
        Devolve None se a fonte não puder ser usada (ausente ou de outra revisão)."""
        if not p:
            self.pendencia("EVIDENCIAS_AUSENTES", self.cfg["arquivos"].get("evidencias"),
                           "arquivo de EVIDÊNCIAS ausente: zona, status de posição e ID da auditoria foram PRESERVADOS "
                           "(não atualizados) e o agrupamento não foi conferido")
            return None
        nome = os.path.basename(p).upper()
        m2, m1 = re.search(r"REV(\d+)_06\.2", nome), re.search(r"REV(\d+)_06\.1", nome)
        rev_con, rev_red = revisao_do_nome(os.path.basename(p_conexoes)), revisao_do_nome(os.path.basename(p_redes))
        if m2 and m1:
            difs = [f"{a} da evidência × {b} da fonte" for a, b in ((f"REV{m2.group(1)} (06.2)", rev_con),
                                                                     (f"REV{m1.group(1)} (06.1)", rev_red))
                    if a.split()[0] != b]
            if difs:
                self.pendencia("EVIDENCIAS_REVISAO_INCOMPATIVEL", os.path.basename(p),
                               "; ".join(difs) + " — evidências não usadas; zona/status preservados", bloqueante=True)
                return None
        else:
            self.pendencia("EVIDENCIAS_REVISAO_NAO_IDENTIFICADA", os.path.basename(p),
                           "o nome não indica as revisões da 06.2/06.1; a correspondência fica a cargo do conjunto de "
                           "controles aprovado para estes arquivos (hash)")
        self.fonte("evidencias", p)
        out = {}
        with planilha(p) as wb:
            for row in wb["ELEMENTOS_ZONA"].iter_rows(min_row=2, values_only=True):
                if row[0]:
                    g_ = frozenset(s.strip() for s in str(row[1]).split("|") if s.strip())
                    out[g_] = {"id": txt(row[0]), "zona": txt(row[5]), "status": txt(row[8])}
        return out

    def conexoes(self, p, p_evidencias=None, p_redes=None):
        fid = self.fonte("conexoes", p)
        linhas = self._ler_conexoes(p)
        juntada = {r[0]: r[1] for r in self.c.execute("SELECT item, unidade FROM bm_itens WHERE planilha='JUNTADA' AND vigente=1")}

        # elemento físico: handle CAD; sem handle, o número do projeto (identidade provisória)
        grupos = defaultdict(list)
        for k, r in linhas:
            nom = txt(r["nomenclatura"])
            chave = f"H:{txt(r['handle'])}" if txt(r["handle"]) else f"N:{numero_projeto(nom)}"
            grupos[chave].append((k, r))
        nom_chave = defaultdict(set)
        for chave, rs in grupos.items():
            for _, r in rs:
                nom_chave[txt(r["nomenclatura"])].add(chave)
        for nom, ch in nom_chave.items():
            if len(ch) > 1:
                self.pendencia("NOMENCLATURA_EM_MAIS_DE_UM_ELEMENTO", nom, f"chaves {sorted(ch)}", bloqueante=True)

        ev = self._evidencias(p_evidencias, p, p_redes)
        if ev is not None:
            meus = {frozenset(txt(r["nomenclatura"]) for _, r in rs) for rs in grupos.values()}
            iguais = sum(1 for g_ in ev if g_ in meus)
            self.conferencias["agrupamento_vs_evidencias"] = {"evidencias": len(ev), "iguais": iguais}
            for g_ in ev:
                if g_ not in meus:
                    self.pendencia("AGRUPAMENTO_DIVERGE_EVIDENCIAS", " | ".join(sorted(g_)), "grupo das EVIDÊNCIAS sem par na 06.2")

        info = {}
        for chave, rs in grupos.items():
            noms = sorted({txt(r["nomenclatura"]) for _, r in rs})
            coords = {(num(r["utm_e"]), num(r["utm_n"]), num(r["lat"]), num(r["lon"])) for _, r in rs}
            numero = numero_projeto(noms[0])
            if len(coords) > 1:
                self.pendencia("COORDENADAS_DIVERGENTES", numero, f"{len(coords)} coordenadas distintas no mesmo elemento")
            info[chave] = {"noms": noms, "numero": numero,
                           "coord": sorted(coords, key=lambda x: tuple(v is None for v in x))[0]}

        el_ids = self._resolver_elementos(grupos, info)
        at_exist = {r["id"]: dict(r) for r in self.c.execute("SELECT * FROM atividades")}   # antes de nomes temporários
        at_ids = self._resolver_atividades(grupos, info, el_ids)

        # aplica: elementos → atividades → itens/avisos
        sem_ev = 0
        el_exist = {r["id"]: dict(r) for r in self.c.execute("SELECT * FROM elementos")}
        for chave, rs in grupos.items():
            eid, inf = el_ids[chave], info[chave]
            ue, un, la, lo = inf["coord"]
            campos = dict(chave_natural=chave, numero=inf["numero"], tipo=tipo_elemento(inf["numero"]), vigente=1,
                          handle=txt(rs[0][1]["handle"]),
                          pb=inf["noms"][0].split("_")[1] if inf["noms"][0].count("_") >= 3 else None,
                          utm_e=ue, utm_n=un, lat=la, lon=lo)
            dados_ev = ev.get(frozenset(inf["noms"])) if ev is not None else None
            if dados_ev is not None:
                campos.update(zona=dados_ev["zona"], status_posicao=dados_ev["status"], id_auditoria=dados_ev["id"])
            else:
                sem_ev += 1          # zona/status/ID da auditoria preservados (não viram NULL)
                if ev is not None:
                    self.pendencia("ELEMENTO_SEM_GRUPO_NAS_EVIDENCIAS", inf["numero"],
                                   f"{' | '.join(inf['noms'])}: zona/status preservados")
            self.upsert("elementos", "elemento", {"id": eid}, campos, atual=el_exist.get(eid))
            self.prov("elemento", eid, fid, "CONEXÕES", rs[0][0], {"handle": txt(rs[0][1]["handle"]), "nomenclaturas": inf["noms"]})
        self.conferencias["elementos_com_zona_preservada"] = sem_ev

        for (chave, tipo_frente), (aid, linhas_at) in at_ids.items():
            noms_at = sorted({txt(r["nomenclatura"]) for _, r in linhas_at})
            if len(noms_at) > 1:
                self.pendencia("ATIVIDADE_COM_VARIAS_NOMENCLATURAS", info[chave]["numero"], " | ".join(noms_at))
            k0, r0 = linhas_at[0]
            self.upsert("atividades", "atividade", {"id": aid}, dict(
                elemento_id=el_ids[chave], tipo_frente=tipo_frente, nomenclatura=noms_at[0], frente_id=txt(r0["local"]),
                equipe=txt(r0["equipe"]), vigente=1), atual=at_exist.get(aid))
            self.nome_hist("atividade", aid, noms_at[0], fid)
            self.prov("atividade", aid, fid, "CONEXÕES", k0, {"nomenclatura": noms_at[0], "local": txt(r0["local"])})
            self._itens(aid, noms_at[0], linhas_at, juntada, fid)
        self.retirar_nao_vistos("avisos_origem", "aviso", self._av_vistos)

    def _resolver_elementos(self, grupos, info):
        """chave natural → ID estável. Promove N:<número> para H:<handle> sem perder ID nem histórico."""
        exist = {r["id"]: dict(r) for r in self.c.execute("SELECT * FROM elementos")}
        por_chave = {e["chave_natural"]: i for i, e in exist.items()}
        noms_ex = defaultdict(set)
        for r in self.c.execute("SELECT elemento_id, nomenclatura FROM atividades WHERE vigente=1"):
            noms_ex[r[0]].add(r[1])
        hist_at = self.historico("ATIVIDADE")
        hist = {r[0] for r in self.c.execute("SELECT elemento_id, id FROM atividades") if r[1] in hist_at}
        conc = self.concil["ELEMENTO"]
        atrib, usados = {}, set()
        for chave in grupos:                                             # 1) de-para explícito
            e = conc["usar"].get(chave)
            if e:
                if e["usar_id"] not in exist or e["usar_id"] in usados:
                    self.pendencia("CONCILIACAO_INVALIDA", chave, f"usar_id {e['usar_id']} inexistente ou já usado",
                                   bloqueante=True)
                else:
                    atrib[chave] = e["usar_id"]
                    usados.add(e["usar_id"])
                    self._concil_aplicada(e, e["usar_id"], f"{chave} → {e['usar_id']}")
        for chave in grupos:                                             # 2) mesma chave natural
            if chave not in atrib and por_chave.get(chave) and por_chave[chave] not in usados:
                atrib[chave] = por_chave[chave]
                usados.add(por_chave[chave])
        for chave, inf in info.items():                                  # 3) identidade provisória / handle redesenhado
            if chave in atrib:
                continue
            cands = [i for i, e in exist.items() if e["vigente"] and i not in usados and e["numero"] == inf["numero"]
                     and e["chave_natural"] not in grupos]
            if len(cands) != 1:
                continue
            e = exist[cands[0]]
            mesmos_nomes = noms_ex[e["id"]] <= set(inf["noms"])
            if chave.startswith("H:") and e["chave_natural"].startswith("N:") and mesmos_nomes:
                atrib[chave] = e["id"]
                usados.add(e["id"])
                self.pendencia("IDENTIDADE_PROMOVIDA", e["id"],
                               f"{inf['numero']}: {e['chave_natural']} → {chave} (mesmo nº e mesmas nomenclaturas; ID e "
                               f"histórico mantidos; posição vinda da fonte: {inf['coord'][2:] if inf['coord'][2] else 'sem posição'})")
            elif (chave.startswith("H:") and e["chave_natural"].startswith("H:") and mesmos_nomes
                  and _mesma_posicao((e["utm_e"], e["utm_n"]), inf["coord"][:2])):
                atrib[chave] = e["id"]
                usados.add(e["id"])
                self.pendencia("ELEMENTO_HANDLE_ALTERADO", e["id"],
                               f"{inf['numero']}: {e['chave_natural']} → {chave} (mesmo nº, nomenclaturas e posição)")
        for chave in grupos:                                             # 4) novos
            if chave not in atrib:
                atrib[chave] = self.novo_id("EL-", "elementos", 4)
        for eid, e in exist.items():                                     # retirada
            if not e["vigente"] or eid in usados:
                continue
            if eid in conc["retirar"]:
                self._concil_aplicada(conc["retirar"][eid], eid, f"retirada confirmada de {e['numero']}")
            elif eid in hist:
                self.pendencia("ELEMENTO_COM_HISTORICO_SEM_CORRESPONDENCIA", eid,
                               f"{e['numero']} ({e['chave_natural']}) tem histórico e não tem correspondência na nova "
                               f"revisão. Registre o de-para (usar_id) ou confirme a retirada (retirar_id) em "
                               f"config/conciliacoes.json.", bloqueante=True)
            self.retirar("elementos", "elemento", eid)
        return atrib

    def _resolver_atividades(self, grupos, info, el_ids):
        exist = {r["id"]: dict(r) for r in self.c.execute("SELECT * FROM atividades")}
        por_chave = {(a["elemento_id"], a["tipo_frente"]): i for i, a in exist.items()}
        out = {}
        for chave, rs in grupos.items():
            por_frente = defaultdict(list)
            for k, r in rs:
                nom = txt(r["nomenclatura"])
                por_frente[nom.split("_")[2] if nom.count("_") >= 3 else "?"].append((k, r))
            for tf, linhas_at in por_frente.items():
                aid = por_chave.get((el_ids[chave], tf)) or self.novo_id("AT-", "atividades", 4)
                out[(chave, tf)] = (aid, linhas_at)
        vistos = {aid for aid, _ in out.values()}
        hist = self.historico("ATIVIDADE")
        conc = self.concil["ATIVIDADE"]
        for aid, a in exist.items():
            if not a["vigente"] or aid in vistos:
                continue
            if aid in conc["retirar"]:
                self._concil_aplicada(conc["retirar"][aid], aid, f"retirada confirmada de {a['nomenclatura']}")
            elif aid in hist and a["elemento_id"] not in self.concil["ELEMENTO"]["retirar"]:
                self.pendencia("ATIVIDADE_COM_HISTORICO_SEM_CORRESPONDENCIA", aid,
                               f"{a['nomenclatura']} tem histórico e saiu da 06.2. Confirme a retirada (retirar_id) ou "
                               f"concilie o elemento em config/conciliacoes.json.", bloqueante=True)
            self.retirar("atividades", "atividade", aid)
        # nomes temporários para permitir troca de nomenclatura entre atividades (índice único vigente)
        novos_nomes = {aid: sorted({txt(r["nomenclatura"]) for _, r in la})[0] for aid, la in out.values()}
        for aid, nome in novos_nomes.items():
            if aid in exist and exist[aid]["nomenclatura"] != nome:
                self.c.execute("UPDATE atividades SET nomenclatura=? WHERE id=?", ("⁣" + aid, aid))
        return out

    def _itens(self, aid, nomenclatura, linhas_at, juntada, fid):
        """Itens por assinatura técnica completa. Ordem das linhas não muda identidade.
        Correspondência ambígua → itens novos + pendência (nada é transferido por presunção)."""
        novos = []
        for k, r in linhas_at:
            codigo = txt(r["codigo_bm"])
            if codigo in juntada and juntada[codigo]:
                unidade, fonte_un = R.normalizar_unidade(juntada[codigo]), f"BM Juntada {codigo}"
            else:
                unidade, fonte_un = None, f"PENDENTE: código '{codigo}' sem item/unidade na Juntada"
            campos = dict(descricao=txt(r["descricao"]) or "(sem descrição)", pressao=txt(r["pressao"]),
                          resina=txt(r["resina"]), tipo_conexao=txt(r["tipo_conexao"]), solda=txt(r["solda"]),
                          diametro=txt(r["diametro"]), material=txt(r["material"]), grupo=txt(r["grupo"]),
                          quantidade=num(r["quantidade"]), unidade=unidade, unidade_fonte=fonte_un, codigo_bm=codigo,
                          enquadramento_bm=txt(r["enquadramento"]), obs_enquadramento_bm=txt(r["obs_enquadramento"]),
                          trechos=txt(r["trechos"]), tipo_entroncamento=txt(r["tipo_entroncamento"]),
                          status_conexao=txt(r["status_conexao"]))
            novos.append({"k": k, "r": r, "campos": campos, "sig": assinatura_item(campos)})
        exist = [dict(x) for x in self.c.execute("SELECT * FROM itens_projeto WHERE atividade_id=? ORDER BY id", (aid,))]
        seq = Counter()
        for x in exist:
            sig, _, n = x["chave_ocorrencia"].rpartition("#")
            seq[sig] = max(seq[sig], int(n) if n.isdigit() else 0)
        usados, atrib = set(), [None] * len(novos)
        iguais = lambda x, nv: x["chave_ocorrencia"].rpartition("#")[0] == nv["sig"] and all(
            x[c] == v for c, v in nv["campos"].items())
        for vig in (1, 0):                                   # 1) correspondência exata (vigentes, depois reativação)
            for n, nv in enumerate(novos):
                if atrib[n] is None:
                    x = next((x for x in exist if x["vigente"] == vig and x["id"] not in usados and iguais(x, nv)), None)
                    if x:
                        atrib[n] = x
                        usados.add(x["id"])
        resto_n = defaultdict(list)                          # 2) mesma assinatura, atributo mutável alterado
        for n, nv in enumerate(novos):
            if atrib[n] is None:
                resto_n[nv["sig"]].append(n)
        for sig, ns in resto_n.items():
            ex = [x for x in exist if x["vigente"] and x["id"] not in usados and x["chave_ocorrencia"].rpartition("#")[0] == sig]
            if len(ns) == 1 and len(ex) == 1:
                atrib[ns[0]] = ex[0]
                usados.add(ex[0]["id"])
            elif ex:
                self.pendencia("ITEM_SEM_CORRESPONDENCIA_INEQUIVOCA", nomenclatura,
                               f"{sig}: {len(ex)} item(ns) anterior(es) × {len(ns)} na revisão, com atributos diferentes — "
                               f"criados itens novos; os anteriores ficam não vigentes (sem transferência presumida)")
        for n, nv in enumerate(novos):
            x = atrib[n]
            if x is None:
                seq[nv["sig"]] += 1
                iid, chave_oc, atual = self.novo_id("IP-", "itens_projeto", 6), f"{nv['sig']}#{seq[nv['sig']]}", None
            else:
                iid, chave_oc, atual = x["id"], x["chave_ocorrencia"], x
            self.upsert("itens_projeto", "item_projeto", {"id": iid},
                        dict(atividade_id=aid, chave_ocorrencia=chave_oc, vigente=1, **nv["campos"]), atual=atual)
            self.prov("item_projeto", iid, fid, "CONEXÕES", nv["k"], {self.colunas_06_2[c]: bruto(v) for c, v in nv["r"].items()})
            usados.add(iid)
            for tipo, campo in (("PROBLEMA", "consideracao"), ("SUGESTAO", "sugestao")):
                for texto in [t.strip() for t in str(nv["r"][campo] or "").split(" | ") if t.strip()]:
                    h = _hash_texto(texto)
                    res = self.upsert("avisos_origem", "aviso", {"item_id": iid, "tipo": tipo, "texto_hash": h},
                                      dict(prioridade=txt(nv["r"]["prioridade"]), texto=texto, vigente=1))
                    avid = self.c.execute("SELECT id FROM avisos_origem WHERE item_id=? AND tipo=? AND texto_hash=?",
                                          (iid, tipo, h)).fetchone()[0]
                    self._av_vistos.add(avid)
                    if res == "novo":
                        self.prov("aviso", str(avid), fid, "CONEXÕES", nv["k"], {self.colunas_06_2[campo]: texto})
        for x in exist:
            if x["vigente"] and x["id"] not in usados:
                self.retirar("itens_projeto", "item_projeto", x["id"])

    def materiais(self):
        """Tubos PEAD da rede (fornecimento Caesb, SL0327 §6) por DE; itens MATERIAL do BM (aquisição Ankara);
        consumo diário. Nunca apaga material (há movimentos ligados)."""
        for r in self.c.execute("SELECT de_mm, SUM(extensao_m) ext FROM trechos WHERE vigente=1 GROUP BY de_mm ORDER BY de_mm").fetchall():
            de = int(r["de_mm"])
            self.upsert("materiais", "material", {"id": f"CAESB-PEAD-DE{de}"}, dict(
                descricao=f"TUBO PEAD DE{de} PN10 — fornecimento Caesb", unidade="M", propriedade="CAESB",
                preco_unitario=0, qtd_projeto=round(r["ext"], 2)))
        for b in self.c.execute("SELECT * FROM bm_itens WHERE planilha='VERTICAL' AND classificacao='MATERIAL' AND vigente=1").fetchall():
            self.upsert("materiais", "material", {"id": b["item"]}, dict(
                descricao=b["descricao"], unidade=b["unidade"] or "UN", propriedade="ANKARA",
                preco_unitario=b["preco_unitario"] or 0, bm_item_id=b["id"], qtd_projeto=b["quantidade"]))
        for m in self.cfg["materiais_consumo"]:
            self.c.execute("INSERT OR IGNORE INTO materiais(id,descricao,unidade,propriedade) VALUES (?,?,?,'CONSUMO')",
                           (m["id"], m["descricao"], m["unidade"]))

    def decisoes(self):
        """Só INSERE decisões, escopos e vínculos novos; nunca altera o que já foi registrado.
        O padrão de texto (padrao_sugestao) só serve para sugerir avisos parecidos para revisão."""
        p = os.path.join(BASE, "config", "decisoes.json")
        if not os.path.exists(p):
            return 0
        with open(p, encoding="utf-8") as f:
            lista = json.load(f)["decisoes"]
        n = 0
        agora = S._agora()
        for d in lista:
            n += self.c.execute(
                """INSERT OR IGNORE INTO decisoes(id,assunto,objeto_tipo,objeto_ref,status,decisao,responsavel,evidencia,data,
                     padrao_aviso,registrada_em) VALUES (?,?,?,?,?,?,?,?,?,?,?)""",
                (d["id"], d["assunto"], d.get("objeto_tipo"), d.get("objeto_ref"), d["status"], d["decisao"],
                 d.get("responsavel"), d["evidencia"], d.get("data"), d.get("padrao_sugestao"), agora)).rowcount
            for ref in d.get("escopo", {}).get("elementos", []):
                self.c.execute("INSERT OR IGNORE INTO decisao_objetos VALUES (?, 'ELEMENTO', ?)", (d["id"], ref))
            for v in d.get("vinculos_avisos", []):
                self.c.execute("INSERT OR IGNORE INTO decisao_vinculos VALUES (?,?,?,?)",
                               (d["id"], v["elemento"], v["texto_hash"], agora))
        for dec, numero, h in self.c.execute("SELECT decisao_id, elemento_numero, texto_hash FROM decisao_vinculos").fetchall():
            achou = self.c.execute(
                """SELECT 1 FROM avisos_origem x JOIN itens_projeto i ON i.id=x.item_id JOIN atividades a ON a.id=i.atividade_id
                   JOIN elementos e ON e.id=a.elemento_id WHERE e.numero=? AND x.texto_hash=? AND x.vigente=1 LIMIT 1""",
                (numero, h)).fetchone()
            if not achou:
                self.pendencia("VINCULO_DE_DECISAO_SEM_AVISO_VIGENTE", dec,
                               f"{numero} · aviso {h[:12]} não está na revisão atual (o vínculo é mantido)")
        return n

    # -------------------------------------------------------------- execução
    def executar(self):
        cfg = self.cfg
        obrig = {k: caminho(cfg, k, self.pasta) for k in ("redes", "conexoes", "bm", "reajuste")}
        faltam = [cfg["arquivos"][k] for k, v in obrig.items() if not v]
        if faltam:
            raise RuntimeError("Faltam planilhas obrigatórias em planilhas/: " + ", ".join(faltam))
        for e, motivo in self._concil_invalidas:
            self.pendencia("CONCILIACAO_INVALIDA", json.dumps(e, ensure_ascii=False)[:200], motivo, bloqueante=True)
        self.parametro("contrato", cfg["contrato"], "config/fontes.json")
        vc = cfg["valor_contrato"]
        self.parametro("valor_contrato", vc["valor"], vc["fonte"])
        self.parametro("medicao_oficial_liberada", "NAO", "trava das etapas 0/1 (alterar só após validação)", sobrescrever=False)
        self.equipes()
        self.bm(obrig["bm"])
        self.reajuste(obrig["reajuste"])
        self.trechos(obrig["redes"])
        self.conexoes(obrig["conexoes"], caminho(cfg, "evidencias", self.pasta), obrig["redes"])
        self.materiais()
        self.decisoes_novas = self.decisoes()


# ------------------------------------------------------------------ migração de esquema
TABELAS_V1_RENOMEAR = ["trechos", "elementos", "conjuntos", "pecas", "medicoes", "medicao_lancamentos", "parametros"]


def _add_coluna(conn, tabela, coluna, tipo):
    if coluna not in {r[1] for r in conn.execute(f"PRAGMA table_info({tabela})")}:
        conn.execute(f"ALTER TABLE {tabela} ADD COLUMN {coluna} {tipo}")


def preparar_v1(conn):
    """Renomeia as tabelas v1 de estrutura incompatível para v1_* (preservadas para auditoria) e cria o esquema atual.

    Deve rodar com foreign_keys=OFF e legacy_alter_table=ON (as referências das tabelas operacionais
    continuam apontando para os nomes trechos/medicoes, que passam a ser as tabelas novas).
    """
    for t in TABELAS_V1_RENOMEAR:
        conn.execute(f'ALTER TABLE "{t}" RENAME TO "v1_{t}"')
    for ix in ("ix_pecas_conj", "ix_conj_elem", "ix_lanc_med"):
        conn.execute(f"DROP INDEX IF EXISTS {ix}")
    _add_coluna(conn, "bm_itens", "vigente", "INTEGER NOT NULL DEFAULT 1")
    S.aplicar_esquema(conn)
    conn.execute("INSERT OR IGNORE INTO parametros(chave,valor,fonte) SELECT chave, valor, 'migrado do esquema v1' "
                 "FROM v1_parametros WHERE chave NOT IN ('p0')")
    conn.execute("INSERT INTO medicoes(id,numero,tipo,periodo_ini,periodo_fim,status) "
                 "SELECT id, numero, 'TESTE', periodo_ini, periodo_fim, status FROM v1_medicoes")
    conn.execute(
        """INSERT INTO medicao_lancamentos(id,medicao_id,gaveta,servico,bm_item_id,alvo_tipo,alvo_id,nf_id,material_id,
             unidade,quantidade,status,motivo,observacao)
           SELECT id, medicao_id, gaveta, servico, bm_item_id, alvo_tipo, alvo_id, nf_id, material_id, unidade, quantidade,
             CASE WHEN servico LIKE 'MATERIAL%' OR bm_item_id IS NULL THEN 'PENDENTE' ELSE 'PROPOSTO' END,
             CASE WHEN servico LIKE 'MATERIAL%' THEN 'regime de faturamento pendente (matriz de regimes não aprovada)'
                  WHEN bm_item_id IS NULL THEN 'sem item de BM com fonte' END,
             observacao
           FROM v1_medicao_lancamentos""")


def preparar_v3(conn):
    """Esquema v2 → v3, sem perda: colunas novas, assinatura técnica completa nos itens já importados
    (mantendo o nº de cada um) e tabelas de escopo das decisões."""
    _add_coluna(conn, "pendencias_importacao", "bloqueante", "INTEGER NOT NULL DEFAULT 0")
    _add_coluna(conn, "medicoes", "congelada_em", "TEXT")
    _add_coluna(conn, "medicoes", "resumo_congelado", "TEXT")
    for c, t in (("preco_congelado", "REAL"), ("classificacao_congelada", "TEXT"), ("item_congelado", "TEXT")):
        _add_coluna(conn, "medicao_lancamentos", c, t)
    for r in conn.execute("SELECT * FROM itens_projeto").fetchall():
        n = r["chave_ocorrencia"].rsplit("|", 1)[-1]
        conn.execute("UPDATE itens_projeto SET chave_ocorrencia=? WHERE id=?", (f"{assinatura_item(dict(r))}#{n}", r["id"]))
    S.aplicar_esquema(conn)


def preparar_v4(conn):
    """Esquema v3 → v4 (etapa 3, estoque), sem perda: colunas de responsável/autorização nos movimentos e as
    tabelas de reservas, estornos e saldo declarado com a equipe. Movimentos existentes ficam como estão
    (BAIXA_TEORICA antiga continua registrada, mas deixa de reduzir o saldo físico)."""
    _add_coluna(conn, "movimentos", "responsavel", "TEXT")
    _add_coluna(conn, "movimentos", "autorizacao", "TEXT")
    S.aplicar_esquema(conn)


def preparar_v5(conn):
    """Esquema v4 → v5: vínculo do movimento com a retificação e tabelas de retificação e ajuste de reserva."""
    _add_coluna(conn, "movimentos", "retificacao_id", "INTEGER REFERENCES retificacoes(id)")
    S.aplicar_esquema(conn)


def preparar_v6(conn):
    """Esquema v5 → v6 (etapa 4, campo): metadados da foto original, informante do apontamento, intervalos
    por etapa e soldas. Fotos existentes ficam marcadas LEGADO (imagem já carimbada; original não preservado).
    Apontamentos antigos não ganham intervalo inventado: a posição deles é desconhecida."""
    for c, t in (("sha256", "TEXT"), ("bytes", "INTEGER"), ("mime", "TEXT"), ("origem_data", "TEXT"),
                 ("origem_gps", "TEXT"), ("enviada_em", "TEXT")):
        _add_coluna(conn, "fotos", c, t)
    conn.execute("UPDATE fotos SET origem_data='LEGADO' WHERE origem_data IS NULL")
    _add_coluna(conn, "apontamentos", "informado_por", "TEXT")
    S.aplicar_esquema(conn)


def preparar_v7(conn):
    """Esquema v6 → v7, sem perda:
    * intervalos_execucao ganha as origens LEGADO/RETIFICACAO e o vínculo com a conciliação do legado. O CHECK de
      uma tabela SQLite não muda por ALTER: a tabela é recriada e TODAS as linhas copiadas (conferido) antes de
      descartar a cópia antiga;
    * fotos ganha `validacao` (preenchida depois por servicos.revalidar_fotos);
    * tabelas de conciliação do legado.
    Apontamentos antigos sem intervalo continuam como estão: viram "avanço legado sem posição", preservado no
    resumo de execução até a conciliação explícita."""
    _add_coluna(conn, "fotos", "validacao", "TEXT")
    sql = (conn.execute("SELECT sql FROM sqlite_master WHERE type='table' AND name='intervalos_execucao'").fetchone() or [""])[0]
    if sql and "LEGADO" not in sql:
        antes = conn.execute("SELECT COUNT(*) FROM intervalos_execucao").fetchone()[0]
        conn.execute("ALTER TABLE intervalos_execucao RENAME TO intervalos_execucao_v6")
        conn.execute("DROP INDEX IF EXISTS ix_intervalos")
        S.aplicar_esquema(conn)
        cols = [r[1] for r in conn.execute("PRAGMA table_info(intervalos_execucao_v6)")]
        lista = ", ".join(cols)
        conn.execute(f"INSERT INTO intervalos_execucao({lista}) SELECT {lista} FROM intervalos_execucao_v6")
        depois = conn.execute("SELECT COUNT(*) FROM intervalos_execucao").fetchone()[0]
        if depois != antes:
            raise RuntimeError(f"Migração v7: intervalos copiados {depois} ≠ {antes}; nada foi alterado.")
        conn.execute("DROP TABLE intervalos_execucao_v6")
    S.aplicar_esquema(conn)


def preparar_v8(conn):
    """Esquema v7 → v8, só acréscimos (nenhuma linha alterada):
    * intervalos_execucao: `anulado_por` / `retificacao_id` (retificação sem apagar) e `referencia` (registro do
      levantamento de campo);
    * apontamentos: `anulado_por` / `retificacao_id` (o original retificado sai das somas e fica no histórico);
    * tabelas retificacoes_intervalos e revisoes_cadastro (conferência interna Ankara, não é aprovação Caesb)."""
    retif = "INTEGER REFERENCES retificacoes_intervalos(id)"
    for tabela, coluna, tipo in (("intervalos_execucao", "referencia", "TEXT"), ("intervalos_execucao", "anulado_por", retif),
                                 ("intervalos_execucao", "retificacao_id", retif), ("apontamentos", "anulado_por", retif),
                                 ("apontamentos", "retificacao_id", retif)):
        _add_coluna(conn, tabela, coluna, tipo)
    S.aplicar_esquema(conn)


def preparar_v9(conn):
    """Esquema v8 → v9: tabelas de acesso e auditoria (usuários, sessões, auditoria, reenvio). Só acréscimos."""
    S.aplicar_esquema(conn)


def preparar_v10(conn):
    """Esquema v9 → v10: etapas complementares (liberação topográfica, escavação) e ocorrências. Só acréscimos."""
    S.aplicar_esquema(conn)


def _ddl_atual(tabela):
    with open(os.path.join(S.BASE, "schema.sql"), encoding="utf-8") as f:
        texto = f.read()
    m = re.search(rf"CREATE TABLE IF NOT EXISTS {tabela} \(.*?\n\);", texto, re.S)
    if not m:
        raise RuntimeError(f"Tabela {tabela} não está no schema.sql.")
    return m.group(0)


def _recriar_tabela(conn, tabela, marcador):
    """Recria a tabela com o CHECK/colunas do esquema atual (SQLite não altera CHECK por ALTER), sem perda:
    cria a nova com outro nome, copia TODAS as linhas (mesmos ids), confere a contagem, descarta a antiga e dá à nova o
    nome original. A tabela antiga NÃO é renomeada: assim as chaves estrangeiras das tabelas-filhas continuam
    apontando para o nome original (agora a nova). Roda com as chaves estrangeiras desligadas (importar.executar) e a
    integridade é conferida por foreign_key_check antes do COMMIT."""
    sql = (conn.execute("SELECT sql FROM sqlite_master WHERE type='table' AND name=?", (tabela,)).fetchone() or [""])[0]
    if not sql or marcador in sql:
        return 0
    nova = f"{tabela}__nova"
    antes = conn.execute(f"SELECT COUNT(*) FROM {tabela}").fetchone()[0]
    conn.execute(_ddl_atual(tabela).replace(f"CREATE TABLE IF NOT EXISTS {tabela} (", f"CREATE TABLE {nova} (", 1))
    novas = {r[1] for r in conn.execute(f"PRAGMA table_info({nova})")}
    cols = ", ".join(c for c in (r[1] for r in conn.execute(f"PRAGMA table_info({tabela})")) if c in novas)
    conn.execute(f"INSERT INTO {nova}({cols}) SELECT {cols} FROM {tabela}")
    depois = conn.execute(f"SELECT COUNT(*) FROM {nova}").fetchone()[0]
    if depois != antes:
        raise RuntimeError(f"Migração: {tabela} copiada com {depois} ≠ {antes} linhas; nada foi alterado.")
    conn.execute(f"DROP TABLE {tabela}")
    conn.execute(f"ALTER TABLE {nova} RENAME TO {tabela}")
    S.aplicar_esquema(conn)                          # índices da tabela
    return antes


def preparar_v11(conn):
    """Esquema v10 → v11 (compras e recebimento), sem perda:
    * notas_fiscais e movimentos são recriadas para aceitar os novos estados/tipos (cópia conferida, mesmos ids);
      notas já existentes ficam com origem DIRETA (registro anterior à conferência por item);
    * nf_itens ganha o vínculo com o item do pedido e a linha da nota; tabelas de compras, recebimento e inventário."""
    S.aplicar_esquema(conn)                          # cria fornecedores/pedidos/documentos (referenciados pelas novas)
    _recriar_tabela(conn, "notas_fiscais", "A_CONFERIR")
    _recriar_tabela(conn, "movimentos", "TRANSF_SAIDA")
    for coluna, tipo in (("pedido_item_id", "INTEGER REFERENCES pedido_itens(id)"), ("linha_id", "INTEGER REFERENCES nf_linhas(id)")):
        _add_coluna(conn, "nf_itens", coluna, tipo)
    conn.execute("UPDATE notas_fiscais SET origem='DIRETA' WHERE origem IS NULL")
    S.aplicar_esquema(conn)


def preparar_v12(conn):
    """Esquema v11 → v12: programação das equipes (planejamento). Só acréscimos."""
    S.aplicar_esquema(conn)


def preparar_v13(conn):
    """Esquema v12 → v13 (medição): chave da parcela em cada lançamento, quem fechou a medição e a conferência interna.
    Lançamentos existentes recebem a chave na ordem em que foram gravados; se a mesma parcela já aparecer antes (dupla
    medição antiga), o lançamento repetido fica sem chave — nada é apagado — e aparece no relatório de duplicidades."""
    for coluna in ("fechada_por TEXT", "fechada_em TEXT", "conferencia_id INTEGER"):
        _add_coluna(conn, "medicoes", *coluna.split(" ", 1))
    _add_coluna(conn, "medicao_lancamentos", "chave", "TEXT")
    vistas = {r[0] for r in conn.execute("SELECT chave FROM medicao_lancamentos WHERE chave IS NOT NULL")}
    sem_chave = 0
    for l in conn.execute("""SELECT l.id, m.tipo, l.gaveta, l.servico, l.alvo_tipo, l.alvo_id, b.item, l.unidade, l.nf_id,
                                    l.material_id FROM medicao_lancamentos l JOIN medicoes m ON m.id=l.medicao_id
                             LEFT JOIN bm_itens b ON b.id=l.bm_item_id WHERE l.chave IS NULL ORDER BY l.id""").fetchall():
        ch = S.chave_parcela(*tuple(l)[1:])
        if ch in vistas:
            sem_chave += 1
            continue
        vistas.add(ch)
        conn.execute("UPDATE medicao_lancamentos SET chave=? WHERE id=?", (ch, l[0]))
    S.aplicar_esquema(conn)
    return sem_chave


def congelar_fechadas(conn):
    """Medições já FECHADAS antes do congelamento: grava os valores do cadastro presente no banco ANTES
    desta importação, para que a nova revisão do BM não as recalcule."""
    ids = [r[0] for r in conn.execute("SELECT id FROM medicoes WHERE status='FECHADA' AND congelada_em IS NULL")]
    for mid in ids:
        S.congelar_medicao(conn, mid, "migração de esquema (cadastro do BM existente antes desta importação)")
    return ids


REFS_ALVO = ["movimentos", "checkins", "apontamentos", "fotos", "execucao", "medicao_lancamentos"]
REFS_TRECHO = [("cadastro", "trecho_id"), ("cadastro_historico", "trecho_id"), ("ligacoes_prediais", "trecho_id")]


def remapear_v1(conn, imp):
    """Troca as referências antigas (PONTO / nomenclatura) pelas identidades estáveis. Nada é apagado;
    referência sem correspondência fica com o valor antigo e vira pendência."""
    agora = S._agora()
    tr = {r[0]: r[1] for r in conn.execute("SELECT ponto, id FROM trechos WHERE vigente=1")}
    at = {r[0]: r[1] for r in conn.execute("SELECT nomenclatura, id FROM atividades WHERE vigente=1")}
    resumo = Counter()
    for tabela in REFS_ALVO:
        for tipo, antigo in conn.execute(f"SELECT DISTINCT alvo_tipo, alvo_id FROM {tabela} WHERE alvo_id IS NOT NULL").fetchall():
            if tipo == "TRECHO":
                novo_tipo, novo = "TRECHO", tr.get(antigo)
            elif tipo in ("CONJUNTO", "ELEMENTO"):
                novo_tipo, novo = "ATIVIDADE", at.get(antigo)
            else:
                continue
            conn.execute("INSERT OR REPLACE INTO migracao_mapa VALUES (?,?,?,?,?)",
                         (tabela, "alvo", f"{tipo}:{antigo}", f"{novo_tipo}:{novo}" if novo else None, agora))
            if novo:
                conn.execute(f"UPDATE {tabela} SET alvo_tipo=?, alvo_id=? WHERE alvo_tipo=? AND alvo_id=?",
                             (novo_tipo, novo, tipo, antigo))
                resumo[f"{tabela}:mapeados"] += 1
            else:
                conn.execute("INSERT INTO pendencias_importacao(importacao_id,tipo,referencia,detalhe) VALUES (?,?,?,?)",
                             (imp, "REFERENCIA_V1_SEM_PAR", f"{tabela}:{tipo}:{antigo}", "valor antigo preservado"))
                resumo[f"{tabela}:sem_par"] += 1
    for tabela, coluna in REFS_TRECHO:
        for (antigo,) in conn.execute(f"SELECT DISTINCT {coluna} FROM {tabela}").fetchall():
            novo = tr.get(antigo)
            conn.execute("INSERT OR REPLACE INTO migracao_mapa VALUES (?,?,?,?,?)", (tabela, coluna, antigo, novo, agora))
            if novo:
                conn.execute(f"UPDATE {tabela} SET {coluna}=? WHERE {coluna}=?", (novo, antigo))
                resumo[f"{tabela}:mapeados"] += 1
            else:
                conn.execute("INSERT INTO pendencias_importacao(importacao_id,tipo,referencia,detalhe) VALUES (?,?,?,?)",
                             (imp, "REFERENCIA_V1_SEM_PAR", f"{tabela}.{coluna}:{antigo}", "valor antigo preservado"))
                resumo[f"{tabela}:sem_par"] += 1
    return dict(resumo)


# ------------------------------------------------------------------ auditoria pós-carga (antes do COMMIT)
def obter_contagens(conn, conferencias):
    q = lambda sql, *a: conn.execute(sql, a).fetchone()[0]
    unidades = {r[0] or "PENDENTE": r[1] for r in conn.execute(
        "SELECT unidade, COUNT(*) FROM itens_projeto WHERE vigente=1 GROUP BY unidade")}
    out = {
        "trechos": q("SELECT COUNT(*) FROM trechos WHERE vigente=1"),
        "extensao_m": round(q("SELECT SUM(extensao_m) FROM trechos WHERE vigente=1") or 0, 2),
        "elementos": q("SELECT COUNT(*) FROM elementos WHERE vigente=1"),
        "nos": q("SELECT COUNT(*) FROM elementos WHERE vigente=1 AND tipo='NÓ'"),
        "caixas": q("SELECT COUNT(*) FROM elementos WHERE vigente=1 AND tipo<>'NÓ'"),
        "atividades": q("SELECT COUNT(*) FROM atividades WHERE vigente=1"),
        "itens_projeto": q("SELECT COUNT(*) FROM itens_projeto WHERE vigente=1"),
        "bm_vertical": q("SELECT COUNT(*) FROM bm_itens WHERE planilha='VERTICAL' AND vigente=1"),
        "bm_juntada": q("SELECT COUNT(*) FROM bm_itens WHERE planilha='JUNTADA' AND vigente=1"),
        "sem_posicao": sorted(r[0] for r in conn.execute(
            "SELECT numero FROM elementos WHERE vigente=1 AND (lat IS NULL OR lon IS NULL)")),
        "unidades_linhas": dict(sorted(unidades.items())),
        "quantidade_m": round(q("SELECT COALESCE(SUM(quantidade),0) FROM itens_projeto WHERE vigente=1 AND unidade='M'"), 4),
    }
    ag = conferencias.get("agrupamento_vs_evidencias")
    if ag:
        out["agrupamento_evidencias_iguais"] = ag["iguais"]
    return out


def controle_para(cfg, fontes):
    """Conjunto de controles para EXATAMENTE estes arquivos (hash de cada fonte). Não diz se é válido/aprovado."""
    atual = {k: v["sha256"] for k, v in fontes.items()}
    return next((c for c in cfg.get("controles", []) if c.get("fontes") == atual), None)


# Conferências que todo conjunto precisa trazer (tipo esperado de cada valor)
CONTROLES_OBRIGATORIOS = {"trechos": "inteiro", "extensao_m": "número", "elementos": "inteiro", "atividades": "inteiro",
                          "itens_projeto": "inteiro", "bm_vertical": "inteiro", "bm_juntada": "inteiro",
                          "sem_posicao": "lista de nº de projeto", "unidades_linhas": "unidade → nº de linhas"}
CONTROLES_OPCIONAIS = {"nos": "inteiro", "caixas": "inteiro", "quantidade_m": "número",
                       "agrupamento_evidencias_iguais": "inteiro"}
TEXTO_PROVISORIO = re.compile(r"\bA CONFIRMAR\b|\bPREENCHER\b|\bTODO\b|\bXXX\b|^\s*(PENDENTE|A DEFINIR|-+|\?+)\s*$",
                              re.I)


def _tipo_ok(tipo, v):
    inteiro = lambda x: isinstance(x, int) and not isinstance(x, bool) and x >= 0
    if tipo == "inteiro":
        return inteiro(v)
    if tipo == "número":
        return isinstance(v, (int, float)) and not isinstance(v, bool) and v >= 0
    if tipo == "lista de nº de projeto":
        return isinstance(v, list) and all(isinstance(x, str) and x.strip() for x in v)
    if tipo == "unidade → nº de linhas":
        return isinstance(v, dict) and bool(v) and all(isinstance(k, str) and k.strip() and inteiro(x) for k, x in v.items())
    return False


def validar_controle(c, perfil):
    """(problemas de estrutura, problemas de aprovação) de um conjunto de controles.

    Aprovação efetiva = bloco `aprovacao` com status "APROVADO", nome de quem aprovou e data ISO, sem
    texto provisório. O sistema nunca preenche a aprovação: ela é escrita pela pessoa em config/controles.json.
    """
    texto_ok = lambda v: isinstance(v, str) and bool(v.strip()) and not TEXTO_PROVISORIO.search(v)
    estrutura, aprovacao = [], []
    if not texto_ok(c.get("id")):
        estrutura.append("campo 'id' ausente ou provisório")
    if c.get("perfil") != perfil:
        estrutura.append(f"perfil do conjunto = {c.get('perfil')!r}; esta importação exige perfil {perfil!r}")
    fontes = c.get("fontes")
    if not isinstance(fontes, dict) or not fontes or not all(
            isinstance(h, str) and re.fullmatch(r"[0-9a-f]{64}", h) for h in fontes.values()):
        estrutura.append("campo 'fontes' ausente, vazio ou com hash SHA-256 inválido")
    esperado = c.get("esperado")
    if not isinstance(esperado, dict) or not esperado:
        estrutura.append("campo 'esperado' ausente ou vazio: nenhuma conferência de contagem seria feita")
        esperado = {}
    obrig = dict(CONTROLES_OBRIGATORIOS)
    if isinstance(fontes, dict) and "evidencias" in fontes:
        obrig["agrupamento_evidencias_iguais"] = "inteiro"
    faltam = [k for k in obrig if k not in esperado]
    if faltam and esperado:
        estrutura.append("conferências obrigatórias ausentes em 'esperado': " + ", ".join(faltam))
    tipos = {**CONTROLES_OPCIONAIS, **obrig}
    for k, v in esperado.items():
        if k not in tipos:
            estrutura.append(f"'esperado.{k}' não é uma conferência conhecida")
        elif not _tipo_ok(tipos[k], v):
            estrutura.append(f"'esperado.{k}' = {v!r}: tipo inválido (esperado {tipos[k]})")
    if not texto_ok(c.get("origem_dos_valores")):
        estrutura.append("campo 'origem_dos_valores' ausente ou provisório: indique o documento que confirma os valores")
    ap = c.get("aprovacao")
    if not isinstance(ap, dict):
        aprovacao.append("bloco 'aprovacao' ausente" + (" ('aprovado_por' solto não vale como aprovação)"
                                                         if "aprovado_por" in c else ""))
    else:
        if ap.get("status") != "APROVADO":
            aprovacao.append(f"aprovacao.status = {ap.get('status')!r} (precisa ser \"APROVADO\")")
        if not texto_ok(ap.get("por")):
            aprovacao.append("aprovacao.por ausente ou provisório: nome de quem conferiu e aprovou")
        try:
            date.fromisoformat(str(ap.get("em")))
        except ValueError:
            aprovacao.append(f"aprovacao.em = {ap.get('em')!r}: data AAAA-MM-DD da aprovação")
    return estrutura, aprovacao


def auditar(conn, cfg, imp, fontes, conferencias, modo="APLICACAO"):
    """Conferências da carga. `bloqueios` lista o que impede o COMMIT.

    Controle inexistente, incompleto ou sem aprovação efetiva bloqueia a APLICAÇÃO. Na SIMULAÇÃO
    vira aviso (`avisos_controle`): a simulação serve justamente para conferir e preparar a aprovação.
    """
    q = lambda sql, *a: conn.execute(sql, a).fetchone()[0]
    bloqueios, avisos_controle = [], []
    obtido = obter_contagens(conn, conferencias)
    controle = controle_para(cfg, fontes)
    perfil = cfg.get("perfil", "REAL")
    estrutura, aprovacao = validar_controle(controle, perfil) if controle else ([], [])
    if controle is None:
        avisos_controle.append({"tipo": "SEM_CONTROLES_APROVADOS", "referencia": "config/controles.json",
                                "detalhe": "não há conjunto de controles para estes arquivos (hash). Confira os valores "
                                           "obtidos contra as fontes e registre o conjunto em config/controles.json "
                                           "(modelo no relatório)."})
    if estrutura:
        avisos_controle.append({"tipo": "CONTROLE_INCOMPLETO", "referencia": controle.get("id"),
                                "detalhe": "; ".join(estrutura)})
    if aprovacao:
        avisos_controle.append({"tipo": "CONTROLE_SEM_APROVACAO", "referencia": controle.get("id"),
                                "detalhe": "; ".join(aprovacao) + ". A aprovação é escrita pela pessoa responsável em "
                                           "config/controles.json depois de conferir os valores; o sistema não a preenche."})
    if modo == "APLICACAO":
        bloqueios += avisos_controle
    esperado = controle.get("esperado") if controle else None
    if not isinstance(esperado, dict):
        contagens = {}
    else:
        contagens = {k: {"esperado": v, "obtido": obtido.get(k), "ok": obtido.get(k) == v}
                     for k, v in esperado.items()}
        for k, c in contagens.items():
            if not c["ok"]:
                bloqueios.append({"tipo": "CONTROLE_DIVERGENTE", "referencia": k,
                                  "detalhe": f"esperado {c['esperado']} × obtido {c['obtido']} ({controle.get('id')})"})

    # BM: a totalização do sistema reproduz o rodapé da própria planilha (mesma ordem de arredondamento)?
    bdi_m, bdi_s, k = (float(S.parametro(conn, x, 0)) for x in ("bdi_material", "bdi_servico", "coef_k"))
    itens = [(r[0], r[1] or 0, r[2] or 0) for r in conn.execute(
        "SELECT classificacao, quantidade, preco_unitario FROM bm_itens WHERE planilha='VERTICAL' AND vigente=1")]
    tot = R.totalizar_bm(itens, bdi_m, bdi_s, k, modo="contrato")
    ref = {x: float(S.parametro(conn, x, "nan")) for x in ("bm_total_material", "bm_total_servico", "p0_bm")}
    bm_ok = (abs(tot["material"] - ref["bm_total_material"]) < 0.005 and abs(tot["servico"] - ref["bm_total_servico"]) < 0.005
             and abs(tot["total_exato"] - ref["p0_bm"]) < 1e-4)
    totalizacao = {"sistema": {"material": tot["material"], "servico": tot["servico"], "p0": tot["total_exato"]},
                   "celulas_bm": ref, "ok": bm_ok}
    if not bm_ok:
        bloqueios.append({"tipo": "TOTALIZACAO_BM_DIVERGE_DAS_CELULAS", "referencia": "BM VERTICAL",
                          "detalhe": f"sistema {totalizacao['sistema']} × células {ref}: quantidades/preços não "
                                     f"reproduzem o rodapé da própria planilha"})

    vc = float(S.parametro(conn, "valor_contrato"))
    p0 = ref["p0_bm"]
    total_contratual = {"contrato_3_1": vc, "fonte_contrato": cfg["valor_contrato"]["fonte"], "bm_p0": p0,
                        "diferenca": round(p0 - vc, 7),
                        "situacao": "DIVERGÊNCIA REGISTRADA — conciliar (não ajustado)" if abs(p0 - vc) > 0.004 else "IGUAL"}

    unidades = {r[0] or "PENDENTE": {"linhas": r[1], "quantidade": round(r[2] or 0, 4)} for r in conn.execute(
        "SELECT unidade, COUNT(*), SUM(quantidade) FROM itens_projeto WHERE vigente=1 GROUP BY unidade ORDER BY unidade")}
    enquad = {r[0] or "(vazio)": r[1] for r in conn.execute(
        "SELECT enquadramento_bm, COUNT(*) FROM itens_projeto WHERE vigente=1 GROUP BY 1 ORDER BY 2 DESC")}
    avisos = {f"{r[0]}/{r[1] or 'sem prioridade'}": r[2] for r in conn.execute(
        "SELECT tipo, prioridade, COUNT(*) FROM avisos_origem WHERE vigente=1 GROUP BY 1,2 ORDER BY 1,2")}
    with open(os.path.join(BASE, "config", "servicos_rede.json"), encoding="utf-8") as f:
        servs = json.load(f)["gaveta1"]
    pendencias = {r[0]: r[1] for r in conn.execute(
        "SELECT tipo, COUNT(*) FROM pendencias_importacao WHERE importacao_id=? GROUP BY 1", (imp,))}
    return {
        "controle_aplicado": controle.get("id") if controle else None,
        "controle_situacao": ("INEXISTENTE" if controle is None else "INCOMPLETO" if estrutura
                              else "SEM_APROVACAO" if aprovacao else "APROVADO"),
        "controle_aprovacao": controle.get("aprovacao") if controle else None,
        "avisos_controle": avisos_controle,
        "contagens": contagens,
        # só há confirmação de conferência com conjunto completo, aprovado e sem divergência
        "contagens_ok": bool(controle) and not estrutura and not aprovacao and all(c["ok"] for c in contagens.values()),
        "obtido": obtido,
        "modelo_controles": None if controle else {
            "id": "PREENCHER", "perfil": perfil, "fontes": {k: v["sha256"] for k, v in fontes.items()},
            "arquivos": {k: v["nome"] for k, v in fontes.items()},
            "esperado": obtido, "origem_dos_valores": "PREENCHER: documento/fonte que confirma cada valor",
            "aprovacao": {"status": "PENDENTE", "por": None, "em": None},
            "_atencao": "valores 'esperado' copiados do obtido nesta simulação: confira cada um na fonte antes de aprovar"},
        "totalizacao_bm": totalizacao,
        "total_contratual": total_contratual,
        "unidades_06_2": unidades,
        "itens_sem_quantidade": q("SELECT COUNT(*) FROM itens_projeto WHERE vigente=1 AND quantidade IS NULL"),
        "enquadramento_bm_BA": enquad,
        "avisos_AA_AB_AW": avisos,
        "decisoes_registradas": q("SELECT COUNT(*) FROM decisoes"),
        "servicos_06_1_sem_item_confirmado": [s["servico"] for s in servs if s["status"] != "CONFIRMADO"],
        "trechos_mnd": [r[0] for r in conn.execute("SELECT ponto FROM trechos WHERE vigente=1 AND metodologia='MND' ORDER BY 1")],
        "pendencias_importacao": pendencias,
        "bloqueios": bloqueios,
    }
