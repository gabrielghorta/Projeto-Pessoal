"""Exportação da medição para .xlsx (conferência interna: gavetas, regras, elegibilidade, parcelas).

Esta planilha NÃO é o modelo oficial: o pacote no modelo oficial do BM (juntada, nª MC, cronograma FF e gráficos) sai de
pacote_medicao.py. Enquanto a medição oficial estiver bloqueada (D-56), toda exportação sai marcada TESTE / NÃO OFICIAL. Valores seguem a ordem do BM (ROUND por item; BDI por classe; K no total).
"""
import io
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import medicao as M
import regras as R
import servicos as S

LARANJA = "E85D22"
GRAFITE = "17202B"
CAB = PatternFill("solid", fgColor=GRAFITE)
SUB = PatternFill("solid", fgColor="FBE3D6")
FINO = Side(style="thin", color="E2E5EA")
BORDA = Border(left=FINO, right=FINO, top=FINO, bottom=FINO)
MOEDA = 'R$ #,##0.00'
QTD = '#,##0.000'
AVISO_TESTE = "TESTE — NÃO OFICIAL · conferência interna; regras financeiras e acesso pendentes de validação"


def _cabecalho(ws, titulo, sub, colunas, larguras, teste):
    ws.append([f"{titulo}  [{AVISO_TESTE}]" if teste else titulo])
    ws["A1"].font = Font(bold=True, size=13, color="C92A2A" if teste else LARANJA)
    ws.append([sub])
    ws["A2"].font = Font(italic=True, color="555555")
    ws.append([])
    ws.append(colunas)
    for i, w in enumerate(larguras, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
        c = ws.cell(row=4, column=i)
        c.font, c.fill, c.alignment, c.border = Font(bold=True, color="FFFFFF"), CAB, Alignment(wrap_text=True, vertical="center"), BORDA
    ws.freeze_panes = "A5"


def _fmt(ws, linha, formatos):
    for col, f in formatos.items():
        ws.cell(row=linha, column=col).number_format = f
    for c in ws[linha]:
        c.border = BORDA


def planilha_medicao(conn, medicao_id):
    r = S.resumo_medicao(conn, medicao_id)
    med, teste = r["medicao"], r["teste"]
    sub = f"{S.contrato(conn)} · {med['numero']}ª medição ({med['tipo']}) · {med['periodo_ini']} a {med['periodo_fim']} · {med['status']}"
    wb = Workbook()

    ws = wb.active
    ws.title = "RESUMO"
    _cabecalho(ws, "Medição — resumo", sub, ["Gaveta", "Descrição", "Lançamentos", "Pendentes", "Valor P0 (c/ BDI e K)"],
               [10, 52, 14, 12, 22], teste)
    for g, d in r["gavetas"].items():
        ws.append([g, d["nome"], len(d["lancamentos"]), d["pendentes"], d["valor"]])
        _fmt(ws, ws.max_row, {5: MOEDA})
    t = r["totalizacao"]
    for rot, val, f in (("Σ MATERIAL (ROUND por item)", t["material"], MOEDA), ("Σ SERVIÇO (ROUND por item)", t["servico"], MOEDA),
                        ("BDI material", t["bdi_material"], MOEDA), ("BDI serviço", t["bdi_servico"], MOEDA),
                        ("K", t["k"], MOEDA), ("Subtotal P0", r["valor_pi"], MOEDA),
                        ("(−) Excedente sobre teto de subgrupo", -r["excedente_teto"], MOEDA),
                        ("Valor P0 limitado ao teto", r["valor_pi_limitado"], MOEDA),
                        ("Índice de reajuste", r["indice_reajuste"] if r["indice_reajuste"] is not None else "PENDENTE", "0.00%"),
                        ("Reajuste", r["valor_reajuste"] if r["valor_reajuste"] is not None else "PENDENTE", MOEDA),
                        ("TOTAL", r["valor_total"] if r["valor_total"] is not None else "PENDENTE", MOEDA)):
        ws.append(["", rot, "", "", val])
        ws.cell(row=ws.max_row, column=2).font = Font(bold=True)
        _fmt(ws, ws.max_row, {5: f})
    for linha in filter(None, [("Valores CONGELADOS no fechamento da medição." if r.get("congelada") else None),
                               r.get("alerta_recalculo"), r["aviso_reajuste"],
                               f"{r['pendentes']} lançamento(s) PENDENTE(s): fora do valor até haver item/regra com fonte."
                               if r["pendentes"] else None]
                        + [f"SIMULAÇÃO — hipótese não validada: {h}" for h in r.get("hipoteses", [])]):
        ws.append(["", linha])
        ws.cell(row=ws.max_row, column=2).font = Font(bold=True, color="C92A2A")

    # preço usado nesta medição (congelado, se fechada)
    preco_med = {l["bm_item_id"]: l["preco_unitario"] for d in r["gavetas"].values() for l in d["lancamentos"]
                 if l["status"] == "PROPOSTO" and l["bm_item_id"] is not None}
    # quantidades PROPOSTAS por item (esta medição e anteriores do mesmo tipo)
    atual, anterior = defaultdict(float), defaultdict(float)
    for l in conn.execute(
            """SELECT l.bm_item_id, l.quantidade, m.numero FROM medicao_lancamentos l JOIN medicoes m ON m.id=l.medicao_id
               WHERE l.status='PROPOSTO' AND m.tipo=? AND m.numero <= ?""", (med["tipo"], med["numero"])):
        (atual if l["numero"] == med["numero"] else anterior)[l["bm_item_id"]] += l["quantidade"]

    ws = wb.create_sheet("BM VERTICAL")
    _cabecalho(ws, "BM — Vertical (conferência)", sub,
               ["Subgrupo", "Item", "Descrição", "Classif.", "Unid.", "Quant. contrato", "Preço unit. (s/ BDI)",
                "Quant. anterior", "Quant. nesta medição", "Quant. acumulada", "Saldo", "Total nesta medição (s/ BDI) = ROUND(q×p;2)"],
               [10, 13, 60, 11, 8, 15, 16, 15, 15, 15, 15, 22], teste)
    sg_atual = None
    for b in conn.execute("SELECT * FROM bm_itens WHERE planilha='VERTICAL' AND vigente=1 ORDER BY linha_origem"):
        if b["subgrupo"] != sg_atual:
            sg_atual = b["subgrupo"]
            ws.append([sg_atual, "", b["subgrupo_nome"] or ""])
            for c in ws[ws.max_row]:
                c.fill, c.font = SUB, Font(bold=True)
        qa, qm = anterior[b["id"]], atual[b["id"]]
        ws.append([b["subgrupo"], b["item"], b["descricao"], b["classificacao"], b["unidade"], b["quantidade"], b["preco_unitario"],
                   qa, qm, qa + qm, (b["quantidade"] or 0) - qa - qm,
                   float(R.produto_excel(qm, preco_med.get(b["id"], b["preco_unitario"])))])
        _fmt(ws, ws.max_row, {6: QTD, 7: MOEDA, 8: QTD, 9: QTD, 10: QTD, 11: QTD, 12: MOEDA})

    for nome, filtro in (("MC REDE", lambda l: l["alvo_tipo"] == "TRECHO"), ("MC NÓS E NF", lambda l: l["alvo_tipo"] != "TRECHO")):
        ws = wb.create_sheet(nome)
        _cabecalho(ws, f"Memória de cálculo — {nome[3:].lower()}", sub,
                   ["Gaveta", "Trecho / atividade / NF", "Serviço", "Unid.", "Quantidade", "Item BM", "Status", "Motivo / observação"],
                   [8, 36, 34, 8, 14, 13, 12, 60], teste)
        for g, d in r["gavetas"].items():
            for l in d["lancamentos"]:
                if not filtro(l):
                    continue
                ref = l["alvo_nome"] or (f"NF {l['nf_id']}" if l["nf_id"] else "")
                ws.append([g, ref, l["servico"], l["unidade"], l["quantidade"], l["item"] or "—", l["status"],
                           " · ".join(filter(None, [l["motivo"], l["observacao"]]))])
                _fmt(ws, ws.max_row, {5: QTD})
                if l["status"] == "PENDENTE":
                    ws.cell(row=ws.max_row, column=7).font = Font(bold=True, color="C92A2A")

    _abas_conferencia(wb, conn, medicao_id, sub, teste)

    buf = io.BytesIO()
    wb.save(buf)
    prefixo = "TESTE_" if teste else ""
    return f"{prefixo}MEDICAO_{med['numero']:02d}_MORRO_DA_CRUZ.xlsx", buf.getvalue()


def _abas_conferencia(wb, conn, medicao_id, sub, teste):
    """Conferência interna (retrato gravado; ou o retrato atual, se ainda não conferida), regras, elegibilidade e parcelas."""
    conf = S.conferencia_vigente(conn, medicao_id)
    ret = conf["retrato"] if conf else M.retrato(conn, medicao_id)
    ws = wb.create_sheet("CONFERÊNCIA")
    _cabecalho(ws, "Conferência interna da medição (Ankara — não é aprovação da Caesb)", sub,
               ["Trecho / atividade", "OK interno em", "OK interno por", "Fotos válidas", "Fotos faltando", "Verificação",
                "Soldas aprovadas / informadas", "Ocorrências abertas", "Cadastro (etapa · GDOC · ESET)"],
               [36, 14, 16, 12, 28, 16, 18, 12, 46], teste)
    if conf:
        cab = (f"Conferência #{conf['id']}: {conf['resultado']} por {conf['por']} em {conf['em']}"
               + ("" if conf["vale"] else " — NÃO VALE MAIS: os lançamentos mudaram depois dela"))
    else:
        cab = "SEM CONFERÊNCIA INTERNA registrada: abaixo, o retrato atual (o fechamento exige conferência CONFORME)."
    ws.append([cab])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True, color=LARANJA if conf and conf["vale"] else "C92A2A")
    if conf:
        for k, v in conf["checklist"].items():
            ws.append([f"  [{'x' if v else ' '}] {R.CHECKLIST_CONFERENCIA_MEDICAO.get(k, k)}"])
        if conf.get("obs"):
            ws.append([f"  Observação: {conf['obs']}"])
    for a in ret["alvos"]:
        cad = a.get("cadastro") or {}
        ws.append([a["nome"], a["ok_interno_em"], a["ok_interno_por"], a["fotos_validas"], ", ".join(a["fotos_faltando"]) or "—",
                   a.get("verificacao") or "—", f"{a['soldas_aprovadas']:g} / {a['soldas_informadas']:g}", a["ocorrencias_abertas"],
                   " · ".join(str(x) for x in (cad.get("etapa"), cad.get("protocolo_gdoc"), cad.get("memorando_eset")) if x) or "—"])
        _fmt(ws, ws.max_row, {})
    for t in ret["impeditivos"]:
        ws.append([f"IMPEDITIVO: {t}"])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True, color="C92A2A")
    for t in ret["alertas"]:
        ws.append([f"Alerta: {t}"])
        ws.cell(row=ws.max_row, column=1).font = Font(color="9A5B00")

    ws = wb.create_sheet("REGRAS")
    _cabecalho(ws, "Regras da medição — confirmadas, provisórias e bloqueadas", sub,
               ["Regra", "Situação", "Efeito no cálculo", "Fonte", "Decisão"], [70, 20, 70, 40, 14], teste)
    for r in M.regras(conn):
        ws.append([r["regra"], r["situacao"], r["efeito"], r["fonte"] or "—", r["decisao"] or "—"])
        _fmt(ws, ws.max_row, {})
        if r["situacao"] != "CONFIRMADA":
            ws.cell(row=ws.max_row, column=2).font = Font(bold=True, color="C92A2A")

    el = M.elegibilidade(conn, medicao_id)
    ws = wb.create_sheet("ELEGIBILIDADE")
    _cabecalho(ws, "Elegibilidade do período", f"{sub} · {el['regra']}",
               ["Situação", "Trecho / atividade", "Frente", "Extensão (m)", "Assentado no período (m)", "Verificado no período (m)",
                "Assentado total (m)", "OK interno", "Impedimentos", "Alertas"],
               [22, 36, 10, 12, 14, 14, 14, 12, 60, 50], teste)
    for l in el["linhas"]:
        ws.append([l["situacao"], l["nome"], l["frente_id"], l["extensao_m"], l["assentado_periodo_m"], l["verificado_periodo_m"],
                   l.get("assentado_total_m"), l["ok_interno_em"] or "—", "; ".join(l["impedimentos"]) or "—", "; ".join(l["alertas"]) or "—"])
        _fmt(ws, ws.max_row, {4: QTD, 5: QTD, 6: QTD, 7: QTD})
    ws.append([])
    ws.append(["Material no período (por propriedade) — BLOQUEADO: regime não aprovado (D-53 · D-54); recebido não é faturável"])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True, color="C92A2A")
    ws.append(["Material", "Descrição", "Propriedade", "Unid.", "Recebido/aceito", "Aplicado", "Situação"])
    for m in el["materiais"]:
        ws.append([m["material_id"], m["descricao"], m["propriedade"], m["unidade"], m["recebido"], m["aplicado"], m["situacao"]])
        _fmt(ws, ws.max_row, {5: QTD, 6: QTD})

    ws = wb.create_sheet("PARCELAS")
    _cabecalho(ws, "Parcelas medidas por item do BM (medições do mesmo tipo)", sub,
               ["Item", "Descrição", "Unid.", "Quant. contrato", "Anterior", "Nesta", "Acumulado", "Saldo", "Acima do contrato"],
               [13, 60, 8, 15, 14, 14, 14, 14, 12], teste)
    for i in M.acumulado_itens(conn, medicao_id):
        ws.append([i["item"], i["descricao"], i["unidade"], i["contrato"], i["anterior"], i["nesta"], i["acumulado"], i["saldo"],
                   "SIM" if i["excede_contrato"] else ""])
        _fmt(ws, ws.max_row, {4: QTD, 5: QTD, 6: QTD, 7: QTD, 8: QTD})
    for d in M.duplicidades(conn, medicao_id):
        ws.append([f"DUPLICIDADE: {d['servico']} de {d['alvo_nome']} nas medições {d['medicoes']}"])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True, color="C92A2A")
