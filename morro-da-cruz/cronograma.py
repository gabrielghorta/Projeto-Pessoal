"""Cronograma da obra vindo do MS Project (Planejamento › Cronograma da obra).

O .mpp do Project é um formato fechado: o sistema lê o que o próprio Project exporta —
  * XML do Project (Arquivo › Salvar como › "Formato XML (*.xml)"): o mais completo (EDT, datas, duração, % concluída,
    predecessoras, recursos, linha de base e anotações);
  * planilha do Excel (.xlsx) ou CSV (Arquivo › Exportar), com uma linha de títulos (em português ou em inglês).
O arquivo enviado fica guardado como veio em <pasta do banco>/documentos/cronograma/ (o backup copia a pasta inteira) e a
leitura vira a versão atual (cronograma.json); as versões anteriores continuam guardadas (historico.json). Nada no banco
muda: o cronograma é a referência para programar as equipes. O planejado do dia continua sendo a programação (Dados da
obra) e o feito, os registros de campo.
Cada tarefa é ligada à frente (Zona 01…11, Adutora, CH) e aos trechos citados no nome ("TR014", "TR001 a TR010"), com a
zona e a PB herdadas das tarefas-resumo de cima quando o nome não traz; para as tarefas com trechos ligados, a consulta
mostra o assentado registrado no sistema ao lado do % concluído do Project.
"""
import base64
import csv
import hashlib
import io
import json
import os
import re
import unicodedata
import xml.etree.ElementTree as ET
import zipfile
from datetime import date, datetime, timedelta

import compras as C
import regras as R
import relogio as RL
import servicos as S
from servicos import RegraViolada, _agora

LIMITE = 25 * 1024 * 1024
LIMITE_DESCOMPACTADO = 250 * 1024 * 1024
LIMITE_TAREFAS = 20000
NS = "{http://schemas.microsoft.com/project}"
PASTA = "cronograma"
ATUAL, HISTORICO = "cronograma.json", "historico.json"
EXT = {"XML": "xml", "XLSX": "xlsx", "CSV": "csv"}
FORMATO = {"XML": "XML do Project", "XLSX": "planilha do Excel exportada do Project", "CSV": "CSV exportado do Project"}
COMO_EXPORTAR = ("No Project: Arquivo › Salvar como › tipo \"Formato XML (*.xml)\" (ou Arquivo › Exportar para Excel) e "
                 "envie esse arquivo.")
# vínculo das predecessoras (sigla do Project em português): término-início, início-início, término-término, início-término
TIPO_XML = {"0": "TT", "1": "TI", "2": "IT", "3": "II"}
TIPO_EN = {"FS": "TI", "SS": "II", "FF": "TT", "SF": "IT"}
COLUNAS = {
    "id": "id", "identificador": "id", "identificacao": "id", "unique id": "uid", "id exclusivo": "uid",
    "identificacao exclusiva": "uid", "nome": "nome", "nome da tarefa": "nome", "name": "nome", "task name": "nome",
    "tarefa": "nome", "edt": "edt", "wbs": "edt", "numero da estrutura de topicos": "estrutura", "outline number": "estrutura",
    "nivel da estrutura de topicos": "nivel", "outline level": "nivel", "nivel": "nivel", "duracao": "duracao",
    "duration": "duracao", "inicio": "inicio", "start": "inicio", "termino": "fim", "finish": "fim", "fim": "fim",
    "% concluida": "concluido", "% concluido": "concluido", "% complete": "concluido", "porcentagem concluida": "concluido",
    "% completo": "concluido", "predecessoras": "predecessoras", "predecessors": "predecessoras",
    "nomes dos recursos": "recursos", "resource names": "recursos", "recursos": "recursos", "resumo": "resumo",
    "summary": "resumo", "marco": "marco", "milestone": "marco", "critica": "critica", "critical": "critica",
    "anotacoes": "notas", "notes": "notas", "notas": "notas", "inicio da linha de base": "base_inicio",
    "baseline start": "base_inicio", "termino da linha de base": "base_fim", "baseline finish": "base_fim",
}
SIM = {"sim", "yes", "s", "y", "1", "true", "verdadeiro"}


def _sem_acento(s):
    return "".join(c for c in unicodedata.normalize("NFD", str(s or "")) if unicodedata.category(c) != "Mn").lower().strip()


# ------------------------------------------------------------------ valores
def _data(v):
    """Data do Project (ISO do XML, célula de data do Excel ou texto 'seg 21/10/26 08:00') → AAAA-MM-DD, ou None."""
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    s = str(v).strip()
    m = re.search(r"(\d{4})-(\d{2})-(\d{2})", s)
    if m:
        a, mo, d = int(m[1]), int(m[2]), int(m[3])
    else:
        m = re.search(r"(\d{1,2})/(\d{1,2})/(\d{2,4})", s)
        if not m:
            return None
        d, mo, a = int(m[1]), int(m[2]), int(m[3])
        a += 2000 if a < 100 else 0
    try:
        return date(a, mo, d).isoformat()
    except ValueError:
        return None


def _duracao_iso(v, min_dia):
    """Duração do XML (PT120H0M0S = horas de trabalho) → dias de trabalho."""
    m = re.fullmatch(r"P(?:(\d+(?:\.\d+)?)D)?(?:T(?:(\d+(?:\.\d+)?)H)?(?:(\d+(?:\.\d+)?)M)?(?:(\d+(?:\.\d+)?)S)?)?", str(v or "").strip())
    if not m or not any(m.groups()):
        return None
    d, h, mi, s = (float(x or 0) for x in m.groups())
    return round(d + (h * 60 + mi + s / 60) / min_dia, 2)


def _duracao_txt(v):
    """'5 dias', '3,5d', '16 hrs', '2 sems', '1 ms', '5 days?', '5 edias' (decorridos) → dias de trabalho."""
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = _sem_acento(v).replace(",", ".").replace("?", "")
    m = re.match(r"\s*([+-]?[\d.]+)\s*([a-z]*)", s)
    if not m:
        return None
    try:
        n = float(m[1])
    except ValueError:
        return None
    u = m[2][1:] if m[2].startswith("e") and len(m[2]) > 1 else m[2]
    if u.startswith("mi"):
        return round(n / 480, 2)
    if u.startswith("h"):
        return round(n / 8, 2)
    if u.startswith(("s", "w")):
        return round(n * 5, 2)
    if u.startswith(("ms", "me", "mo")):
        return round(n * 20, 2)
    return round(n, 2)


def _pct(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        x = float(v)
        return round(x * 100 if 0 < x < 1 else x, 1)
    m = re.search(r"[\d.,]+", str(v))
    if not m:
        return None
    try:
        return round(float(m[0].replace(",", ".")), 1)
    except ValueError:
        return None


def _sim(v):
    return _sem_acento(v) in SIM if not isinstance(v, bool) else v


def _predecessoras_txt(v):
    """'3;5TI+2 dias;7II' ou '3,5FS+2 days' → [{"id", "tipo", "atraso_dias"}]."""
    out = []
    for parte in re.split(r"[;,](?![\d]+\s*(?:d|h|dia|day))", str(v or "")):
        m = re.match(r"\s*(\d+)\s*(TI|II|TT|IT|FS|SS|FF|SF)?\s*([+-]\s*[\d.,]+\s*\w*)?", parte.strip(), re.I)
        if m:
            tipo = (m[2] or "TI").upper()
            out.append({"id": int(m[1]), "tipo": TIPO_EN.get(tipo, tipo),
                        "atraso_dias": _duracao_txt(m[3].replace(" ", "")) if m[3] else 0})
    return out


# ------------------------------------------------------------------ leitores
def _ler_xml(dados):
    if b"<!DOCTYPE" in dados.upper() or b"<!ENTITY" in dados.upper():
        raise RegraViolada("XML com DOCTYPE/ENTITY não é aceito (o XML do Project não usa).")
    try:
        raiz = ET.fromstring(dados)
    except ET.ParseError as e:
        raise RegraViolada(f"XML inválido: {e}") from None
    if raiz.tag != NS + "Project":
        raise RegraViolada("O XML não é de um projeto do MS Project (falta o elemento Project). " + COMO_EXPORTAR)
    t = lambda el, tag: (el.findtext(NS + tag) or "").strip()
    min_dia = float(t(raiz, "MinutesPerDay") or 480) or 480
    recursos = {t(r, "UID"): t(r, "Name") for r in raiz.iter(NS + "Resource") if t(r, "Name")}
    por_tarefa = {}
    for a in raiz.iter(NS + "Assignment"):
        nome = recursos.get(t(a, "ResourceUID"))
        if nome:
            por_tarefa.setdefault(t(a, "TaskUID"), []).append(nome)
    tarefas, id_de = [], {}
    for el in raiz.iter(NS + "Task"):
        uid, nivel = t(el, "UID"), int(t(el, "OutlineLevel") or 0)
        if t(el, "IsNull") == "1" or nivel == 0 or not t(el, "Name"):
            continue                                   # a tarefa 0 é o resumo do projeto inteiro
        base = next((b for b in el.findall(NS + "Baseline") if (t(b, "Number") or "0") == "0"), None)
        id_de[uid] = int(t(el, "ID") or len(tarefas) + 1)
        tarefas.append({
            "id": id_de[uid], "nome": t(el, "Name")[:300], "edt": t(el, "WBS") or t(el, "OutlineNumber"), "nivel": nivel,
            "inicio": _data(t(el, "Start")), "fim": _data(t(el, "Finish")), "duracao_dias": _duracao_iso(t(el, "Duration"), min_dia),
            "concluido": _pct(t(el, "PercentComplete")), "resumo": t(el, "Summary") == "1", "marco": t(el, "Milestone") == "1",
            "critica": t(el, "Critical") == "1", "inativa": t(el, "Active") == "0",
            "recursos": por_tarefa.get(uid, []), "notas": t(el, "Notes")[:500],
            "linha_base": {"inicio": _data(t(base, "Start")), "fim": _data(t(base, "Finish"))} if base is not None else None,
            "_pred": [(t(p, "PredecessorUID"), TIPO_XML.get(t(p, "Type"), "TI"), float(t(p, "LinkLag") or 0) / 10 / min_dia)
                      for p in el.findall(NS + "PredecessorLink")]})
        if len(tarefas) > LIMITE_TAREFAS:
            raise RegraViolada(f"Cronograma com mais de {LIMITE_TAREFAS} tarefas.")
    for x in tarefas:
        x["predecessoras"] = [{"id": id_de[u], "tipo": tp, "atraso_dias": round(lag, 2)} for u, tp, lag in x.pop("_pred") if u in id_de]
    projeto = {"nome": t(raiz, "Title") or t(raiz, "Name"), "inicio": _data(t(raiz, "StartDate")),
               "fim": _data(t(raiz, "FinishDate")), "data_status": _data(t(raiz, "StatusDate"))}
    return projeto, tarefas


def _ler_tabela(linhas):
    """Linhas (listas de valores) de uma planilha/CSV exportada: acha a linha de títulos e monta as tarefas."""
    cab, tarefas = None, []
    for k, row in enumerate(linhas):
        if cab is None:
            campos = {i: COLUNAS.get(_sem_acento(v)) for i, v in enumerate(row) if v not in (None, "")}
            if "nome" in campos.values():
                cab = {i: c for i, c in campos.items() if c}
            elif k > 20:
                break
            continue
        reg = {c: row[i] for i, c in cab.items() if i < len(row)}
        bruto = str(reg.get("nome") or "")
        nome = bruto.strip()
        if not nome:
            continue
        estrutura = str(reg.get("estrutura") or reg.get("edt") or "").strip()
        try:
            nivel = int(float(reg.get("nivel")))
        except (TypeError, ValueError):
            nivel = len(estrutura.split(".")) if re.fullmatch(r"\d+(\.\d+)*", estrutura) else 1 + (len(bruto) - len(bruto.lstrip())) // 3
        try:
            tid = int(float(reg.get("id")))
        except (TypeError, ValueError):
            tid = len(tarefas) + 1
        tarefas.append({
            "id": tid, "nome": nome[:300], "edt": str(reg.get("edt") or estrutura or "").strip(), "nivel": max(nivel, 1),
            "inicio": _data(reg.get("inicio")), "fim": _data(reg.get("fim")), "duracao_dias": _duracao_txt(reg.get("duracao")),
            "concluido": _pct(reg.get("concluido")), "resumo": _sim(reg.get("resumo")) if "resumo" in reg else None,
            "marco": _sim(reg.get("marco")) if "marco" in reg else False, "critica": _sim(reg.get("critica")) if "critica" in reg else False,
            "inativa": False, "recursos": [r.strip() for r in re.split(r"[;,]", str(reg.get("recursos") or "")) if r.strip()],
            "notas": str(reg.get("notas") or "")[:500], "predecessoras": _predecessoras_txt(reg.get("predecessoras")),
            "linha_base": ({"inicio": _data(reg.get("base_inicio")), "fim": _data(reg.get("base_fim"))}
                           if reg.get("base_inicio") or reg.get("base_fim") else None)})
        if len(tarefas) > LIMITE_TAREFAS:
            raise RegraViolada(f"Cronograma com mais de {LIMITE_TAREFAS} tarefas.")
    if cab is None:
        raise RegraViolada("Não achei a linha de títulos: a planilha precisa de uma coluna \"Nome\" (ou \"Nome da tarefa\"). "
                           + COMO_EXPORTAR)
    for i, x in enumerate(tarefas):                     # sem a coluna Resumo: é resumo quem tem tarefa de nível maior logo abaixo
        if x["resumo"] is None:
            x["resumo"] = i + 1 < len(tarefas) and tarefas[i + 1]["nivel"] > x["nivel"]
    return tarefas


def _ler_xlsx(dados):
    try:
        with zipfile.ZipFile(io.BytesIO(dados)) as z:
            if sum(i.file_size for i in z.infolist()) > LIMITE_DESCOMPACTADO:
                raise RegraViolada("Planilha grande demais depois de descompactada.")
    except zipfile.BadZipFile:
        raise RegraViolada("Arquivo .xlsx inválido.") from None
    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(dados), read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001 — openpyxl levanta vários tipos para arquivo corrompido
        raise RegraViolada(f"Não consegui abrir a planilha: {e}") from None
    try:
        erro = None
        for ws in wb.worksheets:
            try:
                return _ler_tabela(ws.iter_rows(values_only=True, max_row=LIMITE_TAREFAS + 30))
            except RegraViolada as e:
                erro = e
        raise erro or RegraViolada("Planilha sem abas.")
    finally:
        wb.close()


def _ler_csv(dados):
    for cod in ("utf-8-sig", "cp1252"):
        try:
            texto = dados.decode(cod)
            break
        except UnicodeDecodeError:
            continue
    primeira = texto.split("\n", 1)[0]
    sep = max((";", ",", "\t"), key=primeira.count)
    return _ler_tabela(csv.reader(io.StringIO(texto), delimiter=sep))


def ler(dados, nome=""):
    """(formato, projeto, tarefas) do arquivo exportado pelo Project; recusa o .mpp e o que não for cronograma."""
    if not dados:
        raise RegraViolada("Arquivo vazio.")
    if len(dados) > LIMITE:
        raise RegraViolada("Arquivo maior que 25 MB.")
    if dados[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        raise RegraViolada("Este é o arquivo .mpp do Project, que o sistema não consegue abrir. " + COMO_EXPORTAR)
    if dados[:2] == b"PK":
        return "XLSX", {}, _ler_xlsx(dados)
    if dados.lstrip()[:1] == b"<" or dados[:3] == b"\xef\xbb\xbf" and dados[3:].lstrip()[:1] == b"<":
        return ("XML", *_ler_xml(dados))
    if str(nome).lower().endswith((".csv", ".txt")) or b";" in dados[:2000] or b"," in dados[:2000]:
        return "CSV", {}, _ler_csv(dados)
    raise RegraViolada("Formato não reconhecido. " + COMO_EXPORTAR)


# ------------------------------------------------------------------ ligação com a obra
def _frente(nome):
    s = _sem_acento(nome)
    m = re.search(r"\bzona\s*-?\s*0?(\d{1,2})\b", s)
    if m and 1 <= int(m[1]) <= 11:
        return f"Zona{int(m[1]):02d}"
    if re.search(r"\badutora\b", s):
        return "ADUTORA"
    if re.search(r"\bch\b|caixas? hidraulica", s):
        return "CH"
    return None


def _pb(nome):
    m = re.search(r"\bPB\s*-?\s*(\d{1,2})\b", str(nome), re.I)
    return f"PB{int(m[1]):02d}" if m else None


def _faixas_tr(nome):
    out = []
    for m in re.finditer(r"\bTR\s*-?\s*(\d{1,4})(?:\s*(?:a|ate|até|ao|-|–)\s*(?:TR\s*-?\s*)?(\d{1,4}))?", str(nome), re.I):
        a, b = int(m[1]), int(m[2] or m[1])
        if a <= b <= a + 300:
            out.append((a, b))
    return out


def ligar(conn, tarefas):
    """Frente, PB e trechos de cada tarefa (o que o nome cita, herdando zona e PB das tarefas-resumo de cima)."""
    pontos = [r[0] for r in conn.execute("SELECT ponto FROM trechos WHERE vigente=1")]
    completos = set(pontos)
    por_num = {}
    for p in pontos:
        partes = p.split("_")
        n = re.sub(r"\D", "", partes[-1])
        if len(partes) >= 3 and n:
            por_num.setdefault((partes[0].upper(), partes[1].upper(), int(n)), []).append(p)
    pilha = []
    for x in tarefas:
        while pilha and pilha[-1][0] >= x["nivel"]:
            pilha.pop()
        frente = _frente(x["nome"]) or (pilha[-1][1] if pilha else None)
        pb = _pb(x["nome"]) or (pilha[-1][2] if pilha else None)
        alvos = [m[0] for m in re.finditer(r"(?:Zona\d\d|ADUTORA)_PB\w+?_TR\d+", x["nome"], re.I) if m[0] in completos]
        if frente and pb:
            for a, b in _faixas_tr(x["nome"]):
                for n in range(a, b + 1):
                    alvos += por_num.get((frente.upper(), pb.upper(), n), [])
        x["frente"], x["pb"], x["alvos"] = frente, pb, sorted(set(alvos))[:500]
        pilha.append((x["nivel"], frente, pb))
    return tarefas


def _executado(conn, tarefas):
    """% assentado registrado no sistema nos trechos ligados a cada tarefa (registros efetivos, como na ficha)."""
    pontos = sorted({p for x in tarefas for p in x.get("alvos") or []})
    if not pontos:
        return
    ext, ass = {}, {}
    for i in range(0, len(pontos), 500):
        lote = pontos[i:i + 500]
        ph = ",".join("?" * len(lote))
        ids = {r[0]: (r[1], float(r[2])) for r in conn.execute(f"SELECT id, ponto, extensao_m FROM trechos WHERE vigente=1 AND ponto IN ({ph})", lote)}
        ext.update({p: e for p, e in ids.values()})
        cob = {}
        for tid, a, b in conn.execute(f"""SELECT trecho_id, ini_m, fim_m FROM intervalos_execucao WHERE {S.EFETIVO} AND etapa='ASSENTADO'
                                          AND trecho_id IN ({','.join('?' * len(ids))})""", list(ids)):
            cob.setdefault(tid, []).append((a, b))
        ass.update({ids[tid][0]: R.comprimento(R.unir_intervalos(lst)) for tid, lst in cob.items()})
    for x in tarefas:
        if x.get("alvos"):
            total = sum(ext.get(p, 0) for p in x["alvos"])
            x["executado_m"] = round(sum(min(ass.get(p, 0), ext.get(p, 0)) for p in x["alvos"]), 2)
            x["executado_pct"] = round(100 * x["executado_m"] / total, 1) if total else None


# ------------------------------------------------------------------ guardar e consultar
def _pasta(conn):
    return os.path.join(C.documento_dir(conn), PASTA)


def _ler_json(caminho, padrao=None):
    try:
        with open(caminho, encoding="utf-8") as f:
            return json.load(f)
    except (OSError, ValueError):
        return padrao


def _gravar_json(caminho, dados):
    tmp = caminho + ".tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(dados, f, ensure_ascii=False)
    os.replace(tmp, caminho)


def importar(conn, conteudo_b64, nome, por):
    """Guarda o arquivo como veio e a leitura dele como versão atual do cronograma. Devolve o resumo da versão."""
    try:
        dados = base64.b64decode(conteudo_b64 or "", validate=True)
    except (ValueError, TypeError):
        raise RegraViolada("Conteúdo do arquivo inválido.") from None
    formato, projeto, tarefas = ler(dados, nome)
    tarefas = [x for x in tarefas if x["nome"]]
    if not tarefas:
        raise RegraViolada("O arquivo não tem nenhuma tarefa. " + COMO_EXPORTAR)
    ligar(conn, tarefas)
    datas = [d for x in tarefas for d in (x["inicio"], x["fim"]) if d]
    projeto = {**{"nome": None, "inicio": None, "fim": None, "data_status": None}, **projeto}
    projeto["inicio"] = projeto["inicio"] or (min(datas) if datas else None)
    projeto["fim"] = projeto["fim"] or (max(datas) if datas else None)
    pasta = _pasta(conn)
    os.makedirs(pasta, exist_ok=True)
    historico = _ler_json(os.path.join(pasta, HISTORICO), [])
    h = hashlib.sha256(dados).hexdigest()
    if historico and historico[-1]["sha256"] == h:
        raise RegraViolada("Este arquivo é igual à versão atual do cronograma: nada mudou.")
    arquivo = f"{_agora().replace(':', '-')}_{h[:8]}.{EXT[formato]}"
    with open(os.path.join(pasta, arquivo), "wb") as f:
        f.write(dados)
    sem_data = sum(1 for x in tarefas if not x["inicio"] or not x["fim"])
    meta = {"versao": len(historico) + 1, "arquivo": arquivo, "nome_original": os.path.basename(str(nome or ""))[:200] or arquivo,
            "formato": FORMATO[formato], "sha256": h, "bytes": len(dados), "enviado_por": por, "enviado_em": _agora(),
            "n_tarefas": len(tarefas), "ligadas_a_trechos": sum(1 for x in tarefas if x["alvos"]),
            "com_frente": sum(1 for x in tarefas if x["frente"]), "sem_data": sem_data, "projeto": projeto}
    _gravar_json(os.path.join(pasta, ATUAL), {**meta, "tarefas": tarefas})
    _gravar_json(os.path.join(pasta, HISTORICO), historico + [meta])
    return meta


def consultar(conn):
    """Versão atual (com o executado do sistema nas tarefas ligadas a trechos) e o histórico das versões enviadas."""
    pasta = _pasta(conn)
    atual = _ler_json(os.path.join(pasta, ATUAL))
    versoes = list(reversed(_ler_json(os.path.join(pasta, HISTORICO), [])))
    if not atual:
        return {"vazio": True, "versoes": versoes, "como_exportar": COMO_EXPORTAR}
    _executado(conn, atual["tarefas"])
    hoje = RL.hoje().isoformat()
    atual["hoje"] = hoje
    atual["em_andamento"] = sum(1 for x in atual["tarefas"] if not x["resumo"] and x["inicio"] and x["fim"]
                                and x["inicio"] <= hoje <= x["fim"])
    atual["atrasadas"] = sum(1 for x in atual["tarefas"] if not x["resumo"] and x["fim"] and x["fim"] < hoje
                             and (x["concluido"] or 0) < 100)
    return {**atual, "versoes": versoes, "como_exportar": COMO_EXPORTAR,
            "proximos_14_dias": (RL.hoje() + timedelta(days=14)).isoformat()}
