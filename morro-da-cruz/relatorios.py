"""Relatórios de operação (módulo F), todos calculados a partir dos registros efetivos (sem estornos, sem intervalos anulados).

* Ciclo do material: previsto (projeto) × solicitado × comprado × a receber × recebido × reservado × com as equipes ×
  aplicado × perda × medido. Material não entra em medição enquanto o regime estiver bloqueado (D-53, D-54): a coluna
  "medido" mostra isso, sem zero inventado.
* Avanço por frente, equipe, período e etapa, a partir dos intervalos efetivos. O avanço legado sem posição aparece à
  parte (não tem data por intervalo nem posição).
* Pendências por frente: ocorrências, verificação, reaterro sem cadastro, legado, aguardando OK interno, cadastro,
  programação vencida; e as gerais (reservas vencidas, notas a conferir, quarentena).
"""
import io
from collections import defaultdict
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Font

import compras as C
import planejamento as P
import regras as R
import servicos as S
from servicos import _rows

ETAPAS = ("ASSENTADO", "CADASTRO_CAMPO", "REATERRADO")


# ------------------------------------------------------------------ ciclo do material
def materiais(conn):
    def soma(sql, args=()):
        return {r[0]: r[1] for r in conn.execute(sql, args)}
    solicitado = soma("""SELECT i.material_id, SUM(i.quantidade) FROM solicitacao_itens i JOIN solicitacoes_compra s
                         ON s.id=i.solicitacao_id WHERE s.status='APROVADA' GROUP BY 1""")
    comprado = soma("""SELECT i.material_id, SUM(i.quantidade) FROM pedido_itens i JOIN pedidos_compra p ON p.id=i.pedido_id
                       WHERE p.status<>'CANCELADO' GROUP BY 1""")
    mov = defaultdict(lambda: defaultdict(float))
    for r in conn.execute(f"""SELECT material_id, tipo, COALESCE(status,'ACEITO') st, SUM(quantidade) q FROM movimentos
                              WHERE {S.VALIDO} GROUP BY 1, 2, 3"""):
        mov[r[0]][(r[1], r[2])] += r[3]
    reservado = soma("SELECT material_id, SUM(quantidade) FROM reservas WHERE status='ATIVA' GROUP BY 1")
    out = []
    for m in _rows(conn.execute("SELECT * FROM materiais ORDER BY propriedade, id")):
        mv = mov[m["id"]]
        t = lambda tipo, st="ACEITO": mv.get((tipo, st), 0.0)
        recebido = t("ENTRADA_NF") + t("ROMANEIO_CAESB") + t("ENTRADA_CONSUMO") - t("DEVOLUCAO_FORNECEDOR")
        com_equipes = (t("CUSTODIA") + t("TRANSF_ENTRADA") - t("APLICADO") - t("SOBRA") - t("PERDA") - t("TRANSF_SAIDA"))
        saldo, quar = S.saldo_almox(conn, m["id"])
        linha = {"material_id": m["id"], "descricao": m["descricao"], "unidade": m["unidade"], "propriedade": m["propriedade"],
                 "previsto": m["qtd_projeto"], "solicitado": solicitado.get(m["id"], 0.0), "comprado": comprado.get(m["id"], 0.0),
                 "a_receber": P.a_receber(conn, m["id"]) if m["id"] in comprado else 0.0, "quarentena": quar,
                 "recebido": recebido, "saldo_almox": saldo, "reservado": reservado.get(m["id"], 0.0),
                 "com_equipes": com_equipes, "aplicado": t("APLICADO"), "perda": t("PERDA"),
                 "emprestado": t("EMPRESTIMO_OUTRA_OBRA"), "medido": None,
                 "medido_situacao": "material não entra na medição: regime bloqueado (D-53 · D-54)"}
        for k in ("solicitado", "comprado", "a_receber", "quarentena", "recebido", "saldo_almox", "reservado", "com_equipes",
                  "aplicado", "perda", "emprestado"):
            linha[k] = round(linha[k], 4)
        linha["movimentado"] = any(linha[k] for k in ("solicitado", "comprado", "recebido", "reservado", "aplicado", "quarentena",
                                                       "com_equipes", "perda", "emprestado"))
        if linha["movimentado"] or m["qtd_projeto"]:
            out.append(linha)
    return out


# ------------------------------------------------------------------ avanço
def avanco(conn, ini=None, fim=None, frente=None, equipe=None):
    """Metros por frente × equipe × etapa no período (intervalos efetivos) e a situação acumulada de cada frente."""
    ini, fim = ini or "0000-01-01", fim or "9999-12-31"
    onde, args = f"i.{S.EFETIVO} AND i.data BETWEEN ? AND ?", [ini, fim]
    if frente:
        onde, args = onde + " AND t.frente_id=?", args + [frente]
    if equipe:
        onde, args = onde + " AND i.equipe_id=?", args + [equipe]
    base = f"FROM intervalos_execucao i JOIN trechos t ON t.id=i.trecho_id WHERE {onde}"
    linhas = defaultdict(lambda: {e: 0.0 for e in ETAPAS})
    trechos = defaultdict(set)
    for r in conn.execute(f"SELECT t.frente_id, COALESCE(i.equipe_id, '—'), i.etapa, i.trecho_id, i.fim_m - i.ini_m {base}", args):
        linhas[(r[0], r[1])][r[2]] += r[4]
        trechos[(r[0], r[1])].add(r[3])
    por_equipe = [{"frente_id": f, "equipe_id": e, **{k: round(v, 3) for k, v in d.items()}, "trechos": len(trechos[(f, e)])}
                  for (f, e), d in sorted(linhas.items())]
    por_dia = {r[0]: round(r[1], 3) for r in conn.execute(
        f"SELECT i.data, SUM(i.fim_m - i.ini_m) {base} AND i.etapa='ASSENTADO' GROUP BY 1 ORDER BY 1", args)}
    # acumulado por frente (todas as datas), sobre os trechos vigentes
    acum = defaultdict(lambda: {e: 0.0 for e in ETAPAS})
    for r in conn.execute(f"""SELECT t.frente_id, i.etapa, SUM(i.fim_m - i.ini_m) FROM intervalos_execucao i
                              JOIN trechos t ON t.id=i.trecho_id WHERE i.{S.EFETIVO} AND t.vigente=1 GROUP BY 1, 2"""):
        acum[r[0]][r[1]] += r[2]
    ext = {r[0]: (r[1], r[2]) for r in conn.execute("SELECT frente_id, COUNT(*), SUM(extensao_m) FROM trechos WHERE vigente=1 GROUP BY 1")}
    legado = defaultdict(float)
    for (tid,) in conn.execute(f"SELECT DISTINCT ap.alvo_id {S._SQL_LEGADO}").fetchall():
        t = S._one(conn, "SELECT * FROM trechos WHERE id=?", tid)
        lg = S.legado_pendente(conn, t) if t else None
        if lg:
            legado[t["frente_id"]] += lg["ASSENTADO"]
    nos = defaultdict(lambda: [0, 0])
    for r in conn.execute("""SELECT a.frente_id, COUNT(*), SUM(COALESCE(x.concluido,0)) FROM atividades a
                             JOIN elementos e ON e.id=a.elemento_id LEFT JOIN execucao x ON x.alvo_tipo='ATIVIDADE' AND x.alvo_id=a.id
                             WHERE a.vigente=1 AND e.vigente=1 GROUP BY 1"""):
        nos[r[0]] = [r[1], r[2] or 0]
    frentes = []
    for f in R.FRENTES:
        if frente and f != frente:
            continue
        n, e = ext.get(f, (0, 0.0))
        a = acum[f]
        frentes.append({"frente_id": f, "trechos": n, "extensao_m": round(e or 0, 3),
                        **{k.lower() + "_m": round(v, 3) for k, v in a.items()},
                        "assentado_pct": round(100 * a["ASSENTADO"] / e, 2) if e else None,
                        "legado_sem_posicao_m": round(legado[f], 3), "nos": nos[f][0], "nos_concluidos": nos[f][1]})
    tot = {e: round(sum(d[e] for d in linhas.values()), 3) for e in ETAPAS}
    return {"periodo": {"ini": ini, "fim": fim}, "por_equipe": por_equipe, "total_periodo": tot,
            "por_dia": por_dia, "frentes": frentes}


# ------------------------------------------------------------------ pendências
def pendencias(conn, frente=None, hoje=None):
    hoje = hoje or date.today().isoformat()
    fr_trecho = {r[0]: r[1] for r in conn.execute("SELECT id, frente_id FROM trechos")}
    fr_ativ = {r[0]: r[1] for r in conn.execute("SELECT id, frente_id FROM atividades")}
    fr = lambda tipo, aid: fr_trecho.get(aid) if tipo == "TRECHO" else fr_ativ.get(aid)
    cont = defaultdict(lambda: defaultdict(int))
    for r in conn.execute("SELECT alvo_tipo, alvo_id, tipo FROM ocorrencias WHERE resolvida_em IS NULL"):
        cont[fr(r[0], r[1])]["impedimentos" if r[2] == "IMPEDIMENTO" else "nao_conformidades"] += 1
    for r in conn.execute(f"SELECT trecho_id, verificado_em, sem_cadastro_campo FROM intervalos_execucao WHERE {S.EFETIVO}"):
        if not r[1]:
            cont[fr_trecho.get(r[0])]["intervalos_sem_verificacao"] += 1
        if r[2]:
            cont[fr_trecho.get(r[0])]["reaterro_sem_cadastro"] += 1
    for (tid,) in conn.execute(f"SELECT DISTINCT ap.alvo_id {S._SQL_LEGADO}"):
        cont[fr_trecho.get(tid)]["legado_sem_posicao"] += 1
    for r in conn.execute("SELECT alvo_tipo, alvo_id FROM execucao WHERE concluido=1 AND ok_fiscal_em IS NULL"):
        f = fr(r[0], r[1])
        cont[f]["aguardando_ok_interno"] += 1
        if S.fotos_faltando(conn, r[0], r[1]):
            cont[f]["concluidos_sem_fotos"] += 1
    for r in conn.execute("SELECT trecho_id, etapa FROM cadastro"):
        if r[1] in ("1_PRONTO", "5A_NEGADO", "3_RECEBIDO"):
            cont[fr_trecho.get(r[0])]["cadastro_" + r[1].split("_", 1)[1].lower()] += 1
    for r in conn.execute("SELECT alvo_tipo, alvo_id FROM programacoes WHERE status='PLANEJADA' AND data < ?", (hoje,)):
        cont[fr(r[0], r[1])]["programacao_vencida"] += 1
    frentes = [{"frente_id": f, **cont.get(f, {}), "total": sum(cont.get(f, {}).values())}
               for f in R.FRENTES if not frente or f == frente]
    gerais = {"reservas_vencidas": len(C.reservas_vencidas(conn, hoje)),
              "notas_a_conferir": conn.execute("SELECT COUNT(*) FROM notas_fiscais WHERE status='A_CONFERIR'").fetchone()[0],
              "itens_em_quarentena": conn.execute(f"SELECT COUNT(*) FROM movimentos WHERE tipo='ENTRADA_NF' AND status='QUARENTENA' "
                                                  f"AND {S.VALIDO}").fetchone()[0],
              "sem_frente": dict(cont.get(None, {}))}
    return {"frentes": frentes, "gerais": gerais, "hoje": hoje}


def painel(conn, ini, fim, frente=None):
    return {"materiais": materiais(conn), "avanco": avanco(conn, ini, fim, frente), "pendencias": pendencias(conn, frente),
            "faltantes": P.faltantes(conn, ini, fim, frente)}


# ------------------------------------------------------------------ exportação
ROTULOS_PEND = {"impedimentos": "Impedimentos abertos", "nao_conformidades": "Não conformidades abertas",
                "intervalos_sem_verificacao": "Intervalos sem verificação", "reaterro_sem_cadastro": "Reaterro sem cadastro de campo",
                "legado_sem_posicao": "Trechos com legado sem posição", "aguardando_ok_interno": "Concluídos aguardando OK interno",
                "concluidos_sem_fotos": "Concluídos sem as fotos obrigatórias", "cadastro_pronto": "Cadastro pronto, não enviado",
                "cadastro_recebido": "Cadastro recebido, a conferir/protocolar", "cadastro_negado": "Cadastro negado pela Caesb",
                "programacao_vencida": "Programação vencida sem encerramento"}


def planilha(conn, ini, fim, frente=None):
    d = painel(conn, ini, fim, frente)
    wb = Workbook()
    neg = Font(bold=True)

    def aba(ws, titulo, cab, linhas):
        ws.append([titulo])
        ws["A1"].font = Font(bold=True, size=12)
        ws.append([f"{S.contrato(conn)} · período {ini} a {fim}" + (f" · frente {frente}" if frente else " · obra inteira")])
        ws.append([])
        ws.append(cab)
        for c in ws[4]:
            c.font = neg
        for l in linhas:
            ws.append(l)
        ws.freeze_panes = "A5"

    ws = wb.active
    ws.title = "MATERIAIS"
    aba(ws, "Ciclo do material — previsto × comprado × recebido × reservado × aplicado × medido",
        ["Material", "Descrição", "Unid.", "Propriedade", "Previsto (projeto)", "Solicitado (aprovado)", "Comprado", "A receber",
         "Quarentena", "Recebido (aceito − devolvido)", "Saldo almox.", "Reservado", "Com as equipes", "Aplicado", "Perda", "Medido"],
        [[m["material_id"], m["descricao"], m["unidade"], m["propriedade"], m["previsto"], m["solicitado"], m["comprado"],
          m["a_receber"], m["quarentena"], m["recebido"], m["saldo_almox"], m["reservado"], m["com_equipes"], m["aplicado"],
          m["perda"], m["medido_situacao"]] for m in d["materiais"]])
    ws = wb.create_sheet("AVANÇO")
    av = d["avanco"]
    aba(ws, "Avanço no período por frente e equipe (intervalos efetivos, m)",
        ["Frente", "Equipe", "Assentado", "Cadastro de campo", "Reaterrado", "Trechos"],
        [[l["frente_id"], l["equipe_id"], l["ASSENTADO"], l["CADASTRO_CAMPO"], l["REATERRADO"], l["trechos"]] for l in av["por_equipe"]]
        + [["TOTAL", "", av["total_periodo"]["ASSENTADO"], av["total_periodo"]["CADASTRO_CAMPO"], av["total_periodo"]["REATERRADO"], ""]])
    ws.append([])
    ws.append(["Acumulado por frente (todas as datas)", "Trechos", "Extensão (m)", "Assentado (m)", "% assentado", "Cadastro de campo (m)",
               "Reaterrado (m)", "Legado sem posição (m)", "Nós", "Nós concluídos"])
    for c in ws[ws.max_row]:
        c.font = neg
    for f in av["frentes"]:
        ws.append([f["frente_id"], f["trechos"], f["extensao_m"], f["assentado_m"], f["assentado_pct"], f["cadastro_campo_m"],
                   f["reaterrado_m"], f["legado_sem_posicao_m"], f["nos"], f["nos_concluidos"]])
    ws = wb.create_sheet("PENDÊNCIAS")
    pend = d["pendencias"]
    aba(ws, "Pendências por frente", ["Frente"] + list(ROTULOS_PEND.values()) + ["Total"],
        [[f["frente_id"]] + [f.get(k, 0) for k in ROTULOS_PEND] + [f["total"]] for f in pend["frentes"]])
    ws.append([])
    g = pend["gerais"]
    for rot, v in (("Reservas vencidas", g["reservas_vencidas"]), ("Notas a conferir", g["notas_a_conferir"]),
                   ("Itens em quarentena", g["itens_em_quarentena"])):
        ws.append([rot, v])
    ws = wb.create_sheet("FALTANTES")
    aba(ws, "Materiais faltantes para a programação do período",
        ["Material", "Descrição", "Unid.", "Previsto", "Reservado", "Disponível", "A receber", "Faltante"],
        [[f["material_id"], f["descricao"], f["unidade"], f["previsto"], f["reservado"], f["disponivel"], f["a_receber"], f["faltante"]]
         for f in d["faltantes"]])
    buf = io.BytesIO()
    wb.save(buf)
    return f"RELATORIOS_{ini}_{fim}{'_' + frente if frente else ''}.xlsx", buf.getvalue()
