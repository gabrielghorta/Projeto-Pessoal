"""Gera planilhas SINTÉTICAS mínimas com a mesma estrutura das oficiais, para testar a importação
sem depender dos arquivos do contrato (que não são versionados). Os valores aqui não são dados da obra."""
import hashlib
import json
import os

from openpyxl import Workbook
from openpyxl.utils import column_index_from_string as col

ARQ = {
    "redes": "06.1 - SINTETICA_REDES_REV01.xlsx",
    "conexoes": "06.2 - SINTETICA_CONEXOES_REV08.xlsm",
    "bm": "BM_SINTETICO_REV04.xlsx",
    "reajuste": "BM_SINTETICO_REV03_REAJUSTE.xlsx",
    "evidencias": "EVIDENCIAS_SINTETICAS_REV08_06.2_e_REV01_06.1.xlsx",
}

TRECHOS = [
    # ponto, local, DE, extensão, método, código BM, handle, utm_ini, utm_fim, nó ini, nó fim
    ("Zona01_PB01_Rede_TR001", "Zona01", 63, 76.55, "VCA", "5.8.44", "5FD12", (10, 10), (20, 20), "Zona01_PB01_Rede_NÓ1.8", "Zona01_PB01_Rede_NÓ2.7"),
    ("Zona01_PB01_Rede_TR002", "Zona01", 63, 75.07, "VCA", "5.8.44", "5FD13", (20, 20), (30, 30), "Zona01_PB01_Rede_NÓ2.7", "Zona01_PB01_Rede_DF-33"),
    ("Zona02_PB04_Rede_TR036", "Zona02", 63, 17.95, "MND", "5.8.44", "6A001", (40, 40), (50, 50), "Zona02_PB04_Rede_NÓ9.1", "Zona02_PB04_Rede_NÓ9.2"),
]

# nomenclatura, descrição, diâmetro, pressão, quant., código BM, handle, AA, AB, AW, BA, BB, BC
CONEXOES = [
    ("Zona01_PB01_Rede_NÓ1.8", "LUVA ELETROFUSÃO", "63", "PN10/16", 1, "1.14.27", "5FD4F", "P1 aviso A | P1 aviso B", "sugestão A", "P1",
     "CONTRATUAL CONFIRMADO", None, "Zona01_PB01_Rede_TR001"),
    ("Zona01_PB01_Rede_NÓ1.8", "TÊ 90º ELETROFUSÃO INJETADO", "63", "PN10", 1, "1.14.53", "5FD4F", None, None, None,
     "CONTRATUAL CONFIRMADO", None, "Zona01_PB01_Rede_TR001"),
    ("Zona01_PB01_Rede_DF-33", "TÊ 90º ELETROFUSÃO INJETADO", "63", "PN10", 1, "1.14.53", "5FDEA", "DF-05 duplicidade histórica", None, "P2",
     "CONTRATUAL CONFIRMADO", None, "Zona01_PB01_Rede_TR002"),
    ("Zona01_PB01_CH_DF-33", "LUVA ELETROFUSÃO", "63", "PN10/16", 1, "1.14.27", "5FDEA", None, None, None, "CONTRATUAL CONFIRMADO", None, None),
    ("Zona01_PB01_CH_DF-33", "LUVA ELETROFUSÃO", "63", "PN10/16", 1, "1.14.27", "5FDEA", None, None, None, "CONTRATUAL CONFIRMADO", None, None),
    ("Zona01_PB01_CH_DF-33", "TUBO", "63", "PN10", 0.71, "1.14.77", "5FDEA", None, None, None, "CONTRATUAL CONFIRMADO", None, None),
    ("Zona01_PB01_CH_DF-33", "TÊ 90º DE REDUÇÃO", "90", "PN16", 2, "1.14.57", "5FDEA", None, None, None, "CONTRATUAL — VALIDAR", "PN diferente", None),
    ("Zona01_PB01_CH_DF-33", "TUBO C/ FLANGE E PONTA PN10 L=0,36 M", "100", "PN10", 1, "APROVAR", "5FDEA", None, None, None,
     "APROVAR — SEM ITEM NO BM", "sem item", None),
    ("Zona11_PBND_CH_MV-11", "MEDIDOR DE VAZÃO", "150", "PN10", 1, "3.1.1", None, None, None, None, "CONTRATUAL CONFIRMADO", None, None),
    # textos reais dos avisos (06.2 REV08) para testar o escopo das decisões
    ("Zona11_PB15_Rede_DF-05", "TÊ 90º ELETROFUSÃO INJETADO", "63", "PN10", 1, "1.14.53", "6DF05",
     "P1 — DF-05 contada em duas pranchas (16 e 19).", "Tirar a DF-05 do resumo e do quadro de uma das pranchas.", "P1",
     "CONTRATUAL CONFIRMADO", None, None),
    ("Zona01_PB02_Rede_DF-32", "TÊ 90º ELETROFUSÃO INJETADO", "63", "PN10", 1, "1.14.53", "6DF32",
     "P2 — DF-32 no canto da curva 2 (6C232): o tê da descarga pode fazer a função da curva (possível duplicidade).",
     "Possível duplicidade tê × curva: definir a montagem com a fiscalização (F9) antes de comprar.", "P2",
     "CONTRATUAL CONFIRMADO", None, None),
]

JUNTADA = [  # item, unidade, preço, classe
    ("1.14.27", "UN", 10.0, "MATERIAL"), ("1.14.53", "UN", 20.0, "MATERIAL"), ("1.14.57", "pç", 30.0, "MATERIAL"),
    ("1.14.77", "M", 17.06, "MATERIAL"), ("3.1.1", "UN", 11167.22, "MATERIAL"), ("5.8.44", "M", 0.78, "SERVIÇO"),
    ("5.2.1", "M", 2.61, "SERVIÇO"), ("5.17.3", "M3", 47.5, "SERVIÇO"),
]
VERTICAL = [  # chave, subgrupo, item, quantidade
    ("1.2_1.14.27", "1.2", "1.14.27", 100), ("1.2_1.14.53", "1.2", "1.14.53", 50), ("1.2_1.14.57", "1.2", "1.14.57", 10),
    ("1.1_1.14.77", "1.1", "1.14.77", 15), ("5.1_1.14.77", "5.1", "1.14.77", 10), ("5.1_3.1.1", "5.1", "3.1.1", 3),
    ("2.2_5.8.44", "2.2", "5.8.44", 70458.38), ("2.2_5.2.1", "2.2", "5.2.1", 111219.2), ("4.2_5.17.3", "4.2", "5.17.3", 6026.69),
]


def _rede(p, trechos):
    wb = Workbook()
    ws = wb.active
    ws.title = "REDE"
    for i, t in enumerate(trechos, 6):
        ponto, local, de, ext, met, cod, h, a, b, ni, nf = t
        valores = {"A": i - 5, "G": "ELETROFUSÃO", "H": de, "I": "PEAD", "K": local, "M": ext, "N": met, "R": ponto,
                   "S": "TERCIÁRIA", "U": "002/037", "X": "ST BAIRRO", "AB": f"TRECHO DE {ni}", "AC": cod,
                   "AE": 1.2, "AF": 1.363, "AG": 0.65, "AH": 0.1, "AS": 0.585 * ext, "AW": 0.6 * ext, "BA": 0.08 * ext,
                   "BM": f"EQUIPE {local}", "BN": "Rede", "BO": ponto.split("_")[1], "BP": ni, "BQ": nf, "BR": h,
                   "BS": a[0], "BT": a[1], "BU": b[0], "BV": b[1], "BW": -15.9, "BX": -47.7}
        if met == "MND":
            valores.update({"AF": 0, "AS": 0, "AW": 0, "BA": 0})     # como na 06.1 real: MND sem vala, profundidade 0
        for c, v in valores.items():
            ws.cell(row=i, column=col(c), value=v)
    wb.save(p)


def _conexoes(p, linhas):
    wb = Workbook()
    ws = wb.active
    ws.title = "CONEXÕES"
    cab = {"B": "DESCRIÇÃO", "C": "PRESSÃO", "D": "RESINA", "E": "TIPO_CONEXÃO", "F": "SOLDA", "G": "DIÂMETRO", "H": "MATERIAL",
           "I": "GRUPO", "J": "QUANT.", "K": "NOMENCLATURA (OBRA)", "M": "LOCAL", "R": "CÓDIGO BM", "AA": "CONSIDERAÇÃO",
           "AB": "SUGESTÃO DE CORREÇÃO", "AE": "EQUIPE", "AG": "TIPO ENTRONCAMENTO", "AH": "STATUS CONEXÃO", "AR": "HANDLE CAD",
           "AS": "COORD. E (UTM)", "AT": "COORD. N (UTM)", "AU": "LATITUDE", "AV": "LONGITUDE", "AW": "PRIORIDADE DO ERRO",
           "BA": "ENQUADRAMENTO BM", "BB": "OBS. ENQUADRAMENTO BM", "BC": "TRECHO (06.1)"}
    for c, v in cab.items():
        ws.cell(row=5, column=col(c), value=v)
    for i, (nom, d, diam, pn, q, cod, h, aa, ab, aw, ba, bb, bc) in enumerate(linhas, 6):
        local = "CH" if "_CH_" in nom else nom.split("_")[0]
        posicao = (203552.5, 8239199.1, -15.9087, -47.7686) if h else (None, None, None, None)
        for c, v in {"B": d, "C": pn, "D": "PE100", "E": d.split()[0], "F": "ELETROFUSÃO", "G": diam, "H": "PEAD", "I": "HIDROMECÂNICO",
                     "J": q, "K": nom, "M": local, "R": cod, "AA": aa, "AB": ab, "AE": f"EQUIPE {local}", "AG": "A DEFINIR",
                     "AH": "CONEXÃO PREVISTA", "AR": h, "AS": posicao[0], "AT": posicao[1], "AU": posicao[2], "AV": posicao[3],
                     "AW": aw, "BA": ba, "BB": bb, "BC": bc}.items():
            ws.cell(row=i, column=col(c), value=v)
    wb.save(p)


def _bm(p, quantidade_override=None, precos=None):
    wb = Workbook()
    j = wb.active
    j.title = "BM MORRO DA CRUZ"
    for i, (it, un, preco, cl) in enumerate(JUNTADA, 6):
        preco = (precos or {}).get(it, preco)
        for k, v in enumerate([it, f"desc {it}", "N/A", "N/A", "N/A", "N/A", "N/A", "N/A", "G", "G", "000", cl, f"DESCRIÇÃO {it}",
                               un, 1, preco, preco], 1):
            j.cell(row=i, column=k, value=v)
    v_ = wb.create_sheet("BM MORRO DA CRUZ - VERTICAL")
    precos = {it: (un, (precos or {}).get(it, pr), cl) for it, un, pr, cl in JUNTADA}
    tot = {"MATERIAL": 0, "SERVIÇO": 0}
    linha = 6
    for chave, sg, it, q in VERTICAL:
        q = (quantidade_override or {}).get(chave, q)
        un, pr, cl = precos[it]
        tot[cl] += q * pr
        for k, val in enumerate([1, sg, f"GRUPO {sg}", chave, it, f"desc {it}", "N/A", "N/A", "N/A", "N/A", "N/A", "N/A", "G", "G",
                                 "000", cl, f"DESCRIÇÃO {it}", un, q, pr, round(q * pr, 2)], 1):
            v_.cell(row=linha, column=k, value=val)
        linha += 1
    from decimal import ROUND_HALF_UP, Decimal
    r2 = lambda x: float(Decimal(format(x, ".15g")).quantize(Decimal("0.01"), ROUND_HALF_UP))
    m, s = r2(tot["MATERIAL"]), r2(tot["SERVIÇO"])
    bdi_m, bdi_s, k = 0.1528, 0.2148, -1 + (1 - 0.1404)
    t02 = m + s + m * bdi_m + s * bdi_s
    linha += 1
    for rot, vals in (("MATERIAL", [m]), ("SERVIÇO", [s]), ("BDI MATERIAL =>", [bdi_m, m * bdi_m]),
                      ("BDI SERVIÇO =>", [bdi_s, s * bdi_s]), ("COEF. - K =>", [k, t02 * k]), ("TOTAL 01 + 02 + BDI -K", [t02 * (1 + k)])):
        if rot == "MATERIAL":
            v_.cell(row=linha, column=1, value="RESUMO PRINCIPAL")
        v_.cell(row=linha, column=18, value=rot)
        for c, val in enumerate(vals, 19):
            v_.cell(row=linha, column=c, value=val)
        linha += 1
    wb.save(p)
    return t02 * (1 + k)


def _reajuste(p):
    from datetime import datetime
    wb = Workbook()
    ws = wb.active
    ws.title = "REAJUSTE"
    for i, (rot, v) in enumerate([("Data-base do reajuste - conforme CEFR", datetime(2025, 1, 31)),
                                  ("Percentual do 1º reajuste informado pela CEFR", 0.0329),
                                  ("1º aniversário da data-base", datetime(2026, 1, 31)),
                                  ("Início da execução (OS, página 1):", datetime(2026, 10, 1))], 1):
        ws.cell(row=i, column=1, value=rot)
        ws.cell(row=i, column=2, value=v)
    wb.save(p)


def _evidencias(p, linhas):
    wb = Workbook()
    ws = wb.active
    ws.title = "ELEMENTOS_ZONA"
    grupos = {}
    for l in linhas:
        chave = l[6] or l[0].split("_")[-1]
        grupos.setdefault(chave, set()).add(l[0])
    for i, (chave, noms) in enumerate(grupos.items(), 2):
        num = sorted(noms)[0].split("_")[-1]
        ws.cell(row=i, column=1, value=f"SUBDMC1_{num}")
        ws.cell(row=i, column=2, value=" | ".join(sorted(noms)))
        ws.cell(row=i, column=6, value="Zona01")
        ws.cell(row=i, column=9, value="COERENTE")
    wb.save(p)


def sha(p):
    with open(p, "rb") as f:
        return hashlib.sha256(f.read()).hexdigest()


def conjunto_controles(pasta, esperado, ident="CTRL-SINTETICO", arquivos=None):
    """Conjunto de controles para os arquivos presentes em `pasta` (hash de cada um)."""
    arquivos = arquivos or ARQ
    fontes = {k: sha(os.path.join(pasta, n)) for k, n in arquivos.items() if os.path.exists(os.path.join(pasta, n))}
    return {"id": ident, "perfil": "TESTE", "fontes": fontes, "esperado": esperado,
            "origem_dos_valores": "listas do gerador sintético (perfil de TESTE)",
            "aprovacao": {"status": "APROVADO", "por": "teste automatizado (perfil sintético)", "em": "2026-10-16"}}


def esperado_de(trechos, conexoes, evidencias=True):
    """Controles calculados das listas de entrada (independentes do importador)."""
    elementos = {l[6] or l[0].split("_")[-1] for l in conexoes}
    unidades = {it: un.upper() for it, un, _, _ in JUNTADA}
    linhas = {}
    for l in conexoes:
        u = unidades.get(l[5], "PENDENTE")
        linhas[u] = linhas.get(u, 0) + 1
    e = {"trechos": len(trechos), "extensao_m": round(sum(t[3] for t in trechos), 2), "elementos": len(elementos),
         "atividades": len({l[0] for l in conexoes}), "itens_projeto": len(conexoes), "bm_vertical": len(VERTICAL),
         "bm_juntada": len(JUNTADA), "sem_posicao": sorted({l[0].split("_")[-1] for l in conexoes if not l[6]}),
         "unidades_linhas": dict(sorted(linhas.items()))}
    if evidencias:
        e["agrupamento_evidencias_iguais"] = len(elementos)
    return e


def gerar(pasta, trechos=None, conexoes=None, bm_override=None, precos=None, evidencias=True):
    os.makedirs(pasta, exist_ok=True)
    trechos = TRECHOS if trechos is None else trechos
    conexoes = CONEXOES if conexoes is None else conexoes
    _rede(os.path.join(pasta, ARQ["redes"]), trechos)
    _conexoes(os.path.join(pasta, ARQ["conexoes"]), conexoes)
    p0 = _bm(os.path.join(pasta, ARQ["bm"]), bm_override, precos)
    _reajuste(os.path.join(pasta, ARQ["reajuste"]))
    pe = os.path.join(pasta, ARQ["evidencias"])
    if evidencias:
        _evidencias(pe, conexoes)
    elif os.path.exists(pe):
        os.remove(pe)
    return {
        "contrato": "CT SINTÉTICO", "obra": "teste", "perfil": "TESTE",
        "valor_contrato": {"valor": "1.00", "fonte": "teste"},
        "arquivos": dict(ARQ, mapa="inexistente.html"),
        "controles": [conjunto_controles(pasta, esperado_de(trechos, conexoes, evidencias))],
        "conciliacoes": [],
        "materiais_consumo": [{"id": "CONS-AREIA", "descricao": "Areia", "unidade": "M3"}],
        "_p0": p0,
    }


if __name__ == "__main__":
    import sys
    print(json.dumps(gerar(sys.argv[1]), ensure_ascii=False, indent=2))
