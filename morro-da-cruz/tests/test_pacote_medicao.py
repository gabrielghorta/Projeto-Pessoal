"""Medição no modelo oficial do BM (pacote_medicao.py e xlsx_modelo.py), sobre as planilhas reais (BM REV04).

O pacote sai do próprio arquivo oficial: a cadeia de fórmulas MC → vertical → juntada → cronograma FF continua valendo,
com uma folha por item (fotos e justificativa), a administração local proporcional, reajuste R$ 0,00 não pedido, e tudo
marcado TESTE – NÃO OFICIAL. Com o LibreOffice instalado, o pacote é recalculado e os totais conferidos com o sistema."""
import io
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import unittest
import zipfile

from openpyxl import load_workbook

AQUI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(AQUI))
sys.path.insert(0, AQUI)
import medicao as MD  # noqa: E402
import pacote_medicao as PM  # noqa: E402
import regras as R  # noqa: E402
import servicos as S  # noqa: E402
import xlsx_modelo as X  # noqa: E402
from base_real import BaseReal  # noqa: E402
from fotos_teste import jpeg  # noqa: E402

TRS = ["Zona01_PB01_Rede_TR001", "Zona01_PB01_Rede_TR002", "Zona01_PB01_Rede_TR003"]
TUDO = dict.fromkeys(R.CHECKLIST_CONFERENCIA_MEDICAO, True)
SOFFICE = shutil.which("soffice") or shutil.which("libreoffice")


class TestEditorXlsx(unittest.TestCase):
    """O editor por dentro do arquivo: fórmula compartilhada vira explícita e a cópia de bloco desloca as referências."""

    XML = ('<?xml version="1.0" encoding="UTF-8"?><worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
           '<dimension ref="A1:C3"/><cols><col min="1" max="5" width="10" customWidth="1"/></cols><sheetData>'
           '<row r="1"><c r="A1"><v>1</v></c><c r="B1" s="3"><f t="shared" ref="B1:B3" si="0">A1*2+$A$1</f><v>3</v></c></row>'
           '<row r="2"><c r="A2"><v>2</v></c><c r="B2" s="3"><f t="shared" si="0"/><v>5</v></c></row>'
           '<row r="3"><c r="A3"><v>3</v></c><c r="B3" s="3"><f t="shared" si="0"/><v>7</v></c>'
           '<c r="C3"><f>SUM(\'OUTRA ABA\'!$W$1:$W$9,B3)</f></c></row>'
           '</sheetData><pageMargins left="0.7" right="0.7" top="0.7" bottom="0.7" header="0.3" footer="0.3"/>'
           '<pageSetup paperSize="9" scale="50" orientation="landscape"/></worksheet>')

    def test_compartilhada_vira_explicita_e_bloco_desloca(self):
        a = X.Aba(self.XML)
        a.expandir_compartilhadas()
        self.assertEqual([a.celula(f"B{r}").f for r in (1, 2, 3)], ["A1*2+$A$1", "A2*2+$A$1", "A3*2+$A$1"])
        self.assertEqual(a.celula("B2").fatr, {})
        a.copiar_bloco(1, 3, 11)
        self.assertEqual(a.celula("B12").f, "A12*2+$A$11")                    # absoluta da aba também anda (cópia de folha)
        self.assertEqual(a.celula("C13").f, "SUM('OUTRA ABA'!$W$1:$W$9,B13)")  # a de outra aba fica
        self.assertEqual(a.celula("B12").atr["s"], "3")
        a.sem_valores_guardados()
        a.ocultar_colunas([2, 3])
        a.uma_pagina_de_largura()
        a.cabecalho("TESTE – NÃO OFICIAL")
        xml = a.xml()
        self.assertIn('<col min="2" max="3" width="10" customWidth="1" hidden="1"/>', xml)
        self.assertIn('<pageSetUpPr fitToPage="1"/>', xml)
        self.assertIn('fitToWidth="1" fitToHeight="0"', xml)
        self.assertNotIn('scale="50"', xml)
        self.assertLess(xml.index("<pageSetup"), xml.index("<headerFooter"))   # ordem do esquema
        self.assertIn('<dimension ref="A1:C13"/>', xml)
        self.assertNotIn("<v>5</v>", xml)

    def test_ida_e_volta_do_modelo_oficial_sem_mudar_nada(self):
        caminho = PM.caminho_modelo()
        if not caminho:
            self.skipTest("BM oficial ausente em planilhas/")
        with open(caminho, "rb") as f:
            p = X.Pacote(f.read())
        a = p.aba(PM.VERTICAL)
        a.expandir_compartilhadas()
        self.assertTrue(a.celula("V31").f.startswith("IFERROR(SUMIFS('01ª MC'!$W$1:$W$1840"))
        self.assertEqual(a.celula("W31").f, "ROUND(V31*$T31,2)")
        p.recalcular_ao_abrir()
        wb = load_workbook(io.BytesIO(p.salvar()))
        self.assertEqual(wb.sheetnames, [n for n, _ in p.abas()])
        self.assertEqual(wb[PM.VERTICAL]["W546"].value, "=ROUND(V546*$T546,2)")


class BasePacote(BaseReal):
    LIBERAR_TUDO = True

    def setUp(self):
        super().setUp()
        self.mid = S.criar_medicao(self.c, 1, "2026-10-01", "2026-10-15")
        for t in ("MANHA", "TARDE"):
            for d in ("2026-10-10", "2026-10-20"):
                S.registrar_clima(self.c, d, "Zona01", t, "SOL")

    def executar(self, ponto, data="2026-10-10"):
        ext = S.alvo(self.c, "TRECHO", ponto)["extensao_m"]
        S.apontar(self.c, data, "Zona01", "TRECHO", ponto, intervalos_novos=[{"etapa": "ASSENTADO", "ini_m": 0, "fim_m": ext},
                                                                             {"etapa": "REATERRADO", "ini_m": 0, "fim_m": ext}])
        if not (S.execucao(self.c, "TRECHO", self.trecho(ponto)) or {}).get("concluido"):
            S.apontar(self.c, data, "Zona01", "TRECHO", ponto, concluido_100=True)
        for i, cat in enumerate(("SINALIZACAO", "EXECUCAO", "CONCLUIDO")):
            S.registrar_foto(self.c, "TRECHO", ponto, cat, jpeg(), data_hora=f"{data}T0{7 + i}:00:00", dir_fotos=self.fotos)
        S.verificar_trecho(self.c, ponto, "sup1")
        S.ok_fiscalizacao(self.c, "TRECHO", ponto, "eng1", data)

    def gerar(self, mid=None):
        nome, dados, r = PM.gerar(self.c, mid or self.mid, dir_fotos=self.fotos)
        return nome, dados, r, load_workbook(io.BytesIO(dados))

    def recalcular(self, dados):
        """O pacote recalculado pelo LibreOffice (valores das fórmulas)."""
        pasta = tempfile.mkdtemp(dir=self.tmp)
        arq = os.path.join(pasta, "pacote.xlsx")
        with open(arq, "wb") as f:
            f.write(dados)
        saida = os.path.join(pasta, "lo")
        subprocess.run([SOFFICE, "--headless", "--calc", "--convert-to", "xlsx", "--outdir", saida, arq],
                       capture_output=True, timeout=300, env={**os.environ, "HOME": pasta})
        return load_workbook(os.path.join(saida, "pacote.xlsx"), data_only=True)


class TestPacote(BasePacote):
    def test_folhas_fotos_sugeridas_e_administracao_local(self):
        for p in TRS:
            self.executar(p)
        r = PM.resumo(self.c, self.mid)
        itens = [f["item"] for f in r["folhas"]]
        linhas = [S._bm_item(self.c, k)["linha_origem"] for k in itens]
        self.assertEqual(linhas, sorted(linhas))                                # na ordem do BM vertical
        self.assertIn("2.2_5.8.44", itens)
        ass = next(f for f in r["folhas"] if f["item"] == "2.2_5.8.44")
        ext = sum(S.alvo(self.c, "TRECHO", p)["extensao_m"] for p in TRS)
        self.assertAlmostEqual(ass["quantidade"], ext, places=6)
        self.assertEqual(len(ass["linhas"]), 3)
        self.assertTrue(all(l["estacas"] and l["estacas"][:2] == (0, 0.0) for l in ass["linhas"]))
        l1 = ass["linhas"][0]
        self.assertAlmostEqual(l1["estacas"][2] * 20 + l1["estacas"][3], l1["quantidade"], places=6)
        # 6 fotos sugeridas: as mais recentes dos trechos do item
        self.assertTrue(ass["fotos_sugeridas"])
        self.assertEqual(len(ass["fotos"]), 6)
        recentes = [x["id"] for x in sorted(ass["candidatas"], key=lambda x: (x["data_hora"], x["id"]), reverse=True)[:6]]
        self.assertEqual(ass["fotos"], recentes)
        # ADM proporcional: % ADM = ADM ÷ (subtotal 01 do contrato − ADM); faturamento sem a ADM
        adm = r["adm"]
        self.assertAlmostEqual(adm["total_contrato"], 1304685.94, places=2)
        self.assertAlmostEqual(adm["quantidade"], adm["faturamento"] * adm["percentual"] / adm["total_contrato"], places=12)
        self.assertGreater(adm["valor"], 0)
        self.assertTrue(any("Reajuste R$ 0,00" in a for a in r["avisos"]))
        self.assertTrue(any("PENDENTE" in a for a in r["avisos"]))

    def test_pacote_no_modelo_oficial_com_formulas_e_marca_de_teste(self):
        for p in TRS:
            self.executar(p)
        nome, dados, r, wb = self.gerar()
        self.assertEqual(nome, "BM_MORRO_DA_CRUZ_01a_MEDICAO_TESTE_NAO_OFICIAL.xlsx")
        self.assertEqual(wb.sheetnames, [PM.JUNTADA, PM.VERTICAL, "01ª MC", PM.FF, PM.REAJUSTE, "GRÁFICOS",
                                         "CONFERÊNCIA CONTRATO", PM.AVISOS])
        z = zipfile.ZipFile(io.BytesIO(dados))
        self.assertNotIn("xl/calcChain.xml", z.namelist())
        self.assertIn('fullCalcOnLoad="1"', z.read("xl/workbook.xml").decode())
        self.assertEqual(len([n for n in z.namelist() if n.startswith("xl/charts/chart")]), 4)    # os 4 gráficos ficam
        self.assertTrue(any(n.startswith("xl/media/mc_foto_") for n in z.namelist()))
        # memória de cálculo: uma folha por item + a da ADM, cada uma numa página
        mc = wb["01ª MC"]
        n_folhas = len(r["folhas"]) + 1
        self.assertEqual([mc[f"B{1 + 92 * k}"].value for k in range(n_folhas)], [f"FOLHA - {k + 1:03d}" for k in range(n_folhas)])
        self.assertEqual([mc[f"C{7 + 92 * k}"].value for k in range(len(r["folhas"]))], [f["item"] for f in r["folhas"]])
        self.assertIn("TESTE – NÃO OFICIAL", mc["C1"].value)
        adm_lin = 92 * len(r["folhas"]) + 7
        self.assertEqual(mc[f"C{adm_lin}"].value, PM.CHAVE_ADM)
        self.assertEqual(mc[f"W{adm_lin}"].value, f"=SUM($W{adm_lin + 20}:$W{adm_lin + 85})")
        self.assertIn("'BM MORRO DA CRUZ - VERTICAL'!$W$7:$W$542", mc[f"R{adm_lin + 20}"].value)
        self.assertEqual(len(mc.row_breaks.brk), n_folhas - 1)
        self.assertIsNone(mc.legacy_drawing)                                    # as instruções de preenchimento manual saem
        # memorial do assentamento: estaca 0 até o fim do trecho, sem encadear
        k = next(i for i, f in enumerate(r["folhas"]) if f["item"] == "2.2_5.8.44")
        lin = 27 + 92 * k
        self.assertEqual((mc[f"I{lin}"].value, mc[f"J{lin}"].value, mc[f"U{lin}"].value), (0, 0, 1))
        self.assertEqual(mc[f"E{lin}"].value, TRS[0])
        self.assertEqual((mc[f"E{lin + 3}"].value, mc[f"I{lin + 3}"].value), (None, None))   # linha vazia: sem estaca herdada
        # vertical: o mês puxa da MC; a ADM da folha da ADM
        v = wb[PM.VERTICAL]
        ass = S._bm_item(self.c, "2.2_5.8.44")["linha_origem"]
        self.assertEqual(v[f"V{ass}"].value, f"=IFERROR(SUMIFS('01ª MC'!$W$1:$W${92 * n_folhas},'01ª MC'!$C$1:$C${92 * n_folhas},$D{ass}),0)")
        self.assertEqual(v["V546"].value, f"='01ª MC'!$W${adm_lin}")
        self.assertEqual(v[f"X{ass}"].value, 0)
        self.assertIn("TESTE – NÃO OFICIAL", v["E1"].value)
        # juntada: só os itens medidos à vista; as colunas das outras medições ocultas
        j = wb[PM.JUNTADA]
        jl = {b["item"]: b["linha_origem"] for b in self.c.execute(
            "SELECT item, linha_origem FROM bm_itens WHERE planilha='JUNTADA' AND vigente=1").fetchall()}
        self.assertFalse(j.row_dimensions[jl["5.8.44"]].hidden)
        self.assertFalse(j.row_dimensions[jl["4.1.1"]].hidden)                 # a ADM entra quando há faturamento
        self.assertTrue(j.row_dimensions[jl["1.1.1"]].hidden)
        self.assertTrue(j.column_dimensions["T"].hidden)
        self.assertFalse(j.column_dimensions["R"].hidden)
        self.assertTrue(j.sheet_properties.pageSetUpPr.fitToPage)
        # reajuste: fração elegível 0, R$ 0,00 não pedido
        rj = wb[PM.REAJUSTE]
        self.assertEqual((rj["E24"].value, rj["F24"].value), (0, 0))
        self.assertIn("NÃO PEDIDO", rj["I24"].value)
        self.assertIn("1ª MEDIÇÃO", z.read("xl/worksheets/sheet1.xml").decode())        # cabeçalho de impressão
        avisos = [c.value for c in wb[PM.AVISOS]["A"] if c.value]
        self.assertTrue(avisos[0].startswith("MEDIÇÃO DE TESTE – NÃO OFICIAL"))

    @unittest.skipUnless(SOFFICE, "LibreOffice ausente: o recálculo das fórmulas não é conferido")
    def test_recalculado_bate_com_o_sistema(self):
        for p in TRS:
            self.executar(p)
        _, dados, r, _ = self.gerar()
        wb = self.recalcular(dados)
        v, j, ff = wb[PM.VERTICAL], wb[PM.JUNTADA], wb[PM.FF]
        self.assertAlmostEqual(j["R411"].value, r["totais"]["juntada"], places=2)     # P0 com BDI e K
        self.assertAlmostEqual(v["V556"].value, r["totais"]["vertical"], places=2)
        self.assertAlmostEqual(v["W546"].value, r["adm"]["valor"], places=2)          # ADM proporcional
        self.assertAlmostEqual(j["R417"].value, r["totais"]["juntada"], places=2)     # total geral = P0 (reajuste 0)
        self.assertAlmostEqual(ff["F129"].value, r["totais"]["vertical"], places=2)   # cronograma FF executado
        sem_adm = R.totalizar_bm([(f["classificacao"], f["quantidade"], S._bm_item(self.c, f["item"])["preco_unitario"])
                                  for f in r["folhas"]], *S._fatores(self.c))["total"]
        self.assertAlmostEqual(S.resumo_medicao(self.c, self.mid)["valor_pi"], sem_adm, places=2)

    def test_fotos_e_justificativa_da_folha(self):
        for p in TRS:
            self.executar(p)
        f = next(x for x in PM.folhas(self.c, self.mid) if x["item"] == "2.2_5.8.44")
        cand = [x["id"] for x in f["candidatas"]]
        PM.salvar_folha(self.c, self.mid, "2.2_5.8.44", "eng1", fotos=[cand[-1], cand[0]], justificativa="  Trecho conforme NS.  ")
        f = next(x for x in PM.folhas(self.c, self.mid) if x["item"] == "2.2_5.8.44")
        self.assertEqual((f["fotos"], f["fotos_sugeridas"], f["justificativa"]), ([cand[-1], cand[0]], False, "Trecho conforme NS."))
        _, dados, r, wb = self.gerar()
        k = [x["item"] for x in r["folhas"]].index("2.2_5.8.44")
        self.assertEqual(wb["01ª MC"][f"X{9 + 92 * k}"].value, "Trecho conforme NS.")
        dr = zipfile.ZipFile(io.BytesIO(dados)).read("xl/drawings/drawing3.xml").decode()
        self.assertEqual(len(re.findall(rf"<xdr:row>{92 * k + 8}</xdr:row><xdr:rowOff>", dr)), 2)   # 2 fotos na galeria
        for fotos, msg in (([999999], "não são dos trechos"), (cand[:7], "no máximo 6"), ([cand[0], cand[0]], "duas vezes")):
            with self.assertRaisesRegex(S.RegraViolada, msg):
                PM.salvar_folha(self.c, self.mid, "2.2_5.8.44", "eng1", fotos=fotos)
        with self.assertRaisesRegex(S.RegraViolada, "não tem quantidade proposta"):
            PM.salvar_folha(self.c, self.mid, "2.2_9.9.9", "eng1", fotos=[])
        PM.salvar_folha(self.c, self.mid, "2.2_5.8.44", "eng1", fotos=None, justificativa="")     # volta à sugestão
        f = next(x for x in PM.folhas(self.c, self.mid) if x["item"] == "2.2_5.8.44")
        self.assertTrue(f["fotos_sugeridas"])
        self.assertIsNone(f["justificativa"])

    def test_acima_de_100_so_fecha_com_justificativa_e_sai_aditivo_pendente(self):
        self.executar(TRS[0])
        b = S._bm_item(self.c, "2.2_5.2.7")                                  # locação: o lançamento passa do contrato
        self.c.execute("UPDATE medicao_lancamentos SET quantidade=? WHERE medicao_id=? AND bm_item_id=?",
                       (b["quantidade"] + 10, self.mid, b["id"]))
        f = next(x for x in PM.folhas(self.c, self.mid) if x["item"] == "2.2_5.2.7")
        self.assertTrue(f["excede"])
        MD.conferir(self.c, self.mid, "CONFORME", TUDO, "eng1")
        with self.assertRaisesRegex(S.RegraViolada, r"acima de 100 % do contrato sem justificativa.*5\.2\.7"):
            S.fechar_medicao(self.c, self.mid, "eng1")
        _, _, r, wb = self.gerar()
        k = [x["item"] for x in r["folhas"]].index("2.2_5.2.7")
        self.assertTrue(wb["01ª MC"][f"X{9 + 92 * k}"].value.startswith("ADITIVO PENDENTE"))
        self.assertTrue(any("ADITIVO PENDENTE" in a and "falta a justificativa" in a for a in r["avisos"]))
        PM.salvar_folha(self.c, self.mid, "2.2_5.2.7", "eng1", justificativa=PM.FRASES[3])
        S.fechar_medicao(self.c, self.mid, "eng1")
        with self.assertRaisesRegex(S.RegraViolada, "Medição fechada"):
            PM.salvar_folha(self.c, self.mid, "2.2_5.2.7", "eng1", justificativa="outra")

    def test_segunda_medicao_congela_a_primeira_e_puxa_da_02a_mc(self):
        self.executar(TRS[0])
        MD.conferir(self.c, self.mid, "CONFORME", TUDO, "eng1")
        S.fechar_medicao(self.c, self.mid, "eng1")
        m2 = S.criar_medicao(self.c, 2, "2026-10-16", "2026-10-31")
        self.executar(TRS[1], "2026-10-20")
        nome, dados, r, wb = self.gerar(m2)
        self.assertEqual(nome, "BM_MORRO_DA_CRUZ_02a_MEDICAO_TESTE_NAO_OFICIAL.xlsx")
        self.assertIn("02ª MC", wb.sheetnames)
        self.assertNotIn("01ª MC", wb.sheetnames)
        self.assertEqual(wb["02ª MC"].print_area, "'02ª MC'!$B$1:$X$" + str(92 * (len(r["folhas"]) + 1)))
        v = wb[PM.VERTICAL]
        ass = S._bm_item(self.c, "2.2_5.8.44")["linha_origem"]
        self.assertAlmostEqual(v[f"V{ass}"].value, S.alvo(self.c, "TRECHO", TRS[0])["extensao_m"], places=6)   # congelada
        self.assertTrue(v[f"X{ass}"].value.startswith("=IFERROR(SUMIFS('02ª MC'!$W$1:"))
        self.assertIsInstance(v["V546"].value, float)
        self.assertTrue(v["X546"].value.startswith("='02ª MC'!$W$"))
        self.assertEqual(wb["02ª MC"]["T2"].value, "='BM MORRO DA CRUZ - VERTICAL'!$X$4")
        rj = wb[PM.REAJUSTE]
        self.assertEqual([(rj[f"E{x}"].value, rj[f"F{x}"].value) for x in (24, 25)], [(0, 0), (0, 0)])
        self.assertTrue(any("anteriores" in a for a in r["avisos"]))
        if SOFFICE:
            wb = self.recalcular(dados)
            self.assertAlmostEqual(wb[PM.JUNTADA]["T411"].value, r["totais"]["juntada"], places=2)

    def test_modelo_diferente_do_importado_e_recusado(self):
        self.executar(TRS[0])
        falso = os.path.join(self.tmp, "outro.xlsx")
        with open(PM.caminho_modelo(), "rb") as f, open(falso, "wb") as g:
            g.write(f.read() + b"\0")
        with self.assertRaisesRegex(S.RegraViolada, "não é o que o sistema importou"):
            PM.gerar(self.c, self.mid, dir_fotos=self.fotos, modelo=falso)


class TestPacoteHTTP(BasePacote):
    def test_consultar_escolher_e_baixar(self):
        self.executar(TRS[0])
        eng = self.cli("eng1")
        r = self.ok(eng.get(f"/api/medicao/pacote?id={self.mid}"))
        f = next(x for x in r["folhas"] if x["item"] == "2.2_5.8.44")
        self.assertEqual(f["linhas"], 1)
        self.assertEqual(set(f["candidatas"][0]), {"id", "arquivo", "data_hora", "categoria", "alvo_tipo", "alvo_id", "alvo"})
        self.assertEqual(r["frases"], PM.FRASES)
        escolha = [f["candidatas"][1]["id"]]
        self.ok(eng.post("/api/medicao/folha", {"medicao_id": self.mid, "item": "2.2_5.8.44", "fotos": escolha,
                                                "justificativa": PM.FRASES[0], "por": "outro"}))
        self.assertEqual([tuple(x) for x in self.c.execute("SELECT por, fotos FROM medicao_folhas")],
                         [("eng1", json.dumps(escolha))])                       # autor = sessão, não o digitado
        st, xlsx = eng.get(f"/api/medicao/pacote/baixar?id={self.mid}")
        self.assertEqual(st, 200)
        wb = load_workbook(io.BytesIO(xlsx))
        self.assertIn("01ª MC", wb.sheetnames)
        self.assertEqual(self.c.execute("SELECT COUNT(*) FROM auditoria WHERE rota='/api/medicao/pacote/baixar'").fetchone()[0], 1)
        enc = self.cli("enc3")                                                  # equipe: sem a tela de medição
        self.recusa(enc.get(f"/api/medicao/pacote?id={self.mid}"), 403, "")
        st, _ = enc.get(f"/api/medicao/pacote/baixar?id={self.mid}")
        self.assertEqual(st, 403)
        self.recusa(enc.post("/api/medicao/folha", {"medicao_id": self.mid, "item": "2.2_5.8.44", "fotos": None}), 403, "")


if __name__ == "__main__":
    unittest.main()
