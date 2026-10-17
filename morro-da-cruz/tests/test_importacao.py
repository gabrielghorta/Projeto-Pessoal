"""Etapas 0 e 1: backup verificável, carga não destrutiva, importação idempotente, proveniência,
identidade estável, unidades com fonte, BA/BB/AW, decisões preservadas e migração de esquema.

Inclui uma regressão por achado da auditoria independente do commit b63b5db (A01–A12), reproduzindo
exatamente o cenário descrito. Usa planilhas SINTÉTICAS (tests/planilhas_sinteticas.py).

Cada teste falha se deixar arquivo aberto dentro da pasta temporária (no Windows isso impede apagar a
pasta: achado A12). A verificação lê /proc/self/fd e só roda no Linux.
"""
import gc
import hashlib
import json
import os
import shutil
import sqlite3
import subprocess
import sys
import tempfile
import unittest

AQUI = os.path.dirname(os.path.abspath(__file__))
APP = os.path.dirname(AQUI)
sys.path.insert(0, APP)
sys.path.insert(0, AQUI)
import backup  # noqa: E402
import importacao as I  # noqa: E402
import importar  # noqa: E402
import planilhas_sinteticas as PS  # noqa: E402
import servicos as S  # noqa: E402
from fotos_teste import jpeg  # noqa: E402

D = "2026-10-15"
TR001, TR002 = "Zona01_PB01_Rede_TR001", "Zona01_PB01_Rede_TR002"


def sha(p):
    with open(p, "rb") as f:
        return hashlib.sha256(f.read()).hexdigest()


def arquivos_abertos_em(pasta):
    d = "/proc/self/fd"
    if not os.path.isdir(d):
        return None
    out = []
    for f in os.listdir(d):
        try:
            alvo = os.readlink(os.path.join(d, f))
        except OSError:
            continue
        if alvo.startswith(pasta):
            out.append(alvo)
    return sorted(out)


class Base(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.planilhas = os.path.join(self.tmp, "planilhas")
        self.cfg = PS.gerar(self.planilhas)
        self.dados = os.path.join(self.tmp, "dados")
        self.fotos = os.path.join(self.dados, "fotos")
        os.makedirs(self.fotos)
        self.db = os.path.join(self.dados, "obra.db")
        self.rels = os.path.join(self.tmp, "relatorios")
        self._conns = []

    def tearDown(self):
        for c in self._conns:
            c.close()
        abertos = arquivos_abertos_em(self.tmp)       # A12: nada pode ficar aberto (sem depender de gc)
        shutil.rmtree(self.tmp)
        if abertos:
            self.fail(f"arquivos ainda abertos ao fim do teste: {abertos}")

    def aplicar(self, **kw):
        return importar.aplicar(self.db, self.planilhas, cfg=self.cfg, dir_relatorios=self.rels, **kw)

    def regerar(self, **kw):
        self.cfg = PS.gerar(self.planilhas, **kw)
        return self.cfg

    def reprovar(self, *tipos):
        antes = sha(self.db) if os.path.exists(self.db) else None
        with self.assertRaises(importar.ImportacaoReprovada) as cm:
            self.aplicar()
        rel = cm.exception.relatorio
        achados = [b["tipo"] for b in rel["bloqueios"]]
        for t in tipos:
            self.assertIn(t, achados, rel["bloqueios"])
        self.assertEqual(sha(self.db) if os.path.exists(self.db) else None, antes)   # banco byte a byte igual (ou inexistente)
        self.assertTrue(os.path.exists(rel["arquivo_relatorio"]))    # relatório persistente da reprovação
        with open(rel["arquivo_relatorio"], encoding="utf-8") as f:
            self.assertEqual(json.load(f)["status"], "REPROVADA")
        return rel

    def conectar(self):
        c = S.conectar(self.db)
        self._conns.append(c)
        return c

    def q(self, sql, *a):
        c = sqlite3.connect(self.db)
        try:
            return c.execute(sql, a).fetchall()
        finally:
            c.close()

    def tid(self, ponto):
        return self.q("SELECT id FROM trechos WHERE ponto=? AND vigente=1", ponto)[0][0]


# ====================================================================== etapa 0
class TestBackup(Base):
    def test_criar_verificar_e_detectar_adulteracao(self):
        self.aplicar()
        S.registrar_foto(self.conectar(), "TRECHO", TR001, "EXECUCAO", jpeg(), data_hora=f"{D}T10:00:00", dir_fotos=self.fotos)
        pasta = backup.criar(self.db, os.path.join(self.dados, "backups"), self.fotos)
        self.assertEqual(backup.verificar(pasta), (True, []))
        nome = self.q("SELECT arquivo FROM fotos")[0][0]
        with open(os.path.join(pasta, "fotos", nome), "ab") as f:
            f.write(b"x")
        ok, erros = backup.verificar(pasta)
        self.assertFalse(ok)
        self.assertIn("hash", erros[0])

    def test_conteudo_adulterado_no_banco_copiado_e_detectado(self):
        self.aplicar()
        pasta = backup.criar(self.db, os.path.join(self.dados, "backups"), self.fotos)
        c = sqlite3.connect(os.path.join(pasta, "obra.db"))
        c.execute("UPDATE trechos SET extensao_m = extensao_m + 1 WHERE id='TR-00001'")   # mesma contagem de linhas
        c.commit()
        c.close()
        ok, erros = backup.verificar(pasta)
        self.assertFalse(ok)
        self.assertTrue(any("conteúdo" in e or "hash" in e for e in erros), erros)


class TestA03BackupCompletude(Base):
    """A03: foto referenciada no banco e ausente não pode dar backup verificado."""

    def _foto_fantasma(self):
        self.aplicar()
        c = sqlite3.connect(self.db)
        c.execute("INSERT INTO fotos(alvo_tipo,alvo_id,categoria,arquivo,data_hora,carimbo) "
                  "VALUES ('TRECHO','TR-00001','EXECUCAO','foto_ausente.jpg','x','c')")
        c.commit()
        c.close()

    def test_reproducao_auditoria_foto_ausente_reprova(self):
        self._foto_fantasma()
        with self.assertRaises(backup.BackupIncompleto) as cm:
            backup.criar(self.db, os.path.join(self.tmp, "bk"), os.path.join(self.tmp, "nao_existe"))
        pasta = cm.exception.pasta
        self.assertTrue(os.path.isdir(pasta))                    # a cópia fica guardada para recuperação
        ok, erros = backup.verificar(pasta)
        self.assertFalse(ok)
        self.assertIn("foto_ausente.jpg", erros[0])
        det = backup.verificar_detalhado(pasta)
        self.assertEqual((det["integro"], det["completo"], det["ausentes"]), (True, False, ["fotos/foto_ausente.jpg"]))
        with open(os.path.join(pasta, "manifesto.json"), encoding="utf-8") as f:
            self.assertEqual(json.load(f)["situacao"], "INCOMPLETO")

    def test_incompleto_so_com_permissao_explicita(self):
        self._foto_fantasma()
        pasta = backup.criar(self.db, os.path.join(self.tmp, "bk"), self.fotos, permitir_incompleto=True)
        self.assertFalse(backup.verificar(pasta)[0])
        with self.assertRaises(backup.BackupIncompleto):
            backup.restaurar(pasta, self.db, self.fotos)         # restaurar incompleto exige permissão explícita

    def test_importacao_nao_comeca_sem_backup_completo(self):
        self._foto_fantasma()
        antes = sha(self.db)
        with self.assertRaises(backup.BackupIncompleto):
            self.aplicar()
        self.assertEqual(sha(self.db), antes)

    def test_backup_completo_com_foto_real(self):
        self.aplicar()
        S.registrar_foto(self.conectar(), "TRECHO", TR001, "SINALIZACAO", jpeg(), data_hora=f"{D}T09:00:00", dir_fotos=self.fotos)
        pasta = backup.criar(self.db, os.path.join(self.tmp, "bk"), self.fotos)
        det = backup.verificar_detalhado(pasta)
        self.assertTrue(det["integro"] and det["completo"])

    def test_trava_de_manutencao_exclui_operacao_concorrente(self):
        self.aplicar()
        with backup.trava_manutencao(self.dados, "teste"):
            self.assertTrue(backup.em_manutencao(self.dados))     # o servidor recusa POST com a trava (503)
            with self.assertRaises(backup.EmManutencao):
                backup.criar(self.db, os.path.join(self.tmp, "bk"), self.fotos)
            with self.assertRaises(backup.EmManutencao):
                self.aplicar()
        self.assertFalse(backup.em_manutencao(self.dados))


class TestA04Restauracao(Base):
    """A04: a cópia do estado atual antes de restaurar precisa incluir o que está no WAL e as fotos."""

    def test_restaurar_com_wal_e_fotos(self):
        self.aplicar()
        c0 = self.conectar()
        S.registrar_foto(c0, "TRECHO", TR001, "EXECUCAO", jpeg(), data_hora=f"{D}T10:00:00", dir_fotos=self.fotos)
        f1 = self.q("SELECT arquivo FROM fotos")[0][0]
        pasta = backup.criar(self.db, os.path.join(self.dados, "backups"), self.fotos)
        # estado posterior ao backup, confirmado mas ainda no WAL (conexão aberta, sem checkpoint)
        c = sqlite3.connect(self.db, isolation_level=None)
        self._conns.append(c)
        c.execute("PRAGMA journal_mode=WAL")
        c.execute("PRAGMA wal_autocheckpoint=0")
        c.execute("CREATE TABLE marca(x)")
        c.execute("INSERT INTO marca VALUES ('antes')")
        c.execute("INSERT INTO marca VALUES ('posterior')")
        S.registrar_foto(c0, "TRECHO", TR002, "EXECUCAO", jpeg(), data_hora=f"{D}T11:00:00", dir_fotos=self.fotos)
        f2 = [r[0] for r in self.q("SELECT arquivo FROM fotos") if r[0] != f1][0]
        self.assertTrue(os.path.getsize(self.db + "-wal") > 0)
        os.remove(os.path.join(self.fotos, f1))                   # foto perdida depois do backup

        seguranca = backup.restaurar(pasta, self.db, self.fotos)

        # 1) o estado substituído foi guardado pela API do SQLite (WAL incluído), com as fotos existentes
        det = backup.verificar_detalhado(seguranca)
        self.assertEqual((det["integro"], det["ausentes"]), (True, ["fotos/" + f1]))   # f1 já tinha sumido: marcado INCOMPLETO
        g = sqlite3.connect(os.path.join(seguranca, "obra.db"))
        try:
            self.assertEqual([r[0] for r in g.execute("SELECT x FROM marca ORDER BY rowid")], ["antes", "posterior"])
            self.assertEqual(g.execute("SELECT COUNT(*) FROM fotos").fetchone()[0], 2)
        finally:
            g.close()
        self.assertTrue(os.path.exists(os.path.join(seguranca, "fotos", f2)))
        # 2) o banco voltou exatamente ao conteúdo do backup
        with open(os.path.join(pasta, "manifesto.json"), encoding="utf-8") as f:
            self.assertEqual(backup.conteudo(self.db), json.load(f)["conteudo_por_tabela"])
        self.assertEqual(self.q("SELECT COUNT(*) FROM sqlite_master WHERE name='marca'")[0][0], 0)
        # 3) foto do backup devolvida; foto posterior mantida
        self.assertTrue(os.path.exists(os.path.join(self.fotos, f1)))
        self.assertTrue(os.path.exists(os.path.join(self.fotos, f2)))

    def test_restaurar_volta_ao_estado_e_guarda_o_atual(self):
        self.aplicar()
        pasta = backup.criar(self.db, os.path.join(self.dados, "backups"), self.fotos)
        S.criar_medicao(self.conectar(), 1, "2026-10-01", "2026-10-31")
        seguranca = backup.restaurar(pasta, self.db, self.fotos)
        self.assertEqual(self.q("SELECT COUNT(*) FROM medicoes")[0][0], 0)
        self.assertTrue(backup.verificar(seguranca)[0])
        g = sqlite3.connect(os.path.join(seguranca, "obra.db"))
        self.assertEqual(g.execute("SELECT COUNT(*) FROM medicoes").fetchone()[0], 1)
        g.close()


class TestCargaNaoDestrutiva(Base):
    def test_simulacao_nao_altera_o_banco(self):
        self.aplicar()
        antes = sha(self.db)
        rel = importar.simular(self.db, self.planilhas, cfg=self.cfg, dir_relatorios=self.rels)
        self.assertEqual(rel["modo"], "SIMULACAO")
        self.assertEqual(sha(self.db), antes)

    def test_aplicar_faz_backup_verificado_antes(self):
        self.aplicar()
        rel = self.aplicar()
        self.assertTrue(rel["backup"] and backup.verificar(rel["backup"])[0])

    def test_erro_desfaz_tudo_e_preserva_backup(self):
        self.aplicar()
        S.criar_medicao(self.conectar(), 1, "2026-10-01", "2026-10-31")
        antes = sha(self.db)
        from openpyxl import Workbook
        wb = Workbook()
        wb.active.title = "CONEXÕES"
        wb.save(os.path.join(self.planilhas, PS.ARQ["conexoes"]))     # 06.2 corrompida (sem cabeçalho)
        with self.assertRaises(RuntimeError):
            self.aplicar()
        self.assertEqual(sha(self.db), antes)
        self.assertEqual(self.q("SELECT COUNT(*) FROM medicoes")[0][0], 1)
        self.assertTrue(os.listdir(os.path.join(self.dados, "backups")))
        self.assertTrue(any("_falhou_" in n for n in os.listdir(self.rels)))   # relatório da falha persistido

    def test_seed_antigo_nao_apaga_mais(self):
        r = subprocess.run([sys.executable, os.path.join(APP, "seed.py")], capture_output=True)
        self.assertNotEqual(r.returncode, 0)
        self.assertIn("importar.py", r.stderr.decode("utf-8"))      # A12: codificação definida pelo programa


# ====================================================================== A01 — conferências impeditivas
class TestA01ConferenciasAntesDoCommit(Base):
    def test_reproducao_item_retirado_com_controles_antigos(self):
        self.aplicar()
        cfg_antigo = self.cfg
        self.regerar(conexoes=[l for l in PS.CONEXOES if l[1] != "TUBO"])       # um item a menos
        self.cfg = cfg_antigo                                                    # controles da revisão anterior
        rel = self.reprovar("SEM_CONTROLES_APROVADOS")
        self.assertEqual(self.q("SELECT COUNT(*) FROM itens_projeto WHERE vigente=1")[0][0], len(PS.CONEXOES))
        self.assertIsNotNone(rel["auditoria"]["modelo_controles"])               # modelo para aprovar a nova revisão

    def test_controle_da_revisao_com_contagem_errada_reprova(self):
        self.aplicar()
        cfg = self.regerar(conexoes=[l for l in PS.CONEXOES if l[1] != "TUBO"])
        cfg["controles"][0]["esperado"]["itens_projeto"] = len(PS.CONEXOES)      # valor esperado errado
        self.reprovar("CONTROLE_DIVERGENTE")

    def test_reproducao_preco_999_sem_atualizar_o_total(self):
        self.aplicar()
        import openpyxl
        p = os.path.join(self.planilhas, PS.ARQ["bm"])
        wb = openpyxl.load_workbook(p)
        wb["BM MORRO DA CRUZ - VERTICAL"].cell(row=6, column=20, value=999)
        wb.save(p)
        wb.close()
        self.cfg["controles"] = [PS.conjunto_controles(self.planilhas, self.cfg["controles"][0]["esperado"])]
        self.reprovar("TOTALIZACAO_BM_DIVERGE_DAS_CELULAS")
        self.assertEqual(self.q("SELECT preco_unitario FROM bm_itens WHERE item='1.2_1.14.27'"), [(10.0,)])

    def test_nova_revisao_legitima_com_controles_proprios(self):
        self.aplicar()
        self.regerar(conexoes=[l for l in PS.CONEXOES if l[1] != "TUBO"])
        rel = self.aplicar()
        self.assertEqual(rel["status"], "CONCLUIDA")
        self.assertEqual(self.q("SELECT vigente FROM itens_projeto WHERE descricao='TUBO'"), [(0,)])


# ====================================================================== controles: completos e com aprovação efetiva
class TestControlesAprovacao(Base):
    """Controle vazio, incompleto ou sem aprovação efetiva não pode virar 'conferência confirmada'."""

    def ctl(self, **mudancas):
        c = self.cfg["controles"][0]
        for k, v in mudancas.items():
            if v is None:
                c.pop(k, None)
            else:
                c[k] = v
        return c

    def test_sonda_auditoria_esperado_vazio_e_sem_aprovador(self):
        self.ctl(esperado={}, aprovacao=None)
        rel = self.reprovar("CONTROLE_INCOMPLETO", "CONTROLE_SEM_APROVACAO")
        self.assertFalse(rel["auditoria"]["contagens_ok"])
        self.assertFalse(os.path.exists(self.db) and self.q("SELECT name FROM sqlite_master WHERE name='trechos'"))

    def test_sonda_auditoria_sem_aprovador(self):
        self.ctl(aprovacao=None)
        rel = self.reprovar("CONTROLE_SEM_APROVACAO")
        self.assertEqual(rel["auditoria"]["controle_situacao"], "SEM_APROVACAO")

    def test_aprovado_por_solto_ou_provisorio_nao_vale(self):
        self.ctl(aprovacao=None, aprovado_por="Fulano")
        self.reprovar("CONTROLE_SEM_APROVACAO")
        self.ctl(aprovacao={"status": "APROVADO", "por": "A CONFIRMAR pelo usuário", "em": "2026-10-16"})
        self.reprovar("CONTROLE_SEM_APROVACAO")
        self.ctl(aprovacao={"status": "PENDENTE", "por": "Fulano", "em": "2026-10-16"})
        self.reprovar("CONTROLE_SEM_APROVACAO")
        self.ctl(aprovacao={"status": "APROVADO", "por": "Fulano", "em": "ontem"})
        self.reprovar("CONTROLE_SEM_APROVACAO")

    def test_campo_obrigatorio_ausente(self):
        self.ctl()["esperado"].pop("trechos")
        rel = self.reprovar("CONTROLE_INCOMPLETO")
        b = next(x for x in rel["bloqueios"] if x["tipo"] == "CONTROLE_INCOMPLETO")
        self.assertIn("trechos", b["detalhe"])                                   # relatório legível: diz o que falta

    def test_tipo_invalido_origem_ausente_e_conferencia_desconhecida(self):
        c = self.ctl(origem_dos_valores="PREENCHER")
        c["esperado"]["trechos"] = "3"
        c["esperado"]["qualquer_coisa"] = 1
        b = next(x for x in self.reprovar("CONTROLE_INCOMPLETO")["bloqueios"] if x["tipo"] == "CONTROLE_INCOMPLETO")
        for trecho in ("esperado.trechos", "origem_dos_valores", "qualquer_coisa"):
            self.assertIn(trecho, b["detalhe"])

    def test_perfil_de_teste_nao_vale_no_perfil_real(self):
        self.cfg["perfil"] = "REAL"
        self.reprovar("CONTROLE_INCOMPLETO")

    def test_aprovacao_valida_conclui_e_registra_quem_aprovou(self):
        self.ctl(aprovacao={"status": "APROVADO", "por": "Eng. Responsável", "em": "2026-10-16"})
        rel = self.aplicar()
        a = rel["auditoria"]
        self.assertEqual((rel["status"], a["controle_situacao"], a["contagens_ok"]), ("CONCLUIDA", "APROVADO", True))
        self.assertEqual(a["controle_aprovacao"]["por"], "Eng. Responsável")

    def test_simulacao_continua_disponivel_sem_aprovacao(self):
        self.aplicar()
        antes = sha(self.db)
        self.ctl(aprovacao=None)
        rel = importar.simular(self.db, self.planilhas, cfg=self.cfg, dir_relatorios=self.rels)
        self.assertEqual(rel["status"], "SIMULADA_SEM_CONTROLE_APROVADO")
        self.assertFalse(rel["auditoria"]["contagens_ok"])                       # sem falsa confirmação
        self.assertTrue(all(c["ok"] for c in rel["auditoria"]["contagens"].values()))
        self.assertEqual([x["tipo"] for x in rel["auditoria"]["avisos_controle"]], ["CONTROLE_SEM_APROVACAO"])
        self.assertEqual(rel["bloqueios"], [])
        self.assertEqual(sha(self.db), antes)
        self.assertIsNone(self.cfg["controles"][0].get("aprovacao"))             # nada preenchido em lugar da pessoa

    def test_simulacao_sem_conjunto_traz_modelo_com_aprovacao_pendente(self):
        self.cfg["controles"] = []
        rel = importar.simular(self.db, self.planilhas, cfg=self.cfg, dir_relatorios=self.rels)
        m = rel["auditoria"]["modelo_controles"]
        self.assertEqual(m["aprovacao"], {"status": "PENDENTE", "por": None, "em": None})
        self.assertEqual(I.validar_controle(m, "TESTE")[1] != [], True)          # o modelo sozinho não aprova nada

    def test_conjunto_real_versionado_esta_completo_e_sem_aprovacao_inventada(self):
        real = I.carregar_config()
        self.assertEqual(real["perfil"], "REAL")
        for c in real["controles"]:
            estrutura, _ = I.validar_controle(c, "REAL")
            self.assertEqual(estrutura, [], c["id"])
            ap = c.get("aprovacao") or {}
            self.assertTrue(ap.get("status") != "APROVADO" or "teste" not in str(ap.get("por")).lower())


class TestServidorParado(Base):
    """Importação e restauração exigem o servidor parado (limitação de concorrência desta versão)."""

    def test_aplicar_e_restaurar_recusam_com_servidor_respondendo(self):
        import socket
        self.aplicar()
        pasta = backup.criar(self.db, os.path.join(self.dados, "backups"), self.fotos)
        srv = socket.socket()
        srv.bind(("127.0.0.1", 0))
        srv.listen()
        try:
            marca = backup.marcar_servidor(self.dados, "127.0.0.1", srv.getsockname()[1])
            antes = sha(self.db)
            with self.assertRaises(backup.ServidorAtivo):
                self.aplicar()
            with self.assertRaises(backup.ServidorAtivo):
                backup.restaurar(pasta, self.db, self.fotos)
            self.assertEqual(sha(self.db), antes)
            importar.simular(self.db, self.planilhas, cfg=self.cfg, dir_relatorios=self.rels)   # simulação liberada
        finally:
            srv.close()
        self.assertEqual(self.aplicar()["status"], "SEM_ALTERACOES")              # marca órfã não bloqueia
        os.remove(marca)


# ====================================================================== A02 — identidade física dos trechos
class TestA02IdentidadeTrechos(Base):
    def _checkin_tr001(self):
        self.aplicar()
        S.checkin(self.conectar(), D, "Zona01", "TRECHO", TR001)
        return self.tid(TR001)

    def test_reproducao_nomes_trocados_historico_fica_na_geometria(self):
        tid = self._checkin_tr001()
        t = [list(x) for x in PS.TRECHOS]
        t[0][0], t[1][0] = t[1][0], t[0][0]
        self.regerar(trechos=[tuple(x) for x in t])
        rel = self.aplicar()
        ck = self.q("SELECT c.alvo_id, t.ponto, t.handle, t.utm_ini FROM checkins c JOIN trechos t ON t.id=c.alvo_id")
        self.assertEqual(ck, [(tid, TR002, "5FD12", "[10.0, 10.0]")])           # mesma linha física, novo nome
        self.assertEqual(self.q("SELECT handle FROM trechos WHERE ponto=? AND vigente=1", TR001), [("5FD13",)])
        renomeados = [p for p in rel["pendencias"] if p["tipo"] == "TRECHO_RENOMEADO"]
        self.assertEqual(len(renomeados), 2)
        self.assertEqual(sorted(r[0] for r in self.q("SELECT nome FROM nomes_historico WHERE entidade_id=?", tid)),
                         [TR001, TR002])

    def test_divisao_de_trecho_com_historico_exige_de_para(self):
        tid = self._checkin_tr001()
        div = [("Zona01_PB01_Rede_TR001A", "Zona01", 63, 40.0, "VCA", "5.8.44", "5FD12", (10, 10), (15, 15), "a", "b"),
               ("Zona01_PB01_Rede_TR001B", "Zona01", 63, 36.55, "VCA", "5.8.44", "5FD12", (15, 15), (20, 20), "b", "c")]
        self.regerar(trechos=div + list(PS.TRECHOS[1:]))
        rel = self.reprovar("TRECHO_COM_HISTORICO_SEM_CORRESPONDENCIA")
        b = next(x for x in rel["bloqueios"] if x["tipo"] == "TRECHO_COM_HISTORICO_SEM_CORRESPONDENCIA")
        self.assertIn("TR001A", b["detalhe"])                                     # sucessores apontados, não escolhidos
        self.assertEqual(self.q("SELECT ponto FROM trechos WHERE id=?", tid), [(TR001,)])
        # conciliação explícita: a parte A herda o ID (e o histórico); a parte B nasce com ID novo
        self.cfg["conciliacoes"] = [{"tipo": "TRECHO", "ponto": "Zona01_PB01_Rede_TR001A", "handle": "5FD12",
                                     "usar_id": tid, "responsavel": "Eng. teste", "data": "2026-10-16", "motivo": "divisão"}]
        self.assertEqual(self.aplicar()["status"], "CONCLUIDA")
        self.assertEqual(self.q("SELECT t.ponto FROM checkins c JOIN trechos t ON t.id=c.alvo_id"),
                         [("Zona01_PB01_Rede_TR001A",)])
        self.assertNotEqual(self.tid("Zona01_PB01_Rede_TR001B"), tid)

    def test_uniao_de_trechos_com_historico_exige_de_para(self):
        self._checkin_tr001()
        uniao = [("Zona01_PB01_Rede_TR001", "Zona01", 63, 151.62, "VCA", "5.8.44", "5FD12", (10, 10), (30, 30), "a", "c")]
        self.regerar(trechos=uniao + [PS.TRECHOS[2]])
        self.reprovar("TRECHO_COM_HISTORICO_SEM_CORRESPONDENCIA")

    def test_mesmo_nome_outra_geometria_sem_historico_ganha_id_novo(self):
        self.aplicar()
        antigo = self.tid(TR002)                                                   # sem histórico
        t = list(PS.TRECHOS)
        t[1] = (TR002, "Zona01", 63, 80.0, "VCA", "5.8.44", "7XXXX", (21, 21), (31, 31), "x", "y")
        self.regerar(trechos=t)
        rel = self.aplicar()
        self.assertNotEqual(self.tid(TR002), antigo)                               # nome não decide identidade
        self.assertEqual(self.q("SELECT vigente FROM trechos WHERE id=?", antigo), [(0,)])
        self.assertIn("TRECHO_SUBSTITUIDO", [p["tipo"] for p in rel["pendencias"]])

    def test_trecho_renomeado_mantem_id_e_historico(self):
        self.aplicar()
        tid = self.tid(TR002)
        self.regerar(trechos=[("Zona01_PB01_Rede_TR099",) + t[1:] if t[0] == TR002 else t for t in PS.TRECHOS])
        self.aplicar()
        self.assertEqual(self.q("SELECT ponto FROM trechos WHERE id=?", tid), [("Zona01_PB01_Rede_TR099",)])


# ====================================================================== A05 — escopo das decisões
class TestA05Decisoes(Base):
    def test_decisao_df05_so_no_escopo_explicito(self):
        self.aplicar()
        c = self.conectar()
        d05 = S.detalhe(c, "ATIVIDADE", "Zona11_PB15_Rede_DF-05")
        ligados = {a["texto"]: a["decisoes"] for a in d05["avisos"]}
        self.assertEqual(ligados["P1 — DF-05 contada em duas pranchas (16 e 19)."], ["DEC-0001"])
        self.assertEqual(ligados["Tirar a DF-05 do resumo e do quadro de uma das pranchas."], ["DEC-0001"])
        # reprodução da auditoria: "possível duplicidade" da DF-32 não recebe a decisão encerrada da DF-05
        d32 = S.detalhe(c, "ATIVIDADE", "Zona01_PB02_Rede_DF-32")
        self.assertTrue(d32["avisos"])
        self.assertTrue(all(a["decisoes"] == [] for a in d32["avisos"]))
        self.assertNotIn("DEC-0001", [d["id"] for d in d32["decisoes"]])
        # texto que cita a DF-05 em OUTRO elemento: só sugestão para revisão, nunca decisão aplicada
        d33 = S.detalhe(c, "ATIVIDADE", "Zona01_PB01_Rede_DF-33")
        aviso = next(a for a in d33["avisos"] if "DF-05" in a["texto"])
        self.assertEqual((aviso["decisoes"], aviso["sugestoes"]), ([], ["DEC-0001"]))
        self.assertEqual([d["id"] for d in d33["decisoes_sugeridas"]], ["DEC-0001"])
        self.assertEqual(d33["decisoes"], [])

    def test_reimportacao_nao_altera_decisao_nem_aviso(self):
        self.aplicar()
        c = sqlite3.connect(self.db)
        c.execute("UPDATE decisoes SET decisao='texto revisto pelo usuário' WHERE id='DEC-0001'")
        c.commit()
        c.close()
        self.aplicar()
        self.assertEqual(self.q("SELECT decisao FROM decisoes WHERE id='DEC-0001'"), [("texto revisto pelo usuário",)])
        self.assertEqual(self.q("SELECT COUNT(*) FROM decisao_vinculos WHERE decisao_id='DEC-0001'"), [(2,)])
        self.assertEqual(self.q("SELECT DISTINCT vigente FROM avisos_origem WHERE texto LIKE '%DF-05%'"), [(1,)])

    def test_decisao_por_objeto_sem_posicao(self):
        self.aplicar()
        d = S.detalhe(self.conectar(), "ATIVIDADE", "Zona11_PBND_CH_MV-11")
        self.assertEqual([x["id"] for x in d["decisoes"]], ["DEC-0004"])


# ====================================================================== A06 — evidências ausentes/incompatíveis
class TestA06Evidencias(Base):
    def test_reproducao_sem_evidencias_preserva_zonas(self):
        self.aplicar()
        antes = self.q("SELECT id, zona, status_posicao, id_auditoria FROM elementos ORDER BY id")
        self.assertTrue(all(r[1] for r in antes))
        self.regerar(evidencias=False)                                # controle aprovado SEM evidências
        rel = self.aplicar()
        self.assertEqual(self.q("SELECT id, zona, status_posicao, id_auditoria FROM elementos ORDER BY id"), antes)
        self.assertIn("EVIDENCIAS_AUSENTES", [p["tipo"] for p in rel["pendencias"]])

    def test_sem_evidencias_com_controle_que_as_exige_reprova(self):
        self.aplicar()
        os.remove(os.path.join(self.planilhas, PS.ARQ["evidencias"]))
        self.reprovar("SEM_CONTROLES_APROVADOS")

    def test_evidencias_de_outra_revisao_reprova(self):
        self.aplicar()
        antes = self.q("SELECT zona FROM elementos ORDER BY id")
        p = os.path.join(self.planilhas, PS.ARQ["evidencias"])
        outra = os.path.join(self.planilhas, "EVIDENCIAS_SINTETICAS_REV07_06.2_e_REV01_06.1.xlsx")
        os.rename(p, outra)
        self.cfg["arquivos"]["evidencias"] = os.path.basename(outra)
        self.reprovar("EVIDENCIAS_REVISAO_INCOMPATIVEL")
        self.assertEqual(self.q("SELECT zona FROM elementos ORDER BY id"), antes)


# ====================================================================== A07 — identidade provisória
class TestA07IdentidadeProvisoria(Base):
    MV = "Zona11_PBND_CH_MV-11"

    def test_reproducao_mv11_recebe_handle_sem_perder_id_e_historico(self):
        self.aplicar()
        el, at = self.q("SELECT e.id, a.id FROM elementos e JOIN atividades a ON a.elemento_id=e.id WHERE e.numero='MV-11'")[0]
        S.checkin(self.conectar(), D, "CH", "ATIVIDADE", self.MV)
        self.regerar(conexoes=[l[:6] + ("7AB01",) + l[7:] if "MV-11" in l[0] else l for l in PS.CONEXOES])
        rel = self.aplicar()
        self.assertEqual(rel["status"], "CONCLUIDA")
        self.assertEqual(self.q("SELECT id, chave_natural, vigente FROM elementos WHERE numero='MV-11'"), [(el, "H:7AB01", 1)])
        self.assertEqual(self.q("SELECT alvo_id FROM checkins"), [(at,)])
        self.assertIn("IDENTIDADE_PROMOVIDA", [p["tipo"] for p in rel["pendencias"]])
        self.assertIsNotNone(self.q("SELECT lat FROM elementos WHERE numero='MV-11'")[0][0])   # posição veio da fonte

    def test_vtf_sem_handle_continua_sem_coordenada(self):
        vtf = ("Zona03_PBND_CH_VTF-12", "VENTOSA", "50", "PN10", 1, "3.1.1", None, None, None, None,
               "CONTRATUAL CONFIRMADO", None, None)
        self.regerar(conexoes=PS.CONEXOES + [vtf])
        self.aplicar()
        self.assertEqual(self.q("SELECT chave_natural, lat, lon, utm_e FROM elementos WHERE numero='VTF-12'"),
                         [("N:VTF-12", None, None, None)])
        est = S.estado_mapa(self.conectar())
        self.assertEqual(sorted(e["numero"] for e in est["sem_posicao"]), ["MV-11", "VTF-12"])


# ====================================================================== A08 — identidade dos itens
class TestA08IdentidadeItens(Base):
    LUVA = ("Zona01_PB01_Rede_NÓ1.8", "LUVA X", "63", "PN10", 1, "1.14.27", "5FD4F", None, None, None,
            "CONTRATUAL CONFIRMADO", None, None)

    def _resinas(self, ordem):
        import openpyxl
        p = os.path.join(self.planilhas, PS.ARQ["conexoes"])
        wb = openpyxl.load_workbook(p)
        ws = wb["CONEXÕES"]
        for r, rs in zip([r for r in range(6, ws.max_row + 1) if ws.cell(r, 2).value == "LUVA X"], ordem):
            ws.cell(r, 4, rs)
        wb.save(p)
        wb.close()
        self.cfg["controles"] = [PS.conjunto_controles(self.planilhas, self.cfg["controles"][0]["esperado"])]

    def test_reproducao_resinas_reordenadas_nao_trocam_identidade(self):
        self.regerar(conexoes=PS.CONEXOES + [self.LUVA, self.LUVA])
        self._resinas(["PE80", "PE100"])
        self.aplicar()
        antes = dict(self.q("SELECT resina, id FROM itens_projeto WHERE descricao='LUVA X'"))
        self._resinas(["PE100", "PE80"])
        rel = self.aplicar()
        self.assertEqual(dict(self.q("SELECT resina, id FROM itens_projeto WHERE descricao='LUVA X'")), antes)
        self.assertEqual(rel["alteracoes"].get("item_projeto", {}).get("alterados", 0), 0)

    def test_linha_inserida_nao_troca_identidade(self):
        self.aplicar()
        antes = dict(self.q("SELECT chave_ocorrencia || '@' || atividade_id, id FROM itens_projeto"))
        nova = ("Zona01_PB01_Rede_NÓ1.8", "CURVA 90º", "63", "PN10", 1, "1.14.27", "5FD4F", None, None, None,
                "CONTRATUAL CONFIRMADO", None, None)
        self.regerar(conexoes=[nova] + PS.CONEXOES)                    # todas as linhas descem uma posição
        rel = self.aplicar()
        depois = dict(self.q("SELECT chave_ocorrencia || '@' || atividade_id, id FROM itens_projeto"))
        for k, v in antes.items():
            self.assertEqual(depois[k], v)
        self.assertEqual(rel["alteracoes"]["item_projeto"]["novos"], 1)
        self.assertEqual(rel["alteracoes"]["item_projeto"].get("alterados", 0), 0)

    def test_duplicatas_identicas_reordenadas_sem_alteracao(self):
        luva2 = self.LUVA[:4] + (2,) + self.LUVA[5:]
        self.regerar(conexoes=PS.CONEXOES + [self.LUVA, luva2])
        self.aplicar()
        antes = dict(self.q("SELECT quantidade, id FROM itens_projeto WHERE descricao='LUVA X'"))
        self.regerar(conexoes=PS.CONEXOES + [luva2, self.LUVA])
        rel = self.aplicar()
        self.assertEqual(rel["status"], "SEM_ALTERACOES")
        self.assertEqual(dict(self.q("SELECT quantidade, id FROM itens_projeto WHERE descricao='LUVA X'")), antes)

    def test_correspondencia_ambigua_vira_pendencia_sem_transferencia(self):
        luva2 = self.LUVA[:4] + (2,) + self.LUVA[5:]
        self.regerar(conexoes=PS.CONEXOES + [self.LUVA, luva2])
        self.aplicar()
        ids_antes = {r[0] for r in self.q("SELECT id FROM itens_projeto WHERE descricao='LUVA X'")}
        self.regerar(conexoes=PS.CONEXOES + [self.LUVA[:4] + (3,) + self.LUVA[5:], self.LUVA[:4] + (4,) + self.LUVA[5:]])
        rel = self.aplicar()
        self.assertIn("ITEM_SEM_CORRESPONDENCIA_INEQUIVOCA", [p["tipo"] for p in rel["pendencias"]])
        self.assertEqual(self.q("SELECT DISTINCT vigente FROM itens_projeto WHERE id IN (?,?)", *sorted(ids_antes)), [(0,)])

    def test_quantidade_alterada_registra_antes_e_depois(self):
        self.aplicar()
        self.regerar(conexoes=[l[:4] + (5,) + l[5:] if l[1] == "TUBO" else l for l in PS.CONEXOES])
        rel = self.aplicar()
        self.assertEqual(rel["alteracoes"]["item_projeto"]["alterados"], 1)
        self.assertEqual(self.q("SELECT antes, depois FROM alteracoes WHERE campo='quantidade'"), [("0.71", "5.0")])


# ====================================================================== A09 — medição fechada congelada
class TestA09MedicaoCongelada(Base):
    def _medicao(self, c, numero, ini, fim):
        mid = S.criar_medicao(c, numero, ini, fim)
        bid = c.execute("SELECT id FROM bm_itens WHERE item='2.2_5.2.1'").fetchone()[0]
        c.execute("INSERT INTO medicao_lancamentos(medicao_id,gaveta,servico,bm_item_id,unidade,quantidade,status) "
                  "VALUES (?,3,'CADASTRO',?,'M',2.5,'PROPOSTO')", (mid, bid))
        return mid

    def test_reproducao_preco_dobrado_nao_muda_medicao_fechada(self):
        self.aplicar()
        c = self.conectar()
        m1 = self._medicao(c, 1, "2026-10-01", "2026-10-31")
        S.fechar_medicao(c, m1)
        m2 = self._medicao(c, 2, "2026-11-01", "2026-11-30")              # aberta
        v1, v2 = S.resumo_medicao(c, m1)["valor_pi"], S.resumo_medicao(c, m2)["valor_pi"]
        self.regerar(precos={"5.2.1": 2.61 * 2})                           # BM coerente, preço dobrado
        self.assertTrue(self.aplicar()["auditoria"]["totalizacao_bm"]["ok"])
        r1 = S.resumo_medicao(c, m1)
        self.assertEqual(r1["valor_pi"], v1)                               # fechada: congelada
        self.assertTrue(r1["congelada"])
        self.assertIn("Mantido o valor congelado", r1["alerta_recalculo"])
        self.assertGreater(S.resumo_medicao(c, m2)["valor_pi"], v2)        # aberta: usa o cadastro atual
        with self.assertRaises(S.RegraViolada):
            S.fechar_medicao(c, m1)                                        # não refecha

    def test_migracao_v2_congela_medicao_fechada_antes_do_novo_bm(self):
        c = sqlite3.connect(self.db)
        with open(os.path.join(AQUI, "fixtures", "schema_v2.sql"), encoding="utf-8") as f:
            c.executescript(f.read())
        c.execute("INSERT INTO schema_versao VALUES (2, 'x', 'v2')")
        c.executemany("INSERT INTO parametros VALUES (?,?,'v2')", [("bdi_material", "0.1528"), ("bdi_servico", "0.2148"),
                                                                    ("coef_k", "-0.1404")])
        c.execute("INSERT INTO bm_itens(id,planilha,item,item_bm,subgrupo,classificacao,descricao,unidade,quantidade,"
                  "preco_unitario) VALUES (1,'VERTICAL','2.2_5.2.1','5.2.1','2.2','SERVIÇO','cadastro','M',111219.2,2.61)")
        c.execute("INSERT INTO medicoes(id,numero,tipo,periodo_ini,periodo_fim,status) "
                  "VALUES (1,1,'TESTE','2026-10-01','2026-10-31','FECHADA')")
        c.execute("INSERT INTO medicao_lancamentos(medicao_id,gaveta,servico,bm_item_id,unidade,quantidade,status) "
                  "VALUES (1,3,'CADASTRO',1,'M',2.5,'PROPOSTO')")
        # item importado no v2, com a chave antiga (descrição|DE|PN|nº)
        c.execute("INSERT INTO equipes VALUES ('Zona01','EQUIPE ZONA 01',NULL)")
        c.execute("INSERT INTO elementos(id,chave_natural,numero,tipo,handle) VALUES ('EL-0001','H:5FD4F','NÓ1.8','NÓ','5FD4F')")
        c.execute("INSERT INTO atividades(id,elemento_id,tipo_frente,nomenclatura,frente_id) "
                  "VALUES ('AT-0001','EL-0001','Rede','Zona01_PB01_Rede_NÓ1.8','Zona01')")
        c.execute("INSERT INTO itens_projeto(id,atividade_id,chave_ocorrencia,descricao,pressao,resina,tipo_conexao,solda,"
                  "diametro,material,quantidade) VALUES ('IP-000001','AT-0001','LUVA ELETROFUSÃO|63|PN10/16|1',"
                  "'LUVA ELETROFUSÃO','PN10/16','PE100','LUVA','ELETROFUSÃO','63','PEAD',1)")
        c.commit()
        c.close()
        self.regerar(precos={"5.2.1": 5.22})
        rel = self.aplicar()
        self.assertEqual((rel["esquema_antes"], rel["medicoes_congeladas_na_migracao"]), (2, [1]))
        self.assertEqual(self.q("SELECT preco_unitario FROM bm_itens WHERE id=1"), [(5.22,)])
        r = S.resumo_medicao(self.conectar(), 1)
        self.assertEqual(r["totalizacao"]["servico"], 6.53)                 # ROUND(2,5 × 2,61; 2) — preço do fechamento
        self.assertEqual(self.q("SELECT preco_congelado FROM medicao_lancamentos"), [(2.61,)])
        self.assertEqual(self.q("SELECT id FROM itens_projeto WHERE descricao='LUVA ELETROFUSÃO' AND atividade_id='AT-0001'"),
                         [("IP-000001",)])                                 # identidade preservada na migração


# ====================================================================== A10 — objeto retirado
class TestA10ObjetoRetirado(Base):
    def test_reproducao_trecho_retirado_nao_aceita_checkin(self):
        self.aplicar()
        tid = self.tid("Zona02_PB04_Rede_TR036")
        self.regerar(trechos=PS.TRECHOS[:2])
        self.aplicar()
        c = self.conectar()
        self.assertEqual(self.q("SELECT vigente FROM trechos WHERE id=?", tid), [(0,)])
        with self.assertRaisesRegex(S.RegraViolada, "retirado"):
            S.checkin(c, "2026-10-16", "Zona02", "TRECHO", tid)
        self.assertEqual(S.detalhe(c, "TRECHO", tid)["id"], tid)          # consulta histórica continua possível
        for f in (lambda: S.registrar_foto(c, "TRECHO", tid, "EXECUCAO", jpeg(), dir_fotos=self.fotos),
                  lambda: S.ok_fiscalizacao(c, "TRECHO", tid, "x", D)):
            with self.assertRaisesRegex(S.RegraViolada, "retirado"):
                f()

    def test_checkin_aberto_em_trecho_retirado_fica_para_tratamento_explicito(self):
        self.aplicar()
        c = self.conectar()
        tid = self.tid(TR002)
        ck = S.checkin(c, D, "Zona01", "TRECHO", TR002)
        S.registrar_clima(c, D, "Zona01", "MANHA", "SOL")
        S.registrar_clima(c, D, "Zona01", "TARDE", "SOL")
        S.apontar(c, D, "Zona01", "TRECHO", TR002, intervalos_novos=[{"etapa": "ASSENTADO", "ini_m": 0, "fim_m": 10}])
        self.regerar(trechos=[PS.TRECHOS[0], PS.TRECHOS[2]])
        self.reprovar("TRECHO_COM_HISTORICO_SEM_CORRESPONDENCIA")         # retirada com histórico exige confirmação
        self.cfg["conciliacoes"] = [{"tipo": "TRECHO", "retirar_id": tid, "responsavel": "Eng. teste", "data": "2026-10-16",
                                     "motivo": "trecho excluído do projeto"}]
        self.aplicar()
        est = S.estado_mapa(c, data=D)
        self.assertNotIn(TR002, est["status"])                             # fora do estado corrente do mapa
        self.assertEqual(est["pins"], [])
        (ret,) = est["retirados_com_historico"]                            # um registro por objeto, com as situações
        self.assertEqual(ret["alvo_id"], tid)
        self.assertEqual(len(ret["situacoes"]), 2)                          # execução sem OK + check-in aberto
        self.assertIn("check-in", ret["historico"])
        S.checkout(c, ck, [])                                              # movimento antigo pode ser encerrado
        self.assertEqual(self.q("SELECT COUNT(*) FROM apontamentos WHERE alvo_id=?", tid), [(1,)])   # nada apagado


# ====================================================================== A11 — MND no detalhe
class TestA11MND(Base):
    def test_mnd_detalhe_e_calculo_com_a_mesma_regra(self):
        self.aplicar()
        c = self.conectar()
        d = S.detalhe(c, "TRECHO", "Zona02_PB04_Rede_TR036")
        self.assertEqual((d["metodologia"], d["profundidade_m"]), ("MND", 0.0))
        self.assertIsNone(d["perfil_bloqueado"])
        self.assertEqual(S.previsto(c, "TRECHO", "Zona02_PB04_Rede_TR036", 1.0),
                         [{"servico": "ASSENTAMENTO", "unidade": "M", "previsto": 1.0}])
        self.assertTrue(d["mnd"])                                          # enquadramento financeiro segue pendente


# ====================================================================== idempotência, unidades, migração v1
class TestIdempotencia(Base):
    def test_reimportar_mesmos_arquivos_sem_alteracoes(self):
        r1 = self.aplicar()
        self.assertEqual(r1["status"], "CONCLUIDA")
        self.assertTrue(r1["auditoria"]["contagens_ok"], r1["auditoria"]["contagens"])
        self.assertTrue(r1["auditoria"]["totalizacao_bm"]["ok"])
        ids = self.q("SELECT id FROM itens_projeto ORDER BY id")
        antes = backup.conteudo(self.db)
        r2 = self.aplicar()
        self.assertEqual(r2["status"], "SEM_ALTERACOES")
        depois = backup.conteudo(self.db)
        self.assertEqual(sorted(t for t in antes if antes[t] != depois[t]), ["importacoes"])   # conteúdo, não só contagem
        self.assertEqual(r2["alteracoes_registradas"], 0)
        self.assertEqual(self.q("SELECT id FROM itens_projeto ORDER BY id"), ids)
        self.assertEqual(self.q("SELECT COUNT(*) FROM fontes_arquivo")[0][0], 5)

    def test_proveniencia_com_hash_aba_linha(self):
        self.aplicar()
        p = self.q("""SELECT p.aba, p.linha, f.sha256, f.revisao FROM proveniencia p JOIN fontes_arquivo f ON f.id=p.fonte_id
                      JOIN itens_projeto i ON i.id=p.entidade_id WHERE p.entidade='item_projeto' AND i.descricao='TUBO'""")
        self.assertEqual(len(p), 1)
        self.assertEqual(p[0][:2], ("CONEXÕES", 11))
        self.assertEqual(p[0][2], sha(os.path.join(self.planilhas, PS.ARQ["conexoes"])))
        self.assertEqual(p[0][3], "REV08")

    def test_elemento_fisico_unico_com_duas_atividades(self):
        self.aplicar()
        el = self.q("SELECT id, chave_natural FROM elementos WHERE numero='DF-33'")
        self.assertEqual(len(el), 1)
        self.assertEqual(el[0][1], "H:5FDEA")
        self.assertEqual(self.q("SELECT tipo_frente, frente_id FROM atividades WHERE elemento_id=? ORDER BY 1", el[0][0]),
                         [("CH", "CH"), ("Rede", "Zona01")])

    def test_item_retirado_fica_nao_vigente(self):
        self.aplicar()
        self.regerar(conexoes=[l for l in PS.CONEXOES if l[1] != "TUBO"])
        self.aplicar()
        self.assertEqual(self.q("SELECT vigente FROM itens_projeto WHERE descricao='TUBO'"), [(0,)])


class TestUnidadesEAvisos(Base):
    def test_unidade_so_com_fonte(self):
        self.aplicar()
        un = dict(self.q("SELECT descricao, unidade FROM itens_projeto"))
        self.assertEqual(un["TUBO"], "M")
        self.assertEqual(un["TÊ 90º DE REDUÇÃO"], "PÇ")
        self.assertIsNone(un["TUBO C/ FLANGE E PONTA PN10 L=0,36 M"])
        self.assertIn("PENDENTE", self.q("SELECT unidade_fonte FROM itens_projeto WHERE codigo_bm='APROVAR'")[0][0])

    def test_contador_por_unidade_separa_pendente(self):
        self.aplicar()
        c = self.conectar()
        t = S.totais_atividade(c, "Zona01_PB01_CH_DF-33")
        self.assertEqual(dict(t["totais"]), {"UN": 2.0, "M": 0.71, "PÇ": 2.0})
        self.assertEqual(t["pendentes"], 1)
        with self.assertRaisesRegex(Exception, "sem quantidade/unidade"):
            S.previsto(c, "ATIVIDADE", "Zona01_PB01_CH_DF-33", 0)

    def test_ba_bb_aw_e_avisos_separados(self):
        self.aplicar()
        self.assertEqual(self.q("SELECT enquadramento_bm, obs_enquadramento_bm FROM itens_projeto WHERE descricao='TÊ 90º DE REDUÇÃO'"),
                         [("CONTRATUAL — VALIDAR", "PN diferente")])
        av = self.q("SELECT tipo, prioridade, texto FROM avisos_origem a JOIN itens_projeto i ON i.id=a.item_id "
                    "WHERE i.descricao='LUVA ELETROFUSÃO' AND i.atividade_id IN (SELECT id FROM atividades WHERE nomenclatura LIKE '%NÓ1.8') ORDER BY 1,3")
        self.assertEqual(av, [("PROBLEMA", "P1", "P1 aviso A"), ("PROBLEMA", "P1", "P1 aviso B"), ("SUGESTAO", "P1", "sugestão A")])


class TestMigracaoV1(Base):
    def _criar_v1(self):
        c = sqlite3.connect(self.db)
        with open(os.path.join(AQUI, "fixtures", "schema_v1.sql"), encoding="utf-8") as f:
            c.executescript(f.read())
        c.execute("INSERT INTO equipes(id,nome) VALUES ('Zona01','EQUIPE ZONA 01')")
        c.execute("INSERT INTO trechos(id,frente_id,extensao_m,quantidades) VALUES ('Zona01_PB01_Rede_TR001','Zona01',76.55,'{}')")
        c.execute("INSERT INTO trechos(id,frente_id,extensao_m,quantidades) VALUES ('Zona09_PB99_Rede_TR777','Zona01',1,'{}')")
        c.execute("INSERT INTO elementos(id,numero,tipo) VALUES ('SUBDMC1_DF-33','DF-33','DF')")
        c.execute("INSERT INTO conjuntos(id,elemento_id,frente_id) VALUES ('Zona01_PB01_CH_DF-33','SUBDMC1_DF-33','Zona01')")
        c.execute("INSERT INTO materiais(id,descricao,unidade,propriedade) VALUES ('CAESB-PEAD-DE63','tubo','M','CAESB')")
        c.execute("INSERT INTO medicoes(id,numero,periodo_ini,periodo_fim) VALUES (1,1,'2026-10-01','2026-10-31')")
        c.execute("INSERT INTO checkins(id,data,equipe_id,alvo_tipo,alvo_id) VALUES (1,'2026-10-15','Zona01','TRECHO','Zona01_PB01_Rede_TR001')")
        c.execute("INSERT INTO movimentos(data,material_id,tipo,quantidade,alvo_tipo,alvo_id,checkin_id) "
                  "VALUES ('2026-10-15','CAESB-PEAD-DE63','CUSTODIA',80,'TRECHO','Zona01_PB01_Rede_TR001',1)")
        c.execute("INSERT INTO fotos(alvo_tipo,alvo_id,categoria,arquivo,data_hora,carimbo) "
                  "VALUES ('CONJUNTO','Zona01_PB01_CH_DF-33','EXECUCAO','a.jpg','2026-10-15T10:00','c')")
        c.execute("INSERT INTO fotos(alvo_tipo,alvo_id,categoria,arquivo,data_hora,carimbo) "
                  "VALUES ('TRECHO','Zona09_PB99_Rede_TR777','EXECUCAO','b.jpg','2026-10-15T10:00','c')")
        c.execute("INSERT INTO execucao(alvo_tipo,alvo_id,concluido,ok_fiscal_em,medicao_id) VALUES ('TRECHO','Zona01_PB01_Rede_TR001',1,'2026-10-15',1)")
        c.execute("INSERT INTO cadastro(trecho_id,etapa,atualizado_em) VALUES ('Zona01_PB01_Rede_TR001','1_PRONTO','x')")
        c.execute("INSERT INTO medicao_lancamentos(medicao_id,gaveta,servico,alvo_tipo,alvo_id,unidade,quantidade,valor) "
                  "VALUES (1,1,'MATERIAL 20%','TRECHO','Zona01_PB01_Rede_TR001','UN',1,99.9)")
        c.commit()
        c.close()
        for n in ("a.jpg", "b.jpg"):                                       # fotos referenciadas existem no disco
            with open(os.path.join(self.fotos, n), "wb") as f:
                f.write(n.encode())

    def test_migra_preservando_movimentos_e_remapeando(self):
        self._criar_v1()
        with self.assertRaises(S.EsquemaAntigo):
            S.conectar(self.db)                                            # servidor recusa v1 sem migração
        rel = self.aplicar()
        self.assertEqual(rel["esquema_antes"], 1)
        tr = self.tid(TR001)
        at = self.q("SELECT id FROM atividades WHERE nomenclatura='Zona01_PB01_CH_DF-33'")[0][0]
        self.assertEqual(self.q("SELECT alvo_tipo, alvo_id, quantidade FROM movimentos"), [("TRECHO", tr, 80.0)])
        self.assertEqual(self.q("SELECT alvo_tipo, alvo_id FROM checkins"), [("TRECHO", tr)])
        self.assertEqual(self.q("SELECT trecho_id FROM cadastro"), [(tr,)])
        self.assertEqual(self.q("SELECT alvo_id, ok_fiscal_em FROM execucao"), [(tr, "2026-10-15")])
        fotos = dict(self.q("SELECT arquivo, alvo_tipo || ':' || alvo_id FROM fotos"))
        self.assertEqual(fotos, {"a.jpg": f"ATIVIDADE:{at}", "b.jpg": "TRECHO:Zona09_PB99_Rede_TR777"})
        self.assertTrue(self.q("SELECT 1 FROM pendencias_importacao WHERE referencia LIKE '%TR777%'"))
        self.assertEqual(self.q("SELECT tipo FROM medicoes"), [("TESTE",)])
        self.assertEqual(self.q("SELECT status FROM medicao_lancamentos"), [("PENDENTE",)])
        self.assertEqual(self.q("SELECT valor FROM v1_medicao_lancamentos"), [(99.9,)])
        self.assertEqual(rel["registros_operacionais_preservados"]["fotos"], 2)
        self.assertTrue(backup.verificar(rel["backup"])[0])
        S.conectar(self.db).close()


# ====================================================================== A12 — arquivos fechados, codificação
class TestA12ArquivosECodificacao(Base):
    def test_importacao_fecha_planilhas_sem_coleta_de_lixo(self):
        gc.disable()
        try:
            self.aplicar()
            self.assertEqual(arquivos_abertos_em(self.planilhas) or [], [])
        finally:
            gc.enable()

    def test_saida_do_programa_em_utf8_mesmo_com_console_cp1252(self):
        self.aplicar()
        pasta = backup.criar(self.db, os.path.join(self.tmp, "bk"), self.fotos)
        env = dict(os.environ, PYTHONIOENCODING="cp1252", PYTHONUTF8="0")
        r = subprocess.run([sys.executable, os.path.join(APP, "backup.py"), "verificar", pasta], capture_output=True, env=env)
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertIn("Integridade: OK", r.stdout.decode("utf-8"))


if __name__ == "__main__":
    unittest.main()
