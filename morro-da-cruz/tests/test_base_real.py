"""Verificações sobre as planilhas REAIS do contrato (planilhas/, fora do Git), sempre em cópias descartáveis.

Pulados quando as planilhas não estão presentes. Conferem conteúdo e vínculos, não só contagens:
  * carga real + reimportação idempotente com o conteúdo de todas as tabelas inalterado;
  * troca de nomes TR001 ↔ TR002 numa cópia da 06.1 real (os dois têm o mesmo handle 5FD12; só as pontas UTM
    os distinguem): o check-in fica na geometria física original;
  * migração de um banco v1 com movimentos sobre a base real;
  * restauração com banco em WAL e fotos;
  * decisões só no escopo explícito (DF-05) e MND sem bloqueio de vala no detalhe.
"""
import json
import os
import shutil
import sqlite3
import sys
import tempfile
import unittest

AQUI = os.path.dirname(os.path.abspath(__file__))
APP = os.path.dirname(AQUI)
sys.path.insert(0, APP)
sys.path.insert(0, AQUI)
import backup  # noqa: E402
import importacao as I  # noqa: E402
import importar  # noqa: E402
import servicos as S  # noqa: E402
from fotos_teste import jpeg  # noqa: E402
from perfil_teste import cfg_teste_real  # noqa: E402

CFG = I.carregar_config()                 # perfil REAL (aprovação do usuário)
CFGT = cfg_teste_real()                   # perfil TESTE: mesmos controles, aprovados pelo teste automatizado
TEM_DADOS = all(I.caminho(CFG, k) for k in ("redes", "conexoes", "evidencias", "bm", "reajuste"))
TR001, TR002 = "Zona01_PB01_Rede_TR001", "Zona01_PB01_Rede_TR002"
D = "2026-10-15"


@unittest.skipUnless(TEM_DADOS, "planilhas oficiais ausentes em planilhas/")
class BaseReal(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.raiz = tempfile.mkdtemp(prefix="base_real_")
        cls.db0 = os.path.join(cls.raiz, "carga", "obra.db")
        os.makedirs(os.path.join(cls.raiz, "carga", "fotos"))
        cls.rel0 = importar.aplicar(cls.db0, I.DIR_PLANILHAS, cfg=CFGT, dir_relatorios=os.path.join(cls.raiz, "rel"))

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.raiz)

    def setUp(self):
        self.tmp = tempfile.mkdtemp(dir=self.raiz)
        self.dados = os.path.join(self.tmp, "dados")
        self.fotos = os.path.join(self.dados, "fotos")
        os.makedirs(self.fotos)
        self.db = os.path.join(self.dados, "obra.db")
        self.rels = os.path.join(self.tmp, "rel")
        self._conns = []

    def tearDown(self):
        for c in self._conns:
            c.close()

    def copia(self):
        backup.copiar_banco(self.db0, self.db)

    def conectar(self):
        c = S.conectar(self.db)
        self._conns.append(c)
        return c

    def q(self, sql, *a):
        c = sqlite3.connect(self.db)
        try:
            return c.execute(sql, a).fetchall()
        finally:
            c.close()

    # ---------------------------------------------------------------- carga e reimportação
    def test_carga_real_controles_e_totais(self):
        a = self.rel0["auditoria"]
        self.assertEqual(self.rel0["status"], "CONCLUIDA")
        self.assertTrue(a["contagens_ok"], a["contagens"])
        self.assertTrue(a["totalizacao_bm"]["ok"])
        self.assertEqual(a["obtido"]["unidades_linhas"], {"M": 73, "PENDENTE": 361, "PÇ": 54, "UN": 6051})
        self.assertEqual(a["obtido"]["quantidade_m"], 72.95)
        self.assertEqual(a["total_contratual"]["diferenca"], 0.0975869)
        self.assertEqual(self.rel0["bloqueios"], [])

    def test_reimportacao_real_nao_muda_conteudo(self):
        self.copia()
        antes = backup.conteudo(self.db)
        rel = importar.aplicar(self.db, I.DIR_PLANILHAS, cfg=CFGT, dir_relatorios=self.rels)
        self.assertEqual(rel["status"], "SEM_ALTERACOES")
        depois = backup.conteudo(self.db)
        mudou = sorted(t for t in antes if antes[t] != depois[t])
        self.assertEqual(mudou, ["importacoes"])                 # só o registro da própria importação

    def test_perfil_real_sem_aprovacao_simula_mas_nao_aplica(self):
        estrutura, aprovacao = I.validar_controle(CFG["controles"][0], "REAL")
        self.assertEqual(estrutura, [])                                  # conjunto real completo
        if not aprovacao:
            self.skipTest("o conjunto real já foi aprovado pelo usuário")
        self.copia()
        antes = backup.conteudo(self.db)
        sim = importar.simular(self.db, I.DIR_PLANILHAS, cfg=CFG, dir_relatorios=self.rels)
        self.assertEqual(sim["status"], "SIMULADA_SEM_CONTROLE_APROVADO")
        self.assertFalse(sim["auditoria"]["contagens_ok"])
        self.assertTrue(all(c["ok"] for c in sim["auditoria"]["contagens"].values()))   # valores conferem, sem confirmação
        with self.assertRaises(importar.ImportacaoReprovada) as cm:
            importar.aplicar(self.db, I.DIR_PLANILHAS, cfg=CFG, dir_relatorios=self.rels)
        self.assertIn("CONTROLE_SEM_APROVACAO", [b["tipo"] for b in cm.exception.relatorio["bloqueios"]])
        self.assertEqual(backup.conteudo(self.db), antes)

    # ---------------------------------------------------------------- nomes trocados na 06.1 real
    def test_nomes_trocados_na_06_1_real(self):
        import openpyxl
        self.copia()
        c = self.conectar()
        tid1 = c.execute("SELECT id FROM trechos WHERE ponto=?", (TR001,)).fetchone()[0]
        geo1 = c.execute("SELECT handle, utm_ini, utm_fim, extensao_m FROM trechos WHERE id=?", (tid1,)).fetchone()
        S.checkin(c, D, "Zona01", "TRECHO", TR001)
        S.registrar_foto(c, "TRECHO", TR001, "SINALIZACAO", jpeg(), data_hora=f"{D}T08:00:00", dir_fotos=self.fotos)
        pl = os.path.join(self.tmp, "planilhas")
        shutil.copytree(I.DIR_PLANILHAS, pl)
        p = os.path.join(pl, CFG["arquivos"]["redes"])
        wb = openpyxl.load_workbook(p, data_only=True)            # valores (o importador lê valores)
        ws = wb["REDE"]
        linhas = {ws.cell(r, 18).value: r for r in range(6, ws.max_row + 1) if ws.cell(r, 18).value in (TR001, TR002)}
        ws.cell(linhas[TR001], 18, TR002)
        ws.cell(linhas[TR002], 18, TR001)
        wb.save(p)
        wb.close()
        cfg = json.loads(json.dumps(CFGT))
        ctl = dict(cfg["controles"][0], id="CTRL-TESTE nomes trocados")
        ctl["fontes"] = dict(ctl["fontes"], redes=backup.sha256(p))
        cfg["controles"] = [ctl]
        rel = importar.aplicar(self.db, pl, cfg=cfg, dir_relatorios=self.rels)
        self.assertEqual(rel["status"], "CONCLUIDA")
        ck = self.q("""SELECT t.id, t.ponto, t.handle, t.utm_ini, t.utm_fim, t.extensao_m FROM checkins c
                       JOIN trechos t ON t.id=c.alvo_id""")
        self.assertEqual(ck, [(tid1, TR002, *geo1)])                # mesma geometria física, nome novo
        self.assertEqual(self.q("SELECT alvo_id FROM fotos"), [(tid1,)])
        self.assertEqual(sorted(p["referencia"] for p in rel["pendencias"] if p["tipo"] == "TRECHO_RENOMEADO"),
                         sorted([tid1, self.q("SELECT id FROM trechos WHERE ponto=? AND vigente=1", TR001)[0][0]]))
        self.assertTrue(rel["auditoria"]["contagens_ok"])

    # ---------------------------------------------------------------- migração v1 com movimentos
    def test_migracao_v1_real_com_movimentos(self):
        c = sqlite3.connect(self.db)
        with open(os.path.join(AQUI, "fixtures", "schema_v1.sql"), encoding="utf-8") as f:
            c.executescript(f.read())
        c.execute("INSERT INTO equipes(id,nome) VALUES ('Zona01','EQUIPE ZONA 01'), ('CH','EQUIPE CH')")
        c.execute("INSERT INTO trechos(id,frente_id,extensao_m,quantidades) VALUES (?,'Zona01',76.55,'{}')", (TR001,))
        c.execute("INSERT INTO elementos(id,numero,tipo) VALUES ('SUBDMC1_DF-33','DF-33','DF')")
        c.execute("INSERT INTO conjuntos(id,elemento_id,frente_id) VALUES ('Zona01_PB01_CH_DF-33','SUBDMC1_DF-33','CH')")
        c.execute("INSERT INTO materiais(id,descricao,unidade,propriedade) VALUES ('CAESB-PEAD-DE63','tubo','M','CAESB')")
        c.execute("INSERT INTO medicoes(id,numero,periodo_ini,periodo_fim,status) VALUES (1,1,'2026-10-01','2026-10-31','FECHADA')")
        c.execute("INSERT INTO checkins(id,data,equipe_id,alvo_tipo,alvo_id) VALUES (1,?, 'Zona01','TRECHO',?)", (D, TR001))
        for t, q in (("ROMANEIO_CAESB", 120), ("CUSTODIA", 80), ("APLICADO", 76.55), ("SOBRA", 3.45)):
            c.execute("INSERT INTO movimentos(data,material_id,tipo,quantidade,alvo_tipo,alvo_id,checkin_id) "
                      "VALUES (?,'CAESB-PEAD-DE63',?,?,'TRECHO',?,1)", (D, t, q, TR001))
        c.execute("INSERT INTO fotos(alvo_tipo,alvo_id,categoria,arquivo,data_hora,carimbo) "
                  "VALUES ('CONJUNTO','Zona01_PB01_CH_DF-33','EXECUCAO','a.jpg',?,'c')", (D,))
        c.execute("INSERT INTO execucao(alvo_tipo,alvo_id,concluido,ok_fiscal_em,medicao_id) VALUES ('TRECHO',?,1,?,1)", (TR001, D))
        c.commit()
        c.close()
        with open(os.path.join(self.fotos, "a.jpg"), "wb") as f:
            f.write(b"a")
        rel = importar.aplicar(self.db, I.DIR_PLANILHAS, cfg=CFGT, dir_relatorios=self.rels)
        self.assertEqual((rel["esquema_antes"], rel["status"]), (1, "CONCLUIDA"))
        self.assertTrue(rel["auditoria"]["contagens_ok"])
        tr = self.q("SELECT id FROM trechos WHERE ponto=? AND vigente=1", TR001)[0][0]
        at = self.q("SELECT id FROM atividades WHERE nomenclatura='Zona01_PB01_CH_DF-33' AND vigente=1")[0][0]
        self.assertEqual(self.q("SELECT tipo, quantidade, alvo_id FROM movimentos ORDER BY id"),
                         [("ROMANEIO_CAESB", 120.0, tr), ("CUSTODIA", 80.0, tr), ("APLICADO", 76.55, tr), ("SOBRA", 3.45, tr)])
        self.assertEqual(self.q("SELECT alvo_tipo, alvo_id FROM fotos"), [("ATIVIDADE", at)])
        self.assertEqual(self.q("SELECT alvo_id, ok_fiscal_em FROM execucao"), [(tr, D)])
        self.assertEqual(rel["medicoes_congeladas_na_migracao"], [1])
        c = self.conectar()
        self.assertEqual(S.totais_atividade(c, at)["texto"], "17 UN · 0,71 m")

    # ---------------------------------------------------------------- restauração com WAL e fotos
    def test_restauracao_real_wal_e_fotos(self):
        self.copia()
        c0 = self.conectar()
        S.registrar_foto(c0, "TRECHO", TR001, "SINALIZACAO", jpeg(), data_hora=f"{D}T08:00:00", dir_fotos=self.fotos)
        pasta = backup.criar(self.db, os.path.join(self.dados, "backups"), self.fotos)
        w = sqlite3.connect(self.db, isolation_level=None)
        self._conns.append(w)
        w.execute("PRAGMA journal_mode=WAL")
        w.execute("PRAGMA wal_autocheckpoint=0")
        S.criar_medicao(c0, 1, "2026-10-01", "2026-10-31")
        S.checkin(c0, D, "Zona01", "TRECHO", TR002)
        seguranca = backup.restaurar(pasta, self.db, self.fotos)
        g = sqlite3.connect(os.path.join(seguranca, "obra.db"))
        try:
            self.assertEqual(g.execute("SELECT COUNT(*) FROM medicoes").fetchone()[0], 1)    # estado do WAL preservado
            self.assertEqual(g.execute("SELECT COUNT(*) FROM checkins").fetchone()[0], 1)
        finally:
            g.close()
        self.assertTrue(backup.verificar(seguranca)[0])
        with open(os.path.join(pasta, "manifesto.json"), encoding="utf-8") as f:
            self.assertEqual(backup.conteudo(self.db), json.load(f)["conteudo_por_tabela"])

    # ---------------------------------------------------------------- decisões e MND na base real
    def test_decisao_df05_so_no_escopo_e_mnd_sem_bloqueio_de_vala(self):
        self.copia()
        c = self.conectar()
        fora, df05 = [], []
        for aid, numero in c.execute("SELECT a.id, e.numero FROM atividades a JOIN elementos e ON e.id=a.elemento_id "
                                     "WHERE a.vigente=1").fetchall():
            for av in S.avisos_e_decisoes(c, aid)["avisos"]:
                if "DEC-0001" in av["decisoes"]:
                    (df05 if numero == "DF-05" else fora).append(av["texto"])
        self.assertEqual(fora, [])                                   # eram 18 vínculos indevidos (9 × aviso/sugestão)
        self.assertEqual(sorted(set(df05)), ["P1 — DF-05 contada em duas pranchas (16 e 19).",
                                             "Tirar a DF-05 do resumo e do quadro de uma das pranchas."])
        d = S.detalhe(c, "TRECHO", "Zona02_PB04_Rede_TR036")
        self.assertEqual((d["metodologia"], d["perfil_bloqueado"], d["mnd"]), ("MND", None, True))


if __name__ == "__main__":
    unittest.main()
