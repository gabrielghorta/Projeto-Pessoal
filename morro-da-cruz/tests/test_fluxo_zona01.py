"""Piloto Zona01 (TESTE) sobre as planilhas oficiais reais em planilhas/ (não versionadas).

Importa num banco temporário e percorre TR001, NÓ1.8, DF-33 e um trecho MND, conferindo as travas das
etapas 0/1: nenhuma medição oficial, nenhum regime de material aplicado, MND sem substituição automática,
pendências fora do valor e totalização do BM na ordem das células.
"""
import os
import shutil
import sqlite3
import sys
import tempfile
import unittest

AQUI = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, AQUI)
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import exportar  # noqa: E402
import importacao as I  # noqa: E402
import importar  # noqa: E402
import regras as R  # noqa: E402
import servicos as S  # noqa: E402
from fotos_teste import jpeg  # noqa: E402
from perfil_teste import cfg_teste_real  # noqa: E402

CFG = I.carregar_config()
TEM_DADOS = all(I.caminho(CFG, k) for k in ("redes", "conexoes", "bm", "reajuste"))
TR = "Zona01_PB01_Rede_TR001"
MND = "Zona02_PB04_Rede_TR036"
NO18 = "Zona01_PB01_Rede_NÓ1.8"
D = "2026-10-15"


@unittest.skipUnless(TEM_DADOS, "planilhas oficiais ausentes em planilhas/")
class PilotoZona01(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.mkdtemp()
        cls.db = os.path.join(cls.tmp, "piloto.db")
        cls.rel = importar.executar(cls.db, "APLICACAO", cfg=cfg_teste_real(), dir_relatorios=os.path.join(cls.tmp, "relatorios"))
        cls.c = S.conectar(cls.db)
        cls.fotos = os.path.join(cls.tmp, "fotos")
        cls.med = S.criar_medicao(cls.c, 1, "2026-10-01", "2026-10-31")
        for t in ("MANHA", "TARDE"):
            for eq in ("Zona01", "Zona02"):
                S.registrar_clima(cls.c, D, eq, t, "SOL")

    @classmethod
    def tearDownClass(cls):
        cls.c.close()
        shutil.rmtree(cls.tmp)

    def foto(self, tipo, alvo, cat):
        S.registrar_foto(self.c, tipo, alvo, cat, jpeg(), data_hora=f"{D}T10:00:00", dir_fotos=self.fotos)

    def lanc(self, alvo_id):
        return S.resumo_medicao(self.c, self.med), [dict(r) for r in self.c.execute(
            "SELECT l.*, b.item FROM medicao_lancamentos l LEFT JOIN bm_itens b ON b.id=l.bm_item_id WHERE alvo_id=?", (alvo_id,))]

    # ---------------------------------------------------------------- carga real
    def test_01_auditoria_da_carga_real(self):
        a = self.rel["auditoria"]
        self.assertTrue(a["contagens_ok"], a["contagens"])
        self.assertTrue(a["totalizacao_bm"]["ok"])
        self.assertAlmostEqual(a["totalizacao_bm"]["sistema"]["p0"], 49878748.6575869, places=6)
        self.assertEqual(a["total_contratual"]["contrato_3_1"], 49878748.56)
        self.assertAlmostEqual(a["total_contratual"]["diferenca"], 0.0975869, places=6)
        self.assertIn("DIVERGÊNCIA", a["total_contratual"]["situacao"])
        u = a["unidades_06_2"]
        self.assertEqual((u["M"]["linhas"], u["M"]["quantidade"]), (73, 72.95))
        self.assertEqual(u["PÇ"]["linhas"], 54)
        self.assertEqual(u["PENDENTE"]["linhas"], 361)
        self.assertEqual(a["enquadramento_bm_BA"]["APROVAR — SEM ITEM NO BM"], 361)
        self.assertEqual(len(a["trechos_mnd"]), 4)
        self.assertEqual(self.rel["conferencias"]["agrupamento_vs_evidencias"], {"evidencias": 1191, "iguais": 1191})

    def test_02_df33_um_elemento_duas_atividades(self):
        el = self.c.execute("SELECT id FROM elementos WHERE numero='DF-33' AND handle='5FDEA'").fetchall()
        self.assertEqual(len(el), 1)
        noms = sorted(r[0] for r in self.c.execute("SELECT nomenclatura FROM atividades WHERE elemento_id=?", (el[0][0],)))
        self.assertEqual(noms, ["Zona01_PB01_CH_DF-33", "Zona01_PB01_Rede_DF-33"])
        self.assertEqual(S.totais_atividade(self.c, "Zona01_PB01_CH_DF-33")["texto"], "17 UN · 0,71 m")

    # ---------------------------------------------------------------- piloto
    def test_03_tr001_proposta_sem_regime_e_pendencias_fora_do_valor(self):
        c = self.c
        S.entrada_material(c, D, "CAESB-PEAD-DE63", 100, "ROMANEIO_CAESB", "ROM-1")
        ck = S.checkin(c, D, "Zona01", "TRECHO", TR, materiais=[{"material_id": "CAESB-PEAD-DE63", "quantidade": 80}])
        S.apontar(c, D, "Zona01", "TRECHO", TR, intervalos_novos=[{"etapa": "ASSENTADO", "ini_m": 0, "fim_m": 40}, {"etapa": "REATERRADO", "ini_m": 0, "fim_m": 30}])
        S.apontar(c, D, "Zona01", "TRECHO", TR, concluido_100=True)
        for cat in ("SINALIZACAO", "EXECUCAO", "CONCLUIDO"):
            self.foto("TRECHO", TR, cat)
        S.checkout(c, ck, [{"material_id": "CAESB-PEAD-DE63", "aplicado": 76.55, "sobra": 3.45}])
        sid = S.registrar_soldas(c, D, "Zona01", "TRECHO", TR, "ELETROFUSAO", 11, "Soldador X", "Encarregado")
        S.registrar_soldas(c, D, "Zona01", "TRECHO", TR, "ELETROFUSAO", 2, "Soldador X", "Encarregado")   # sem verificação
        S.verificar_soldas(c, sid, "Fiscal", "APROVADA")
        S.verificar_trecho(c, TR, "Fiscal")                              # o OK interno exige a verificação (v20)
        S.ok_fiscalizacao(c, "TRECHO", TR, "eng", D)
        r, ls = self.lanc(S.alvo(c, "TRECHO", TR)["id"])
        por = {l["servico"]: l for l in ls}
        self.assertEqual((por["ASSENTAMENTO"]["item"], por["ASSENTAMENTO"]["status"]), ("2.2_5.8.44", "PROPOSTO"))
        self.assertAlmostEqual(por["ASSENTAMENTO"]["quantidade"], 76.55)
        # soldas: só as verificadas e aprovadas (11), não a estimativa ⌈76,55/6⌉ = 13; regra contratual pendente
        self.assertEqual((por["SOLDA"]["status"], por["SOLDA"]["quantidade"]), ("PENDENTE", 11))
        self.assertIn("APROVADAS", por["SOLDA"]["motivo"])
        self.assertEqual(por["LASTRO"]["status"], "PENDENTE")
        self.assertTrue(all(l["material_id"] is None for l in ls))        # nenhum regime de material
        self.assertTrue(r["teste"])
        # valor = só PROPOSTOS, na ordem do BM
        itens = {}
        for l in ls:
            if l["status"] == "PROPOSTO":
                b = c.execute("SELECT classificacao, preco_unitario FROM bm_itens WHERE id=?", (l["bm_item_id"],)).fetchone()
                itens.setdefault(l["bm_item_id"], [b[0], 0, b[1]])[1] += l["quantidade"]
        esperado = R.totalizar_bm([tuple(v) for v in itens.values()], 0.1528, 0.2148, -1 + (1 - 0.1404))["total"]
        self.assertAlmostEqual(r["gavetas"][1]["valor"], esperado, places=2)

    def test_04_mnd_sem_substituicao_automatica(self):
        c = self.c
        S.apontar(c, D, "Zona02", "TRECHO", MND, concluido_100=True)
        for cat in ("SINALIZACAO", "EXECUCAO", "CONCLUIDO"):
            self.foto("TRECHO", MND, cat)
        S.verificar_trecho(c, MND, "Fiscal")
        S.ok_fiscalizacao(c, "TRECHO", MND, "eng", D)
        _, ls = self.lanc(S.alvo(c, "TRECHO", MND)["id"])
        ass = [l for l in ls if l["servico"] == "ASSENTAMENTO"][0]
        self.assertEqual(ass["status"], "PENDENTE")
        self.assertIsNone(ass["bm_item_id"])
        self.assertIn("MND", ass["motivo"])

    def test_05_no18_material_fica_pendente(self):
        c = self.c
        prev = S.previsto(c, "ATIVIDADE", NO18, 0)
        self.assertEqual({p["unidade"] for p in prev}, {"UN"})
        S.apontar(c, D, "Zona01", "ATIVIDADE", NO18, concluido_100=True)
        self.assertEqual(S.fotos_faltando(c, "ATIVIDADE", NO18), ["EXECUCAO", "CONCLUIDO"])   # herdou sinalização do TR001
        for cat in ("EXECUCAO", "CONCLUIDO"):
            self.foto("ATIVIDADE", NO18, cat)
        S.ok_fiscalizacao(c, "ATIVIDADE", NO18, "eng", D)
        _, ls = self.lanc(S.alvo(c, "ATIVIDADE", NO18)["id"])
        self.assertTrue(ls)
        self.assertTrue(all(l["status"] == "PENDENTE" and "regime" in l["motivo"] for l in ls))

    def test_06_nf_aceita_cria_estoque_sem_lancamento(self):
        c = self.c
        antes = c.execute("SELECT COUNT(*) FROM medicao_lancamentos").fetchone()[0]
        mat = c.execute("SELECT id FROM materiais WHERE propriedade='ANKARA' AND unidade='UN' LIMIT 1").fetchone()[0]
        nf = S.registrar_nf(c, "N1", "F", D, [{"material_id": mat, "quantidade": 10, "valor_unitario": 1}], faturamento_etapas=True)
        S.aceitar_nf(c, nf, "laudo.pdf")
        self.assertEqual(S.saldo_almox(c, mat)[0], 10)
        self.assertEqual(c.execute("SELECT COUNT(*) FROM medicao_lancamentos").fetchone()[0], antes)

    def test_07_medicao_oficial_bloqueada_e_export_marcado(self):
        with self.assertRaisesRegex(R.RegraViolada, "OFICIAL bloqueada"):
            S.criar_medicao(self.c, 1, "2026-10-01", "2026-10-31", tipo="OFICIAL")
        nome, dados = exportar.planilha_medicao(self.c, self.med)
        self.assertTrue(nome.startswith("TESTE_"))
        import io
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(dados))
        self.assertTrue(all("NÃO OFICIAL" in str(ws["A1"].value) for ws in wb.worksheets))

    def test_08_reimportacao_real_sem_alteracoes(self):
        copia = os.path.join(self.tmp, "copia.db")
        o, d = sqlite3.connect(self.db), sqlite3.connect(copia)
        o.backup(d)
        d.close()
        o.close()
        rel = importar.executar(copia, "SIMULACAO", cfg=cfg_teste_real(), dir_relatorios=os.path.join(self.tmp, "relatorios"))
        self.assertEqual(rel["status"], "SEM_ALTERACOES")


if __name__ == "__main__":
    unittest.main()
