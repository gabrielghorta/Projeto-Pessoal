"""Módulo F — relatórios: ciclo do material, avanço por frente/equipe/período/etapa e pendências por frente, conferidos
contra os movimentos e intervalos efetivos. Planilhas sintéticas, bancos temporários."""
import io
import os
import sys
import unittest

from openpyxl import load_workbook

AQUI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(AQUI))
sys.path.insert(0, AQUI)
import compras as C  # noqa: E402
import planejamento as P  # noqa: E402
import relatorios as RL  # noqa: E402
import servicos as S  # noqa: E402
from test_acesso import BaseHTTP  # noqa: E402
from test_campo import TR, BaseCampo, D  # noqa: E402
from test_compras import MAT, BaseCompras  # noqa: E402

D2 = "2026-10-16"
TR2 = "Zona01_PB01_Rede_TR002"
MND = "Zona02_PB04_Rede_TR036"


class TestCicloDoMaterial(BaseCompras):
    def test_colunas_batem_com_os_movimentos(self):
        pid = self.pedido(100)
        nf = self.importar(601, [("L1", "Luva eletrofusão", "CX", 10, 100.0)])
        (l1, _), = self.linhas(nf)
        pi = self.q("SELECT id FROM pedido_itens WHERE pedido_id=?", pid)[0][0]
        C.conferir_nf(self.c, nf, pid, [{"linha_id": l1, "material_id": MAT, "fator": 10, "pedido_item_id": pi}], "almox2")
        ni = self.q("SELECT id FROM nf_itens WHERE nf_id=?", nf)[0][0]
        C.receber(self.c, nf, D2, [{"nf_item_id": ni, "quantidade": 60}], "almox1")
        m = {x["material_id"]: x for x in RL.materiais(self.c)}[MAT]
        self.assertEqual((m["solicitado"], m["comprado"], m["quarentena"], m["recebido"], m["a_receber"]), (100, 100, 60, 0, 40))
        ri = self.q("SELECT id FROM recebimento_itens")[0][0]
        C.inspecionar(self.c, ri, 50, 10, "LAUDO-1", "almox1", D2, "10 luvas ovalizadas")
        C.devolver_ao_fornecedor(self.c, ni, D2, 5, "defeito de fábrica", "almox1", "NF dev 1")
        for t in ("MANHA", "TARDE"):
            S.registrar_clima(self.c, D2, "Zona01", t, "SOL")
        S.reservar(self.c, D2, "Zona01", MAT, 8)
        S.reservar(self.c, D2, "Zona02", MAT, 3)
        ck = S.checkin(self.c, D2, "Zona01", "TRECHO", TR, materiais=[{"material_id": MAT, "quantidade": 20}])
        S.checkout(self.c, ck, [{"material_id": MAT, "aplicado": 12, "sobra": 3, "perda": 1, "motivo": "luva caiu na vala",
                                 "fica_com_equipe": 4}])
        m = {x["material_id"]: x for x in RL.materiais(self.c)}[MAT]
        esperado = {"comprado": 100, "a_receber": 55, "quarentena": 0, "recebido": 45, "saldo_almox": 28, "reservado": 3,
                    "com_equipes": 4, "aplicado": 12, "perda": 1}
        self.assertEqual({k: m[k] for k in esperado}, esperado)
        self.assertIsNone(m["medido"])                                                # nada de zero inventado
        self.assertIn("D-53", m["medido_situacao"])
        # recebido líquido = almoxarifado + com as equipes + aplicado + perda + emprestado
        self.assertAlmostEqual(m["recebido"], m["saldo_almox"] + m["com_equipes"] + m["aplicado"] + m["perda"] + m["emprestado"])
        self.assertEqual(m["saldo_almox"], S.saldo_almox(self.c, MAT)[0])
        # estorno sai do relatório
        mv = self.q("SELECT id FROM movimentos WHERE tipo='DEVOLUCAO_FORNECEDOR'")[0][0]
        S.estornar(self.c, mv, "devolução lançada em duplicidade", "almox2")
        self.assertEqual({x["material_id"]: x for x in RL.materiais(self.c)}[MAT]["recebido"], 50)


class TestAvancoEPendencias(BaseCampo):
    def setUp(self):
        super().setUp()
        for eq in ("Zona01", "Zona02"):
            for t in ("MANHA", "TARDE"):
                S.registrar_clima(self.c, D2, eq, t, "SOL")

    def test_avanco_por_frente_equipe_periodo_e_etapa(self):
        self.ap(("ASSENTADO", 0, 30), ("CADASTRO_CAMPO", 0, 30))
        self.ap(("REATERRADO", 0, 20))
        ap = S.apontar(self.c, D2, "Zona02", "TRECHO", TR2, intervalos_novos=[{"etapa": "ASSENTADO", "ini_m": 0, "fim_m": 12}])
        S.apontar(self.c, D2, "Zona02", "TRECHO", MND, intervalos_novos=[{"etapa": "ASSENTADO", "ini_m": 0, "fim_m": 10}])
        S.retificar_intervalos(self.c, "APONTAMENTO", ap, [{"etapa": "ASSENTADO", "ini_m": 0, "fim_m": 10}], "trena: 10 m", "eng1", D2)
        a = RL.avanco(self.c, D, D2)
        lin = {(l["frente_id"], l["equipe_id"]): l for l in a["por_equipe"]}
        self.assertEqual((lin[("Zona01", "Zona01")]["ASSENTADO"], lin[("Zona01", "Zona01")]["REATERRADO"]), (30, 20))
        self.assertEqual(lin[("Zona01", "Zona02")]["ASSENTADO"], 10)                 # retificado: 10, não 12 nem 22
        self.assertEqual(lin[("Zona02", "Zona02")]["ASSENTADO"], 10)
        self.assertEqual(a["total_periodo"]["ASSENTADO"], sum(l["ASSENTADO"] for l in a["por_equipe"]))   # soma por frente = geral
        self.assertEqual(a["por_dia"], {D: 30, D2: 20})
        fz = {f["frente_id"]: f for f in a["frentes"]}
        self.assertEqual(sum(f["assentado_m"] for f in a["frentes"]), 50)
        self.assertEqual(fz["Zona01"]["assentado_m"], 40)
        self.assertAlmostEqual(fz["Zona01"]["assentado_pct"], round(100 * 40 / (76.55 + 75.07), 2))
        self.assertEqual(RL.avanco(self.c, D2, D2)["total_periodo"]["ASSENTADO"], 20)   # período
        self.assertEqual(RL.avanco(self.c, D, D2, frente="Zona02")["total_periodo"]["ASSENTADO"], 10)
        self.assertEqual(RL.avanco(self.c, D, D2, equipe="Zona01")["total_periodo"]["ASSENTADO"], 30)
        # legado sem posição aparece à parte
        tid = S.alvo(self.c, "TRECHO", TR2)["id"]
        self.c.execute("INSERT INTO apontamentos(data,equipe_id,alvo_tipo,alvo_id,metros_executados) VALUES (?, 'Zona01', 'TRECHO', ?, 7)", (D, tid))
        self.assertEqual({f["frente_id"]: f for f in RL.avanco(self.c, D, D2)["frentes"]}["Zona01"]["legado_sem_posicao_m"], 7)

    def test_pendencias_por_frente(self):
        self.ap(("ASSENTADO", 0, 30), ("REATERRADO", 0, 20))                          # reaterro sem cadastro (alerta)
        S.abrir_ocorrencia(self.c, "TRECHO", TR, "IMPEDIMENTO", "CHUVA", "vala alagada", D, "enc1")
        S.abrir_ocorrencia(self.c, "TRECHO", MND, "NAO_CONFORMIDADE", "EXECUCAO", "furo desviado", D, "eng1")
        S.apontar(self.c, D2, "Zona01", "TRECHO", TR2, concluido_100=True)            # concluído, sem fotos, sem OK
        P.programar(self.c, D, "Zona02", "TRECHO", MND, "ASSENTAMENTO", "eng1")
        p = RL.pendencias(self.c, hoje="2026-10-20")
        fz = {f["frente_id"]: f for f in p["frentes"]}
        z1, z2 = fz["Zona01"], fz["Zona02"]
        # reaterro sem cadastro: TR001 (0–20 m) e TR002 (o 100 % sem levantamento de campo)
        self.assertEqual((z1["impedimentos"], z1["reaterro_sem_cadastro"], z1["aguardando_ok_interno"], z1["concluidos_sem_fotos"]), (1, 2, 1, 1))
        self.assertGreaterEqual(z1["intervalos_sem_verificacao"], 2)
        self.assertEqual((z2["nao_conformidades"], z2["programacao_vencida"]), (1, 1))
        self.assertEqual(fz["Zona03"]["total"], 0)
        self.assertEqual([f["frente_id"] for f in RL.pendencias(self.c, "Zona02", "2026-10-20")["frentes"]], ["Zona02"])


class TestRelatoriosPeloHTTP(BaseHTTP):
    def test_consulta_e_exportacao(self):
        self.assertEqual(self.cliente().get("/api/relatorios")[0], 401)
        enc, eng = self.cliente("enc1"), self.cliente("eng1")                         # Relatórios: tela do Administrativo
        self.assertEqual((enc.get("/api/relatorios")[0], enc.get(f"/api/relatorios/exportar?ini={D}&fim={D2}")[0]), (403, 403))
        st, r = eng.get(f"/api/relatorios?ini={D}&fim={D2}&frente=Zona01")
        self.assertEqual(st, 200, r)
        self.assertEqual(set(r), {"materiais", "avanco", "pendencias", "faltantes"})
        self.assertEqual([f["frente_id"] for f in r["avanco"]["frentes"]], ["Zona01"])
        st, xlsx = eng.get(f"/api/relatorios/exportar?ini={D}&fim={D2}")
        self.assertEqual(st, 200)
        wb = load_workbook(io.BytesIO(xlsx))
        self.assertEqual(wb.sheetnames, ["MATERIAIS", "AVANÇO", "PENDÊNCIAS", "FALTANTES"])
        self.assertEqual(self.q("SELECT COUNT(*) FROM auditoria WHERE rota='/api/relatorios/exportar'")[0][0], 1)


if __name__ == "__main__":
    unittest.main()
