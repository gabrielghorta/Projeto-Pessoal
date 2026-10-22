"""Nota de Serviço sintética para os testes (os PDFs reais têm nome e CREA de pessoas e não ficam no repositório).

Monta o texto no leiaute que o pdftotext dá para o modelo Caesb, o PDF com camada de texto (fonte Courier, uma célula
por trecho de texto, como o Excel imprime; com `girar=True` a página tem /Rotate 90, como a NS assinada) e a planilha
.xlsx e .xls com as mesmas células. As assinaturas levam só marcadores ([NOME], [CREA])."""
import io
import struct
from datetime import date

# colunas (centro, em caracteres) do pdftotext da NS 006: inteira, +, fração, elemento, terreno, GI, profundidade,
# método, material, parcial, acumulada, boca, fundo, visor, cruzeta, observação
COLUNAS = (4, 14, 24, 58, 91, 104, 121, 135, 148, 158, 173, 185, 196, 207, 223, 266)
CABECALHO_TABELA = [
    ' ' * 92 + '{bloco}',
    ' ' * 94 + 'COTAS',
    ' ' * 11 + 'ESTACA' + ' ' * 127 + 'MATERIAL           DISTÂNCIA              ESCAVAÇÃO          VISOR         CRUZETA'
    + ' ' * 32 + 'OBSERVAÇÕES',
    ' ' * 98 + '"GI" GERATRIZ                    MÉTODO',
    ' ' * 54 + 'ELEMENTOS' + ' ' * 52 + 'PROFUNDIDADE',
    ' ' * 87 + 'TERRENO     INFERIOR DO                    EXECUTIVO',
    ' Inteira     +       Fração' + ' ' * 75 + 'TUBO' + ' ' * 38 + '"D" (mm)   PARCIAL     ACUMULADA     BOCA (m)   FUNDO (m)',
]


def br(v, casas=2):
    """1089.55 -> '1.089,550' (casas=3); 17.32 -> '17,32'."""
    if v is None or v == '':
        return ''
    t = f'{float(v):,.{casas}f}'
    return t.replace(',', '§').replace('.', ',').replace('§', '.')


def estaca(m):
    """122.01 -> (6, '2,01'): estaca de 20 m + fração."""
    i = int(round(m, 2) // 20)
    return i, br(round(m - 20 * i, 2))


def _colocar(cols):
    s = ''
    for col, txt in cols:
        txt = str(txt)
        if not txt:
            continue
        ini = max(col - len(txt) // 2, len(s) + 2 if s else 0)
        s = s.ljust(ini) + txt
    return s


def linha_dados(l):
    """l: {estaca_m, elemento, terreno, gi, metodo ('V.C.A'|'MND'), material ('PEAD-160'), parcial, acumulada, boca, fundo,
    visor, cruzeta, obs}; campo ausente = célula vazia. Profundidade = terreno − GI."""
    i, f = estaca(l['estaca_m']) if l.get('estaca_m') is not None else ('', '')
    prof = round(l['terreno'] - l['gi'], 2) if l.get('terreno') is not None and l.get('gi') is not None else None
    vals = [i, '+' if i != '' else '', f, l.get('elemento') or '', br(l.get('terreno'), 3), br(l.get('gi'), 3), br(prof),
            l.get('metodo', 'V.C.A'), l.get('material', 'PEAD-160'), br(l.get('parcial')), br(l.get('acumulada')),
            br(l.get('boca')), br(l.get('fundo')), br(l.get('visor')), br(l.get('cruzeta')), l.get('obs') or '']
    return _colocar(zip(COLUNAS, vals))


def grupo_padrao(titulo, elementos, distancias, obs=None, cota=1012.45, metodo='V.C.A', material='PEAD-160',
                 inicio=0.0, segmentos=True):
    """Grupo com uma linha por elemento: distâncias parciais entre elementos consecutivos (a 1ª linha tem parcial 0).
    `inicio` = acumulada da 1ª linha (continuação de outro grupo). obs = lista com a observação de cada linha."""
    linhas, acc = [], inicio
    for k, el in enumerate(elementos):
        parc = 0.0 if k == 0 else distancias[k - 1]
        acc = round(acc + parc, 2)
        z = round(cota - 0.37 * k, 3)
        linhas.append({'estaca_m': acc, 'elemento': el, 'terreno': z, 'gi': round(z - 1.25, 3), 'metodo': metodo,
                       'material': material, 'parcial': parc, 'acumulada': acc, 'boca': 0.6, 'fundo': 0.6,
                       'obs': (obs or [None] * len(elementos))[k]})
    return {'titulo': titulo, 'linhas': linhas, 'segmentos': segmentos}


def _curto(el):
    return (el or '').split(' (')[0].split(' - ')[0].strip()


def paginas_ns(numero='012', contrato='10239/2026', data='15-out-26', localizacao='REDE DE DISTRIBUIÇÃO DE ÁGUA - MORRO DA '
               'CRUZ - SÃO SEBASTIÃO/DF', titulo='"REDE DE ÁGUA" - ZONA 03 - PB08 - VER CROQUI ANEXO - VER PROJETO EXECUTIVO',
               bloco='TOPOGRAFIA', grupos=(), notas=(), resumo=None, total=None, papeis=('TOPOGRAFIA:', 'CONSTRUTORA ANKARA'),
               folhas=None):
    """Texto das folhas da NS. `folhas`: lista com os índices dos grupos de cada folha (padrão: tudo na folha 1); NOTAS,
    RESUMO, TOTAL e assinaturas vão na última."""
    folhas = folhas or [list(range(len(grupos)))]
    n = len(folhas)
    if resumo is None:
        resumo = [(f'TUBO PEAD {g["linhas"][0].get("material", "PEAD-160").split("-")[-1]}mm',
                   round(sum(l['parcial'] or 0 for l in g['linhas']), 2), g['titulo']) for g in grupos]
    if total is None:
        total = round(sum(r[1] for r in resumo), 2)
    out = []
    for k, idx in enumerate(folhas, 1):
        ls = [' ' * 107 + 'Companhia de Saneamento Ambiental do Distrito Federal - Caesb' + ' ' * 63 + 'NS nº :' + ' ' * 30 + numero,
              '',
              'CONTRATO:           ' + contrato + ' ' * 200 + f'FOLHA:    {k:02d} de {n:02d}',
              'CONTRATADA: CONSTRUTORA ANKARA ENGENHARIA LTDA',
              ' ' * 233 + 'DATA:   ' + data,
              'LOCALIZAÇÃO: ' + localizacao,
              '',
              ' ' * 69 + titulo,
              '']
        ls += [c.format(bloco=bloco) for c in CABECALHO_TABELA] + ['', '', '', '']
        for gi in idx:
            g = grupos[gi]
            ls.append(' ' * 98 + g['titulo'])
            for j, l in enumerate(g['linhas']):
                if j and g.get('segmentos', True):
                    a, b = g['linhas'][j - 1], l
                    ls.append(' ' * 40 + f'{_curto(b["elemento"]) or "?"} ----> {_curto(a["elemento"]) or "?"} = {br(b["parcial"])} metros')
                ls.append(linha_dados(l))
            ls += ['', '']
        if k == n:
            if notas:
                ls += ['', ' ' * 47 + 'NOTAS'] + [' ' * 17 + f'{el:<13}{txt}' for el, txt in notas]
            ls += ['', '', ' ' * 142 + 'RESUMO']
            ls += [' ' * 143 + f'{d:<27}{br(m):>6}     metros            {c}' for d, m, c in resumo]
            ls += ['', ' ' * 197 + f'TOTAL GERAL         =      {br(total)}      metros', '', '']
            ls += [papeis[0] + ' ' * 75 + papeis[1] + ' ' * 62 + 'Fiscalização: CONCREMAT' + ' ' * 33 + 'Fiscalização: CAESB',
                   'Eng°. [NOME]' + ' ' * 74 + 'Eng° [NOME]' + ' ' * 83 + 'Eng°.' + ' ' * 51 + 'Eng°.',
                   'CREA: [CREA]' + ' ' * 74 + 'CREA: [CREA]' + ' ' * 63 + 'CREA:' + ' ' * 51 + 'CREA:']
        out.append('\n'.join(ls))
    return out


# ------------------------------------------------------------------------------------------------ PDF e planilha
def _pdf_texto(s):
    b = s.encode('cp1252', errors='replace')
    return '(' + ''.join(chr(c) if 32 <= c < 127 and chr(c) not in '()\\' else f'\\{c:03o}' for c in b) + ')'


def pdf_de_texto(paginas, tamanho=4.4, girar=False):
    """PDF (bytes) com a camada de texto das folhas: cada célula (texto separado por 2+ espaços) é um trecho de texto
    na posição da sua coluna, em Courier. girar=True: página em pé com /Rotate 90 (lida deitada), como a NS assinada."""
    import re
    larg, alt = 842, 595
    cw, passo = 0.6 * tamanho, tamanho * 1.45
    objs = ['<< /Type /Catalog /Pages 2 0 R >>', None,
            '<< /Type /Font /Subtype /Type1 /BaseFont /Courier /Encoding /WinAnsiEncoding >>']
    kids = []
    for texto in paginas:
        cmds = ['BT', f'/F1 {tamanho} Tf']
        for i, ln in enumerate(texto.split('\n')):
            y = alt - 12 - i * passo
            for m in re.finditer(r'\S+(?: \S+)*', ln):
                x = 12 + m.start() * cw
                tm = f'0 1 -1 0 {alt - y:.2f} {x:.2f}' if girar else f'1 0 0 1 {x:.2f} {y:.2f}'
                cmds.append(f'{tm} Tm {_pdf_texto(m.group())} Tj')
        cmds.append('ET')
        corpo = '\n'.join(cmds).encode('latin-1')
        objs.append(f'<< /Length {len(corpo)} >>\nstream\n'.encode('latin-1') + corpo + b'\nendstream')
        caixa = f'[0 0 {alt} {larg}]' if girar else f'[0 0 {larg} {alt}]'
        objs.append(f'<< /Type /Page /Parent 2 0 R /MediaBox {caixa} {"/Rotate 90 " if girar else ""}'
                    f'/Resources << /Font << /F1 3 0 R >> >> /Contents {len(objs)} 0 R >>')
        kids.append(len(objs))
    objs[1] = f'<< /Type /Pages /Kids [{" ".join(f"{k} 0 R" for k in kids)}] /Count {len(kids)} >>'
    out = io.BytesIO()
    out.write(b'%PDF-1.4\n%\xe2\xe3\xcf\xd3\n')
    pos = []
    for k, o in enumerate(objs, 1):
        pos.append(out.tell())
        out.write(f'{k} 0 obj\n'.encode('latin-1') + (o if isinstance(o, bytes) else o.encode('latin-1')) + b'\nendobj\n')
    xref = out.tell()
    out.write(f'xref\n0 {len(objs) + 1}\n0000000000 65535 f \n'.encode('latin-1'))
    for p in pos:
        out.write(f'{p:010d} 00000 n \n'.encode('latin-1'))
    out.write(f'trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n'.encode('latin-1'))
    return out.getvalue()


def pdf_sem_texto():
    """PDF só com imagem (NS digitalizada)."""
    from PIL import Image
    b = io.BytesIO()
    Image.new('RGB', (120, 80), 'white').save(b, 'PDF', resolution=72.0)
    return b.getvalue()


def xlsx_de_ns(numero='012', contrato='10239/2026', data_ns=date(2026, 10, 15), localizacao='REDE DE DISTRIBUIÇÃO DE ÁGUA '
               '- MORRO DA CRUZ - SÃO SEBASTIÃO/DF', titulo='"REDE DE ÁGUA" - ZONA 03 - PB08 - VER CROQUI ANEXO - VER PROJETO '
               'EXECUTIVO', bloco='TOPOGRAFIA', grupos=(), notas=(), resumo=None, total=None):
    """Planilha .xlsx com as mesmas células da NS (uma aba = uma folha), como a engenharia recebe do Excel da Caesb."""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'NS'
    if resumo is None:
        resumo = [(f'TUBO PEAD {g["linhas"][0].get("material", "PEAD-160").split("-")[-1]}mm',
                   round(sum(l['parcial'] or 0 for l in g['linhas']), 2), g['titulo']) for g in grupos]
    if total is None:
        total = round(sum(r[1] for r in resumo), 2)

    def put(r, c, v, fmt=None):
        cel = ws.cell(row=r, column=c, value=v)
        if fmt:
            cel.number_format = fmt

    put(1, 4, 'Companhia de Saneamento Ambiental do Distrito Federal - Caesb')
    put(1, 14, 'NS nº :')
    put(1, 15, numero)
    put(3, 1, 'CONTRATO:')
    put(3, 2, contrato)
    put(3, 14, 'FOLHA:')
    put(3, 15, '01 de 01')
    put(4, 1, 'CONTRATADA: CONSTRUTORA ANKARA ENGENHARIA LTDA')
    put(5, 14, 'DATA:')
    put(5, 15, data_ns, 'dd/mm/yyyy')
    put(6, 1, 'LOCALIZAÇÃO: ' + localizacao)
    put(8, 4, titulo)
    put(10, 4, bloco)
    put(11, 5, 'COTAS')
    for c, v in ((1, 'ESTACA'), (9, 'MATERIAL'), (10, 'DISTÂNCIA'), (12, 'ESCAVAÇÃO'), (14, 'VISOR'), (15, 'CRUZETA'), (16, 'OBSERVAÇÕES')):
        put(12, c, v)
    for c, v in ((4, 'ELEMENTOS'), (5, 'TERRENO'), (6, '"GI" GERATRIZ INFERIOR DO TUBO'), (7, 'PROFUNDIDADE'),
                 (8, 'MÉTODO EXECUTIVO')):
        put(13, c, v)
    for c, v in ((1, 'Inteira'), (2, '+'), (3, 'Fração'), (9, '"D" (mm)'), (10, 'PARCIAL'), (11, 'ACUMULADA'), (12, 'BOCA (m)'),
                 (13, 'FUNDO (m)')):
        put(14, c, v)
    r = 16
    for g in grupos:
        put(r, 4, g['titulo'])
        r += 1
        for j, l in enumerate(g['linhas']):
            if j and g.get('segmentos', True):
                a = g['linhas'][j - 1]
                put(r, 4, f'{_curto(l["elemento"]) or "?"} ----> {_curto(a["elemento"]) or "?"} = {br(l["parcial"])} metros')
                r += 1
            i, f = estaca(l['estaca_m'])
            prof = round(l['terreno'] - l['gi'], 2)
            vals = [(1, i, '0'), (2, '+', None), (3, round(l['estaca_m'] - 20 * i, 2), '0.00'), (4, l['elemento'], None),
                    (5, l['terreno'], '#,##0.000'), (6, l['gi'], '#,##0.000'), (7, prof, '0.00'), (8, l.get('metodo', 'V.C.A'), None),
                    (9, l.get('material', 'PEAD-160'), None), (10, l['parcial'], '0.00'), (11, l['acumulada'], '0.00'),
                    (12, l.get('boca'), '0.00'), (13, l.get('fundo'), '0.00'), (16, l.get('obs'), None)]
            for c, v, fmt in vals:
                if v not in (None, ''):
                    put(r, c, v, fmt)
            r += 1
        r += 1
    if notas:
        put(r, 4, 'NOTAS')
        r += 1
        for el, txt in notas:
            put(r, 2, el)
            put(r, 4, txt)
            r += 1
    put(r + 1, 9, 'RESUMO')
    r += 2
    for d, m, cc in resumo:
        put(r, 9, d)
        put(r, 10, m, '0.00')
        put(r, 11, 'metros')
        put(r, 12, cc)
        r += 1
    put(r + 1, 12, 'TOTAL GERAL')
    put(r + 1, 13, '=')
    put(r + 1, 14, total, '0.00')
    put(r + 1, 15, 'metros')
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


# ------------------------------------------------------------------------------------------------ .xls (Excel 97-2003)
def xls_de_ns(limite_registro=8224, variado=True, **kw):
    """A mesma planilha de xlsx_de_ns em .xls (BIFF8), como o Excel da Caesb ainda salva. `limite_registro` pequeno
    parte a tabela de textos em vários registros (CONTINUE); `variado` usa os outros registros do Excel: números
    seguidos em MULRK, a ACUMULADA como fórmula (com o resultado guardado), 'metros' como fórmula de texto e as
    observações como LABEL."""
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_de_ns(**kw)))
    abas = [(ws.title, {(c.row - 1, c.column - 1): (c.value, c.number_format) for row in ws.iter_rows() for c in row
                        if c.value is not None}) for ws in wb.worksheets]
    return xls_de_celulas(abas, limite_registro, variado)


def _reg(tipo, dados):
    return struct.pack('<HH', tipo, len(dados)) + dados


def _largo(s):
    return any(ord(ch) > 255 for ch in s)


def _texto_biff(s, cch_curto=False):
    return (struct.pack('<B' if cch_curto else '<H', len(s)) + bytes([_largo(s)])
            + s.encode('utf-16-le' if _largo(s) else 'latin-1'))


def _sst(textos, limite):
    """SST e CONTINUE como o Excel parte: o cabeçalho do texto não se parte; os caracteres, sim, e o pedaço seguinte
    recomeça com o byte da largura (que pode mudar)."""
    regs, cur = [], bytearray(struct.pack('<II', len(textos), len(textos)))
    for s in textos:
        if len(cur) + 3 > limite:
            regs.append(bytes(cur))
            cur = bytearray()
        largo = _largo(s)
        cur += struct.pack('<HB', len(s), largo)
        resto = s
        while resto:
            cabe = (limite - len(cur)) // (2 if largo else 1)
            if cabe == 0:
                regs.append(bytes(cur))
                largo = _largo(resto)
                cur = bytearray([largo])
                continue
            cur += resto[:cabe].encode('utf-16-le' if largo else 'latin-1')
            resto = resto[cabe:]
    regs.append(bytes(cur))
    return _reg(0x00FC, regs[0]) + b''.join(_reg(0x003C, r) for r in regs[1:])


def _rk(v):
    """O número no formato compacto do Excel (None se não couber)."""
    if isinstance(v, int) and -2 ** 29 <= v < 2 ** 29:
        return (v << 2) & 0xFFFFFFFF | 2
    if isinstance(v, float) and abs(v) < 2 ** 28 and round(v * 100) / 100 == v:
        return (round(v * 100) << 2) & 0xFFFFFFFF | 3
    return None


def xls_de_celulas(abas, limite_registro=8224, variado=True):
    """[(nome, {(linha, coluna): (valor, formato)})] → bytes do .xls (arquivo composto com o fluxo Workbook)."""
    from datetime import datetime
    internos = {'General': 0, '0': 1, '0.00': 2, '#,##0': 3, '#,##0.00': 4}
    fmts = sorted({f for _, cel in abas for _, f in cel.values() if f and f != 'General'})
    ifmt = {f: internos.get(f, 164 + i) for i, f in enumerate(fmts)}
    xf = {'General': 0, **{f: i + 1 for i, f in enumerate(fmts)}}
    textos = []
    for _, cel in abas:
        for v, _ in cel.values():
            if isinstance(v, str) and v not in textos:
                textos.append(v)

    def folha(cel):
        pos = sorted(cel)
        n_lin, n_col = (max(r for r, _ in pos) + 1, max(c for _, c in pos) + 1) if pos else (0, 0)
        out = [_reg(0x0809, struct.pack('<HHHHII', 0x0600, 0x0010, 0x0DBB, 0x07CC, 0, 6)),
               _reg(0x0200, struct.pack('<IIHHH', 0, n_lin, 0, n_col, 0))]
        k = 0
        while k < len(pos):
            (r, c), (v, f) = pos[k], cel[pos[k]]
            x = xf.get(f or 'General', 0)
            cab = struct.pack('<HHH', r, c, x)
            if isinstance(v, (datetime, date)):
                v = (datetime(v.year, v.month, v.day) - datetime(1899, 12, 30)).days if not isinstance(v, datetime) else \
                    (v - datetime(1899, 12, 30)).total_seconds() / 86400
                out.append(_reg(0x0203, cab + struct.pack('<d', float(v))))
            elif isinstance(v, str):
                if variado and v == 'metros':
                    out.append(_reg(0x0006, cab + bytes(6) + b'\xff\xff' + struct.pack('<HIH', 0, 0, 0)))
                    out.append(_reg(0x0207, _texto_biff(v)))
                elif variado and c == 15:
                    out.append(_reg(0x0204, cab + _texto_biff(v)))
                else:
                    out.append(_reg(0x00FD, cab + struct.pack('<I', textos.index(v))))
            elif isinstance(v, bool):
                out.append(_reg(0x0205, cab + bytes([v, 0])))
            elif variado and c == 10 and isinstance(v, (int, float)):
                out.append(_reg(0x0006, cab + struct.pack('<dHIH', float(v), 0, 0, 0)))
            else:
                corrida = [k]                    # números seguidos na linha: um MULRK, como o Excel grava
                while variado and corrida[-1] + 1 < len(pos) and pos[corrida[-1] + 1] == (r, c + len(corrida)) \
                        and _rk(cel[pos[corrida[-1] + 1]][0]) is not None and _rk(v) is not None \
                        and not isinstance(cel[pos[corrida[-1] + 1]][0], (str, bool)) and c + len(corrida) != 10:
                    corrida.append(corrida[-1] + 1)
                if len(corrida) > 1:
                    pares = b''.join(struct.pack('<HI', xf.get(cel[pos[j]][1] or 'General', 0), _rk(cel[pos[j]][0]))
                                     for j in corrida)
                    out.append(_reg(0x00BD, struct.pack('<HH', r, c) + pares + struct.pack('<H', c + len(corrida) - 1)))
                    k += len(corrida)
                    continue
                rk = _rk(v)
                out.append(_reg(0x027E, cab + struct.pack('<I', rk)) if rk is not None else
                           _reg(0x0203, cab + struct.pack('<d', float(v))))
            k += 1
        return b''.join(out) + _reg(0x000A, b'')

    folhas = [folha(cel) for _, cel in abas]
    inicio = [_reg(0x0809, struct.pack('<HHHHII', 0x0600, 0x0005, 0x0DBB, 0x07CC, 0, 6)), _reg(0x0022, b'\0\0')]
    inicio += [_reg(0x041E, struct.pack('<H', ifmt[f]) + _texto_biff(f)) for f in fmts if ifmt[f] >= 164]
    inicio += [_reg(0x00E0, struct.pack('<HH', 0, 0) + bytes(16))]
    inicio += [_reg(0x00E0, struct.pack('<HH', 0, ifmt[f]) + bytes(16)) for f in fmts]
    fim = _sst(textos, limite_registro) + _reg(0x000A, b'')
    tam_folhas = sum(len(_reg(0x0085, b'\0' * 6 + _texto_biff(n, True))) for n, _ in abas)
    pos = len(b''.join(inicio)) + tam_folhas + len(fim)
    for (nome, _), f in zip(abas, folhas):
        inicio.append(_reg(0x0085, struct.pack('<IBB', pos, 0, 0) + _texto_biff(nome, True)))
        pos += len(f)
    return composto({'Workbook': b''.join(inicio) + fim + b''.join(folhas)})


def composto(fluxos):
    """Arquivo composto do Office (versão 3, setores de 512 bytes) com os fluxos dados; o fluxo menor que 4096 bytes
    vai para o mini fluxo da raiz, como o Office faz."""
    ss, fim_, livre, setor_fat, nenhum = 512, 0xFFFFFFFE, 0xFFFFFFFF, 0xFFFFFFFD, 0xFFFFFFFF
    setores, fat = [], []

    def cadeia(dados):
        ini, n = len(setores), max(1, -(-len(dados) // ss))
        for k in range(n):
            setores.append(dados[k * ss:(k + 1) * ss].ljust(ss, b'\0'))
            fat.append(ini + k + 1 if k + 1 < n else fim_)
        return ini

    mini, minifat, entradas = bytearray(), [], []
    for nome, dados in fluxos.items():
        if len(dados) >= 4096:
            entradas.append((nome, cadeia(dados), len(dados)))
        else:
            ini, n = len(mini) // 64, max(1, -(-len(dados) // 64))
            mini += dados.ljust(n * 64, b'\0')
            minifat += [ini + k + 1 if k + 1 < n else fim_ for k in range(n)]
            entradas.append((nome, ini, len(dados)))
    raiz = cadeia(bytes(mini)) if mini else fim_
    n0 = len(setores)
    minifat1 = cadeia(struct.pack(f'<{len(minifat)}I', *minifat)) if minifat else fim_
    n_minifat = len(setores) - n0

    def entrada(nome, tipo, ini, tam, direita=nenhum, filho=nenhum):
        n = (nome + '\0').encode('utf-16-le')
        return (n.ljust(64, b'\0') + struct.pack('<HBB', len(n), tipo, 1) + struct.pack('<III', nenhum, direita, filho)
                + bytes(36) + struct.pack('<IQ', ini, tam))
    dire = entrada('Root Entry', 5, raiz, len(mini), filho=1 if entradas else nenhum)
    for i, (nome, ini, tam) in enumerate(entradas):
        dire += entrada(nome, 2, ini, tam, direita=i + 2 if i + 1 < len(entradas) else nenhum)
    while len(dire) % ss:
        dire += bytes(64) + struct.pack('<HBB', 0, 0, 0) + struct.pack('<III', nenhum, nenhum, nenhum) + bytes(48)
    dir1 = cadeia(bytes(dire))
    n_fat = 1
    while n_fat * (ss // 4) < len(setores) + n_fat:
        n_fat += 1
    fat_ini = len(setores)
    fat += [setor_fat] * n_fat
    fat += [livre] * (n_fat * (ss // 4) - len(fat))
    setores += [struct.pack(f'<{ss // 4}I', *fat[k * (ss // 4):(k + 1) * (ss // 4)]) for k in range(n_fat)]
    difat = [fat_ini + k for k in range(n_fat)] + [livre] * (109 - n_fat)
    cab = (b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1' + bytes(16) + struct.pack('<HHHHH', 0x3E, 3, 0xFFFE, 9, 6) + bytes(6)
           + struct.pack('<IIIIIIIII', 0, n_fat, dir1, 0, 4096, minifat1, n_minifat, fim_, 0) + struct.pack('<109I', *difat))
    return cab + b''.join(setores)


# ------------------------------------------------------------------------------------------------ casos desta obra
def ns_zona03_pb08(**kw):
    """NS de água desta obra: NÓ6.1 → NÓ9.1 (TR020, 277,99 m) → CH_VRP-15 (TR007 + TR006, 295,98 m, passando pelo
    cruzamento NÓ113.1). Devolve os argumentos de paginas_ns/xlsx_de_ns."""
    g = grupo_padrao('REDE ZONA 03 - PB08 - PEAD DE 160 (NÓ6.1 - NÓ9.1 - VRP-15)',
                     ['NÓ6.1 - TÊ 160x110', 'NÓ9.1 - TÊ 160x160', 'CH_VRP-15 (CAIXA DA VRP)'], [277.99, 295.98])
    args = {'grupos': [g], 'notas': [('NÓ9.1', 'TÊ PEAD DE 160x160 ELETROFUSÃO'), ('VRP-15', 'CAIXA DE VRP "PADRÃO CAESB"')]}
    args.update(kw)
    return args
