"""Módulo E — medição de TESTE: parcela única entre medições, elegibilidade por período, conferência interna antes do
fechamento, regras bloqueadas e exportação. Planilhas SINTÉTICAS. Nenhuma medição oficial; nenhuma aprovação preenchida."""
import io
import os
import sqlite3
import sys
import unittest

from openpyxl import load_workbook

AQUI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(AQUI))
sys.path.insert(0, AQUI)
import exportar  # noqa: E402
import medicao as M  # noqa: E402
import regras as R  # noqa: E402
import servicos as S  # noqa: E402
from fotos_teste import jpeg  # noqa: E402
from test_acesso import BaseHTTP  # noqa: E402
from test_campo import TR, BaseCampo, D  # noqa: E402

TR2 = "Zona01_PB01_Rede_TR002"
MND = "Zona02_PB04_Rede_TR036"
TUBO = "CAESB-PEAD-DE63"
TUDO = dict.fromkeys(R.CHECKLIST_CONFERENCIA_MEDICAO, True)


class BaseMedicao(BaseCampo):
    def setUp(self):
        super().setUp()
        self.m1 = S.criar_medicao(self.c, 1, "2026-10-01", "2026-10-31")
        self.m2 = S.criar_medicao(self.c, 2, "2026-11-01", "2026-11-30")

    def concluir(self, alvo=TR, data=D, fotos=True):
        S.apontar(self.c, data, "Zona01", "TRECHO", alvo, concluido_100=True)
        S.verificar_trecho(self.c, alvo, "sup1")                        # o OK interno exige a verificação (v20)
        if fotos:
            for cat in ("SINALIZACAO", "EXECUCAO", "CONCLUIDO"):
                S.registrar_foto(self.c, "TRECHO", alvo, cat, jpeg(), dir_fotos=self.fotos)

    def lanc(self, mid=None):
        sql = "SELECT servico, status, medicao_id FROM medicao_lancamentos" + (" WHERE medicao_id=?" if mid else "") + " ORDER BY id"
        return self.q(sql, *([mid] if mid else []))


class TestParcelaUnica(BaseMedicao):
    def test_mesma_parcela_nao_entra_em_duas_medicoes(self):
        self.concluir()
        S.ok_fiscalizacao(self.c, "TRECHO", TR, "eng1", D)
        antes = self.lanc()
        self.assertTrue(antes and all(m == self.m1 for _, _, m in antes))
        tid = S.alvo(self.c, "TRECHO", TR)["id"]
        bm = S._bm_item(self.c, "2.2_5.8.44")
        with self.assertRaisesRegex(S.RegraViolada, "Parcela já medida: ASSENTAMENTO .* 1ª medição"):
            with S.transacao(self.c):
                S._lancar(self.c, self.m2, 1, "ASSENTAMENTO", bm, "M", 76.55, "TRECHO", tid)
        outro = S._bm_item(self.c, "2.2_5.2.1")                       # mesmo alvo × serviço com outro item (revisão do BM)
        with self.assertRaisesRegex(S.RegraViolada, "Parcela já medida"):
            with S.transacao(self.c):
                S._lancar(self.c, self.m2, 1, "ASSENTAMENTO", outro, "M", 76.55, "TRECHO", tid)
        with self.assertRaisesRegex(S.RegraViolada, "já marcado com OK"):
            S.ok_fiscalizacao(self.c, "TRECHO", TR, "eng1", D)              # clique repetido
        ch = self.q("SELECT chave FROM medicao_lancamentos WHERE servico='ASSENTAMENTO'")[0][0]
        with self.assertRaises(sqlite3.IntegrityError):                # trava no banco, além da regra
            self.c.execute("INSERT INTO medicao_lancamentos(medicao_id,gaveta,servico,unidade,quantidade,chave) "
                           "VALUES (?,1,'ASSENTAMENTO','M',1,?)", (self.m2, ch))
        self.assertEqual(self.lanc(), antes)
        self.assertEqual(M.duplicidades(self.c), [])
        acum = {i["item"]: i for i in M.acumulado_itens(self.c, self.m1)}["2.2_5.8.44"]
        self.assertEqual((acum["anterior"], acum["nesta"]), (0, 76.55))
        self.assertAlmostEqual(acum["saldo"], acum["contrato"] - 76.55, places=4)
        self.assertEqual({i["item"]: i for i in M.acumulado_itens(self.c, self.m2)}["2.2_5.8.44"]["anterior"], 76.55)

    def test_lancamento_em_medicao_fechada_e_recusado(self):
        self.concluir()
        S.ok_fiscalizacao(self.c, "TRECHO", TR, "eng1", D)
        M.conferir(self.c, self.m1, "CONFORME", TUDO, "eng2")
        S.fechar_medicao(self.c, self.m1, "eng2")
        with self.assertRaisesRegex(S.RegraViolada, "já está fechada"):
            S.recomposicao_ok(self.c, "TRECHO", TR, "CALÇADA", 10, D)
        self.assertEqual(self.q("SELECT recomposicao_ok FROM execucao WHERE alvo_tipo='TRECHO'")[0][0], 0)   # nada gravado


class TestElegibilidade(BaseMedicao):
    def test_proposto_apto_parcial_e_material_bloqueado(self):
        self.concluir(TR)
        S.ok_fiscalizacao(self.c, "TRECHO", TR, "eng1", D)
        self.concluir(TR2, fotos=False)                                # completo, sem fotos: apto sem OK, com impedimento
        S.apontar(self.c, D, "Zona02", "TRECHO", MND, intervalos_novos=[{"etapa": "ASSENTADO", "ini_m": 0, "fim_m": 10}])
        iv = self.q("SELECT id FROM intervalos_execucao WHERE trecho_id=(SELECT id FROM trechos WHERE ponto=?)", MND)[0][0]
        S.verificar_intervalo(self.c, iv, "eng2")
        S.abrir_ocorrencia(self.c, "TRECHO", MND, "IMPEDIMENTO", "INTERFERENCIA", "rede de esgoto no traçado", D, "enc1")
        S.entrada_material(self.c, D, TUBO, 100, "ROMANEIO_CAESB", "ROM-1", "almox1")
        ck = S.checkin(self.c, D, "Zona01", "TRECHO", TR, materiais=[{"material_id": TUBO, "quantidade": 30}])
        S.checkout(self.c, ck, [{"material_id": TUBO, "aplicado": 24, "sobra": 6}])
        el = M.elegibilidade(self.c, self.m1)
        sit = {l["nome"]: l for l in el["linhas"]}
        self.assertEqual(sit[TR]["situacao"], "PROPOSTO_NESTA")
        self.assertEqual(sit[TR2]["situacao"], "APTO_SEM_OK")
        self.assertIn("fotos obrigatórias", sit[TR2]["impedimentos"][0])
        self.assertEqual(sit[MND]["situacao"], "PARCIAL_NAO_ELEGIVEL")
        self.assertIn("D-30", sit[MND]["impedimentos"][0])
        self.assertEqual((sit[MND]["assentado_periodo_m"], sit[MND]["verificado_periodo_m"]), (10, 10))
        self.assertIn("D-22", sit[MND]["alertas"][0])
        self.assertEqual(el["contagem"], {"PROPOSTO_NESTA": 1, "OK_SEM_MEDICAO": 0, "APTO_SEM_OK": 1, "PARCIAL_NAO_ELEGIVEL": 1,
                                         "MEDIDO_ANTES": 0})
        self.assertEqual(el["resumo"]["proposto_nesta_m"], 76.55)
        mat = {m["material_id"]: m for m in el["materiais"]}[TUBO]
        self.assertEqual((mat["recebido"], mat["aplicado"], mat["situacao"]), (100, 24, "BLOQUEADO"))
        self.assertFalse(self.q("SELECT 1 FROM medicao_lancamentos WHERE material_id IS NOT NULL OR gaveta=2"))   # nada de material
        nov = {l["nome"]: l["situacao"] for l in M.elegibilidade(self.c, self.m2)["linhas"]}
        self.assertNotIn(TR, nov)                                      # medido em outubro, sem atividade em novembro
        self.assertEqual(nov[TR2], "APTO_SEM_OK")                      # continua aguardando o OK


class TestConferenciaEFechamento(BaseMedicao):
    def test_conferencia_obrigatoria_valida_so_para_os_lancamentos_conferidos(self):
        self.concluir()
        S.ok_fiscalizacao(self.c, "TRECHO", TR, "eng1", D)
        with self.assertRaisesRegex(S.RegraViolada, "exige a conferência interna"):
            S.fechar_medicao(self.c, self.m1, "eng2")
        with self.assertRaisesRegex(S.RegraViolada, "CONFORME exige todos os itens"):
            M.conferir(self.c, self.m1, "CONFORME", {"EXECUTADO": True}, "eng2")
        with self.assertRaisesRegex(S.RegraViolada, "Descreva o que não confere"):
            M.conferir(self.c, self.m1, "NAO_CONFORME", {}, "eng2")
        with self.assertRaisesRegex(S.RegraViolada, "Nada lançado"):
            M.conferir(self.c, self.m2, "NAO_CONFORME", {}, "eng2", "vazia")
        M.conferir(self.c, self.m1, "NAO_CONFORME", {"EXECUTADO": True}, "eng2", "escavação em rocha a conferir")
        with self.assertRaisesRegex(S.RegraViolada, "NÃO CONFORME"):
            S.fechar_medicao(self.c, self.m1, "eng2")
        # evidência perdida depois do OK: CONFORME recusado
        self.c.execute("UPDATE fotos SET validacao='ARQUIVO_AUSENTE' WHERE categoria='CONCLUIDO'")
        with self.assertRaisesRegex(S.RegraViolada, "evidência fotográfica obrigatória"):
            M.conferir(self.c, self.m1, "CONFORME", TUDO, "eng2")
        self.c.execute("UPDATE fotos SET validacao='EXIF_OK' WHERE categoria='CONCLUIDO'")
        cid = M.conferir(self.c, self.m1, "CONFORME", TUDO, "eng2")
        ret = S.conferencia_vigente(self.c, self.m1)["retrato"]
        ev = ret["alvos"][0]
        self.assertEqual((ev["nome"], ev["ok_interno_por"], ev["fotos_validas"], ev["fotos_faltando"]), (TR, "eng1", 3, []))
        self.assertGreater(ret["pendentes"], 0)                       # pendentes identificados, fora do valor
        self.assertTrue(any(r["decisao"] == "D-31" for r in ret["regras_bloqueadas"]))
        S.recomposicao_ok(self.c, "TRECHO", TR, "CALÇADA", 10, D)       # lançamento novo depois da conferência
        self.assertFalse(S.conferencia_vigente(self.c, self.m1)["vale"])
        with self.assertRaisesRegex(S.RegraViolada, "mudaram depois da conferência"):
            S.fechar_medicao(self.c, self.m1, "eng2")
        cid2 = M.conferir(self.c, self.m1, "CONFORME", TUDO, "eng2")
        S.fechar_medicao(self.c, self.m1, "eng2")
        self.assertEqual(self.q("SELECT status, fechada_por, conferencia_id FROM medicoes WHERE id=?", self.m1), [("FECHADA", "eng2", cid2)])
        self.assertEqual(self.q("SELECT id, resultado FROM conferencias_medicao ORDER BY id"),
                         [(cid - 1, "NAO_CONFORME"), (cid, "CONFORME"), (cid2, "CONFORME")])   # histórico preservado
        with self.assertRaisesRegex(S.RegraViolada, "antes do fechamento"):
            M.conferir(self.c, self.m1, "CONFORME", TUDO, "eng2")
        self.assertEqual(self.q("SELECT tipo FROM medicoes"), [("TESTE",), ("TESTE",)])

    def test_regras_bloqueadas_e_nada_de_valor_sem_fonte(self):
        rg = {r["decisao"]: r for r in M.regras(self.c) if r["decisao"]}
        for dec in ("D-31", "D-50", "D-52", "D-53 · D-54", "D-56", "D-16"):
            self.assertEqual(rg[dec]["situacao"], "BLOQUEADA", dec)
        for dec in ("D-55", "D-57"):
            self.assertEqual(rg[dec]["situacao"], "HIPÓTESE", dec)
        # o BM sintético não tem os itens de escavação: o vínculo aparece incompleto (com o BM real, HIPÓTESE do rateio)
        self.assertEqual(rg["D-51"]["situacao"], "VÍNCULO INCOMPLETO")
        self.assertIn("fora do BM vigente: PENDENTE", rg["D-51"]["efeito"])
        self.assertIn("nenhum percentual (80/20 ou outro)", rg["D-53 · D-54"]["efeito"])
        self.assertTrue(any(r["situacao"] == "CONFIRMADA" and r["regra"].startswith("ASSENTAMENTO") for r in M.regras(self.c)))
        S.apontar(self.c, D, "Zona02", "TRECHO", MND, concluido_100=True)
        for cat in ("SINALIZACAO", "EXECUCAO", "CONCLUIDO"):
            S.registrar_foto(self.c, "TRECHO", MND, cat, jpeg(), dir_fotos=self.fotos)
        S.verificar_trecho(self.c, MND, "sup1")
        S.ok_fiscalizacao(self.c, "TRECHO", MND, "eng1", D)
        ass = self.q("SELECT status, bm_item_id, motivo FROM medicao_lancamentos WHERE servico='ASSENTAMENTO'")
        self.assertEqual((ass[0][0], ass[0][1]), ("PENDENTE", None))     # MND não vira vala aberta
        self.assertIn("MND", ass[0][2])
        r = S.resumo_medicao(self.c, self.m1)
        self.assertTrue(r["teste"])
        with self.assertRaisesRegex(S.RegraViolada, "OFICIAL bloqueada"):
            S.criar_medicao(self.c, 3, "2026-12-01", "2026-12-31", tipo="OFICIAL")


class TestMigracaoV13(BaseMedicao):
    def test_duplicidade_antiga_preservada_e_identificada(self):
        self.concluir()
        S.ok_fiscalizacao(self.c, "TRECHO", TR, "eng1", D)
        tid = S.alvo(self.c, "TRECHO", TR)["id"]
        n_antes = self.q("SELECT COUNT(*) FROM medicao_lancamentos")[0][0]
        self.c.close()
        # banco no esquema v12, com a mesma parcela gravada em duas medições pela versão anterior (sem a trava)
        raw = sqlite3.connect(self.db)
        raw.execute("DROP INDEX ux_lanc_chave")
        for col in ("chave",):
            raw.execute(f"ALTER TABLE medicao_lancamentos DROP COLUMN {col}")
        for col in ("fechada_por", "fechada_em", "conferencia_id"):
            raw.execute(f"ALTER TABLE medicoes DROP COLUMN {col}")
        raw.execute("DROP TABLE conferencias_medicao")
        raw.execute("DELETE FROM schema_versao WHERE versao>=12")
        raw.execute("INSERT INTO schema_versao VALUES (12, '2026-10-16', 'v12')")
        raw.execute("INSERT INTO medicao_lancamentos(medicao_id,gaveta,servico,bm_item_id,alvo_tipo,alvo_id,unidade,quantidade,status) "
                    "SELECT ?, gaveta, servico, bm_item_id, alvo_tipo, alvo_id, unidade, quantidade, status FROM medicao_lancamentos "
                    "WHERE servico='ASSENTAMENTO'", (self.m2,))
        raw.commit()
        raw.close()
        rel = self.aplicar()
        self.assertEqual(rel["esquema_antes"], 12)
        c = self.conectar()
        self.assertEqual(c.execute("SELECT COUNT(*) FROM medicao_lancamentos").fetchone()[0], n_antes + 1)   # nada apagado
        self.assertEqual(c.execute("SELECT COUNT(*) FROM medicao_lancamentos WHERE chave IS NULL").fetchone()[0], 1)
        dup = M.duplicidades(c)
        self.assertEqual([(d["servico"], d["alvo_id"], d["medicoes"]) for d in dup], [("ASSENTAMENTO", tid, [1, 2])])
        S.recomposicao_ok(c, "TRECHO", TR, "CALÇADA", 10, D)            # a medição 1 segue utilizável
        with self.assertRaisesRegex(S.RegraViolada, "Parcela em mais de uma medição"):
            M.conferir(c, self.m2, "CONFORME", TUDO, "eng2")
        M.conferir(c, self.m2, "NAO_CONFORME", {}, "eng2", "parcela repetida da 1ª medição")


class TestMedicaoPeloHTTP(BaseHTTP):
    def test_permissao_autoria_fechamento_e_exportacao(self):
        mid = S.criar_medicao(self.c, 1, "2026-10-01", "2026-10-31")
        for t in ("MANHA", "TARDE"):
            S.registrar_clima(self.c, D, "Zona01", t, "SOL")
        S.apontar(self.c, D, "Zona01", "TRECHO", TR, concluido_100=True)
        for cat in ("SINALIZACAO", "EXECUCAO", "CONCLUIDO"):
            S.registrar_foto(self.c, "TRECHO", TR, cat, jpeg(), dir_fotos=self.fotos)
        S.verificar_trecho(self.c, TR, "sup1")
        S.ok_fiscalizacao(self.c, "TRECHO", TR, "eng1", D)
        eng, alm = self.cliente("eng2"), self.cliente("almox1")
        corpo = {"medicao_id": mid, "resultado": "CONFORME", "checklist": TUDO, "por": "Fiscal da Caesb"}
        self.assertEqual(alm.post("/api/medicao/conferir", corpo)[0], 403)
        self.assertEqual(eng.post("/api/medicao/fechar", {"medicao_id": mid})[0], 422)          # sem conferência
        st, r = eng.post("/api/medicao/conferir", corpo)
        self.assertEqual(st, 200, r)
        self.assertEqual(self.q("SELECT por FROM conferencias_medicao"), [("eng2",)])        # autor = sessão, não o digitado
        st, m = eng.get(f"/api/medicao?id={mid}")
        self.assertTrue(m["conferencia"]["vale"])
        self.assertTrue(any(r["decisao"] == "D-56" for r in m["regras"]))
        self.assertEqual(m["roteiro_conferencia"], dict(R.CHECKLIST_CONFERENCIA_MEDICAO))
        st, el = eng.get(f"/api/medicao/elegibilidade?id={mid}")
        self.assertEqual(el["contagem"]["PROPOSTO_NESTA"], 1)
        st, r = eng.post("/api/medicao/fechar", {"medicao_id": mid, "por": "outro"})
        self.assertEqual(st, 200, r)
        self.assertEqual(self.q("SELECT fechada_por FROM medicoes"), [("eng2",)])
        self.assertEqual(self.q("SELECT alvo_tipo, alvo_id FROM auditoria WHERE rota='/api/medicao/fechar' AND status=200"),
                         [("MEDICAO", str(mid))])
        st, xlsx = eng.get(f"/api/medicao/exportar?id={mid}")
        self.assertEqual(st, 200)
        wb = load_workbook(io.BytesIO(xlsx))
        for aba in ("CONFERÊNCIA", "REGRAS", "ELEGIBILIDADE", "PARCELAS"):
            self.assertIn(aba, wb.sheetnames)
        self.assertIn("TESTE — NÃO OFICIAL", wb["REGRAS"]["A1"].value)
        self.assertIn("CONFORME por eng2", wb["CONFERÊNCIA"]["A5"].value)
        nome, _ = exportar.planilha_medicao(self.c, mid)
        self.assertTrue(nome.startswith("TESTE_"))


if __name__ == "__main__":
    unittest.main()
