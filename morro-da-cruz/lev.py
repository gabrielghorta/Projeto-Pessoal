"""LEV: as planilhas da obra como tabelas no sistema — cronograma físico-financeiro Caesb, BM (Juntada e Vertical) e os
levantamentos de redes (06.1) e de conexões (06.2).

* BM, redes e conexões vêm do banco (o que a importação já gravou das planilhas, só o vigente).
* O cronograma vem direto da aba CRONOGRAMA FF CAESB da planilha de reajuste (a mais nova; sem ela, a do BM). Os itens
  são os mesmos nas duas; a de reajuste traz também o bloco do reajuste. Fica em memória até o arquivo mudar.
Só leitura: nada aqui grava.
"""
import os
import re
import threading

import apropriacao as AP
import servicos as S
from servicos import RegraViolada, _rows

ABA_CRONOGRAMA = "CRONOGRAMA FF CAESB"
LINHAS = {"FINANC. PROGRAMADO": "fin_prog", "FAT. PROGRAMADO": "fin_prog", "FINANC. EXECUTADO": "fin_exec",
          "FAT. FINANC. EXECUTADO": "fin_exec", "FÍSICO PROGRAMADO": "fis_prog", "FÍSICO EXECUTADO": "fis_exec"}
_CACHE, _TRAVA = {}, threading.Lock()


def _num(v):
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def _texto(v):
    return re.sub(r"\s+", " ", str(v)).strip() if v not in (None, "") else ""


def _ler_cronograma(caminho):
    from openpyxl import load_workbook
    wb = load_workbook(caminho, read_only=True, data_only=True)
    try:
        if ABA_CRONOGRAMA not in wb.sheetnames:
            raise RegraViolada(f"A planilha {os.path.basename(caminho)} não tem a aba {ABA_CRONOGRAMA}.")
        linhas = [list(r) + [None] * 26 for r in wb[ABA_CRONOGRAMA].iter_rows(values_only=True)]
    finally:
        wb.close()
    cab = next((i for i, r in enumerate(linhas) if _texto(r[5]).upper() == "MÊS 01"), None)
    if cab is None:
        raise RegraViolada(f"Não achei o cabeçalho MÊS 01 … MÊS 18 na aba {ABA_CRONOGRAMA}.")
    cols = [j for j in range(5, 26) if re.fullmatch(r"MÊS \d+", _texto(linhas[cab][j]).upper())]
    meses = [{"n": k + 1, "periodo": _texto(linhas[cab + 1][j]), "dias": _num(linhas[cab + 2][j])} for k, j in enumerate(cols)]
    prazo = next((_num(v) for v in linhas[cab + 2][cols[-1] + 1:cols[-1] + 3] if _num(v)), None)
    blocos, base = [], ""
    for r in linhas[cab + 3:]:
        a, b, d, rot = r[0], _texto(r[1]), _texto(r[3]), _texto(r[4]).upper()
        novo = None
        if a not in (None, ""):
            if isinstance(a, float) and 0 < a < 1:               # BDI, desconto e reajuste: a coluna A é a taxa
                base, novo = b, {"tipo": "taxa", "descricao": f"{b} · {d}" if d else b, "taxa": a}
            elif re.fullmatch(r"\d+(\.\d+)*", _texto(a)):
                novo = {"tipo": "item", "item": _texto(a), "descricao": b, "valor": _num(r[2]), "peso": _num(r[3])}
            else:
                base = " ".join(x for x in (_texto(a), b) if x)
                novo = {"tipo": "total", "descricao": f"{base} · {d}" if d else base}
        elif d and rot and blocos and blocos[-1]["tipo"] != "item":   # o mesmo total, outra medida (ACUMULADO - %)
            novo = {"tipo": "total", "descricao": f"{base} · {d}"}
        if novo:
            novo.setdefault("item", "")
            novo["nivel"] = novo["item"].count(".")
            novo.update(dict.fromkeys(("fin_prog", "fin_exec", "fis_prog", "fis_exec")))
            blocos.append(novo)
        chave = next((v for k, v in LINHAS.items() if rot.startswith(k)), None)
        if chave and blocos:
            blocos[-1][chave] = [_num(r[j]) for j in cols]
    return {"arquivo": os.path.basename(caminho), "meses": meses, "prazo_dias": prazo,
            "itens": [b for b in blocos if b["tipo"] == "item"], "totais": [b for b in blocos if b["tipo"] != "item"]}


def cronograma(caminho):
    """Cronograma físico-financeiro (Caesb) lido da planilha; em memória até o arquivo mudar."""
    if not caminho or not os.path.isfile(caminho):
        raise RegraViolada("A planilha do BM com o cronograma não está no sistema (config/fontes.json).")
    st = os.stat(caminho)
    chave = (caminho, st.st_mtime_ns, st.st_size)
    with _TRAVA:
        if chave not in _CACHE:
            _CACHE.clear()
            _CACHE[chave] = _ler_cronograma(caminho)
        return _CACHE[chave]


def _filtro_texto(q, colunas):
    q = (q or "").strip()
    if not q:
        return "", []
    termo = f"%{q.upper()}%"
    return " AND (" + " OR ".join(f"UPPER(COALESCE({c}, '')) LIKE ?" for c in colunas) + ")", [termo] * len(colunas)


def bm(conn, planilha="JUNTADA", q=None):
    """Itens do BM vigente na ordem da planilha. Juntada: item a item; Vertical: com subgrupo."""
    planilha = (planilha or "JUNTADA").upper()
    if planilha not in ("JUNTADA", "VERTICAL"):
        raise RegraViolada("BM: JUNTADA ou VERTICAL.")
    f, a = _filtro_texto(q, ("item", "item_bm", "descricao", "codigo_caesb", "subgrupo_nome"))
    itens = _rows(conn.execute(f"""SELECT item, item_bm, subgrupo, subgrupo_nome, codigo_caesb, classificacao, descricao, unidade,
                                          quantidade, preco_unitario, valor_total FROM bm_itens
                                   WHERE planilha=? AND vigente=1{f} ORDER BY linha_origem, id""", [planilha, *a]))
    if planilha == "VERTICAL":                       # situação na base de custo (aprovado ou a aprovar) e o código oficial
        sit = {it["item"]: it for it in AP.catalogo(conn, com_reprovados=False) if it["tipo_ref"] == "BM"}
        for i in itens:
            s = sit.get(i["item"]) or {}
            i.update({k: s.get(k) for k in ("situacao", "rotulo", "origem_rotulo", "codigo_oficial", "conciliado")})
    por_classe = {}
    for i in itens:
        por_classe[i["classificacao"] or "—"] = round(por_classe.get(i["classificacao"] or "—", 0) + (i["valor_total"] or 0), 2)
    return {"planilha": planilha, "itens": itens, "total": round(sum(i["valor_total"] or 0 for i in itens), 2),
            "por_classe": por_classe, "bdi": S.parametro(conn, "bdi_material"), "bdi_servico": S.parametro(conn, "bdi_servico")}


def redes(conn, frente=None, q=None, limite=2000):
    """Levantamento de redes (06.1): um trecho por linha e o resumo por diâmetro."""
    f, a = _filtro_texto(q, ("ponto", "pb", "prancha", "codigo_bm", "no_ini", "no_fim"))
    if frente:
        f, a = f + " AND frente_id=?", a + [frente]
    base = f"FROM trechos WHERE vigente=1{f}"
    trechos = _rows(conn.execute(f"""SELECT ponto, frente_id, pb, prancha, de_mm, material, metodologia, classe, extensao_m, codigo_bm,
                                            profundidade_m, no_ini, no_fim {base} ORDER BY item, ponto LIMIT ?""", [*a, int(limite)]))
    resumo = _rows(conn.execute(f"""SELECT de_mm, material, metodologia, COUNT(*) AS trechos, ROUND(SUM(extensao_m), 2) AS extensao_m
                                    {base} GROUP BY de_mm, material, metodologia ORDER BY de_mm, material, metodologia""", a))
    total = conn.execute(f"SELECT COUNT(*), ROUND(COALESCE(SUM(extensao_m), 0), 2) {base}", a).fetchone()
    return {"trechos": trechos, "resumo": resumo, "total_trechos": total[0], "total_m": total[1], "mostrando": len(trechos)}


def conexoes(conn, frente=None, q=None, limite=1500):
    """Levantamento de conexões (06.2): peças por nó/caixa e o resumo por peça."""
    f, a = _filtro_texto(q, ("a.nomenclatura", "i.descricao", "i.codigo_bm", "i.diametro", "e.numero"))
    if frente:
        f, a = f + " AND a.frente_id=?", a + [frente]
    base = f"""FROM itens_projeto i JOIN atividades a ON a.id=i.atividade_id JOIN elementos e ON e.id=a.elemento_id
               WHERE i.vigente=1 AND a.vigente=1{f}"""
    itens = _rows(conn.execute(f"""SELECT a.nomenclatura, a.frente_id, e.numero, e.tipo AS elemento_tipo, i.descricao, i.diametro, i.pressao,
                                          i.material, i.quantidade, i.unidade, i.codigo_bm, i.status_conexao
                                   {base} ORDER BY a.frente_id, a.nomenclatura, i.id LIMIT ?""", [*a, int(limite)]))
    resumo = _rows(conn.execute(f"""SELECT i.descricao, i.diametro, i.unidade, i.codigo_bm, COUNT(DISTINCT a.id) AS nos,
                                           ROUND(SUM(i.quantidade), 3) AS quantidade
                                    {base} GROUP BY i.descricao, i.diametro, i.unidade, i.codigo_bm ORDER BY i.descricao, i.diametro""", a))
    total = conn.execute(f"SELECT COUNT(*), COUNT(DISTINCT a.id) {base}", a).fetchone()
    return {"itens": itens, "resumo": resumo, "total_itens": total[0], "total_nos": total[1], "mostrando": len(itens)}
