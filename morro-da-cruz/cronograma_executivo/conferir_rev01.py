"""Conferências da REV01 antes da entrega (item 11 do pedido de 30/10/2026), sobre o xlsx gerado:

    python3 conferir_rev01.py PASTA_DE_SAIDA

1. cada item do BM fecha em 100% (os que não fecham aparecem com o percentual) e a juntada bate com a soma das equipes;
2. a planilha recalcula sem nenhum erro, com a data do reservatório em branco e com uma data de teste (LibreOffice
   recalcula uma cópia; o arquivo entregue continua com a célula em branco);
3. outubro/2026; 4. os trechos prontos antes de ligados ao reservatório, com o motivo; 5. o quadro REV00 × REV01.
Precisa do LibreOffice (soffice) para o recálculo; sem ele, só confere o que não depende de fórmula."""
import collections
import datetime as dt
import glob
import os
import shutil
import subprocess
import sys
import tempfile

import openpyxl

PASTA = sys.argv[1]
ARQ = glob.glob(os.path.join(PASTA, 'CRONOGRAMA_MORRO_DA_CRUZ_REV01_*.xlsx'))[0]
DATA_TESTE = dt.datetime(2027, 12, 1)
ERROS = ('#VALUE!', '#NAME?', '#REF!', '#DIV/0!', '#N/A', '#NUM!', '#NULL!')
problemas = []


def recalcular(origem, pasta, data=None):
    nome = os.path.join(pasta, ('com_data' if data else 'vazio') + '.xlsx')
    wb = openpyxl.load_workbook(origem)
    wb['PREMISSAS']['C13'] = data
    wb.save(nome)
    saida = os.path.join(pasta, 'recalculado')
    subprocess.run(['soffice', '--headless', '--norestore', '--calc', '--convert-to', 'xlsx', '--outdir', saida, nome],
                   check=True, capture_output=True, timeout=900)
    return openpyxl.load_workbook(os.path.join(saida, os.path.basename(nome)), data_only=True)


def conferir(wb, com_data):
    rot = 'com data de teste' if com_data else 'data em branco'
    erros = collections.Counter()
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and (c.value in ERROS or c.value.startswith('Err:')):
                    erros[ws.title] += 1
    print(f'[{rot}] erros de fórmula:', dict(erros) or 'nenhum')
    if erros:
        problemas.append(f'erros de fórmula ({rot})')
    pr = wb['PROGRAMAÇÃO']
    n = 0
    for row in pr.iter_rows(min_row=4, max_col=18):
        if not isinstance(row[12].value, dt.datetime):
            break
        n += 1
        q, r = row[16].value, row[17].value
        ok = (isinstance(q, dt.datetime) and abs((q - max(DATA_TESTE, row[15].value)).total_seconds()) < 61
              and isinstance(r, (int, float))) if com_data else (q == 'após a operação do reservatório' and r in (None, ''))
        if not ok:
            problemas.append(f'PROGRAMAÇÃO {row[5].value}: ENTRADA EM CARGA {q!r}, DIAS SEM CARGA {r!r} ({rot})')
            break
    print(f'[{rot}] PROGRAMAÇÃO: {n} trechos com ENTRADA EM CARGA e DIAS SEM CARGA conferidos')
    # BM: % de cada item e juntada × soma das equipes
    bm, bc = wb['BM POR EQUIPE'], wb['BM CONSOLIDADO']
    cab = [c.value for c in bm[5]]
    iaq = len(cab) - 1 - cab[::-1].index('QUANT.')
    soma = collections.defaultdict(float)
    for row in bm.iter_rows(min_row=6, max_col=len(cab)):
        if row[3].value and '_' in str(row[3].value) and isinstance(row[18].value, (int, float)):
            soma[str(row[4].value)] += row[iaq].value or 0
    cab = [c.value for c in bc[5]]
    iniv, ipc = cab.index('NÍVEL'), cab.index('% PROGRAMADO')
    iaqc = len(cab) - 1 - cab[::-1].index('QUANT.')
    itens, dif, fora = 0, [], []
    for row in bc.iter_rows(min_row=6, max_col=len(cab)):
        if row[iniv].value == 'ITEM':
            itens += 1
            it = str(row[0].value)
            if abs((row[iaqc].value or 0) - soma.get(it, 0)) > 1e-6:
                dif.append(it)
            if row[14].value and abs((row[ipc].value or 0) - 1) > 1e-6:
                fora.append(f'{it} {row[1].value[:40]}: {100 * row[ipc].value:.2f}%')
    print(f'[{rot}] BM: {itens} itens; juntada diferente da soma das equipes: {len(dif)}; fora de 100%: {fora or "nenhum"}')
    if dif:
        problemas.append(f'juntada diferente da soma das equipes: {dif[:5]}')
    return fora


wb0 = openpyxl.load_workbook(ARQ)
P = wb0['PREMISSAS']
if P['C13'].value is not None:
    problemas.append('a data do reservatório (PREMISSAS!C13) não está em branco')
quadro = any(P.cell(r, 1).value == 'REV00 × REV01 (CENÁRIO B, 90 m/dia)' for r in range(1, 40))
print('PREMISSAS: data do reservatório em branco:', P['C13'].value is None, '| quadro REV00 × REV01:', quadro)
if not quadro:
    problemas.append('falta o quadro REV00 × REV01 nas PREMISSAS')
pr = wb0['PROGRAMAÇÃO']
cats = collections.Counter()
for row in pr.iter_rows(min_row=4, max_col=20):
    s = row[18].value
    if isinstance(s, str) and s.startswith('pronto antes de ligado'):
        cats[s.split(' · ')[-1]] += 1
print('Prontos antes de ligados ao reservatório, por motivo:', dict(cats))
if shutil.which('soffice'):
    with tempfile.TemporaryDirectory() as tmp:
        fora = conferir(recalcular(ARQ, tmp), False)
        wbd = recalcular(ARQ, tmp, DATA_TESTE)
        conferir(wbd, True)
        out = wbd['RESUMO POR EQUIPE']['B13'].value
        print(f'Outubro/2026 (RESUMO POR EQUIPE, total do mês): R$ {out:,.2f}')
        if round(out / 1000) != 103:
            problemas.append(f'outubro/2026 = R$ {out:,.2f}, não R$ 103 mil')
else:
    print('Sem o LibreOffice (soffice): o recálculo não foi conferido.')
    problemas.append('recálculo não conferido (sem LibreOffice)')
print('PROBLEMAS:', problemas or 'nenhum')
sys.exit(1 if problemas else 0)
