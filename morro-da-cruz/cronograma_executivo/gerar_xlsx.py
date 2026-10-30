"""Planilha do cronograma físico executivo base no modelo do BM do Morro da Cruz, dividida por equipe."""
import collections
import datetime as dt
import math

import openpyxl
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as COL
from openpyxl.worksheet.properties import PageSetupProperties

import motor
from motor import CAL, NOME, PRAZO, REAJUSTE

FQ = '_-* #,##0.00_-;\\-* #,##0.00_-;_-* \\-??_-;_-@_-'
FR = '_-"R$ "* #,##0.00_-;"-R$ "* #,##0.00_-;_-"R$ "* \\-??_-;_-@_-'
FP = '0.00%;-0.00%;\\-'
FD = 'dd/mm/yyyy'
FONTE = 'Calibri'
VERDE, CREME, AMARELO, AZUL, LARANJA = '92D050', 'FFFFCC', 'FFFF00', '00B0F0', 'F4B183'
fill = lambda c: PatternFill('solid', fgColor=c)
fino = Side(style='thin', color='BFBFBF')
BORDA = Border(left=fino, right=fino, top=fino, bottom=fino)
OBJETO = ('IMPLANTAÇÃO DE REDES DE DISTRIBUIÇÃO NO MORRO DA CRUZ, SÃO SEBASTIÃO/DF - LC 014/2026 - '
          'JUNTADA 1426 (RED.SSB-D014)')
BLOCOS = [('CANT', 'CANTEIRO DE OBRAS (ANKARA)'), ('ANK1', 'EQUIPE ADUTORA (ANKARA 1)'),
          ('E1A', 'EMPREITEIRO 1 - EQUIPE A'), ('E1B', 'EMPREITEIRO 1 - EQUIPE B'),
          ('E2C', 'EMPREITEIRO 2 - EQUIPE C'), ('E2D', 'EMPREITEIRO 2 - EQUIPE D'),
          ('ANK2', 'EQUIPE CAIXAS E HIDROMECÂNICO (ANKARA 2)'), ('ADM', 'ADMINISTRAÇÃO LOCAL (ANKARA)')]
GRUPOS = ['1.1', '1.2', '2.1', '2.2', '3.1', '3.2', '4.1', '4.2', '5.1', '5.2', '6.2']
MES = {1: 'JAN', 2: 'FEV', 3: 'MAR', 4: 'ABR', 5: 'MAI', 6: 'JUN', 7: 'JUL', 8: 'AGO', 9: 'SET', 10: 'OUT', 11: 'NOV', 12: 'DEZ'}
CORES = ['2a78d6', 'eb6834', '1baf7a', 'eda100', 'e87ba4', '008300', '6250d6', 'e34948']   # paleta validada


def f(bold=False, sz=9, cor='000000', it=False):
    return Font(name=FONTE, size=sz, bold=bold, color=cor, italic=it)


def arred(vals, alvo, nd=4):
    """Arredonda cada valor para nd casas mantendo a soma = alvo (maior resto)."""
    m = 10 ** nd
    alvo_i = round(alvo * m)
    base = [math.floor(v * m + 1e-7) for v in vals]
    falta = alvo_i - sum(base)
    ordem = sorted(range(len(vals)), key=lambda i: -(vals[i] * m - base[i]))
    i = 0
    while falta > 0 and ordem:
        base[ordem[i % len(ordem)]] += 1
        falta -= 1
        i += 1
    ordem_neg = sorted(range(len(vals)), key=lambda i: (vals[i] * m - base[i]))
    i = 0
    while falta < 0 and any(base):
        k = ordem_neg[i % len(ordem_neg)]
        if base[k] > 0:
            base[k] -= 1
            falta += 1
        i += 1
    return [b / m for b in base]


def quantidades(itens, livro):
    """(equipe, chave) -> [q por mês] em 4 casas; a soma das equipes é a quantidade do contrato."""
    nm = len(livro.meses)
    por_item = collections.defaultdict(lambda: collections.defaultdict(lambda: [0.0] * nm))
    for (eq, ch, k), v in livro.q.items():
        por_item[ch][eq][k] += v
    out = {}
    for it in itens:
        if it['grp'] == '6.2' or it['q'] == 0:
            continue
        eqs = por_item[it['chave']]
        nomes = [b for b, _ in BLOCOS if b in eqs]
        programado = sum(sum(eqs[e]) for e in nomes)          # o que o cronograma prevê (o aluguel de contêiner fica abaixo do BM)
        tot_eq = arred([sum(eqs[e]) for e in nomes], min(it['q'], programado))
        for e, t in zip(nomes, tot_eq):
            if t > 0:
                out[(e, it['chave'])] = arred(eqs[e], t)
    return out


def cab_meses(ws, meses, c0, n_prazo=18):
    for k, (n, ini, fim) in enumerate(meses):
        c = c0 + 2 * k
        apos = ini > PRAZO
        ws.cell(1, c, 'MED. Nº.:').font = f(True)
        ws.cell(2, c, f'{n}ª MEDIÇÃO' + (' (APÓS O PRAZO)' if apos else '')).font = f(True, 11 if not apos else 9)
        ws.merge_cells(start_row=2, start_column=c, end_row=2, end_column=c + 1)
        ws.cell(3, c, 'PERÍODO:').font = f(True)
        ws.cell(4, c, f'{ini:%d-%m-%y} A {fim:%d-%m-%y}').font = f(True)
        ws.merge_cells(start_row=4, start_column=c, end_row=4, end_column=c + 1)
        for r in (2, 4):
            ws.cell(r, c).alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            if apos:
                ws.cell(r, c).fill = fill(LARANJA)


def aba_bm(wb, cen, itens, Q, livro, origem):
    ws = wb.create_sheet('BM POR EQUIPE')
    meses = livro.meses
    nm = len(meses)
    c0 = 22                                           # V: primeira coluna QUANT. das medições
    cq = [c0 + 2 * k for k in range(nm)]
    ct = [c + 1 for c in cq]
    cAQ, cAT, cPC, cCR = c0 + 2 * nm, c0 + 2 * nm + 1, c0 + 2 * nm + 2, c0 + 2 * nm + 3
    ws['A1'] = 'CRONOGRAMA FÍSICO EXECUTIVO BASE'
    ws['A1'].font = f(True, 14)
    ws['A2'] = cen['titulo']
    ws['A2'].font = f(True, 11, 'C00000')
    ws['E1'] = 'BM - PLANILHA DE MEDIÇÃO - ORÇAMENTOS PARCIAIS POR EQUIPE (PROGRAMADO)'
    ws['E1'].font = f(True, 16, 'FF0000')
    ws['E3'] = 'OBRA:'
    ws['E3'].font = f(True, 11)
    ws['E4'] = OBJETO
    ws['E4'].font = f(False, 10)
    ws['R3'] = 'DATA BASE - ORÇAMENTO:'
    ws['R3'].font = f(True)
    ws['R4'] = '01/2025'
    ws['R4'].font = f(True, 11, 'FF0000')
    ws['T1'] = 'CONTRATO:'
    ws['T1'].font = f(True, 10)
    ws['T2'] = '10.239/2026'
    ws['T2'].font = f(True)
    ws['T3'] = 'ORIGINAL - P0 (CAESB) - PV'
    ws['T3'].font = f(True)
    ws['T3'].fill = fill(AZUL)
    ws.merge_cells('T3:U3')
    cab_meses(ws, meses, c0)
    ws.cell(2, cAQ, 'ACUMULADO PROGRAMADO').font = f(True, 11)
    ws.merge_cells(start_row=2, start_column=cAQ, end_row=2, end_column=cAT)
    ws.cell(3, cAQ, 'ELAB. POR:').font = f(True, 10)
    ws.cell(4, cAQ, 'ANKARA ENGENHARIA LTDA').font = f(False, 10)
    for r in (3, 4):
        ws.cell(r, cAQ).fill = fill(CREME)
    cab = ['SEQUÊNCIA', 'CÓD. \nCRON. FF', 'DESCRIÇÃO DO ORÇAMENTO', 'ITEM - BM VERTICAL', 'ITEM - BM', 'DESCRIÇÃO',
           'PRESSÃO', 'RESINA', 'TIPO_CONEXÃO', 'SOLDA', 'DIÂMETRO', 'MATERIAL', 'GRUPO', 'GRUPO ORÇAMENTO ANKARA',
           'CÓDIGO CAESB', 'CLASSIFICAÇÃO', 'DESCRIÇÃO SUMÁRIA', 'UNID.', 'QUANT. DA EQUIPE', 'VALOR UNITÁRIO', 'TOTAL']
    for c, t in enumerate(cab, 1):
        ws.cell(5, c, t)
    for k in range(nm):
        ws.cell(5, cq[k], 'QUANT.')
        ws.cell(5, ct[k], 'TOTAL')
    ws.cell(5, cAQ, 'QUANT.')
    ws.cell(5, cAT, 'TOTAL ACUMULADO')
    ws.cell(5, cPC, '% DA EQUIPE PROGRAMADO')
    ws.cell(5, cCR, 'CRITÉRIO DE DISTRIBUIÇÃO (EQUIPE E MÊS)')
    for c in range(1, cCR + 1):
        cel = ws.cell(5, c)
        cel.font = f(True)
        cel.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cel.border = BORDA
        if c <= 3:
            cel.fill = fill(AMARELO)
        elif c in (cAQ, cAT):
            cel.fill = fill(CREME)
    ws.row_dimensions[5].height = 42
    por_it = {it['chave']: it for it in itens}
    ordem_it = {it['chave']: n for n, it in enumerate(itens)}
    adm = next(it for it in itens if it['grp'] == '6.2')

    r = 6
    blocos_linha = {}
    linhas_item = []
    for seq, (eq, nome) in enumerate(BLOCOS, 1):
        rb = r
        blocos_linha[eq] = rb
        for c in range(1, cCR + 1):
            ws.cell(rb, c).fill = fill(VERDE)
            ws.cell(rb, c).font = f(True)
        ws.cell(rb, 1, seq)
        ws.cell(rb, 3, f'{seq} - {nome}')
        ws.cell(rb, 17, f'{seq} - {nome}')
        r += 1
        if eq == 'ADM':
            subgrupos = [(('6.2', 'ADMINISTRAÇÃO LOCAL'), [adm['chave']])]
        else:
            chaves = sorted((ch for (e, ch) in Q if e == eq), key=lambda ch: ordem_it[ch])
            sg = collections.OrderedDict()
            for ch in sorted(chaves, key=lambda ch: (GRUPOS.index(por_it[ch]['grp']), ordem_it[ch])):
                it = por_it[ch]
                sg.setdefault((it['grp'], it['orc']), []).append(ch)
            subgrupos = list(sg.items())
        linhas_sg = []
        for (grp, orc), chaves in subgrupos:
            rs = r
            linhas_sg.append(rs)
            for c in range(1, cCR + 1):
                ws.cell(rs, c).fill = fill(CREME)
                ws.cell(rs, c).font = f(True, 9, 'FF0000')
            ws.cell(rs, 1, seq)
            ws.cell(rs, 3, f'{grp} - {orc}')
            ws.cell(rs, 17, f'{grp} - {orc}')
            r += 1
            r0 = r
            for ch in chaves:
                it = por_it[ch]
                linhas_item.append(r)
                ws.cell(r, 1, seq)
                ws.cell(r, 2, it['grp'])
                ws.cell(r, 3, it['orc'])
                ws.cell(r, 4, f"{eq}_{ch}")
                ws.cell(r, 5, it['item'])
                for j, v in enumerate(it['attrs']):
                    ws.cell(r, 6 + j, v)
                ws.cell(r, 15, str(it['cod'] or '')).number_format = '@'
                ws.cell(r, 16, it['classe'])
                ws.cell(r, 17, it['desc'])
                ws.cell(r, 18, it['un'])
                if eq == 'ADM':
                    ws.cell(r, 19, it['q'])
                else:
                    qs = Q[(eq, ch)]
                    ws.cell(r, 19, round(sum(qs), 4))
                    for k, v in enumerate(qs):
                        if v:
                            ws.cell(r, cq[k], v)
                ws.cell(r, 20, it['vu'])
                ws.cell(r, 21, f'=ROUND(T{r}*$S{r},2)')
                for k in range(nm):
                    ws.cell(r, ct[k], f'=ROUND({COL(cq[k])}{r}*$T{r},2)')
                faixa = f'${COL(c0)}{r}:${COL(ct[-1])}{r}'
                cab_f = f'${COL(c0)}$5:${COL(ct[-1])}$5'
                ws.cell(r, cAQ, f'=SUMIFS({faixa},{cab_f},"QUANT.")')
                ws.cell(r, cAT, f'=SUMIFS({faixa},{cab_f},"TOTAL")')
                ws.cell(r, cPC, f'=IFERROR({COL(cAQ)}{r}/S{r},0)')
                ws.cell(r, cCR, origem.get(ch, ''))
                r += 1
            r1 = r - 1
            for c in [21] + ct + [cAT]:
                ws.cell(rs, c, f'=SUM({COL(c)}{r0}:{COL(c)}{r1})')
        for c in [21] + ct + [cAT]:
            ws.cell(rb, c, '=' + '+'.join(f'{COL(c)}{x}' for x in linhas_sg))
    ultimo = r - 1
    # administração local: a fração do mês é o faturamento do mês (sem BDI, sem a ADM) × % ADM ÷ valor da ADM
    r_adm = next(x for x in linhas_item if ws.cell(x, 2).value == '6.2')
    outros = [blocos_linha[b] for b, _ in BLOCOS if b != 'ADM']
    for k in range(nm):
        soma = '+'.join(f'{COL(ct[k])}{x}' for x in outros)
        ws.cell(r_adm, cq[k], f"=ROUND(({soma})*PREMISSAS!$C$12/$T{r_adm},6)")
    ws.cell(r_adm, 19, 1)
    # formatos das linhas
    for x in range(6, ultimo + 1):
        for c in range(1, cCR + 1):
            cel = ws.cell(x, c)
            if not cel.font.b:
                cel.font = f()
            cel.border = BORDA
        ws.cell(x, 19).number_format = FQ
        ws.cell(x, 20).number_format = FR
        ws.cell(x, 21).number_format = FR
        for k in range(nm):
            ws.cell(x, cq[k]).number_format = FQ
            ws.cell(x, ct[k]).number_format = FR
        ws.cell(x, cAQ).number_format = FQ
        ws.cell(x, cAT).number_format = FR
        ws.cell(x, cPC).number_format = FP
    # resumo principal (como no rodapé do BM)
    r = ultimo + 2
    rotulos = ['MATERIAL', 'SERVIÇO', 'SUB-TOTAL 01', 'BDI MATERIAL =>', 'BDI SERVIÇO =>', 'SUB-TOTAL 02',
               'TOTAL 01 + 02 =>', 'COEF. - K =>', 'TOTAL 01 + 02 + BDI -K (P0)', '% - PRINCIPAL',
               'FATURAMENTO ACUMULADO (P0)', '% - ACUMULADO']
    rr = {t: r + i for i, t in enumerate(rotulos)}
    ws.cell(r, 16, 'RESUMO PRINCIPAL').font = f(True, 10)
    faixa_p = f'$P$6:$P${ultimo}'
    for t, x in rr.items():
        ws.cell(x, 17, t).font = f(True)
        ws.cell(x, 17).alignment = Alignment(horizontal='right')
    ws.cell(rr['BDI MATERIAL =>'], 20, '=PREMISSAS!$C$7')
    ws.cell(rr['BDI SERVIÇO =>'], 20, '=PREMISSAS!$C$8')
    ws.cell(rr['COEF. - K =>'], 20, '=PREMISSAS!$C$9')
    for x in (rr['BDI MATERIAL =>'], rr['BDI SERVIÇO =>'], rr['COEF. - K =>']):
        ws.cell(x, 20).number_format = '0.00%'
    for c in [21] + ct + [cAT]:
        L = COL(c)
        ws.cell(rr['MATERIAL'], c, f'=SUMIFS({L}$6:{L}${ultimo},{faixa_p},"MATERIAL")')
        ws.cell(rr['SERVIÇO'], c, f'=SUMIFS({L}$6:{L}${ultimo},{faixa_p},"SERVIÇO")')
        ws.cell(rr['SUB-TOTAL 01'], c, f'={L}{rr["MATERIAL"]}+{L}{rr["SERVIÇO"]}')
        ws.cell(rr['BDI MATERIAL =>'], c, f'={L}{rr["MATERIAL"]}*$T${rr["BDI MATERIAL =>"]}')
        ws.cell(rr['BDI SERVIÇO =>'], c, f'={L}{rr["SERVIÇO"]}*$T${rr["BDI SERVIÇO =>"]}')
        ws.cell(rr['SUB-TOTAL 02'], c, f'={L}{rr["BDI MATERIAL =>"]}+{L}{rr["BDI SERVIÇO =>"]}')
        ws.cell(rr['TOTAL 01 + 02 =>'], c, f'={L}{rr["SUB-TOTAL 01"]}+{L}{rr["SUB-TOTAL 02"]}')
        ws.cell(rr['COEF. - K =>'], c, f'={L}{rr["TOTAL 01 + 02 =>"]}*$T${rr["COEF. - K =>"]}')
        ws.cell(rr['TOTAL 01 + 02 + BDI -K (P0)'], c, f'={L}{rr["TOTAL 01 + 02 =>"]}+{L}{rr["COEF. - K =>"]}')
        ws.cell(rr['% - PRINCIPAL'], c, f'={L}{rr["TOTAL 01 + 02 + BDI -K (P0)"]}/PREMISSAS!$C$10')
        for t in rotulos:
            ws.cell(rr[t], c).number_format = FP if t.startswith('%') else FR
            ws.cell(rr[t], c).font = f(True)
    p0 = rr['TOTAL 01 + 02 + BDI -K (P0)']
    for k, c in enumerate(ct):
        ant = f'+{COL(ct[k - 1])}{rr["FATURAMENTO ACUMULADO (P0)"]}' if k else ''
        ws.cell(rr['FATURAMENTO ACUMULADO (P0)'], c, f'={COL(c)}{p0}{ant}')
        ws.cell(rr['% - ACUMULADO'], c, f'={COL(c)}{rr["FATURAMENTO ACUMULADO (P0)"]}/PREMISSAS!$C$10')
    for t in ('TOTAL 01 + 02 + BDI -K (P0)', 'FATURAMENTO ACUMULADO (P0)'):
        for c in range(17, cAT + 1):
            ws.cell(rr[t], c).fill = fill(CREME)
    # larguras, congelamento, colunas técnicas agrupadas
    larg = {1: 9, 2: 8, 3: 30, 4: 18, 5: 9, 6: 34, 15: 16, 16: 12, 17: 50, 18: 8, 19: 13, 20: 14, 21: 16}
    for c in range(1, cCR + 1):
        ws.column_dimensions[COL(c)].width = larg.get(c, 9)
    for c in cq:
        ws.column_dimensions[COL(c)].width = 11
    for c in ct:
        ws.column_dimensions[COL(c)].width = 15
    ws.column_dimensions[COL(cAQ)].width = 12
    ws.column_dimensions[COL(cAT)].width = 16
    ws.column_dimensions[COL(cPC)].width = 11
    ws.column_dimensions[COL(cCR)].width = 70
    ws.column_dimensions.group('G', 'N', hidden=True, outline_level=1)
    ws.freeze_panes = 'G6'
    ws.sheet_view.zoomScale = 85
    ws.auto_filter.ref = f'A5:{COL(cCR)}{ultimo}'
    return {'ws': ws, 'ultimo': ultimo, 'cq': cq, 'ct': ct, 'blocos': blocos_linha, 'p0': p0,
            'acum': rr['FATURAMENTO ACUMULADO (P0)'], 'cAT': cAT}


def aba_consolidado(wb, itens, bm, livro):
    """BM horizontal (juntada) somando as equipes pelo ITEM - BM."""
    ws = wb.create_sheet('BM CONSOLIDADO')
    src = openpyxl.load_workbook(motor.BM, read_only=True, data_only=True)['BM MORRO DA CRUZ']
    linhas = [r for r in src.iter_rows(min_row=6, max_row=402, values_only=True)]
    meses = livro.meses
    nm = len(meses)
    c0 = 18
    cq = [c0 + 2 * k for k in range(nm)]
    ct = [c + 1 for c in cq]
    cAQ, cAT, cPC = c0 + 2 * nm, c0 + 2 * nm + 1, c0 + 2 * nm + 2
    ws['A1'] = 'BM - PLANILHA DE MEDIÇÃO - JUNTADA (PROGRAMADO, SOMA DAS EQUIPES)'
    ws['A1'].font = f(True, 16, 'FF0000')
    ws['A3'] = 'OBRA:'
    ws['A3'].font = f(True, 11)
    ws['A4'] = OBJETO
    ws['A4'].font = f(False, 10)
    ws['N3'] = 'DATA BASE - ORÇAMENTO:'
    ws['N3'].font = f(True)
    ws['N4'] = '01/2025'
    ws['N4'].font = f(True, 11, 'FF0000')
    ws['P1'] = 'CONTRATO:'
    ws['P1'].font = f(True, 10)
    ws['P2'] = '10.239/2026'
    ws['P2'].font = f(True)
    ws['P3'] = 'ORIGINAL - P0 (CAESB) - PV'
    ws['P3'].font = f(True)
    ws['P3'].fill = fill(AZUL)
    ws.merge_cells('P3:Q3')
    cab_meses(ws, meses, c0)
    ws.cell(2, cAQ, 'ACUMULADO PROGRAMADO').font = f(True, 11)
    ws.merge_cells(start_row=2, start_column=cAQ, end_row=2, end_column=cAT)
    cab = ['ITEM - BM', 'DESCRIÇÃO', 'PRESSÃO', 'RESINA', 'TIPO_CONEXÃO', 'SOLDA', 'DIÂMETRO', 'MATERIAL', 'GRUPO',
           'GRUPO ORÇAMENTO ANKARA', 'CÓDIGO CAESB', 'CLASSIFICAÇÃO', 'DESCRIÇÃO', 'UNID.', 'QUANT.', 'VALOR UNITÁRIO',
           'TOTAL']
    for c, t in enumerate(cab, 1):
        ws.cell(5, c, t)
    for k in range(nm):
        ws.cell(5, cq[k], 'QUANT.')
        ws.cell(5, ct[k], 'TOTAL')
    ws.cell(5, cAQ, 'QUANT.')
    ws.cell(5, cAT, 'TOTAL ACUMULADO')
    ws.cell(5, cPC, '% PROGRAMADO')
    for c in range(1, cPC + 1):
        cel = ws.cell(5, c)
        cel.font = f(True)
        cel.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cel.border = BORDA
    ws.row_dimensions[5].height = 30
    V = "'BM POR EQUIPE'"
    u = bm['ultimo']
    r = 6
    grupos = []                                         # (linha, prefixo)
    item_rows = []
    for lin in linhas:
        a = lin[0]
        if a is None:
            continue
        eh_item = lin[1] not in (None, '')
        for c, v in enumerate(lin[:17], 1):
            if c == 17:
                continue
            ws.cell(r, c, v)
        ws.cell(r, 11).number_format = '@'
        if eh_item:
            item_rows.append(r)
            ws.cell(r, 17, f'=ROUND(P{r}*O{r},2)')
            for k in range(nm):
                cv = COL(bm['cq'][k])
                ws.cell(r, cq[k], f'=SUMPRODUCT(({V}!$E$6:$E${u}=$A{r})*({V}!{cv}$6:{cv}${u}))')
                ws.cell(r, ct[k], f'=ROUND({COL(cq[k])}{r}*$P{r},2)')
            faixa = f'${COL(c0)}{r}:${COL(ct[-1])}{r}'
            cab_f = f'${COL(c0)}$5:${COL(ct[-1])}$5'
            ws.cell(r, cAQ, f'=SUMIFS({faixa},{cab_f},"QUANT.")')
            ws.cell(r, cAT, f'=SUMIFS({faixa},{cab_f},"TOTAL")')
            ws.cell(r, cPC, f'=IFERROR({COL(cAQ)}{r}/O{r},0)')
        else:
            grupos.append((r, str(a)))
            nivel = str(a).count('.')
            for c in range(1, cPC + 1):
                ws.cell(r, c).fill = fill(VERDE if nivel == 0 else CREME)
                ws.cell(r, c).font = f(True, 9, '000000' if nivel == 0 else 'FF0000')
        r += 1
    ultimo = r - 1
    ws.cell(5, cPC + 1, 'NÍVEL').font = f(True)
    for x in item_rows:
        ws.cell(x, cPC + 1, 'ITEM')
    ws.column_dimensions[COL(cPC + 1)].hidden = True
    for gr, pref in grupos:
        for c in [17] + ct + [cAT]:
            L = COL(c)
            ws.cell(gr, c, f'=SUMIFS({L}$6:{L}${ultimo},$A$6:$A${ultimo},"{pref}.*",'
                           f'${COL(cPC + 1)}$6:${COL(cPC + 1)}${ultimo},"ITEM")')
    for x in range(6, ultimo + 1):
        for c in range(1, cPC + 1):
            cel = ws.cell(x, c)
            if not cel.font.b:
                cel.font = f()
            cel.border = BORDA
        ws.cell(x, 15).number_format = FQ
        ws.cell(x, 16).number_format = FR
        ws.cell(x, 17).number_format = FR
        for k in range(nm):
            ws.cell(x, cq[k]).number_format = FQ
            ws.cell(x, ct[k]).number_format = FR
        ws.cell(x, cAQ).number_format = FQ
        ws.cell(x, cAT).number_format = FR
        ws.cell(x, cPC).number_format = FP
    # resumo
    r = ultimo + 2
    rot = ['MATERIAL', 'SERVIÇO', 'SUB-TOTAL 01', 'BDI MATERIAL =>', 'BDI SERVIÇO =>', 'SUB-TOTAL 02',
           'TOTAL 01 + 02 =>', 'COEF. - K =>', 'TOTAL 01 + 02 + BDI -K', '% - PRINCIPAL']
    rr = {t: r + i for i, t in enumerate(rot)}
    ws.cell(r, 12, 'RESUMO PRINCIPAL').font = f(True, 10)
    niv = f'${COL(cPC + 1)}$6:${COL(cPC + 1)}${ultimo}'
    for t, x in rr.items():
        ws.cell(x, 13, t).font = f(True)
        ws.cell(x, 13).alignment = Alignment(horizontal='right')
    ws.cell(rr['BDI MATERIAL =>'], 16, '=PREMISSAS!$C$7')
    ws.cell(rr['BDI SERVIÇO =>'], 16, '=PREMISSAS!$C$8')
    ws.cell(rr['COEF. - K =>'], 16, '=PREMISSAS!$C$9')
    for x in (rr['BDI MATERIAL =>'], rr['BDI SERVIÇO =>'], rr['COEF. - K =>']):
        ws.cell(x, 16).number_format = '0.00%'
    for c in [17] + ct + [cAT]:
        L = COL(c)
        ws.cell(rr['MATERIAL'], c, f'=SUMIFS({L}$6:{L}${ultimo},$L$6:$L${ultimo},"MATERIAL",{niv},"ITEM")')
        ws.cell(rr['SERVIÇO'], c, f'=SUMIFS({L}$6:{L}${ultimo},$L$6:$L${ultimo},"SERVIÇO",{niv},"ITEM")')
        ws.cell(rr['SUB-TOTAL 01'], c, f'={L}{rr["MATERIAL"]}+{L}{rr["SERVIÇO"]}')
        ws.cell(rr['BDI MATERIAL =>'], c, f'={L}{rr["MATERIAL"]}*$P${rr["BDI MATERIAL =>"]}')
        ws.cell(rr['BDI SERVIÇO =>'], c, f'={L}{rr["SERVIÇO"]}*$P${rr["BDI SERVIÇO =>"]}')
        ws.cell(rr['SUB-TOTAL 02'], c, f'={L}{rr["BDI MATERIAL =>"]}+{L}{rr["BDI SERVIÇO =>"]}')
        ws.cell(rr['TOTAL 01 + 02 =>'], c, f'={L}{rr["SUB-TOTAL 01"]}+{L}{rr["SUB-TOTAL 02"]}')
        ws.cell(rr['COEF. - K =>'], c, f'={L}{rr["TOTAL 01 + 02 =>"]}*$P${rr["COEF. - K =>"]}')
        ws.cell(rr['TOTAL 01 + 02 + BDI -K'], c, f'={L}{rr["TOTAL 01 + 02 =>"]}+{L}{rr["COEF. - K =>"]}')
        ws.cell(rr['% - PRINCIPAL'], c, f'={L}{rr["TOTAL 01 + 02 + BDI -K"]}/PREMISSAS!$C$10')
        for t in rot:
            ws.cell(rr[t], c).number_format = FP if t.startswith('%') else FR
            ws.cell(rr[t], c).font = f(True)
    larg = {1: 9, 2: 34, 11: 16, 12: 12, 13: 50, 14: 8, 15: 12, 16: 14, 17: 16}
    for c in range(1, cPC + 1):
        ws.column_dimensions[COL(c)].width = larg.get(c, 9)
    for c in cq:
        ws.column_dimensions[COL(c)].width = 11
    for c in ct:
        ws.column_dimensions[COL(c)].width = 15
    ws.column_dimensions[COL(cAT)].width = 16
    ws.column_dimensions.group('C', 'J', hidden=True, outline_level=1)
    ws.freeze_panes = 'C6'
    ws.sheet_view.zoomScale = 85
    return {'ultimo': ultimo, 'p0': rr['TOTAL 01 + 02 + BDI -K']}


def aba_ff(wb, bm, livro):
    ws = wb.create_sheet('CRONOGRAMA FF')
    meses = livro.meses
    nm = len(meses)
    V, u = "'BM POR EQUIPE'", bm['ultimo']
    ws['A1'] = 'CLIENTE:'
    ws['A2'] = 'COMPANHIA DE SANEAMENTO AMBIENTAL DO DISTRITO FEDERAL'
    ws['F1'] = 'CONTRATO Nº:'
    ws['F2'] = '10.239/2026'
    ws['H1'] = 'OBJETO:'
    ws['H2'] = OBJETO
    for c in ('A1', 'F1', 'H1'):
        ws[c].font = f(True, 10)
    for c in ('A2', 'F2', 'H2'):
        ws[c].font = f(False, 10)
    for c, t in enumerate(['ITEM', 'DESCRIÇÃO', 'VALORES', '', 'PERÍODO:'], 1):
        ws.cell(4, c, t).font = f(True)
    ws.merge_cells('C4:D4')
    for k, (n, ini, fim) in enumerate(meses):
        c = 6 + k
        ws.cell(5, c, f'MÊS {n:02d}').font = f(True)
        ws.cell(6, c, f'{ini:%d-%m-%y} A {fim:%d-%m-%y}').font = f(False, 8)
        ws.cell(7, c, (fim - ini).days + 1).font = f(False, 8)
        for r in (5, 6, 7):
            ws.cell(r, c).alignment = Alignment(horizontal='center', wrap_text=True)
            if ini > PRAZO:
                ws.cell(r, c).fill = fill(LARANJA)
    cTot = 6 + nm
    ws.cell(5, cTot, 'TOTAL').font = f(True)
    estrut = [('1', 'FORNECIMENTO MATERIAIS', ['1.1', '2.1', '3.1', '4.1', '5.1']),
              ('1.1', 'CANTEIRO DE OBRAS - MATERIAL', ['1.1']), ('2.1', 'REDES - MATERIAL', ['2.1']),
              ('3.1', 'CAIXAS - MATERIAL', ['3.1']), ('4.1', 'INTERLIGAÇÃO E LIGAÇÕES - MATERIAL', ['4.1']),
              ('5.1', 'HIDROMECÂNICO - MATERIAL', ['5.1']),
              ('2', 'EXECUÇÃO DE SERVIÇOS', ['1.2', '2.2', '3.2', '4.2', '5.2', '6.2']),
              ('1.2', 'CANTEIRO DE OBRAS - SERVIÇO', ['1.2']), ('2.2', 'REDES - SERVIÇO', ['2.2']),
              ('3.2', 'CAIXAS - SERVIÇO', ['3.2']), ('4.2', 'INTERLIGAÇÃO E LIGAÇÕES - SERVIÇO', ['4.2']),
              ('5.2', 'HIDROMECÂNICO - SERVIÇO', ['5.2']), ('6.2', 'ADMINISTRAÇÃO LOCAL', ['6.2'])]
    r = 9
    linhas = {}

    def soma_grp(grps, col):
        return '+'.join(f'SUMPRODUCT(({V}!$B$6:$B${u}="{g}")*({V}!{col}$6:{col}${u}))' for g in grps)
    for cod, nome, grps in estrut:
        linhas[cod] = r
        topo = '.' not in cod
        ws.cell(r, 1, cod)
        ws.cell(r, 2, nome)
        ws.cell(r, 3, '=' + soma_grp(grps, 'U'))
        ws.cell(r, 4, f'=C{r}/$C$TOTALGERAL')
        rot = ['FINANC. PROGRAMADO - R$', 'FINANC. EXECUTADO - R$', 'FÍSICO PROGRAMADO - %', 'FÍSICO EXECUTADO - %']
        for i, t in enumerate(rot):
            ws.cell(r + i, 5, t).font = f(True, 8)
        for k in range(nm):
            c = 6 + k
            col_v = COL(bm['ct'][k])
            ws.cell(r, c, '=' + soma_grp(grps, col_v))
            ws.cell(r + 2, c, f'=IFERROR({COL(c)}{r}/$C{r},0)')
            ws.cell(r, c).number_format = FR
            ws.cell(r + 1, c).number_format = FR
            ws.cell(r + 2, c).number_format = FP
            ws.cell(r + 3, c).number_format = FP
            ws.cell(r, c).fill = fill('DDEBF7')
        ws.cell(r, cTot, f'=SUM({COL(6)}{r}:{COL(5 + nm)}{r})').number_format = FR
        ws.cell(r + 2, cTot, f'=SUM({COL(6)}{r + 2}:{COL(5 + nm)}{r + 2})').number_format = FP
        for x in range(r, r + 4):
            for c in (1, 2, 3, 4):
                ws.cell(x, c).font = f(topo)
        ws.cell(r, 3).number_format = FR
        ws.cell(r, 4).number_format = '0.00%'
        if topo:
            for c in range(1, cTot + 1):
                ws.cell(r, c).font = f(True)
        r += 5
    # totais (BDI e desconto por classe), como no modelo
    tot = {}

    def linha(rot, formula_c, formula_mes, fmt=FR):
        nonlocal r
        ws.cell(r, 1, rot).font = f(True)
        ws.cell(r, 3, formula_c).number_format = fmt
        for k in range(nm):
            c = 6 + k
            ws.cell(r, c, formula_mes(COL(c))).number_format = fmt
            ws.cell(r, c).font = f(True)
        ws.cell(r, cTot, f'=SUM({COL(6)}{r}:{COL(5 + nm)}{r})' if fmt == FR else None)
        ws.cell(r, cTot).number_format = fmt
        tot[rot] = r
        r += 1
    m, s = linhas['1'], linhas['2']
    linha('SUB-TOTAL MATERIAL =>', f'=C{m}', lambda L: f'={L}{m}')
    linha('BDI MATERIAL =>', f'=C{r - 1}*PREMISSAS!$C$7', lambda L: f'={L}{r - 1}*PREMISSAS!$C$7')
    linha('DESCONTO CONTRATUAL (MATERIAL) =>', f'=-(C{r - 2}+C{r - 1})*PREMISSAS!$C$9*-1',
          lambda L: f'=(({L}{r - 2})+({L}{r - 1}))*PREMISSAS!$C$9')
    ws.cell(r - 1, 3, f'=(C{r - 3}+C{r - 2})*PREMISSAS!$C$9')
    linha('TOTAL MATERIAL =>', f'=C{r - 3}+C{r - 2}+C{r - 1}', lambda L: f'={L}{r - 3}+{L}{r - 2}+{L}{r - 1}')
    r += 1
    linha('SUB-TOTAL SERVIÇO =>', f'=C{s}', lambda L: f'={L}{s}')
    linha('BDI SERVIÇO =>', f'=C{r - 1}*PREMISSAS!$C$8', lambda L: f'={L}{r - 1}*PREMISSAS!$C$8')
    linha('DESCONTO CONTRATUAL (SERVIÇO) =>', '', lambda L: f'=(({L}{r - 2})+({L}{r - 1}))*PREMISSAS!$C$9')
    ws.cell(r - 1, 3, f'=(C{r - 3}+C{r - 2})*PREMISSAS!$C$9')
    linha('TOTAL SERVIÇO =>', f'=C{r - 3}+C{r - 2}+C{r - 1}', lambda L: f'={L}{r - 3}+{L}{r - 2}+{L}{r - 1}')
    r += 1
    tm, tsv = tot['TOTAL MATERIAL =>'], tot['TOTAL SERVIÇO =>']
    linha('TOTAL GERAL MATERIAL + SERVIÇO (P0) =>', f'=C{tm}+C{tsv}', lambda L: f'={L}{tm}+{L}{tsv}')
    tg = tot['TOTAL GERAL MATERIAL + SERVIÇO (P0) =>']
    linha('FÍSICO PROGRAMADO - % (MÊS)', f'=C{tg}/PREMISSAS!$C$10', lambda L: f'={L}{tg}/PREMISSAS!$C$10', FP)
    ws.cell(r, 1, 'FATURAMENTO PROGRAMADO - R$ (ACUM)').font = f(True)
    for k in range(nm):
        c = 6 + k
        ws.cell(r, c, f'={COL(c)}{tg}' + (f'+{COL(c - 1)}{r}' if k else '')).number_format = FR
        ws.cell(r + 1, c, f'={COL(c)}{r}/PREMISSAS!$C$10').number_format = FP
        ws.cell(r, c).font = f(True)
        ws.cell(r + 1, c).font = f(True)
    ws.cell(r + 1, 1, 'FÍSICO PROGRAMADO - % (ACUM)').font = f(True)
    for x in (tg, r):
        for c in range(1, cTot + 1):
            ws.cell(x, c).fill = fill(CREME)
    # o total geral do C de cada grupo
    for cod, x in linhas.items():
        ws.cell(x, 4, f'=C{x}/$C${tg}')
    ws.column_dimensions['A'].width = 34
    ws.column_dimensions['B'].width = 36
    ws.column_dimensions['C'].width = 17
    ws.column_dimensions['D'].width = 8
    ws.column_dimensions['E'].width = 22
    for k in range(nm + 1):
        ws.column_dimensions[COL(6 + k)].width = 14
    ws.freeze_panes = 'F8'
    ws.sheet_view.zoomScale = 85
    return {'tg': tg, 'acum': r}


def aba_resumo(wb, bm, livro, plano, cen):
    ws = wb.create_sheet('RESUMO POR EQUIPE')
    meses = livro.meses
    nm = len(meses)
    V, u = "'BM POR EQUIPE'", bm['ultimo']
    ws['A1'] = 'FATURAMENTO PROGRAMADO POR EQUIPE (P0, COM BDI E COEFICIENTE K)'
    ws['A1'].font = f(True, 14)
    ws['A2'] = cen['titulo']
    ws['A2'].font = f(True, 11, 'C00000')
    ws.cell(4, 1, 'EQUIPE').font = f(True)
    for k, (n, ini, fim) in enumerate(meses):
        c = 2 + k
        ws.cell(4, c, f'{n}ª MED.\n{MES[ini.month]}/{ini:%y}').font = f(True)
        ws.cell(4, c).alignment = Alignment(horizontal='center', wrap_text=True)
        if ini > PRAZO:
            ws.cell(4, c).fill = fill(LARANJA)
        elif ini < REAJUSTE:
            ws.cell(4, c).fill = fill('E2EFDA')
    cT = 2 + nm
    ws.cell(4, cT, 'TOTAL').font = f(True)
    ws.cell(4, cT + 1, '% DO CONTRATO').font = f(True)
    ws.row_dimensions[4].height = 30
    r = 5
    for seq, (eq, nome) in enumerate(BLOCOS, 1):
        ws.cell(r, 1, nome).font = f()
        for k in range(nm):
            col = COL(bm['ct'][k])
            mat = f'SUMPRODUCT(({V}!$A$6:$A${u}={seq})*({V}!$P$6:$P${u}="MATERIAL")*({V}!{col}$6:{col}${u}))'
            srv = f'SUMPRODUCT(({V}!$A$6:$A${u}={seq})*({V}!$P$6:$P${u}="SERVIÇO")*({V}!{col}$6:{col}${u}))'
            ws.cell(r, 2 + k, f'=({mat}*(1+PREMISSAS!$C$7)+{srv}*(1+PREMISSAS!$C$8))*(1+PREMISSAS!$C$9)')
        ws.cell(r, cT, f'=SUM(B{r}:{COL(cT - 1)}{r})')
        ws.cell(r, cT + 1, f'={COL(cT)}{r}/PREMISSAS!$C$10')
        r += 1
    r_tot = r
    ws.cell(r, 1, 'TOTAL DO MÊS').font = f(True)
    for c in range(2, cT + 1):
        ws.cell(r, c, f'=SUM({COL(c)}5:{COL(c)}{r - 1})').font = f(True)
    ws.cell(r, cT + 1, f'={COL(cT)}{r}/PREMISSAS!$C$10').font = f(True)
    ws.cell(r + 1, 1, 'ACUMULADO').font = f(True)
    ws.cell(r + 2, 1, '% DO CONTRATO NO MÊS').font = f(True)
    ws.cell(r + 3, 1, '% ACUMULADO DO CONTRATO').font = f(True)
    for k in range(nm):
        c = 2 + k
        L = COL(c)
        ws.cell(r + 1, c, f'={L}{r}' + (f'+{COL(c - 1)}{r + 1}' if k else '')).font = f(True)
        ws.cell(r + 2, c, f'={L}{r}/PREMISSAS!$C$10')
        ws.cell(r + 3, c, f'={L}{r + 1}/PREMISSAS!$C$10').font = f(True)
    for x in range(5, r + 4):
        for c in range(2, cT + 2):
            ws.cell(x, c).number_format = FP if (x >= r + 2 or c == cT + 1) else '#,##0'
            ws.cell(x, c).border = BORDA
        ws.cell(x, 1).border = BORDA
    for c in range(1, cT + 2):
        ws.cell(r, c).fill = fill(CREME)
    r_acum_pct = r + 3
    # físico (do motor, em valores)
    r += 6
    ws.cell(r, 1, 'PRODUÇÃO FÍSICA PROGRAMADA').font = f(True, 12)
    r += 1
    T, E = plano['T'], plano['E']
    fis = collections.defaultdict(lambda: [0.0] * nm)
    lig = collections.defaultdict(lambda: [0.0] * nm)
    for i, x in plano['trechos'].items():
        a, b = x['ini'], x['fim']
        for k in range(nm):
            lo, hi = max(a, livro.lim[k]), min(b, livro.lim[k + 1])
            if hi > lo:
                fis[x['eq']][k] += T[i]['L'] * (hi - lo) / (b - a)
        if T[i]['lig']:
            a2, b2 = a + motor.LIG_ATRASO, b + motor.LIG_ATRASO
            for k in range(nm):
                lo, hi = max(a2, livro.lim[k]), min(b2, livro.lim[k + 1])
                if hi > lo:
                    lig[x['eq']][k] += T[i]['lig'] * (hi - lo) / (b2 - a2)
    civ, hid = [0.0] * nm, [0.0] * nm
    for j, c in plano['caixas'].items():
        civ[livro.mes(c['civ_fim'] - 1)] += 1
        hid[livro.mes(c['hid_fim'] - 1)] += 1
    for k, (n, ini, fim) in enumerate(meses):
        ws.cell(r, 2 + k, f'{n}ª MED.\n{MES[ini.month]}/{ini:%y}').font = f(True)
        ws.cell(r, 2 + k).alignment = Alignment(horizontal='center', wrap_text=True)
    ws.cell(r, cT, 'TOTAL').font = f(True)
    ws.row_dimensions[r].height = 30
    r += 1
    secoes = [('REDE ASSENTADA (m)', [(NOME[e], fis[e]) for e in ['ANK1', 'E1A', 'E1B', 'E2C', 'E2D']], '#,##0'),
              ('LIGAÇÕES PREDIAIS (un, estimadas pela rede de DE 63)',
               [(NOME[e], lig[e]) for e in ['E1A', 'E1B', 'E2C', 'E2D', 'ANK1']], '#,##0'),
              ('CAIXAS HIDRÁULICAS (un)', [('Obra civil concluída', civ), ('Hidromecânico montado', hid)], '#,##0')]
    for titulo, series, fmt in secoes:
        ws.cell(r, 1, titulo).font = f(True)
        r += 1
        r0 = r
        for nome, vals in series:
            ws.cell(r, 1, nome).font = f()
            for k, v in enumerate(vals):
                ws.cell(r, 2 + k, round(v, 2)).number_format = fmt
            ws.cell(r, cT, f'=SUM(B{r}:{COL(cT - 1)}{r})').number_format = fmt
            r += 1
        if len(series) > 2:
            ws.cell(r, 1, 'TOTAL').font = f(True)
            for c in range(2, cT + 1):
                ws.cell(r, c, f'=SUM({COL(c)}{r0}:{COL(c)}{r - 1})').number_format = fmt
                ws.cell(r, c).font = f(True)
            r += 1
        r += 1
    ws.column_dimensions['A'].width = 44
    for c in range(2, cT + 2):
        ws.column_dimensions[COL(c)].width = 12
    ws.freeze_panes = 'B5'
    # gráficos (um eixo cada): faturamento mensal empilhado por equipe e acumulado em % do contrato
    ch = BarChart()
    ch.type = 'col'
    ch.grouping = 'stacked'
    ch.overlap = 100
    ch.gapWidth = 40
    ch.title = 'Faturamento programado por mês e por equipe (R$, P0)'
    ch.y_axis.title = 'R$'
    ch.y_axis.numFmt = '#,##0'
    ch.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill='E0E0E0'))
    dados = Reference(ws, min_col=1, max_col=1 + nm, min_row=5, max_row=4 + len(BLOCOS))
    ch.add_data(dados, from_rows=True, titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=2, max_col=1 + nm, min_row=4, max_row=4))
    for s, cor in zip(ch.series, CORES):
        s.graphicalProperties.solidFill = cor
        s.graphicalProperties.line.solidFill = 'FFFFFF'
    ch.legend.position = 'b'
    ch.height, ch.width = 11, 30
    ws.add_chart(ch, f'A{r + 1}')
    ln = LineChart()
    ln.title = 'Faturamento acumulado programado (% do contrato, P0)'
    ln.y_axis.numFmt = '0%'
    ln.y_axis.scaling.min = 0
    ln.y_axis.scaling.max = 1
    ln.add_data(Reference(ws, min_col=1, max_col=1 + nm, min_row=r_acum_pct, max_row=r_acum_pct), from_rows=True,
                titles_from_data=True)
    ln.set_categories(Reference(ws, min_col=2, max_col=1 + nm, min_row=4, max_row=4))
    ln.series[0].graphicalProperties.line.solidFill = CORES[0]
    ln.series[0].graphicalProperties.line.width = 25400
    ln.legend = None
    ln.height, ln.width = 9, 30
    ws.add_chart(ln, f'A{r + 24}')
    return {'tot': r_tot, 'acum_pct': r_acum_pct}


def aba_programacao(wb, plano, livro):
    ws = wb.create_sheet('PROGRAMAÇÃO')
    T, E = plano['T'], plano['E']
    cab = ['ORDEM NA EQUIPE', 'EQUIPE', 'FASE', 'ZONA / FRENTE', 'PB', 'TRECHO (06.1)', 'CLASSIFICAÇÃO', 'DE (mm)',
           'METODOLOGIA', 'EXTENSÃO (m)', 'LIGAÇÕES (est.)', 'INÍCIO', 'TÉRMINO', 'MEDIÇÃO']
    ws['A1'] = 'PROGRAMAÇÃO TRECHO A TRECHO (NOMES DA 06.1 REV01) E CAIXA A CAIXA (06.2 REV08)'
    ws['A1'].font = f(True, 12)
    for c, t in enumerate(cab, 1):
        ws.cell(3, c, t).font = f(True)
        ws.cell(3, c).fill = fill(AMARELO)
        ws.cell(3, c).alignment = Alignment(horizontal='center', wrap_text=True)
    r = 4
    for eq in ['ANK1', 'E1A', 'E1B', 'E2C', 'E2D']:
        for n, i in enumerate(plano['fila'][eq], 1):
            t, x = T[i], plano['trechos'][i]
            vals = [n, NOME[eq], x['fase'], 'Subadutora' if t['sub'] else t['loc'], t['pb'], t['R'],
                    t['rede'].get('CLASSIFICAÇÃO EXECUTIVA'), t['de'], t['met'], t['L'], round(t['lig'], 1),
                    CAL.inicio(x['ini']), CAL.fim(x['fim']), livro.meses[livro.mes(x['fim'] - 1)][0]]
            for c, v in enumerate(vals, 1):
                ws.cell(r, c, v)
            ws.cell(r, 10).number_format = '#,##0.00'
            ws.cell(r, 12).number_format = 'dd/mm/yyyy hh:mm'
            ws.cell(r, 13).number_format = 'dd/mm/yyyy hh:mm'
            r += 1
    ws.auto_filter.ref = f'A3:N{r - 1}'
    r += 2
    ws.cell(r, 1, 'CAIXAS HIDRÁULICAS (DISPOSITIVOS DA 06.2) - EQUIPE CAIXAS E HIDROMECÂNICO (ANKARA 2)').font = f(True, 11)
    r += 1
    cab2 = ['ORDEM', 'CAIXA (06.2)', 'TIPO', 'ZONA', 'PB', 'TRECHOS PRONTOS EM', 'OBRA CIVIL - INÍCIO',
            'OBRA CIVIL - TÉRMINO', 'HIDROMECÂNICO - INÍCIO', 'HIDROMECÂNICO - TÉRMINO']
    for c, t in enumerate(cab2, 1):
        ws.cell(r, c, t).font = f(True)
        ws.cell(r, c).fill = fill(AMARELO)
        ws.cell(r, c).alignment = Alignment(horizontal='center', wrap_text=True)
    r += 1
    cx = plano['caixas']
    for n, j in enumerate(plano['ordem_civil'], 1):
        e = E[j]
        vals = [n, e['noms'][0] if e['noms'] else e['id'], e['tp'], e.get('zona'), e.get('pb'),
                CAL.fim(cx[j]['pronto']), CAL.inicio(cx[j]['civ_ini']), CAL.fim(cx[j]['civ_fim']),
                CAL.inicio(cx[j]['hid_ini']), CAL.fim(cx[j]['hid_fim'])]
        for c, v in enumerate(vals, 1):
            ws.cell(r, c, v)
            if isinstance(v, dt.datetime):
                ws.cell(r, c).number_format = 'dd/mm/yyyy'
        r += 1
    for c, w in enumerate([10, 40, 22, 14, 8, 34, 14, 9, 11, 12, 11, 17, 17, 10], 1):
        ws.column_dimensions[COL(c)].width = w
    for row in ws.iter_rows(min_row=4, max_row=r):
        for cel in row:
            if not cel.font.b:
                cel.font = f()
    ws.freeze_panes = 'A4'


def aba_premissas(wb, cen, plano, livro, comparativo, itens):
    ws = wb.create_sheet('PREMISSAS', 0)
    ws['A1'] = 'CRONOGRAMA FÍSICO EXECUTIVO BASE - CT 10.239/2026 - MORRO DA CRUZ'
    ws['A1'].font = f(True, 16)
    ws['A2'] = cen['titulo']
    ws['A2'].font = f(True, 12, 'C00000')
    ws['A3'] = (f'REV00 de 29/10/2026 · OS em 01/10/2026 · prazo de 540 dias (término contratual em '
                f'{PRAZO:%d/%m/%Y}) · BM REV04, 06.1 REV01 e 06.2 REV08')
    ws['A3'].font = f(False, 10)
    ws['A5'] = 'PARÂMETROS DO CONTRATO (as células amarelas alimentam as fórmulas)'
    ws['A5'].font = f(True, 11)
    par = [('BDI material', 0.1528, '0.00%', 'BM REV04 (resumo principal)'),
           ('BDI serviço', 0.2148, '0.00%', 'BM REV04 (resumo principal)'),
           ('Coeficiente K', -0.1404, '0.00%', 'contrato, cláusula 5.2'),
           ('Valor do contrato (P0)', 49878748.6575869, FR, 'BM REV04, célula T556 do vertical'),
           ('Administração local (item 4.1.1)', 1304685.94, FR, 'BM REV04'),
           ('% da administração local sobre o faturamento', '=C11/(48534308.16-C11)', '0.0000%',
            'TR 14.6: ADM ÷ (subtotal 01 do contrato − ADM); subtotal 01 = R$ 48.534.308,16')]
    for i, (n, v, fmt, fonte) in enumerate(par):
        r = 7 + i
        ws.cell(r, 2, n).font = f()
        ws.cell(r, 3, v).number_format = fmt
        ws.cell(r, 3).font = f(False, 9, '0000FF' if not str(v).startswith('=') else '000000')
        ws.cell(r, 3).fill = fill(AMARELO) if not str(v).startswith('=') else PatternFill()
        ws.cell(r, 4, fonte).font = f(False, 8, '595959', True)
    r = 14
    ws.cell(r, 1, 'RESULTADO DOS DOIS CENÁRIOS').font = f(True, 11)
    r += 1
    for c, t in enumerate(['', 'CENÁRIO A - PRAZO CONTRATUAL', 'CENÁRIO B - RITMO MÉDIO DE 90 m/dia'], 2):
        ws.cell(r, c, t).font = f(True)
        ws.cell(r, c).fill = fill(CREME)
        ws.cell(r, c).alignment = Alignment(wrap_text=True, horizontal='center')
    r += 1
    for nome, a, b in comparativo:
        ws.cell(r, 2, nome).font = f()
        for c, v in ((3, a), (4, b)):
            ws.cell(r, c, v).font = f(True)
            ws.cell(r, c).alignment = Alignment(horizontal='center')
            if isinstance(v, (dt.date, dt.datetime)):
                ws.cell(r, c).number_format = FD
            elif isinstance(v, float) and v < 1.5:
                ws.cell(r, c).number_format = '0.0%'
            elif isinstance(v, float):
                ws.cell(r, c).number_format = FR if v > 1000 else '#,##0.0'
        r += 1
    r += 1
    ws.cell(r, 1, 'REGRAS DO RESPONSÁVEL (29/10/2026)').font = f(True, 11)
    r += 1
    regras = [
        '5 equipes de rede/obra: 1 equipe Ankara (adutora) e 2 empreiteiros com 2 equipes cada nas zonas; a equipe de caixas é uma 2ª equipe Ankara.',
        'Pé alto até o reajuste de janeiro de 2027: até lá só a equipe da adutora, em redes de diâmetro não tão alto (160 mm), produzindo pouco.',
        'O 1º empreiteiro entra em janeiro de 2027 com 2 equipes; o 2º entra em abril de 2027, completando as equipes.',
        'Produtividade média de 80 a 100 m por dia por equipe, em dias sem interferência.',
        'Caixas hidráulicas a partir de março de 2027, com redes já feitas; hidromecânico a partir de julho de 2027 (peças de SP e do exterior).',
        'Depois do reajuste, a equipe da adutora vai para onde for melhor; os empreiteiros não têm zonas definidas.',
        'As ligações prediais vão sendo feitas conforme a rede avança.',
        'Trabalho de segunda a sexta, das 7h às 17h.',
        'Dois cenários: um que fecha no prazo contratual e outro que mantém a média de produtividade.',
    ]
    for t in regras:
        ws.cell(r, 2, '• ' + t).font = f()
        r += 1
    r += 1
    ws.cell(r, 1, 'PREMISSAS ADOTADAS NESTE CRONOGRAMA').font = f(True, 11)
    r += 1
    prem = [
        f"Calendário: seg a sex, 07:00-12:00 e 13:00-17:00 (9 h/dia), sem os feriados nacionais e do DF (Carnaval, Sexta-feira Santa, Corpus Christi, 21/04, 30/11 e os nacionais). Sem recesso de fim de ano.",
        f"Produtividade das frentes de rede depois do reajuste: {cen['p']:.1f} m/dia por equipe ({cen['p_txt']}), igual para todos os diâmetros.",
        f"Pé alto: a Equipe Adutora começa em 16/11/2026 (a NS de topografia vem de 03 a 13/11), a {motor.P_PE_ALTO:.0f} m/dia nos trechos de DE 160 das zonas, na ordem das zonas, até 29/01/2027. Em 01/02/2027 vai para a Subadutora (adutora + DE 225 e 250) e, ao terminá-la, para as zonas.",
        'Ordem das zonas: da mais perto para a mais longe do canteiro (coluna DIST. CANTEIRO da 06.1): ' + ', '.join(plano['zonas']) + '.',
        'As zonas são divididas em blocos por prancha (zona × PB). Cada bloco vai para a primeira equipe que fica livre; dentro do bloco, rede primária, secundária e terciária, do maior para o menor diâmetro.',
        'Ligações prediais: as 11.894 do BM ficam nos trechos de DE 63 (tê de serviço 63 e tê de sela 63 × 20), pelo comprimento, e são feitas pela mesma frente, 5 dias úteis atrás da rede.',
        'Caixas hidráulicas: os 132 dispositivos da 06.2 (45 DF, 26 VTF, 27 VM, 17 VRP, 11 MV e 6 HID). Obra civil a partir de 01/03/2027, depois dos trechos que chegam na caixa: DF 1 dia, VTF 1,5, VM 2, VRP e MV 3, HID 0,5. Hidromecânico a partir de 01/07/2027, depois da obra civil: DF, VTF e HID 0,5 dia, VM 1, VRP e MV 2 (frente de montagem da Equipe Ankara 2).',
        'Peças da 06.2 (grupo hidromecânico do BM): as dos nós entram com o primeiro trecho do nó; as das caixas, na montagem hidromecânica. O item sem peça na 06.2 vai para a montagem.',
        'Topografia (NS) 2 semanas antes de cada trecho; cadastro técnico 30 dias corridos depois; teste hidrostático e desinfecção de 10 dias úteis no fim de cada zona; interligação de 3 dias depois do teste da zona, do teste da Subadutora e da montagem das caixas da zona; desmobilização de 10 dias no fim.',
        'Itens de rede do BM: cada um vai para os trechos pela coluna da aba REDE da 06.1 que o mede (escavação, reaterro, escoramento, transporte etc.); o item com diâmetro, pelos metros daquele DE; os furos direcionais (MND) do BM ficam nos trechos do mesmo DE. O critério de cada linha está na última coluna da aba BM POR EQUIPE.',
        'Canteiro: em out/2026, só o que foi feito (demarcação, alambrado metálico, portão e 2 contêineres, informado em 29/10/2026); as demais instalações de 03/11 a 18/12/2026. Os contêineres chegam com as equipes: 2 em out/26, mais 4 em nov/26, 6 em jan/27 (Empreiteiro 1) e 6 em abr/27 (Empreiteiro 2); o aluguel conta do mês de chegada até o fim da obra, e por isso fica abaixo do BM. Os demais itens mensais (funcionamento do canteiro) em partes iguais por mês, de nov/2026 ao fim da obra.',
        'Administração local medida em proporção ao faturamento do mês (fórmula na aba BM POR EQUIPE).',
        'Material faturado na aplicação. Valores em P0 (sem reajuste). O reajuste de janeiro de 2027 vale para o que começar a partir de 01/02/2027 (contrato 6.1 a 6.4, com pedido da contratada).',
    ]
    for t in prem:
        ws.cell(r, 2, '• ' + t).font = f()
        ws.cell(r, 2).alignment = Alignment(wrap_text=True, vertical='top')
        ws.row_dimensions[r].height = 26 if len(t) > 150 else 14
        r += 1
    r += 1
    ws.cell(r, 1, 'ABAS').font = f(True, 11)
    r += 1
    for n, t in [('BM POR EQUIPE', 'o BM vertical com um bloco por equipe: quantidade e valor de cada item em cada medição'),
                 ('BM CONSOLIDADO', 'o BM horizontal (juntada) somando as equipes, com o resumo de BDI e K'),
                 ('CRONOGRAMA FF', 'o cronograma físico-financeiro no modelo da Caesb, por grupo'),
                 ('RESUMO POR EQUIPE', 'faturamento mensal por equipe, produção física e gráficos'),
                 ('PROGRAMAÇÃO', 'datas de cada trecho e de cada caixa (mesmas do arquivo do Project)')]:
        ws.cell(r, 2, n).font = f(True)
        ws.cell(r, 3, t).font = f()
        r += 1
    ws.column_dimensions['A'].width = 4
    ws.column_dimensions['B'].width = 120
    ws.column_dimensions['C'].width = 26
    ws.column_dimensions['D'].width = 26
    ws.sheet_view.showGridLines = False


def gerar(caminho, cen, plano, livro, origem, itens, D, comparativo):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    Q = quantidades(itens, livro)
    aba_premissas(wb, cen, plano, livro, comparativo, itens)
    bm = aba_bm(wb, cen, itens, Q, livro, origem)
    aba_consolidado(wb, itens, bm, livro)
    aba_ff(wb, bm, livro)
    aba_resumo(wb, bm, livro, plano, cen)
    aba_programacao(wb, plano, livro)
    for ws in wb.worksheets:
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        ws.page_setup.orientation = 'landscape'
    wb.save(caminho)
