"""Motor do cronograma físico executivo base do Morro da Cruz (CT 10.239/2026).

Programa trecho a trecho (06.1), caixa a caixa (06.2) e distribui cada item do BM REV04 pelas equipes e pelos meses.
Tudo em minutos úteis: seg a sex, 7h às 17h com 1 h de almoço (540 min por dia), sem os feriados nacionais e do DF.
"""
import collections
import datetime as dt
import json
import statistics

import os

import openpyxl

APP = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))     # a pasta morro-da-cruz/
HTML = APP + '/mapa/PLANTAS_TRECHOS_E_NOS_MORRO_DA_CRUZ_REV02_EQUIPES.html'
REDES = APP + '/planilhas/06.1 - LEV_DMC_MORRO_DA_CRUZ_REDES_REV01_13-10-26.xlsx'
BM = APP + '/planilhas/BM_MORRO_DA_CRUZ_REV04_IDENTIFICACAO_09-10-26.xlsx'

DIA = 540                                   # minutos úteis por dia (07-12 e 13-17)
OS = dt.date(2026, 10, 1)
PRAZO = dt.date(2028, 3, 23)                # 540 dias corridos contando o dia da OS
REAJUSTE = dt.date(2027, 2, 1)              # aniversário 31/01/2027: o que começa depois entra no reajuste
SUBADUTORA_MM = 200


# ---------------------------------------------------------------- calendário
def pascoa(a):
    """Algoritmo de Meeus/Jones/Butcher."""
    b, c = divmod(a, 100)
    a19 = a % 19
    h = (19 * a19 + b - b // 4 - (b - (b + 8) // 25 + 1) // 3 + 15) % 30
    i, k = divmod(c, 4)
    l = (32 + 2 * (b % 4) + 2 * i - h - k) % 7
    m = (a19 + 11 * h + 22 * l) // 451
    mes, dia = divmod(h + l - 7 * m + 114, 31)
    return dt.date(a, mes, dia + 1)


def feriados(anos):
    f = {}
    for a in anos:
        p = pascoa(a)
        for d, n in [((1, 1), 'Confraternização Universal'), ((4, 21), 'Tiradentes e aniversário de Brasília'),
                     ((5, 1), 'Dia do Trabalho'), ((9, 7), 'Independência'), ((10, 12), 'Nossa Senhora Aparecida'),
                     ((11, 2), 'Finados'), ((11, 15), 'Proclamação da República'), ((11, 20), 'Consciência Negra'),
                     ((11, 30), 'Dia do Evangélico (DF)'), ((12, 25), 'Natal')]:
            f[dt.date(a, *d)] = n
        f[p - dt.timedelta(days=48)] = 'Carnaval'
        f[p - dt.timedelta(days=47)] = 'Carnaval'
        f[p - dt.timedelta(days=2)] = 'Sexta-feira Santa'
        f[p + dt.timedelta(days=60)] = 'Corpus Christi'
    return f


FERIADOS = feriados(range(2026, 2031))


class Calendario:
    def __init__(self, ini=OS, fim=dt.date(2030, 12, 31)):
        self.dias = []
        d = ini
        while d <= fim:
            if d.weekday() < 5 and d not in FERIADOS:
                self.dias.append(d)
            d += dt.timedelta(days=1)
        self.idx = {d: i for i, d in enumerate(self.dias)}

    def w(self, data):
        """Início do primeiro dia útil a partir da data."""
        d = data
        while d not in self.idx:
            d += dt.timedelta(days=1)
        return self.idx[d] * DIA

    def inicio(self, w):
        i, m = divmod(int(w), DIA)
        h = 7 * 60 + m if m < 300 else 13 * 60 + (m - 300)
        return dt.datetime.combine(self.dias[i], dt.time(h // 60, h % 60))

    def fim(self, w):
        w = int(w)
        if w > 0 and w % DIA == 0:
            return dt.datetime.combine(self.dias[w // DIA - 1], dt.time(17, 0))
        i, m = divmod(w, DIA)
        h = 7 * 60 + m if m <= 300 else 13 * 60 + (m - 300)
        return dt.datetime.combine(self.dias[i], dt.time(h // 60, h % 60))

    def data_fim(self, w):
        return self.fim(w).date()


CAL = Calendario()


def meses_do_contrato(ultimo):
    """[(nº, início, fim)] de out/2026 até o mês da data `ultimo` (o 18º mês acaba em 23/03/2028)."""
    out, a, m, n = [], 2026, 10, 1
    while dt.date(a, m, 1) <= ultimo:
        ini = dt.date(a, m, 1)
        prox = dt.date(a + (m == 12), m % 12 + 1, 1)
        fim = prox - dt.timedelta(days=1)
        if n == 18:
            fim = min(fim, PRAZO) if ultimo <= PRAZO else fim
        out.append((n, ini, fim))
        a, m, n = prox.year, prox.month, n + 1
    while len(out) < 18:                     # o BM tem sempre as 18 medições do contrato
        n, ini, fim = len(out) + 1, None, None
        a, m = (out[-1][1].year + (out[-1][1].month == 12), out[-1][1].month % 12 + 1)
        ini = dt.date(a, m, 1)
        fim = dt.date(a + (m == 12), m % 12 + 1, 1) - dt.timedelta(days=1)
        if n == 18:
            fim = PRAZO
        out.append((n, ini, fim))
    return out


# ---------------------------------------------------------------- dados
def ler_dados():
    s = open(HTML, encoding='utf-8').read()
    i = s.find('const DADOS = ') + len('const DADOS = ')
    D, _ = json.JSONDecoder().raw_decode(s[i:])
    wb = openpyxl.load_workbook(REDES, read_only=True, data_only=True)
    rows = list(wb['REDE'].iter_rows(min_row=5, values_only=True))
    h = rows[0]
    rede = {r[h.index('PONTO')]: {k: r[j] for j, k in enumerate(h) if k} for r in rows[1:] if r[h.index('PONTO')]}
    wb = openpyxl.load_workbook(BM, read_only=True, data_only=True)
    itens = []
    for r in wb['BM MORRO DA CRUZ - VERTICAL'].iter_rows(min_row=6, max_row=547, values_only=True):
        if r[1] and r[3]:
            itens.append({'seq': r[0], 'grp': str(r[1]), 'orc': r[2], 'chave': r[3], 'item': str(r[4]),
                          'attrs': list(r[5:14]), 'cod': r[14], 'classe': r[15], 'desc': r[16], 'un': r[17],
                          'q': float(r[18] or 0), 'vu': float(r[19] or 0), 'total': float(r[20] or 0)})
    return D, rede, itens


ehSub = lambda t: t['tipo'] == 'Adutora' or t['de'] >= SUBADUTORA_MM


# ---------------------------------------------------------------- equipes e premissas
EQUIPES = [
    # id, nome, empresa, entrada
    ('ANK1', 'Equipe Adutora (Ankara 1)', 'Ankara', dt.date(2026, 11, 16)),   # a NS de topografia vem antes (03 a 13/11)
    ('E1A', 'Empreiteiro 1 - Equipe A', 'Empreiteiro 1', dt.date(2027, 1, 4)),
    ('E1B', 'Empreiteiro 1 - Equipe B', 'Empreiteiro 1', dt.date(2027, 1, 4)),
    ('E2C', 'Empreiteiro 2 - Equipe C', 'Empreiteiro 2', dt.date(2027, 4, 1)),
    ('E2D', 'Empreiteiro 2 - Equipe D', 'Empreiteiro 2', dt.date(2027, 4, 1)),
    ('ANK2', 'Equipe Caixas e Hidromecânico (Ankara 2)', 'Ankara', dt.date(2027, 3, 1)),
]
NOME = {e[0]: e[1] for e in EQUIPES}
NOME.update({'CANT': 'Canteiro de obras (Ankara)', 'ADM': 'Administração local (Ankara)'})
INICIO_HIDRO = dt.date(2027, 7, 1)
P_PE_ALTO = 60.0                             # m/dia da Equipe Adutora nos 160 mm até o reajuste ("pé alto")
CAIXA_DIAS = {'DF': 1.0, 'VTF': 1.5, 'VM': 2.0, 'VRP': 3.0, 'MV': 3.0, 'HID': 0.5}       # obra civil, dias da equipe
HIDRO_DIAS = {'DF': 0.5, 'VTF': 0.5, 'VM': 1.0, 'VRP': 2.0, 'MV': 2.0, 'HID': 0.5}       # montagem
LIG_ATRASO = 5 * DIA                         # a ligação vem 5 dias úteis atrás da rede, na mesma frente
TOPO_ANTES = 10 * DIA                        # a NS de topografia libera o trecho 2 semanas antes
CADASTRO_DEPOIS = 22 * DIA                   # o cadastro técnico sai até 30 dias corridos depois
TESTE_DIAS, INTERLIG_DIAS, DESMOB_DIAS = 10, 3, 10
# canteiro: em out/2026 só a demarcação, o alambrado, o portão e 2 contêineres (informado em 29/10/2026); o resto em nov-dez
CANT_FEITO = (dt.date(2026, 10, 1), dt.date(2026, 10, 30))
CANT_INST = (dt.date(2026, 11, 3), dt.date(2026, 12, 18))
CANT_ITENS_OUT = {'5.1.2', '1.18.8', '5.14.2', '5.14.3'}     # tapume metálico (alambrado), pilares, base e portão
CONTEINERES = [(dt.date(2026, 10, 1), 2), (dt.date(2026, 11, 3), 4), (dt.date(2027, 1, 4), 6), (dt.date(2027, 4, 1), 6)]
TOPO_INICIO = dt.date(2026, 11, 3)            # a topografia começa em novembro


def base(D, rede):
    T, E = D['trechos'], D['els']
    for i, t in enumerate(T):
        t['i'] = i
        t['sub'] = ehSub(t)
        r = rede.get(t['R'], {})
        t['rede'] = r
        t['dist'] = float(r.get('DIST. CANTEIRO') or 0)
    # zonas da mais perto do canteiro para a mais longe (DIST. CANTEIRO da 06.1)
    dz = collections.defaultdict(list)
    for t in T:
        if not t['sub']:
            dz[t['loc']].append(t['dist'])
    zonas = sorted(dz, key=lambda z: statistics.mean(dz[z]))
    # ligações: só na rede de DE 63 (tê de serviço 63 e tê de sela 63 × 20 do BM), pelo comprimento
    m63 = sum(t['L'] for t in T if t['de'] == 63 and not t['sub'])
    for t in T:
        t['lig'] = 11894 * t['L'] / m63 if (t['de'] == 63 and not t['sub']) else 0.0
    caixas = [j for j, e in enumerate(E) if e['tp'] != 'NÓ']
    return T, E, zonas, caixas


ORDEM_CLASSE = {'PRIMÁRIA': 0, 'SECUNDÁRIA': 1, 'TERCIÁRIA': 2}


def ordem_subadutora(T):
    sub = [t for t in T if t['sub']]
    adj = collections.defaultdict(list)
    for t in sub:
        adj[t['ini']].append(t)
        adj[t['fim']].append(t)
    pontas = [n for n, l in adj.items() if len(l) == 1]
    inicio = max(pontas, key=lambda n: (adj[n][0]['de'], adj[n][0]['tipo'] == 'Adutora'))
    vistos, ordem = set(), []

    def andar(n):
        for t in sorted(adj[n], key=lambda t: -t['de']):
            if t['i'] in vistos:
                continue
            vistos.add(t['i'])
            ordem.append(t)
            andar(t['fim'] if t['ini'] == n else t['ini'])
    import sys
    sys.setrecursionlimit(10000)
    andar(inicio)
    resto = [t for t in sub if t['i'] not in vistos]          # pedaços soltos, na ordem das plantas
    return ordem + sorted(resto, key=lambda t: (t['pb'], t['R']))


# ---------------------------------------------------------------- programação
def programar(D, rede, p):
    """p: produtividade (m/dia) de todas as frentes de rede depois do reajuste. Devolve o plano."""
    T, E, zonas, caixas = base(D, rede)
    ent = {e[0]: CAL.w(e[3]) for e in EQUIPES}
    w_reaj = CAL.w(REAJUSTE)
    plano = {'p': p, 'trechos': {}, 'blocos': [], 'zonas': zonas, 'fila': collections.defaultdict(list)}
    dur = lambda L, prod: max(30, round(L / prod * DIA))       # minutos (mínimo de meia hora)

    def fazer(eq, t, w0, prod, fase):
        w1 = w0 + dur(t['L'], prod)
        plano['trechos'][t['i']] = {'eq': eq, 'ini': w0, 'fim': w1, 'fase': fase}
        plano['fila'][eq].append(t['i'])
        return w1

    # 1) Equipe Adutora no pé alto: os 160 mm das zonas, na ordem das zonas, até o reajuste
    livre = {k: ent[k] for k in ent}
    zona_160 = [t for z in zonas for t in sorted((t for t in T if t['loc'] == z and t['de'] == 160 and not t['sub']),
                                                 key=lambda t: (t['pb'], t['R']))]
    w = livre['ANK1']
    pe_alto = []
    for t in zona_160:
        if w >= w_reaj:
            break
        w = fazer('ANK1', t, w, P_PE_ALTO, '160 mm - pé alto')
        pe_alto.append(t['i'])
    # 2) depois do reajuste, a Subadutora
    w = max(w, w_reaj)
    for t in ordem_subadutora(T):
        w = fazer('ANK1', t, w, p, 'Subadutora')
    livre['ANK1'] = w
    plano['fim_subadutora'] = w
    # 3) blocos (zona, PB) das zonas, para a primeira equipe livre (a Equipe Adutora entra depois da Subadutora)
    feitos = set(pe_alto)
    blocos = []
    for z in zonas:
        por_pb = collections.defaultdict(list)
        for t in T:
            if t['loc'] == z and not t['sub'] and t['i'] not in feitos:
                por_pb[t['pb']].append(t)
        for pb in sorted(por_pb, key=lambda pb: statistics.mean(t['dist'] for t in por_pb[pb])):
            lst = sorted(por_pb[pb], key=lambda t: (ORDEM_CLASSE.get(t['rede'].get('CLASSIFICAÇÃO EXECUTIVA'), 3),
                                                    -t['de'], t['R']))
            blocos.append({'zona': z, 'pb': pb, 'trechos': [t['i'] for t in lst]})
    rede_eqs = ['E1A', 'E1B', 'E2C', 'E2D', 'ANK1']
    ultima_zona = {}
    for b in blocos:
        eq = min(rede_eqs, key=lambda k: (livre[k], ultima_zona.get(k) != b['zona'], rede_eqs.index(k)))
        w = livre[eq]
        for i in b['trechos']:
            w = fazer(eq, T[i], w, p, 'Rede da zona')
        livre[eq] = w
        ultima_zona[eq] = b['zona']
        b['eq'] = eq
        tr = [plano['trechos'][i] for i in b['trechos']]
        b['ini'], b['fim'] = tr[0]['ini'], tr[-1]['fim']
        b['lig'] = sum(T[i]['lig'] for i in b['trechos'])
    # o pé alto também vira bloco (para as ligações e o Project): por zona
    for z in zonas:
        lst = [i for i in pe_alto if T[i]['loc'] == z]
        if lst:
            tr = [plano['trechos'][i] for i in lst]
            blocos.append({'zona': z, 'pb': '160 mm', 'trechos': lst, 'eq': 'ANK1', 'ini': tr[0]['ini'],
                           'fim': tr[-1]['fim'], 'lig': 0.0, 'pe_alto': True})
    plano['blocos'] = blocos

    # 4) caixas: obra civil depois dos trechos ligados a ela; a montagem hidromecânica depois da obra civil
    def pronto(j):
        e = E[j]
        trs = [plano['trechos'][i]['fim'] for i in e['tr'] if i in plano['trechos']]
        if trs:
            return max(trs)
        z = next((n[:6] for n in e['noms'] if n.startswith('Zona')), None)
        return max((plano['trechos'][t['i']]['fim'] for t in T if t['loc'] == z and t['i'] in plano['trechos']),
                   default=ent['ANK2'])
    cx = {}
    resto = set(caixas)
    w = ent['ANK2']
    ordem_civil = []
    while resto:
        j = min(resto, key=lambda j: (pronto(j), j))
        w0 = max(w, pronto(j))
        w1 = w0 + round(CAIXA_DIAS[E[j]['tp']] * DIA)
        cx[j] = {'civ_ini': w0, 'civ_fim': w1, 'pronto': pronto(j)}
        ordem_civil.append(j)
        w = w1
        resto.discard(j)
    w = CAL.w(INICIO_HIDRO)
    ordem_hidro = []
    for j in sorted(caixas, key=lambda j: (cx[j]['civ_fim'], j)):
        w0 = max(w, cx[j]['civ_fim'])
        w1 = w0 + round(HIDRO_DIAS[E[j]['tp']] * DIA)
        cx[j].update(hid_ini=w0, hid_fim=w1)
        ordem_hidro.append(j)
        w = w1
    plano['caixas'], plano['ordem_civil'], plano['ordem_hidro'] = cx, ordem_civil, ordem_hidro

    # 5) fechamento: teste e desinfecção, interligação e cadastro por zona; Subadutora; desmobilização
    zona_cx = collections.defaultdict(list)
    for j in caixas:
        z = next((n[:6] for n in E[j]['noms'] if n.startswith('Zona')), None)
        if any(T[i]['sub'] for i in E[j]['tr']):
            z = 'SUB'
        zona_cx[z].append(j)
    fech = {}
    sub_fim = plano['fim_subadutora']
    fech['SUB'] = {'teste': (sub_fim, sub_fim + TESTE_DIAS * DIA)}
    hid_sub = max((cx[j]['hid_fim'] for j in zona_cx['SUB']), default=0)
    a = max(fech['SUB']['teste'][1], hid_sub)
    fech['SUB']['interlig'] = (a, a + INTERLIG_DIAS * DIA)
    for z in zonas:
        trs = [plano['trechos'][t['i']] for t in T if t['loc'] == z and not t['sub']]
        ini, fim = min(x['ini'] for x in trs), max(x['fim'] for x in trs)
        teste = (fim, fim + TESTE_DIAS * DIA)
        a = max(teste[1], fech['SUB']['teste'][1], max((cx[j]['hid_fim'] for j in zona_cx[z]), default=0))
        prim = min(x['fim'] for x in trs)
        fech[z] = {'ini': ini, 'fim': fim, 'teste': teste, 'interlig': (a, a + INTERLIG_DIAS * DIA),
                   'cadastro': (prim + CADASTRO_DEPOIS, fim + CADASTRO_DEPOIS),
                   'eq': max((plano['trechos'][t['i']] for t in T if t['loc'] == z and not t['sub']),
                             key=lambda x: x['fim'])['eq']}
    sub_trs = [plano['trechos'][t['i']] for t in T if t['sub']]
    fech['SUB'].update(ini=min(x['ini'] for x in sub_trs), fim=sub_fim, eq='ANK1',
                       cadastro=(min(x['fim'] for x in sub_trs) + CADASTRO_DEPOIS, sub_fim + CADASTRO_DEPOIS))
    ultimo = max([f['interlig'][1] for f in fech.values()] + [f['cadastro'][1] for f in fech.values()] +
                 [b['fim'] + LIG_ATRASO for b in blocos if b['lig'] > 0.5] + [b['fim'] for b in blocos] +
                 [cx[j]['hid_fim'] for j in caixas])          # as ligações só atrasam o bloco que tem ligação
    plano['fechamento'] = fech
    plano['zona_cx'] = zona_cx
    plano['desmob'] = (ultimo, ultimo + DESMOB_DIAS * DIA)
    plano['fim'] = ultimo + DESMOB_DIAS * DIA
    plano['canteiro'] = (CAL.w(CANT_FEITO[0]), CAL.w(CANT_FEITO[1] + dt.timedelta(days=1)))
    plano['canteiro_inst'] = (CAL.w(CANT_INST[0]), CAL.w(CANT_INST[1] + dt.timedelta(days=1)))
    plano['T'], plano['E'] = T, E
    return plano


def resolver_prazo(D, rede, alvo=PRAZO):
    """Menor produtividade (m/dia, de 0,5 em 0,5) que ainda termina a obra até o prazo."""
    lim = CAL.w(alvo + dt.timedelta(days=1))
    cabe = lambda k: programar(D, rede, k / 2)['fim'] <= lim      # k em meios de m/dia: a busca anda só de 0,5 em 0,5
    lo, hi = 80, 400                                              # 40 e 200 m/dia
    if not cabe(hi):
        return None
    if cabe(lo):
        return lo / 2
    while hi - lo > 1:                                            # cabe(hi) e não cabe(lo): hi é a menor que cabe
        mid = (lo + hi) // 2
        if cabe(mid):
            hi = mid
        else:
            lo = mid
    return hi / 2


if __name__ == '__main__':
    D, rede, itens = ler_dados()
    for p in (90.0,):
        pl = programar(D, rede, p)
        print('p', p, 'fim da obra', CAL.data_fim(pl['fim']), 'fim subadutora', CAL.data_fim(pl['fim_subadutora']))
        for eq in ['ANK1', 'E1A', 'E1B', 'E2C', 'E2D']:
            f = pl['fila'][eq]
            tr = pl['trechos']
            print(eq, len(f), round(sum(pl['T'][i]['L'] for i in f)), CAL.data_fim(tr[f[-1]]['fim']))
        cx = pl['caixas']
        print('caixas civil fim', CAL.data_fim(max(c['civ_fim'] for c in cx.values())),
              'hidro fim', CAL.data_fim(max(c['hid_fim'] for c in cx.values())))
        for z, f in pl['fechamento'].items():
            print(z, CAL.data_fim(f['fim']), 'interlig', CAL.data_fim(f['interlig'][1]), f['eq'])
    print('prazo: p mínimo =', resolver_prazo(D, rede))
