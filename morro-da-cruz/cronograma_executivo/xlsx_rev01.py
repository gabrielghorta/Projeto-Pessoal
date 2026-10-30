"""Planilha da REV01 (só o cenário B) no modelo da REV00: as abas de gerar_xlsx (BM por equipe, BM consolidado,
cronograma FF e resumo por equipe) e o que a regra nova pede: a data de operação do reservatório nas PREMISSAS (célula
amarela, em branco, com o nome DATA_RESERVATORIO), o quadro REV00 × REV01, as colunas de ligação e de carga na
PROGRAMAÇÃO e a aba CARGA. As datas de carga, os dias sem carga e os km em carga são fórmulas da data do reservatório."""
import collections
import datetime as dt
import math

import openpyxl
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter as COL
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.properties import PageSetupProperties

import analise_rev01 as AN
import gerar_xlsx as GX
import motor_rev01 as M1
import paleta_rev01 as PAL
from gerar_xlsx import AMARELO, BORDA, CREME, FD, FR, LARANJA, f, fill
from motor import CAL, DIA, NOME, PRAZO

NOME_DATA = 'DATA_RESERVATORIO'
CELULA_DATA = 'PREMISSAS!$C$13'
SEM_DATA = 'após a operação do reservatório'
FDH = 'dd/mm/yyyy hh:mm'
FKM = '#,##0.00'
PROG = "'PROGRAMAÇÃO'"
hx = lambda c: c.lstrip('#').upper()
quebra = Alignment(wrap_text=True, vertical='top')
centro = Alignment(horizontal='center', vertical='center', wrap_text=True)


def em_carga(ref):
    return f'=IF({NOME_DATA}="","{SEM_DATA}",MAX({NOME_DATA},{ref}))'


def cabecalho(ws, r, titulos, c0=1, cor=AMARELO, altura=42):
    for k, t in enumerate(titulos):
        cel = ws.cell(r, c0 + k, t)
        cel.font = f(True)
        cel.fill = fill(cor)
        cel.alignment = centro
        cel.border = BORDA
    ws.row_dimensions[r].height = altura


def secao(ws, r, texto, sub=None):
    ws.cell(r, 1, texto).font = f(True, 12)
    if sub:
        ws.cell(r + 1, 1, sub).font = f(False, 9, '595959', True)
        return r + 2
    return r + 1


def linha(ws, r, vals, c0=1, fmts=None):
    for k, v in enumerate(vals):
        cel = ws.cell(r, c0 + k, v)
        cel.font = f()
        cel.border = BORDA
        if isinstance(v, (dt.datetime, dt.date)):
            cel.number_format = FDH if isinstance(v, dt.datetime) and (v.hour or v.minute) else FD
        if fmts and fmts.get(k):
            cel.number_format = fmts[k]


def altura(texto, largura):
    return max(13, 12.5 * math.ceil(len(texto) / max(1, largura * 1.15)))


# ---------------------------------------------------------------- PREMISSAS
def aba_premissas(wb, cen, quadro, textos):
    ws = wb.create_sheet('PREMISSAS', 0)
    ws['A1'] = f"CRONOGRAMA FÍSICO EXECUTIVO - {AN.REV['id']} - CT 10.239/2026 - MORRO DA CRUZ"
    ws['A1'].font = f(True, 16)
    ws['A2'] = cen['titulo']
    ws['A2'].font = f(True, 12, 'C00000')
    ws['A3'] = (f"{AN.REV['id']} de {AN.REV['data']}, que substitui a {AN.REV['ant']} de {AN.REV['ant_data']} · só o "
                f'cenário B · OS em 01/10/2026 · '
                f'prazo de 540 dias (término contratual em {PRAZO:%d/%m/%Y}) · BM REV04, 06.1 REV01 e 06.2 REV08')
    ws['A3'].font = f(False, 10)
    ws['A5'] = 'PARÂMETROS DO CONTRATO E DO RESERVATÓRIO (as células amarelas alimentam as fórmulas)'
    ws['A5'].font = f(True, 11)
    par = [('BDI material', 0.1528, '0.00%', 'BM REV04 (resumo principal)'),
           ('BDI serviço', 0.2148, '0.00%', 'BM REV04 (resumo principal)'),
           ('Coeficiente K', -0.1404, '0.00%', 'contrato, cláusula 5.2'),
           ('Valor do contrato (P0)', 49878748.6575869, FR, 'BM REV04, célula T556 do vertical'),
           ('Administração local (item 4.1.1)', 1304685.94, FR, 'BM REV04'),
           ('% da administração local sobre o faturamento', '=C11/(48534308.16-C11)', '0.0000%',
            'TR 14.6: ADM ÷ (subtotal 01 do contrato − ADM); subtotal 01 = R$ 48.534.308,16'),
           ('DATA DE OPERAÇÃO DO RESERVATÓRIO', None, FD,
            'EM BRANCO: o reservatório ainda vai ser licitado e não tem data. Enquanto estiver em branco, a ENTRADA EM '
            'CARGA mostra "após a operação do reservatório" e os DIAS SEM CARGA ficam em branco. Preencha a data e a '
            'PROGRAMAÇÃO e a CARGA se recalculam sozinhas. Não há carga provisória.')]
    for i, (n, v, fmt, fonte) in enumerate(par):
        r = 7 + i
        ws.cell(r, 2, n).font = f(r == 13)
        c = ws.cell(r, 3, v)
        c.number_format = fmt
        formula = isinstance(v, str) and v.startswith('=')
        c.font = f(r == 13, 10 if r == 13 else 9, '000000' if formula else '0000FF')
        if not formula:
            c.fill = fill(AMARELO)
        c.border = BORDA
        ws.cell(r, 4, fonte).font = f(r == 13, 8, 'C00000' if r == 13 else '595959', True)
        ws.cell(r, 4).alignment = quebra
    ws.row_dimensions[13].height = 52
    r = 15
    ws.cell(r, 1, f"{AN.REV['ant']} × {AN.REV['id']} (CENÁRIO B, 90 m/dia)").font = f(True, 11)
    r += 1
    for c, t in enumerate(['', f"{AN.REV['ant']} ({AN.REV['ant_data']}): {AN.REV['ant_txt']}",
                           f"{AN.REV['id']} ({AN.REV['data']}): {AN.REV['txt']}"], 2):
        ws.cell(r, c, t).font = f(True)
        ws.cell(r, c).fill = fill(CREME)
        ws.cell(r, c).alignment = centro
        ws.cell(r, c).border = BORDA
    ws.row_dimensions[r].height = 30
    r += 1
    for nome, a, b in quadro:
        ws.cell(r, 2, nome).font = f()
        ws.cell(r, 2).alignment = quebra
        for c, v in ((3, a), (4, b)):
            cel = ws.cell(r, c, v)
            cel.font = f(True)
            cel.alignment = quebra
            if isinstance(v, (dt.date, dt.datetime)):
                cel.number_format = FD
        for c in (2, 3, 4):
            ws.cell(r, c).border = BORDA
        ws.row_dimensions[r].height = max(altura(str(a), 44), altura(str(b), 44), 14)
        r += 1
    r += 1
    for titulo, itens in textos:
        ws.cell(r, 1, titulo).font = f(True, 11)
        r += 1
        for t in itens:
            if isinstance(t, tuple):
                ws.cell(r, 2, t[0]).font = f(True)
                ws.cell(r, 3, t[1]).font = f()
                ws.cell(r, 3).alignment = quebra
                ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4)
                ws.row_dimensions[r].height = altura(t[1], 88)
            else:
                ws.cell(r, 2, '• ' + t).font = f()
                ws.cell(r, 2).alignment = quebra
                ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
                ws.row_dimensions[r].height = altura(t, 190)
            r += 1
        r += 1
    ws.column_dimensions['A'].width = 4
    ws.column_dimensions['B'].width = 100
    ws.column_dimensions['C'].width = 44
    ws.column_dimensions['D'].width = 44
    ws.sheet_view.showGridLines = False


# ---------------------------------------------------------------- PROGRAMAÇÃO
def aba_programacao(wb, pl, livro, A):
    ws = wb.create_sheet('PROGRAMAÇÃO')
    T, E, tr = pl['T'], pl['E'], pl['trechos']
    det, mot, lig = A['det'], A['mot'], pl['ligado']
    ws['A1'] = f"PROGRAMAÇÃO TRECHO A TRECHO (NOMES DA 06.1 REV01) E CAIXA A CAIXA (06.2 REV08) - {AN.REV['id']}"
    ws['A1'].font = f(True, 12)
    ws['A2'] = ('LIGADO AO RESERVATÓRIO EM: quando fica pronto o último trecho do caminho até a saída do reservatório. '
                f"PRONTO PARA CARGA: {AN.REV['pronto_txt']}. ENTRADA EM CARGA = a maior entre a data do reservatório "
                '(PREMISSAS!C13) e o PRONTO PARA CARGA. DIAS SEM CARGA = ENTRADA EM CARGA − TÉRMINO (dias corridos).')
    ws['A2'].font = f(False, 9, '595959', True)
    cab = ['ORDEM NA EQUIPE', 'EQUIPE', 'FASE', 'ZONA / FRENTE', 'PB', 'TRECHO (06.1)', 'CLASSIFICAÇÃO', 'DE (mm)',
           'METODOLOGIA', 'EXTENSÃO (m)', 'LIGAÇÕES (est.)', 'INÍCIO', 'TÉRMINO', 'MEDIÇÃO',
           'LIGADO AO RESERVATÓRIO EM', 'PRONTO PARA CARGA (SEM O RESERVATÓRIO)', 'ENTRADA EM CARGA', 'DIAS SEM CARGA',
           'AO FICAR PRONTO', 'O QUE FALTAVA NO CAMINHO ATÉ O RESERVATÓRIO', 'DISTÂNCIA DA SAÍDA PELO TUBO (m)',
           'SETOR (entre registros)']
    cabecalho(ws, 3, cab)
    for c in (15, 16, 17, 18):
        ws.cell(3, c).fill = fill('9DC3E6')
    r = 4
    for eq in ['ANK1', 'E1A', 'E1B', 'E2C', 'E2D']:
        for n, i in enumerate(pl['fila'][eq], 1):
            t, x = T[i], tr[i]
            pronto_antes = lig[i] > x['fim']
            vals = [n, NOME[eq], x['fase'], 'Subadutora' if t['sub'] else t['loc'], t['pb'], t['R'],
                    t['rede'].get('CLASSIFICAÇÃO EXECUTIVA'), t['de'], t['met'], t['L'], round(t['lig'], 1),
                    CAL.inicio(x['ini']), CAL.fim(x['fim']), livro.meses[livro.mes(x['fim'] - 1)][0],
                    CAL.fim(lig[i]), CAL.fim(det[i]['pronto']), em_carga(f'P{r}'),
                    f'=IF({NOME_DATA}="","",ROUND(Q{r}-M{r},0))',
                    (f"pronto antes de ligado: {(lig[i] - x['fim']) / DIA:.1f} dia(s) útil(eis)".replace('.', ',')
                     + f" · {mot[i]['cat']}") if pronto_antes else 'ligado ao ficar pronto',
                    mot[i]['txt'] if pronto_antes else '', round(pl['agua'][i], 1), A['setor_de'][i]]
            for c, v in enumerate(vals, 1):
                ws.cell(r, c, v)
            ws.cell(r, 10).number_format = '#,##0.00'
            ws.cell(r, 21).number_format = '#,##0'
            for c in (12, 13, 15, 16, 17):
                ws.cell(r, c).number_format = FDH
            if pronto_antes:
                for c in (15, 19, 20):
                    ws.cell(r, c).fill = fill(LARANJA)
            r += 1
    r1 = r - 1
    ws.auto_filter.ref = f'A3:{COL(len(cab))}{r1}'
    r += 2
    ws.cell(r, 1, 'CAIXAS HIDRÁULICAS (DISPOSITIVOS DA 06.2) - EQUIPE CAIXAS E HIDROMECÂNICO (ANKARA 2): '
                  'AS DO CAMINHO DA ÁGUA PRIMEIRO').font = f(True, 11)
    r += 1
    cab2 = ['ORDEM', 'CAIXA (06.2)', 'TIPO', 'ZONA', 'PB', 'TRECHOS PRONTOS EM', 'OBRA CIVIL - INÍCIO',
            'OBRA CIVIL - TÉRMINO', 'HIDROMECÂNICO - INÍCIO', 'HIDROMECÂNICO - TÉRMINO', 'NO CAMINHO DA ÁGUA',
            'FRENTES QUE ESPERAM ESTA CAIXA PARA RECEBER ÁGUA']
    cabecalho(ws, r, cab2)
    r += 1
    cx = pl['caixas']
    for n, j in enumerate(pl['ordem_civil'], 1):
        e = E[j]
        dep = ', '.join(AN.zona_txt(z) for z in sorted(A['depende'].get(j, ()), key=lambda z: (z != 'SUB', z)))
        vals = [n, e['noms'][0] if e['noms'] else e['id'], e['tp'], e.get('zona'), e.get('pb'),
                CAL.fim(cx[j]['pronto']), CAL.inicio(cx[j]['civ_ini']), CAL.fim(cx[j]['civ_fim']),
                CAL.inicio(cx[j]['hid_ini']), CAL.fim(cx[j]['hid_fim']),
                'SIM' if j in pl['caixas_caminho'] else '', dep]
        for c, v in enumerate(vals, 1):
            ws.cell(r, c, v)
            if isinstance(v, dt.datetime):
                ws.cell(r, c).number_format = FD
        r += 1
    for c, w in enumerate([10, 40, 22, 14, 8, 34, 14, 9, 11, 12, 11, 17, 17, 10, 17, 19, 22, 11, 44, 70, 13, 10], 1):
        ws.column_dimensions[COL(c)].width = w
    for row in ws.iter_rows(min_row=4, max_row=r):
        for cel in row:
            if not cel.font.b:
                cel.font = f()
    ws.freeze_panes = 'G4'
    return {'r0': 4, 'r1': r1}


# ---------------------------------------------------------------- CARGA
def aba_carga(wb, pl, A, prog):
    ws = wb.create_sheet('CARGA')
    T, E, tr, cx = pl['T'], pl['E'], pl['trechos'], pl['caixas']
    r0, r1 = prog['r0'], prog['r1']
    Q = f'{PROG}!$Q${r0}:$Q${r1}'
    J = f'{PROG}!$J${r0}:$J${r1}'
    ws['A1'] = f"CARGA - A ÁGUA SAI DO RESERVATÓRIO DA ZONA 11 ({AN.REV['id']})"
    ws['A1'].font = f(True, 14)
    ws['A2'] = ('Um trecho entra em carga na maior entre: a operação do reservatório; o caminho até a saída do reservatório '
                'pronto, testado e desinfetado; as VRPs, registros (VM), ventosas (VTF), descargas (DF) e demais caixas do '
                f"caminho montadas; {AN.REV['carga_txt']}. Ligado ao "
                'reservatório = todo o caminho até a saída construído.')
    ws['A2'].font = f(False, 9, '595959', True)
    ws['A3'] = 'DATA DE OPERAÇÃO DO RESERVATÓRIO (PREMISSAS!C13):'
    ws['A3'].font = f(True)
    ws['F3'] = f'=IF({NOME_DATA}="","em branco (a definir)",{NOME_DATA})'
    ws['F3'].number_format = FD
    ws['F3'].font = f(True, 10, 'C00000')
    # 1) zonas
    r = secao(ws, 5, '1. ZONAS NA ORDEM EM QUE A ÁGUA CHEGA: ALIMENTAÇÃO, CAMINHO ATÉ A SAÍDA DO RESERVATÓRIO E DEPENDÊNCIAS',
              'A Subadutora é o tronco que sai do reservatório pelo TR005. "Fora da zona": os trechos de outras frentes '
              'no caminho da água até a entrada da zona. A seção 4 lista esse caminho trecho a trecho.')
    cab = ['ORDEM DA ÁGUA', 'ZONA / FRENTE', 'EXTENSÃO (km)', 'SAÍDAS DO RESERVATÓRIO QUE A ALIMENTAM', 'PONTO DE ALIMENTAÇÃO (1º trecho)',
           'PEÇA DE ENTRADA', 'DISTÂNCIA DA SAÍDA PELO TUBO (km)', 'PONTOS DE ENTRADA DA ÁGUA NA ZONA',
           'CAMINHO FORA DA ZONA (trechos / km)', 'SUBADUTORA NO CAMINHO (km)', 'CAMINHO PRONTO EM',
           'CAIXAS DO CAMINHO ATÉ A ENTRADA', 'CAIXAS DO CAMINHO MONTADAS EM', 'INÍCIO DA REDE DA ZONA',
           'REDE DA ZONA PRONTA EM', 'ZONA TODA LIGADA AO RESERVATÓRIO EM', AN.REV['col_teste'],
           'INTERLIGAÇÃO (término)', 'OUTRAS FRENTES QUE A ZONA ESPERA (teste e interligação)',
           'PRONTA PARA CARGA (sem o reservatório)', 'ENTRADA EM CARGA', 'DIAS ENTRE A REDE PRONTA E A CARGA']
    cabecalho(ws, r, cab, altura=56)
    for c in (20, 21, 22):
        ws.cell(r, c).fill = fill('9DC3E6')
    r += 1
    linhas_zona = {}
    for n, z in enumerate(A['zonas']):
        cxs = ', '.join(E[j]['noms'][0].split('_')[-1] if E[j]['noms'] else E[j]['id'] for j in z['caixas']) or '—'
        outras = ', '.join(AN.zona_txt(x) for x in z['frentes_caminho']) or '—'
        saidas = ', '.join(f"{AN.curto(x)} DE {T[pl['rede'].idx[x]]['de']}" for x in z['saidas']) or '—'
        vals = ['tronco' if z['z'] == 'SUB' else n, AN.zona_txt(z['z']), z['km'], saidas,
                z['entrada'], z['peca'], z['agua_km'], z['n_entradas'],
                f"{len(z['caminho']) - 1} trechos · {z['caminho_km']:.2f} km".replace('.', ','),
                sum(T[k]['L'] for k in z['sub_no_caminho']) / 1000, CAL.fim(z['caminho_pronto']), cxs,
                CAL.fim(z['caixas_montadas']) if z['caixas_montadas'] else '—', CAL.inicio(z['ini']),
                CAL.fim(z['rede_pronta']), CAL.fim(z['toda_ligada']), CAL.fim(z['teste'][1]), CAL.fim(z['interlig'][1]),
                outras, CAL.fim(z['pronta_carga']), em_carga(f'T{r}'),
                f'=IF({NOME_DATA}="","",ROUND(U{r}-O{r},0))']
        linha(ws, r, vals, fmts={2: FKM, 6: FKM, 9: FKM})
        for c in (11, 13, 14, 15, 16, 17, 18, 20, 21):
            ws.cell(r, c).number_format = FD
        for c in (12, 19):
            ws.cell(r, c).alignment = quebra
        ws.row_dimensions[r].height = 26
        linhas_zona[z['z']] = r
        r += 1
    r += 2
    # 2) mês a mês
    r = secao(ws, r, '2. MÊS A MÊS: KM PRONTOS LIGADOS E NÃO LIGADOS AO RESERVATÓRIO NO FIM DO MÊS, E EM CARGA',
              'Km em carga e sem carga: fórmulas da data do reservatório (em branco, nada está em carga). A coluna I lista '
              'os trechos prontos e não ligados no fim do mês, agrupados pelo trecho que falta no caminho, com o motivo.')
    cab = ['MEDIÇÃO', 'MÊS', 'FIM DO MÊS', 'KM PRONTOS (acumulado)', 'KM PRONTOS LIGADOS AO RESERVATÓRIO',
           'KM PRONTOS NÃO LIGADOS', 'KM EM CARGA', 'KM PRONTOS SEM CARGA', 'TRECHOS PRONTOS E NÃO LIGADOS NO FIM DO MÊS (motivo)']
    cabecalho(ws, r, cab, altura=44)
    r += 1
    rm0 = r
    for m in A['mensal']:
        vals = [f"{m['n']}ª", m['rot'], m['fim'], m['km_prontos'], m['km_prontos'] - m['km_nl'], m['km_nl'],
                f'=IF({NOME_DATA}="",0,SUMPRODUCT(--({Q}<C{r}+1),{J})/1000)', f'=D{r}-G{r}', m['texto'] or '—']
        linha(ws, r, vals, fmts={3: FKM, 4: FKM, 5: FKM, 6: FKM, 7: FKM})
        ws.cell(r, 3).number_format = FD
        ws.cell(r, 9).alignment = quebra
        if m['km_nl'] > 0.0005:
            ws.cell(r, 6).fill = fill(LARANJA)
        ws.row_dimensions[r].height = altura(m['texto'] or '—', 110)
        r += 1
    rm1 = r - 1
    ws.cell(r, 2, 'SOMA EM KM × MÊS').font = f(True)
    ws.cell(r, 6, f'=SUM(F{rm0}:F{rm1})').number_format = FKM
    ws.cell(r, 8, f'=SUM(H{rm0}:H{rm1})').number_format = FKM
    for c in (2, 6, 8):
        ws.cell(r, c).font = f(True)
        ws.cell(r, c).fill = fill(CREME)
        ws.cell(r, c).border = BORDA
    ws.cell(r, 9, 'km × mês prontos sem carga: com a data em branco, toda a rede pronta conta (ninguém tem água).').font = f(False, 8, '595959', True)
    linha_soma = r
    r += 3
    # 3) zonas × meses
    frentes = [z['z'] for z in A['zonas']]
    for titulo, chave in (('3a. KM PRONTOS LIGADOS AO RESERVATÓRIO, POR ZONA, NO FIM DE CADA MÊS', 'lig'),
                          ('3b. KM PRONTOS E NÃO LIGADOS AO RESERVATÓRIO, POR ZONA, NO FIM DE CADA MÊS', 'nl')):
        r = secao(ws, r, titulo)
        cabecalho(ws, r, ['ZONA / FRENTE'] + [f"{m['n']}ª\n{m['rot']}" for m in A['mensal']], altura=30)
        r += 1
        for z in frentes:
            vals = [AN.zona_txt(z)]
            for m in A['mensal']:
                pz = [i for i in m['prontos'] if M1.frente_de(T[i]) == z]
                nl = set(m['nao_ligados'])
                km = sum(T[i]['L'] for i in pz if (i in nl) == (chave == 'nl')) / 1000
                vals.append(round(km, 3))
            linha(ws, r, vals, fmts={k: FKM for k in range(1, len(vals))})
            for k in range(1, len(vals)):
                if chave == 'nl' and vals[k] > 0.0005:
                    ws.cell(r, 1 + k).fill = fill(LARANJA)
            r += 1
        r += 2
    # 4) caminho trecho a trecho
    r = secao(ws, r, '4. O CAMINHO DA ÁGUA ATÉ A ENTRADA DE CADA ZONA, TRECHO A TRECHO (DA SAÍDA DO RESERVATÓRIO PARA FORA)',
              'O caminho mais curto, em metros de tubo, da saída do reservatório até o primeiro trecho da zona.')
    cab = ['ZONA', 'Nº NO CAMINHO', 'TRECHO (06.1)', 'FRENTE', 'DE (mm)', 'EXTENSÃO (m)', 'EQUIPE', 'TÉRMINO',
           'CAIXAS NESTE TRECHO (06.2)']
    cabecalho(ws, r, cab, altura=30)
    r += 1
    R = pl['rede']
    for z in A['zonas']:
        for n, k in enumerate(z['caminho'], 1):
            cxs = ', '.join(E[j]['noms'][0] if E[j]['noms'] else E[j]['id'] for j in R.pecas_do_trecho(k)) or ''
            linha(ws, r, [AN.zona_txt(z['z']), n, T[k]['R'], AN.zona_txt(M1.frente_de(T[k])), T[k]['de'], T[k]['L'],
                          NOME[tr[k]['eq']], CAL.fim(tr[k]['fim']), cxs], fmts={5: '#,##0.00'})
            ws.cell(r, 8).number_format = FD
            r += 1
    r += 2
    # 5) caixas de que a carga depende
    r = secao(ws, r, '5. CAIXAS NO CAMINHO DA ÁGUA (VRP, VM, VTF, DF, MV, HID): DE QUE A ENTRADA EM CARGA DEPENDE',
              'Montagem a partir de jul/2027, sem antecipar peças (decisão de 30/10/2026).')
    cab = ['CAIXA (06.2)', 'TIPO', 'ZONA', 'OBRA CIVIL - TÉRMINO', 'HIDROMECÂNICO MONTADO EM', 'FRENTES QUE ESPERAM ESTA CAIXA']
    cabecalho(ws, r, cab, altura=30)
    r += 1
    todas = sorted({j for z in A['zonas'] for j in z['caixas_carga']}, key=lambda j: (cx[j]['hid_fim'], j))
    for j in todas:
        e = E[j]
        dep = ', '.join(AN.zona_txt(z) for z in sorted(A['depende'][j], key=lambda z: (z != 'SUB', z)))
        linha(ws, r, [e['noms'][0] if e['noms'] else e['id'], e['tp'], e.get('zona'), CAL.fim(cx[j]['civ_fim']),
                      CAL.fim(cx[j]['hid_fim']), dep])
        ws.cell(r, 4).number_format = FD
        ws.cell(r, 5).number_format = FD
        r += 1
    r += 2
    # 6) exceções
    nl = A['nl']
    r = secao(ws, r, f'6. TRECHOS PRONTOS ANTES DE LIGADOS AO RESERVATÓRIO: {len(nl)} TRECHOS, '
                     f"{sum(T[x['i']]['L'] for x in nl) / 1000:.2f} KM (EXCEÇÕES À META, COM O MOTIVO)".replace('.', ','),
              'A meta é nenhum trecho pronto fora da ligação, a não ser a frente em execução.')
    cab = ['TRECHO (06.1)', 'ZONA / FRENTE', 'EQUIPE', 'DE (mm)', 'EXTENSÃO (m)', 'TÉRMINO', 'LIGADO EM',
           'ESPERA (dias úteis)', 'CATEGORIA', 'MOTIVO']
    cabecalho(ws, r, cab, altura=30)
    r += 1
    mot = A['mot']
    for x in nl:
        i = x['i']
        linha(ws, r, [T[i]['R'], AN.zona_txt(M1.frente_de(T[i])), NOME[tr[i]['eq']], T[i]['de'], T[i]['L'],
                      CAL.fim(tr[i]['fim']), CAL.fim(pl['ligado'][i]), round(x['espera'] / DIA, 1), mot[i]['cat'],
                      mot[i]['txt']], fmts={4: '#,##0.00', 7: '0.0'})
        ws.cell(r, 6).number_format = FDH
        ws.cell(r, 7).number_format = FDH
        ws.cell(r, 10).alignment = quebra
        ws.row_dimensions[r].height = altura(mot[i]['txt'], 90)
        r += 1
    cats = collections.Counter(mot[x['i']]['cat'] for x in nl)
    km_cat = collections.defaultdict(float)
    for x in nl:
        km_cat[mot[x['i']]['cat']] += T[x['i']]['L'] / 1000
    for cat, n in cats.most_common():
        ws.cell(r, 1, f'{cat}: {n} trechos, {km_cat[cat]:.2f} km'.replace('.', ',')).font = f(True)
        r += 1
    r += 2
    # 7) setores
    r = secao(ws, r, '7. SETORES (A REDE ENTRE REGISTROS VM, VRPs E MACROMEDIDORES)',
              'O teste e a interligação continuam por zona (decisão de 30/10/2026); os setores mostram os pedaços que um '
              'teste por setor isolaria. Setor que passa de uma zona a outra fica com as duas.')
    cab = ['SETOR', 'FRENTES', 'EXTENSÃO (km)', 'TRECHOS', 'ENTRADA DA ÁGUA (peça)', 'DISTÂNCIA DA SAÍDA PELO TUBO (km)',
           'REDE PRONTA EM', 'TODO LIGADO AO RESERVATÓRIO EM', 'PRONTO PARA CARGA (sem o reservatório)', 'ENTRADA EM CARGA']
    cabecalho(ws, r, cab, altura=44)
    ws.cell(r, 10).fill = fill('9DC3E6')
    r += 1
    for s in A['setores']:
        linha(ws, r, [s['n'], ', '.join(AN.zona_txt(z) for z in s['frentes']), s['km'], len(s['membros']), s['entra'],
                      s['agua_km'], CAL.fim(s['pronto']), CAL.fim(s['ligado']), CAL.fim(s['pronta_carga']),
                      em_carga(f'I{r}')], fmts={2: FKM, 5: FKM})
        for c in (7, 8, 9, 10):
            ws.cell(r, c).number_format = FD
        r += 1
    r += 2
    for titulo, chave in (('7a. KM PRONTOS LIGADOS AO RESERVATÓRIO, POR SETOR, NO FIM DE CADA MÊS', 'lig'),
                          ('7b. KM PRONTOS E NÃO LIGADOS AO RESERVATÓRIO, POR SETOR, NO FIM DE CADA MÊS', 'nl')):
        r = secao(ws, r, titulo)
        cabecalho(ws, r, ['SETOR'] + [f"{m['n']}ª\n{m['rot']}" for m in A['mensal']], altura=30)
        r += 1
        for s in A['setores']:
            mb = set(s['membros'])
            vals = [s['n']]
            for m in A['mensal']:
                nlm = set(m['nao_ligados'])
                vals.append(round(sum(T[i]['L'] for i in m['prontos'] if i in mb and (i in nlm) == (chave == 'nl')) / 1000, 3))
            linha(ws, r, vals, fmts={k: FKM for k in range(1, len(vals))})
            for k in range(1, len(vals)):
                if chave == 'nl' and vals[k] > 0.0005:
                    ws.cell(r, 1 + k).fill = fill(LARANJA)
            r += 1
        r += 2
    larg = [12, 22, 13, 20, 34, 30, 13, 12, 22, 12, 13, 30, 13, 13, 13, 14, 13, 13, 22, 14, 18, 13]
    for c, w in enumerate(larg, 1):
        ws.column_dimensions[COL(c)].width = w
    ws.column_dimensions['I'].width = 34
    ws.column_dimensions['J'].width = 40
    ws.freeze_panes = 'C5'
    ws.sheet_view.zoomScale = 85
    return {'linhas_zona': linhas_zona, 'soma': linha_soma}


# ---------------------------------------------------------------- gráficos do resumo com as cores da REV01
def graficos_resumo(wb, livro, res):
    """Troca os gráficos do RESUMO POR EQUIPE: canteiro e administração local numa série cinza e as equipes com as
    cores da REV01 (cada par com CIEDE2000 ≥ 20)."""
    ws = wb['RESUMO POR EQUIPE']
    ws._charts = []
    nm = len(livro.meses)
    r = ws.max_row + 3
    ws.cell(r, 1, 'SÉRIES DO GRÁFICO: FATURAMENTO PROGRAMADO POR MÊS (R$, P0)').font = f(True, 11)
    r += 1
    ws.cell(r, 1, 'SÉRIE').font = f(True)
    for k in range(nm):
        ws.cell(r, 2 + k, f'={COL(2 + k)}4')
        ws.cell(r, 2 + k).alignment = centro
        ws.cell(r, 2 + k).font = f(True)
    r_cab = r
    ws.row_dimensions[r].height = 30
    # linhas do RESUMO: 5 canteiro, 6 ANK1, 7 E1A, 8 E1B, 9 E2C, 10 E2D, 11 ANK2, 12 ADM (ordem de GX.BLOCOS)
    series = [('Canteiro e administração local (Ankara)', lambda L: f'={L}5+{L}12', PAL.CANTEIRO)]
    for n, eq in enumerate(['ANK1', 'E1A', 'E1B', 'E2C', 'E2D', 'ANK2'], 6):
        series.append((NOME[eq], (lambda n: lambda L: f'={L}{n}')(n), PAL.EQUIPES[eq]))
    r0 = r + 1
    for nome, fx, _ in series:
        r += 1
        ws.cell(r, 1, nome).font = f()
        for k in range(nm):
            ws.cell(r, 2 + k, fx(COL(2 + k))).number_format = '#,##0'
    ch = BarChart()
    ch.type = 'col'
    ch.grouping = 'stacked'
    ch.overlap = 100
    ch.gapWidth = 40
    ch.title = f"Faturamento programado por mês e por equipe (R$, P0) - {AN.REV['id']}"
    ch.y_axis.title = 'R$'
    ch.y_axis.numFmt = '#,##0'
    ch.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill='E0E0E0'))
    ch.add_data(Reference(ws, min_col=1, max_col=1 + nm, min_row=r0, max_row=r), from_rows=True, titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=2, max_col=1 + nm, min_row=r_cab, max_row=r_cab))
    for s, (_, _, cor) in zip(ch.series, series):
        s.graphicalProperties.solidFill = hx(cor)
        s.graphicalProperties.line.solidFill = 'FFFFFF'
    ch.legend.position = 'b'
    ch.height, ch.width = 11, 30
    ws.add_chart(ch, f'A{r + 2}')
    ln = LineChart()
    ln.title = f"Faturamento acumulado programado (% do contrato, P0) - {AN.REV['id']}"
    ln.y_axis.numFmt = '0%'
    ln.y_axis.scaling.min = 0
    ln.y_axis.scaling.max = 1
    ln.add_data(Reference(ws, min_col=1, max_col=1 + nm, min_row=res['acum_pct'], max_row=res['acum_pct']),
                from_rows=True, titles_from_data=True)
    ln.set_categories(Reference(ws, min_col=2, max_col=1 + nm, min_row=4, max_row=4))
    ln.series[0].graphicalProperties.line.solidFill = '374151'
    ln.series[0].graphicalProperties.line.width = 25400
    ln.legend = None
    ln.height, ln.width = 9, 30
    ws.add_chart(ln, f'A{r + 25}')


def gerar(caminho, cen, pl, livro, origem, itens, A, quadro, textos):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    Q = GX.quantidades(itens, livro)
    aba_premissas(wb, cen, quadro, textos)
    bm = GX.aba_bm(wb, cen, itens, Q, livro, origem)
    GX.aba_consolidado(wb, itens, bm, livro)
    GX.aba_ff(wb, bm, livro)
    res = GX.aba_resumo(wb, bm, livro, pl, cen)
    graficos_resumo(wb, livro, res)
    prog = aba_programacao(wb, pl, livro, A)
    carga = aba_carga(wb, pl, A, prog)
    for nome, ref in ((NOME_DATA, CELULA_DATA), ('KM_MES_SEM_CARGA', f"CARGA!$H${carga['soma']}")):
        wb.defined_names[nome] = DefinedName(nome, attr_text=ref)
    for ws in wb.worksheets:
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        ws.page_setup.orientation = 'landscape'
    wb.save(caminho)
    return {'bm': bm, 'prog': prog, 'carga': carga}
